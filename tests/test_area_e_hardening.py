"""AREA E hardening: adversarial Excel/HTML layouts (one regression test per issue found).

Fixtures: ``tests/fixtures/excel/adversarial.xlsx`` and ``banner_html.xls``
(built by ``make_fixtures.py``), plus small workbooks built in each test.
"""

from __future__ import annotations

import datetime as dt
import json
import shutil
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from conftest import run_cli
from datanonymous.config import Settings
from datanonymous.generate.rng import stream
from datanonymous.io import excel_html, excelio, registry
from datanonymous.io import excel_read as xr
from datanonymous.io import excel_write as xw
from datanonymous.io.excel_read import SheetScan, XCell
from datanonymous.io.registry import OutTable, ReadRequest, WriteRequest
from datanonymous.model import ColumnProfile, ColumnType, FileLayout, TableProfile

FIX = Path(__file__).resolve().parent / "fixtures" / "excel"
ADV = FIX / "adversarial.xlsx"
BANNER = FIX / "banner_html.xls"
#: invented strings that exist only in titles, side notes, banners, CSS and scripts of the fixtures
SECRETS = ("Quillfeather", "Wilhelmina", "555-0142", "Okafor", "nurse.okafor", "jdoe", "Hartfield", "exportedBy")


# --------------------------------------------------------------------------- helpers
def s(v: str) -> XCell:
    return XCell(v, "s", "General")


def n(v: float) -> XCell:
    return XCell(v, "n", "General")


def rows_of(*rows: list) -> list[tuple[int, dict[int, XCell]]]:
    out = []
    for r, row in enumerate(rows, 1):
        cells = {}
        for c, v in enumerate(row, 1):
            if v is None:
                continue
            cells[c] = v if isinstance(v, XCell) else (s(v) if isinstance(v, str) else n(v))
        out.append((r, cells))
    return out


def scan(rows: list[list], formulas: dict[int, dict[int, str]] | None = None, **kw):
    sc = SheetScan("S", **kw)
    for r, cells in rows_of(*rows):
        sc.feed(r, cells)
    sx = xr.SheetXml()
    for r, cols in sorted((formulas or {}).items()):
        sx.formula_tail.append((r, cols))
    return sc.finish(sx)


def read(path: Path, fmt: str = "xlsx", settings: Settings | None = None, purpose: str = "source"):
    return registry.get(fmt).read(ReadRequest(path, path.name, fmt, None, settings or Settings(mode="fidelity"), purpose=purpose))


def out_tables(rr) -> list[OutTable]:
    tables = []
    for t in rr.tables:
        cols = [ColumnProfile(i, c.name, ColumnType.FACTOR, native=c.native) for i, c in enumerate(t.columns)]
        tp = TableProfile(t.name, t.header, t.n_rows, cols, layout=dict(t.layout))
        tables.append(OutTable(t.name, list(t.header), [c.values for c in t.columns], list(range(len(t.columns))),
                               t.n_rows, dict(t.layout), tp))
    return tables


def write(rr, path: Path, fmt: str = "xlsx", settings: Settings | None = None) -> None:
    req = WriteRequest(path, rr.layout, out_tables(rr), settings or Settings(mode="fidelity"), seed=1,
                       rng_for=lambda purpose: stream(1, "test", purpose))
    registry.get(fmt).write(req)


def all_output_text(*dirs: Path) -> str:
    parts = []
    for d in dirs:
        for p in d.rglob("*"):
            if not p.is_file():
                continue
            if p.suffix == ".xlsx":
                with zipfile.ZipFile(p) as z:
                    parts.extend(z.read(m).decode("utf-8", "replace") for m in z.namelist())
            else:
                parts.append(p.read_bytes().decode("utf-8", "replace"))
    return "\n".join(parts)


# --------------------------------------------------------------------------- footer detection
def test_first_name_max_in_the_last_row_is_data_not_a_footer() -> None:
    # "Max" is a summary word AND a first name: without a formula or a blank row before it, it is data
    rows = [["first_name", "score", "grade"]] + [[f"Zeb{i:02d}", 50 + i, "09"] for i in range(30)]
    rows += [["Max", 77, "10"]]
    out = scan(rows)
    assert out.n_rows == 31 and not out.footer_rows
    assert out.cols[1].values[-1] == "Max"


def test_summary_word_after_a_blank_row_or_a_total_row_is_a_footer() -> None:
    body = [["id", "score"]] + [[f"00{i}", i] for i in range(6)]
    out = scan(body + [[], ["Average", 2.5]])
    assert out.n_rows == 6 and [r for r, _ in out.footer_rows] == [9]
    out2 = scan(body + [["Total", 15], ["Average", 2.5]])
    assert out2.n_rows == 6 and [r for r, _ in out2.footer_rows] == [8, 9]
    out3 = scan(body + [["Max:", 5]])  # summary word right under the data, no formula: data
    assert out3.n_rows == 7 and not out3.footer_rows


def test_row_formulas_are_data_not_a_footer() -> None:
    # '=SUM(B5:D5)' in every row (a per-row total column) used to turn the last 20 data rows into a footer
    rows = [["id", "q1", "q2", "total"]] + [[f"{i:04d}", 1, 2, 3] for i in range(40)]
    formulas = {r: {4: f"=SUM(B{r}:C{r})"} for r in range(2, 42)}
    out = scan(rows, formulas)
    assert out.n_rows == 40 and not out.footer_rows


def test_running_totals_all_the_way_up_are_data() -> None:
    # a cumulative column over the rows above (each row looks like a column total): a long run is data
    rows = [["id", "amount", "running"]] + [[f"{i:04d}", 1, i] for i in range(30)]
    formulas = {r: {3: f"=SUM($B$2:B{r - 1})"} for r in range(3, 32)}
    out = scan(rows, formulas)
    assert out.n_rows == 30 and not out.footer_rows


@pytest.mark.parametrize("formula, row, want", [
    ("=SUM(B2:B6)", 7, True), ("=SUBTOTAL(9,C2:C6)", 7, True), ("=AVERAGE(B:B)", 7, True),
    ("=SUM(B7:D7)", 7, False), ("=SUM($B$2:B7)", 7, False), ("=SUM(Sheet2!B1:B5)", 7, False),
    ("=B2+B3", 7, False), ("=SUM(B2)", 7, True), (None, 7, False), ("=SUM(Sheet2!B:B)", 7, False),
    ("=ROUND(AVERAGE(E2:E6),2)", 7, True), ("=_xlfn.AGGREGATE(9,6,B2:B6)", 7, True), ("=ROUND(B6,2)", 7, False),
])
def test_column_aggregate(formula, row, want) -> None:
    assert xr.is_column_aggregate(formula, row) is want


def test_formulas_filled_down_returning_blank_are_not_data_rows() -> None:
    blank = XCell(None, "", "General", "=")
    rows = [["id", "score", "pass"]] + [[f"{i:03d}", 50 + i, "Y"] for i in range(5)]
    rows += [[None, None, blank] for _ in range(50)]
    out = scan(rows)
    assert out.n_rows == 5 and out.data_end == 6 and not out.footer_rows


# --------------------------------------------------------------------------- header detection
def test_title_in_another_column_is_not_the_header() -> None:
    rows = rows_of(["Title of the report"], [], [None, "col_0_int"], [None, 663], [None, 122], [None, 894])
    assert xr.detect_header_row(rows) == 3


def test_sparse_first_data_row_does_not_push_the_header_down() -> None:
    hdr = [None, "name", "id1", "id2", "score", "id3"]
    rows = rows_of(["Title"], ["Other title"], hdr, [None, None, "015777", None, -10, None],
                   *[[None, f"Name{i}", f"0{i}1", f"0{i}2", 3.5, f"0{i}3"] for i in range(10)])
    assert xr.detect_header_row(rows) == 3


def test_duplicate_labels_after_a_title_row() -> None:
    rows = rows_of(["Attendance for Ms. Okafor"], [], ["student_id", "date", "date", "status"],
                   *[[f"00{i}", 1.0 * i, 2.0 * i, "P"] for i in range(8)])
    assert xr.detect_header_row(rows) == 3


def test_note_beside_the_table_does_not_hide_the_header() -> None:
    body = [[f"00{i}", "09", 50 + i] for i in range(10)]
    body[2] = body[2] + [None, None, "Note: withdrew 9/14"]
    body[3] = body[3] + [None, None, "Call 555-0142"]
    assert xr.detect_header_row(rows_of(["student_id", "grade", "score"], *body)) == 1


def test_single_column_titles_above_typed_data() -> None:
    rows = rows_of(["Immunization follow-up list"], ["Printed 10/01/2024 by nurse.okafor"], ["student_number"],
                   *[[52000 + i] for i in range(10)])
    assert xr.detect_header_row(rows) == 3


def test_single_column_title_blank_row_then_header() -> None:
    rows = rows_of(["Class roster for room 12"], [], ["name"], ["Smith"], ["Jones"], ["Brown"], ["Smith"])
    assert xr.detect_header_row(rows) == 3


def test_single_column_header_blank_row_then_ids_keeps_the_header() -> None:
    # a zero-padded ID is never preferred as a header (it would be copied verbatim)
    rows = rows_of(["student_id"], [], ["000123"], ["000124"], ["000125"])
    assert xr.detect_header_row(rows) == 1


def test_single_column_dated_titles_stacked_on_the_header() -> None:
    # titles directly above the header (no blank row) with plain text data
    rows = rows_of(["Report for room 12 10/02/2024"], [], ["Printed 10/01/2024 by jdoe"],
                   ["Run on 10/05/2024 nightly"], ["c0_lvl"], ["A"], ["B"], [], ["A"], ["C"])
    assert xr.detect_header_row(rows) == 5


def test_single_column_title_block_then_short_label() -> None:
    rows = rows_of(["Students with open referrals"], ["Confidential - do not share"], [], ["name"],
                   ["Smith"], ["Jones"], ["Brown"])
    assert xr.detect_header_row(rows) == 4


def test_single_column_header_blank_then_text_data_keeps_the_header() -> None:
    rows = rows_of(["name"], [], ["Smith"], ["Jones"], ["Brown"], ["Garcia"])
    assert xr.detect_header_row(rows) == 1


def test_single_column_notes_after_data_like_or_short_labels() -> None:
    out = scan([["student_id"]] + [[f"{i:06d}"] for i in range(8)] + [[], ["Source: nightly SIS export"]])
    assert out.n_rows == 8 and [r for r, _ in out.footer_rows] == [11]
    out2 = scan([["grade"]] + [[g] for g in ("A", "B", "C", "A")] + [[], ["Source: nightly SIS export"]])
    assert out2.n_rows == 4 and [r for r, _ in out2.footer_rows] == [7]
    out3 = scan([["name"]] + [[g] for g in ("Smith", "Jones", "Brown", "Lee")] + [[], ["Garcia"]])
    assert out3.n_rows == 6 and not out3.footer_rows  # a short text value after a blank row is data


def test_summary_rows_that_are_all_total_like_are_data() -> None:
    rows = [["category", "count"], ["Total enrolled", 400], ["Total staff", 40], ["Total schools", 4]]
    out = scan(rows)
    assert out.n_rows == 3 and not out.footer_rows
    out2 = scan(rows + [[], ["Source: state report"]])
    assert out2.n_rows == 3 and [r for r, _ in out2.footer_rows] == [6]


def test_note_rows_exclude_data_shaped_values() -> None:
    assert xr.is_note_row({1: s("Source: nightly export")})
    assert not xr.is_note_row({1: s("000123")}) and not xr.is_note_row({1: s("10/01/2024")})
    assert not xr.is_note_row({1: n(5)}) and not xr.is_note_row({1: s("a"), 2: s("b")})


def test_side_columns() -> None:
    assert xr.side_columns({1, 2, 3, 5}, {1, 2, 3}) == [5]
    assert xr.side_columns({1, 2, 4, 5}, {1, 2, 5}) == []  # a gap inside the table: interior
    assert xr.side_columns({1, 3, 4}, {3, 4}) == [1]
    assert xr.side_columns({1, 2, 3}, {1, 2}) == []  # a headerless column next to the table is a column


def test_notes_like() -> None:
    assert xr.notes_like("Class roster - Mr. Okafor", ["Printed 10/01/2024 by jdoe"])
    assert not xr.notes_like("name", ["Smith", "Jones"])
    assert not xr.notes_like("Course title", ["Algebra", "Biology"])
    assert not xr.notes_like("Course list for the fall", ["Algebra I Honors"] * 6)  # too many rows: a table


# --------------------------------------------------------------------------- writer typing
def test_header_cells_keep_number_date_and_bool_types(tmp_path: Path) -> None:
    import pandas as pd

    src = tmp_path / "h.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["k", dt.datetime(2024, 9, 3), True, 2021, 3.5])
    for i in range(6):
        ws.append([f"k{i}", i, i, i, i])
    wb.save(src)
    rr = read(src, settings=Settings(mode="fidelity", header_row=1))
    assert rr.tables[0].header == ["k", "2024-09-03 00:00:00", "TRUE", "2021", "3.5"]  # datetime format
    out = tmp_path / "o.xlsx"
    write(rr, out)
    a = pd.read_excel(src)
    b = pd.read_excel(out)
    assert list(a.columns) == list(b.columns) == ["k", dt.datetime(2024, 9, 3), True, 2021, 3.5]


def test_mixed_text_and_number_column_keeps_leading_zero_text() -> None:
    nat = {"types": {"s": 5, "n": 5}, "main": "n", "formats": {"n": "00000", "s": "00000"}, "errors": [],
           "value_types": {"02134": "s", "2135": "n"}}
    assert xw.typed_value("02139", nat) == ("02139", "s")  # unseen, not how a number cell reads: text
    assert xw.typed_value("2139", nat) == (2139, "n")
    assert xw.typed_value("0.5", nat) == (0.5, "n")
    assert xw.typed_value("1.50", nat) == ("1.50", "s")
    assert [xw.is_canonical_number(v) for v in ("0", "-3", "2.75", "007", "+5", "1e3", "-0", "x")] == [
        True, True, True, False, False, False, False, False]


def test_trailing_blank_synthetic_rows_are_not_lost(tmp_path: Path) -> None:
    import pandas as pd

    src = tmp_path / "g.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "score"])
    for i in range(6):
        ws.append([f"{i:04d}", i] if i % 2 == 0 else [])
    ws.append(["0099", 9])
    wb.save(src)
    rr = read(src)
    assert rr.tables[0].n_rows == 7
    tables = out_tables(rr)
    # synthetic rows that end with two all-blank rows: Excel would drop them
    tables[0].columns = [np.array(["0001", "0002", "", "0003", "0004", "", ""], dtype=object),
                         np.array(["1", "2", "", "3", "4", "", ""], dtype=object)]
    out = tmp_path / "o.xlsx"
    registry.get("xlsx").write(WriteRequest(out, rr.layout, tables, Settings(mode="fidelity"), seed=1,
                                            rng_for=lambda purpose: stream(1, "test", purpose)))
    assert len(pd.read_excel(src)) == len(pd.read_excel(out)) == 7
    assert xw._row_order([np.array(["a", ""], dtype=object)], 2, {}, False) == [1, 0]
    assert list(xw._row_order([np.array(["a", ""], dtype=object)], 2, {}, True)) == [0, 1]  # footer below
    assert list(xw._row_order([np.array(["a", ""], dtype=object)], 2, {1: {}}, False)) == [0, 1]  # side note


def test_placeholder_month_day_without_year_is_a_valid_date() -> None:
    ph = xw.Placeholders(np.random.default_rng(3), keep=False)
    out = ph.text("withdrew 9/14")
    m, d = out.split()[-1].split("/")
    assert len(out) == len("withdrew 9/14") and 1 <= int(m) <= 12 and 1 <= int(d) <= 12


def test_layout_diff_tolerates_a_main_type_tie() -> None:
    class T:
        def __init__(self, nat):
            self.name = "S"
            self.layout = {}
            self.columns = [type("C", (), {"name": "z", "native": nat})()]

    src = {"columns": [["z", "s", "00000", 0.75]]}
    lay = FileLayout(format="xlsx")
    tie = {"main": "n", "types": {"n": 2, "s": 2}, "formats": {"n": "00000", "s": "00000"}}
    assert excelio.layout_diff(lay, lay, [src], [T(tie)]) == []
    worse = {"main": "n", "types": {"n": 3, "s": 1}, "formats": {"n": "00000"}}
    assert excelio.layout_diff(lay, lay, [src], [T(worse)]) == ["sheet 'S': column 'z': cell type s -> n"]


def test_excluded_column_moves_widths_with_the_columns(tmp_path: Path) -> None:
    rr = read(FIX / "multi_sheet.xlsx")
    tables = out_tables(rr)
    t0 = tables[0]
    keep = [i for i in range(len(t0.header)) if i != 1]  # --exclude last_name (column B, width 18)
    tables[0] = OutTable(t0.name, [t0.header[i] for i in keep], [t0.columns[i] for i in keep], keep, t0.n_rows,
                         t0.layout, t0.profile)
    out = tmp_path / "x.xlsx"
    registry.get("xlsx").write(WriteRequest(out, rr.layout, tables, Settings(mode="fidelity"), seed=1,
                                            rng_for=lambda purpose: stream(1, "test", purpose)))
    rr2 = read(out, purpose="selfcheck")
    assert rr.tables[0].layout["widths"] == [[1, 1, 12.5, False], [2, 2, 18.0, False], [4, 4, 13.0, False],
                                             [6, 6, 15.25, False]]
    # enroll_date (D -> C) and attendance_rate (F -> E) keep their widths; B (grade) has none
    assert rr2.tables[0].layout["widths"] == [[1, 1, 12.5, False], [3, 3, 13.0, False], [5, 5, 15.25, False]]
    assert excelio.layout_diff(rr.layout, rr2.layout, [t.layout for t in rr.tables], rr2.tables) == []


def test_damaged_sheet_part_is_a_read_error_not_an_internal_error(tmp_path: Path) -> None:
    from datanonymous.errors import ReadError

    src = tmp_path / "ok.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    wb.save(src)
    bad = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(bad, "w") as zout:
        for item in zin.infolist():
            body = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                body = body[: len(body) // 2]  # truncated XML
            zout.writestr(item, body)
    with pytest.raises(ReadError) as ei:
        read(bad)
    assert "damaged" in ei.value.message and ei.value.fix


# --------------------------------------------------------------------------- the adversarial workbook
def test_adversarial_workbook_reading() -> None:
    rr = read(ADV)
    t = {x.name: x for x in rr.tables}
    roster = t["Roster"]
    assert roster.layout["header_row"] == 3 and roster.n_rows == 30 and roster.layout["n_footer_rows"] == 0
    assert roster.header == ["student_id", "first_name", "grade", "q1", "q2", "total"]
    assert roster.columns[1].values[-1] == "Max"  # the last data row, not a footer label
    side = roster.layout["_side_cells"]
    assert sorted(side) == [1, 2] and all(list(c) == [8] for c in side.values())
    assert t["FillDown"].n_rows == 25
    assert t["Immunization"].layout["header_row"] == 3 and t["Immunization"].header == ["student_number"]
    assert t["Info"].layout["notes_only"] and t["Info"].columns == []
    codes = [m.code for m in rr.messages]
    assert "excel_side_notes" in codes and "excel_notes_sheet" in codes
    for m in rr.messages:  # messages never quote the notes
        assert not any(x in m.text or x in (m.fix or "") for x in SECRETS), m.text


def test_adversarial_workbook_writing_keeps_rows_and_hides_notes(tmp_path: Path) -> None:
    import pandas as pd

    rr = read(ADV)
    out = tmp_path / "adv.xlsx"
    write(rr, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["Roster"]
    assert ws["H5"].value and ws["H6"].value and ws["G5"].value is None
    assert len(ws["H5"].value) == len("Note: Wilhelmina Quillfeather withdrew 9/14")
    with zipfile.ZipFile(out) as z:
        text = "\n".join(z.read(m).decode("utf-8", "replace") for m in z.namelist())
    assert not [x for x in SECRETS if x in text]
    for sheet, header in (("Roster", 2), ("FillDown", 0), ("Immunization", 2)):
        a = pd.read_excel(ADV, sheet_name=sheet, header=header)
        b = pd.read_excel(out, sheet_name=sheet, header=header)
        assert len(a) == len(b) and list(a.columns) == list(b.columns), sheet
    rr2 = read(out, purpose="selfcheck")
    assert excelio.layout_diff(rr.layout, rr2.layout, [x.layout for x in rr.tables], rr2.tables) == []


@pytest.fixture(scope="module")
def adversarial_run(tmp_path_factory: pytest.TempPathFactory):
    base = tmp_path_factory.mktemp("adversarial_run")
    inp = base / "input"
    inp.mkdir()
    for f in (ADV, BANNER):
        shutil.copy(f, inp / f.name)
    out, rep = base / "out", base / "rep"
    res = run_cli(inp, "--out-dir", out, "--report-dir", rep, "--seed", "11")
    assert res.exit_code == 0, res.output
    return inp, out, rep


def test_adversarial_run_never_copies_notes_titles_or_banners(adversarial_run) -> None:
    _, out, rep = adversarial_run
    text = all_output_text(out, rep)
    assert not [x for x in SECRETS if x in text]
    data = json.loads((rep / "report.json").read_text(encoding="utf-8"))
    assert data["totals"]["FAIL"] == 0


def test_adversarial_run_n_rows_in_equals_n_rows_out(adversarial_run) -> None:
    import pandas as pd

    inp, out, _ = adversarial_run
    for sheet, header in (("Roster", 2), ("FillDown", 0), ("Immunization", 2)):
        a = pd.read_excel(inp / ADV.name, sheet_name=sheet, header=header)
        b = pd.read_excel(out / ADV.name, sheet_name=sheet, header=header)
        assert len(a) == len(b) and list(a.columns) == list(b.columns), sheet
    ws = openpyxl.load_workbook(out / ADV.name)["Roster"]
    assert ws.max_row == 33  # 30 data rows under the header, no row lost to a fake footer


def test_adversarial_run_is_byte_identical_for_the_same_seed(adversarial_run, tmp_path: Path) -> None:
    inp, out, _ = adversarial_run
    out2 = tmp_path / "out2"
    res = run_cli(inp, "--out-dir", out2, "--report-dir", tmp_path / "rep2", "--seed", "11")
    assert res.exit_code == 0, res.output
    for f in (ADV.name, BANNER.name):
        assert (out / f).read_bytes() == (out2 / f).read_bytes(), f


# --------------------------------------------------------------------------- HTML
def test_clean_css_drops_comments_text_and_urls() -> None:
    css = ('/* exported for Okafor */ td { color: red; content: "jdoe"; } @import url("x/jdoe.css");'
           ' .x { background: url(file:///C:/Users/jdoe/a.png); mso-number-format:"0\\.00"; }'
           ' @page { mso-header-data:"&LPrinted by jdoe"; margin: 1in; }')
    out = excel_html.clean_css(css)
    assert "jdoe" not in out and "Okafor" not in out and "url(" not in out
    assert 'mso-number-format:"0\\.00"' in out and "color: red" in out and "margin: 1in" in out
    attrs = excel_html._attrs([("style", "background:url(file:///C:/Users/jdoe/a.png); color:red"),
                               ("title", "Okafor")], "td")
    assert attrs == ' style=" color:red"'  # inline CSS cleaned too; title (free text) dropped


def test_banner_table_is_notes_and_script_in_table_is_dropped(tmp_path: Path) -> None:
    rr = read(BANNER, "html_xls")
    assert [t.n_rows for t in rr.tables] == [0, 24]
    assert rr.tables[0].layout["notes_only"] and rr.tables[0].header == []
    assert rr.tables[1].header == ["id", "grp", "score"]
    assert "html_notes_table" in [m.code for m in rr.messages]
    out = tmp_path / "b.xls"
    write(rr, out, "html_xls")
    text = out.read_text(encoding="utf-8")
    assert not [x for x in SECRETS if x in text]
    assert "<script" not in text and "mso-number-format" in text and "<th>id</th>" in text
    rr2 = read(out, "html_xls", purpose="selfcheck")
    assert [t.n_rows for t in rr2.tables] == [0, 24]
    assert excelio.layout_diff(rr.layout, rr2.layout, [x.layout for x in rr.tables], rr2.tables) == []


def test_reader_messages_and_layout_quote_no_title_text(tmp_path: Path) -> None:
    # the public layout (report) carries counts and positions only
    rr = read(ADV)
    public = json.dumps(rr.layout.to_public())
    assert not [x for x in SECRETS if x in public]
