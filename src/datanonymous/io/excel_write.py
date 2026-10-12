"""Fresh-workbook ``.xlsx`` writer (AREA E helper; design 6 "Writing", section 4, section 10).

The output is always built from scratch with openpyxl in write-only mode (never
copy-and-overwrite: pivot caches, comments, properties and other parts can
hold real data). Per sheet, in the source order and visibility:

* column widths (and hidden columns), the default column width, freeze panes,
  the autofilter range (criteria dropped) and the merged ranges that do not
  touch data rows;
* title/notes rows above the header in their source positions with
  **same-shape placeholders** (letters -> random letters of the same case,
  digits -> random digits, punctuation and spacing kept; dates inside the text
  -> synthetic dates in the same format; e-mail addresses, URLs, phone
  numbers, SSNs, IPs and street addresses inside the text -> the shared
  embedded-PII rules, reserved domains and fresh phone bases (D24, D46, R8-06); never a
  real text, phone or SSN of the run, R8-04: :class:`Placeholders`;
  number/date cells -> random values of the same shape); ``--keep-titles``
  copies them verbatim;
* the header row verbatim (number header cells stay numbers);
* the data rows typed per column from the source storage (``native``,
  :mod:`.excel_storage`): text cells stay text (``'000123'`` and strings
  starting with ``=`` too), numbers are numbers, ISO dates become date cells
  (serial numbers in the source's date format; a date cell whose source format
  is no date format, such as ``General``, was stored as ISO text, ``t="d"``,
  and is written so again in that format: K-14, :func:`set_value`),
  ``TRUE``/``FALSE`` become booleans, error tokens seen in the source become
  error cells; the column's mix of cell types and number formats (50 text and
  50 number cells; ``0%`` beside ``0.00``) is kept by exact quota, a level
  keeps its source type. A value that cannot take the column's type is
  written as text (the self-check then reports the cell-type change);
* footer rows: SUM/COUNT/COUNTA/AVERAGE (and SUBTOTAL 1/2/3/9) formulas over
  same-sheet ranges are recomputed from the written synthetic cells and
  written as static values; any other footer number/date/formula is left
  blank; "Total"-vocabulary labels are kept, other footer text becomes a
  placeholder (``--keep-titles``: verbatim).

Number formats are written with their literal text masked (``io.numfmt``: a
format such as ``"Zelmira Quenby "0`` can quote a name; ``0.0" kg"`` is kept).
Formulas never appear in the output (cached values become static values).
Document properties are neutral and the zip is normalised (fixed timestamps),
so the same seed gives byte-identical files (D26).
"""

from __future__ import annotations

import datetime as dt
import io
import math
import random
import re
import zipfile
from functools import lru_cache
from types import SimpleNamespace
from typing import TYPE_CHECKING, Any, Callable

import numpy as np

from ..errors import WriteError
from .excel_read import SUMMARY_WORDS, XCell, canon_number, register_written
from .excel_storage import (_INT_RE, _NUM_RE, ColumnPlan, convert, is_canonical_number,  # noqa: F401 (re-exported)
                            parse_date, plan_column)
from .numfmt import safe_number_format

if TYPE_CHECKING:  # pragma: no cover
    from .registry import OutTable, WriteRequest

_KEEP_LABEL_RE = re.compile(r"^\s*(?:(?:grand|sub)[\s-]*)?totals?\s*:?\s*$|"
                            r"^\s*(?:" + SUMMARY_WORDS + r")\s*:?\s*$", re.I)
_CORE_XML = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
             b'<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
             b'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
             b'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
             b'<dcterms:created xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z</dcterms:created>'
             b'<dcterms:modified xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z</dcterms:modified>'
             b'</cp:coreProperties>')
_ZIP_TIME = (1980, 1, 1, 0, 0, 0)


# --------------------------------------------------------------------------- typing synthetic strings
def _convert(s: str, t: str, kind: str | None) -> Any:
    return convert(s, t, kind) if t != "e" else None


def _guess_type(s: str, nat: dict[str, Any]) -> str:
    types = nat.get("types") or {}
    st = s.strip()
    if "n" in types and _NUM_RE.fullmatch(st) and ("s" not in types or (st == s and is_canonical_number(st))):
        return "n"
    if "d" in types and parse_date(st, nat.get("date_kind")) is not None:
        return "d"
    if "b" in types and st.upper() in ("TRUE", "FALSE"):
        return "b"
    return "s" if "s" in types else (nat.get("main") or "s")


def typed_value(s: Any, nat: dict[str, Any] | None) -> tuple[Any, str]:
    """(Python value, cell type) for one synthetic string under the column's source storage, alone.

    A single-cell view (the value's source type, else its likely type); the
    writer types whole columns with :func:`.excel_storage.plan_column`, which
    keeps the source's mix of types and number formats by exact quota.
    """
    if s is None:
        return None, ""
    if not isinstance(s, str):
        s = str(s)
    if s == "":
        return None, ""
    if not nat:
        return s, "s"
    if s in (nat.get("errors") or ()):
        return s, "e"
    vt = nat.get("value_types")
    if vt is not None:
        t = vt.get(s) or _guess_type(s, nat)
    else:
        t = nat.get("main") or "s"
    v = _convert(s, t, nat.get("date_kind"))
    if v is None:
        return s, "s"
    return v, t


# --------------------------------------------------------------------------- placeholders
_MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october",
           "november", "december"]
_MON_RE = "(?:" + "|".join(m[:3] + "(?:" + m[3:] + ")?" if len(m) > 3 else m for m in _MONTHS) + ")"
_SPAN_PATTERNS = [
    ("iso", re.compile(r"(?<!\d)(\d{4})-(\d{2})-(\d{2})(?:([ T])(\d{2}):(\d{2})(?::(\d{2}))?)?(?!\d)")),
    ("mdy", re.compile(r"(?<!\d)(\d{1,2})([/.-])(\d{1,2})\2(\d{4}|\d{2})(?!\d)")),
    ("mon_d_y", re.compile(r"\b(" + _MON_RE + r")(\.?\s+)(\d{1,2})((?:st|nd|rd|th)?,?\s+)(\d{4})\b", re.I)),
    ("d_mon_y", re.compile(r"(?<!\d)(\d{1,2})(\s+|-)(" + _MON_RE + r")(\.?,?\s+|-)(\d{4}|\d{2})\b", re.I)),
    ("mon_y", re.compile(r"\b(" + _MON_RE + r")(\.?,?\s+)(\d{4})\b", re.I)),
    ("time", re.compile(r"(?<!\d)(\d{1,2}):(\d{2})(?::(\d{2}))?(\s?)([AaPp]\.?[Mm]\.?)?(?!\d)")),
    # month/day without a year ("withdrew 9/14"): both parts become 1-12, valid either way round
    ("md", re.compile(r"(?<![\d/])(\d{1,2})/(\d{1,2})(?![\d/])")),
]


#: Draws of one placeholder text before the last one is taken (a redraw happens only when a draw equals a real text,
#: phone number or SSN of the run).
_PLACEHOLDER_ATTEMPTS = 20
#: The column profile the shared embedded-PII synth is built with for text outside the tables (no column: the
#: default ``--min-cell`` for phone area codes, so a title's area code is never kept).
_NO_COLUMN = SimpleNamespace(params={})


class Placeholders:
    """Same-shape placeholder content for title/notes/footer cells (seeded; one mapping per workbook).

    A text's e-mail addresses, URLs and web domains, phone numbers, SSNs, IPv4 addresses and street addresses are
    built by the shared embedded-PII rules (``generate.embedded.SpanSynth``, R8-06: reserved domains example.org /
    .com / .net, fresh phone bases, SSN area 9xx, documentation IPs, D24), its dates become synthetic dates of the same
    format, and the rest is masked (letters -> random letters of the same case, digits -> random digits). A
    placeholder is never a real text of the run, nor holds a real phone number or SSN of the run (R8-04: the run's
    registries, ``privacy.texts`` and ``privacy.digits``, are active while a file is written): it is drawn again.
    """

    def __init__(self, rng: np.random.Generator, keep: bool) -> None:
        self.rng = rng
        self.keep = keep
        self.cache: dict[Any, Any] = {}

    def _rand_date(self) -> dt.datetime:
        start = dt.date(2000, 1, 1).toordinal()
        d = dt.date.fromordinal(start + int(self.rng.integers(0, 365 * 30)))
        return dt.datetime(d.year, d.month, d.day)

    def text(self, s: str) -> str:
        if self.keep:
            return s
        hit = self.cache.get(("s", s))
        if hit is None:
            hit = self.cache[("s", s)] = self._text(s)
        return hit

    def _text(self, s: str) -> str:
        from ..generate import embedded
        from ..privacy.digits import holds_real as real_digits
        from ..privacy.texts import holds_real as real_text

        # no random number is drawn for a text without a personal-data shape (the synth is None)
        synth = embedded.SpanSynth.build(_NO_COLUMN, [(s, 1)], self.rng)
        out = s
        for _ in range(_PLACEHOLDER_ATTEMPTS):
            out = synth.rewrite(s, self._shape) if synth is not None else self._shape(s)
            # the last check of every generated text (a no-op unless a shape is unsafe): reserved domains, invalid SSNs ...
            out = embedded.guard(out, random.Random(len(out)))
            if not (real_text(out) or real_digits(out)):
                break
        return out

    def _shape(self, s: str) -> str:
        """``s`` with its dates replaced by synthetic dates of the same format and the rest masked."""
        spans: list[tuple[int, int, str]] = []
        taken = [False] * len(s)
        for kind, rx in _SPAN_PATTERNS:
            for m in rx.finditer(s):
                if any(taken[m.start():m.end()]):
                    continue
                spans.append((m.start(), m.end(), self._date_text(kind, m)))
                for i in range(m.start(), m.end()):
                    taken[i] = True
        spans.sort()
        out: list[str] = []
        pos = 0
        for a, b, rep in spans:
            out.append(self._mask(s[pos:a]))
            out.append(rep)
            pos = b
        out.append(self._mask(s[pos:]))
        return "".join(out)

    def _mask(self, s: str) -> str:
        if not s:
            return s
        r = self.rng
        chars = []
        for ch in s:
            if "a" <= ch <= "z":
                chars.append(chr(97 + int(r.integers(0, 26))))
            elif "A" <= ch <= "Z":
                chars.append(chr(65 + int(r.integers(0, 26))))
            elif ch.isdigit():
                chars.append(str(int(r.integers(0, 10))))
            elif ch.isalpha():
                chars.append(chr(65 + int(r.integers(0, 26))) if ch.isupper() else chr(97 + int(r.integers(0, 26))))
            else:
                chars.append(ch)
        return "".join(chars)

    def _date_text(self, kind: str, m: re.Match) -> str:
        d = self._rand_date()
        g = m.groups()

        def pad(val: int, src: str) -> str:
            return f"{val:0{len(src)}d}" if len(src) > 1 else str(val)

        def month_name(src: str) -> str:
            name = _MONTHS[d.month - 1]
            name = name[:len(src)] if len(src) <= 3 else name
            if src.isupper():
                return name.upper()
            if src[0].isupper():
                return name.capitalize()
            return name

        if kind == "iso":
            out = f"{d.year:04d}-{d.month:02d}-{d.day:02d}"
            if g[3]:
                out += f"{g[3]}{int(self.rng.integers(0, 24)):02d}:{int(self.rng.integers(0, 60)):02d}"
                if g[6]:
                    out += f":{int(self.rng.integers(0, 60)):02d}"
            return out
        if kind == "mdy":
            a = int(self.rng.integers(1, 13))
            b = int(self.rng.integers(1, 13))  # <= 12: valid as m/d and d/m
            year = f"{d.year:04d}" if len(g[3]) == 4 else f"{d.year % 100:02d}"
            return f"{pad(a, g[0])}{g[1]}{pad(b, g[2])}{g[1]}{year}"
        if kind == "mon_d_y":
            return f"{month_name(g[0])}{g[1]}{pad(min(d.day, 28), g[2])}{g[3]}{d.year:04d}"
        if kind == "d_mon_y":
            year = f"{d.year:04d}" if len(g[4]) == 4 else f"{d.year % 100:02d}"
            return f"{pad(min(d.day, 28), g[0])}{g[1]}{month_name(g[2])}{g[3]}{year}"
        if kind == "mon_y":
            return f"{month_name(g[0])}{g[1]}{d.year:04d}"
        if kind == "md":
            return f"{pad(int(self.rng.integers(1, 13)), g[0])}/{pad(int(self.rng.integers(1, 13)), g[1])}"
        # time
        ampm = g[4]
        hour = int(self.rng.integers(1, 13)) if ampm else int(self.rng.integers(0, 24))
        out = f"{pad(hour, g[0])}:{int(self.rng.integers(0, 60)):02d}"
        if g[2]:
            out += f":{int(self.rng.integers(0, 60)):02d}"
        if ampm:
            word = "AM" if self.rng.random() < 0.5 else "PM"
            letters = [c for c in ampm if c.isalpha()]
            word = word if letters[0].isupper() else word.lower()
            word = ampm.replace(letters[0], word[0], 1).replace(letters[1], word[1], 1)
            out += g[3] + word
        return out

    def value(self, x: XCell) -> Any:
        """Placeholder for a non-text cell value (same Python type and shape)."""
        v = x.v
        if self.keep or v is None:
            return v
        key = ("v", type(v).__name__, v if not isinstance(v, float) else repr(v))
        hit = self.cache.get(key)
        if hit is not None:
            return hit
        hit = self.cache[key] = self._value(v)
        return hit

    def _value(self, v: Any) -> Any:
        r = self.rng
        if isinstance(v, bool):
            return bool(r.random() < 0.5)
        if isinstance(v, int):
            s = str(abs(v))
            if len(s) == 1:
                out = int(r.integers(0, 10))
            else:
                out = int(r.integers(10 ** (len(s) - 1), 10 ** len(s)))
            return -out if v < 0 else out
        if isinstance(v, float):
            s = canon_number(abs(v))
            out_s = "".join(str(int(r.integers(1 if i == 0 and ch != "0" else 0, 10))) if ch.isdigit() else ch
                            for i, ch in enumerate(s))
            try:
                out = float(out_s)
            except ValueError:
                out = float(r.integers(0, 100))
            return -out if v < 0 else out
        if isinstance(v, dt.datetime):
            d = self._rand_date()
            if v.hour or v.minute or v.second:
                d = d.replace(hour=int(r.integers(0, 24)), minute=int(r.integers(0, 60)))
            return d
        if isinstance(v, dt.date):
            return self._rand_date().date()
        if isinstance(v, dt.time):
            return dt.time(int(r.integers(0, 24)), int(r.integers(0, 60)))
        if isinstance(v, dt.timedelta):
            hours = max(1, int(abs(v.total_seconds()) // 3600))
            return dt.timedelta(hours=int(r.integers(0, hours + 1)), minutes=int(r.integers(0, 60)))
        return v


# --------------------------------------------------------------------------- footer recomputation
_AGG_RE = re.compile(r"^=\s*(?:_xlfn\.)?(SUM|COUNT|COUNTA|AVERAGE|SUBTOTAL)\s*\((.*)\)\s*$", re.I | re.S)
_RANGE_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)(?::\$?([A-Za-z]{1,3})\$?(\d+))?$")
_COLS_RE = re.compile(r"^\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})$")
_SUBTOTAL = {1: "AVERAGE", 2: "COUNT", 3: "COUNTA", 9: "SUM", 101: "AVERAGE", 102: "COUNT", 103: "COUNTA",
             109: "SUM"}
_NO_VALUE = object()


def _col_idx(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def parse_aggregate(formula: str | None) -> tuple[str, list[tuple[int, int, int, int]]] | None:
    """('SUM'|'COUNT'|'COUNTA'|'AVERAGE', [(min_col, min_row, max_col, max_row)]) or None (not recomputable).

    ``max_row`` 0 means a whole-column reference.
    """
    if not formula:
        return None
    m = _AGG_RE.match(formula.strip())
    if not m:
        return None
    func = m.group(1).upper()
    body = m.group(2)
    if "(" in body or "!" in body or '"' in body:
        return None
    args = [a.strip() for a in body.split(",")]
    if func == "SUBTOTAL":
        try:
            code = int(args[0])
        except (ValueError, IndexError):
            return None
        func = _SUBTOTAL.get(code, "")
        if not func:
            return None
        args = args[1:]
    ranges = []
    for a in args:
        rm = _RANGE_RE.match(a)
        if rm:
            c1, r1 = _col_idx(rm.group(1)), int(rm.group(2))
            c2 = _col_idx(rm.group(3)) if rm.group(3) else c1
            r2 = int(rm.group(4)) if rm.group(4) else r1
            ranges.append((min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2)))
            continue
        cm = _COLS_RE.match(a)
        if cm:
            c1, c2 = _col_idx(cm.group(1)), _col_idx(cm.group(2))
            ranges.append((min(c1, c2), 1, max(c1, c2), 0))
            continue
        return None
    return (func, ranges) if ranges else None


def evaluate_aggregate(func: str, values: list[Any]) -> Any:
    """Excel semantics over the referenced cell values; returns a value, an error token, or _NO_VALUE."""
    nums: list[Any] = []
    has_dates = False
    for v in values:
        if v is None:
            continue
        if isinstance(v, str) and v.startswith("#") and func in ("SUM", "AVERAGE"):
            if v in ("#N/A", "#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NUM!", "#NULL!"):
                return v
        if isinstance(v, bool):
            continue
        if isinstance(v, (int, float)):
            nums.append(v)
        elif isinstance(v, (dt.date, dt.time, dt.timedelta)):
            has_dates = True
    if func == "COUNTA":
        return sum(1 for v in values if v is not None and v != "")
    if func == "COUNT":
        return len(nums) + sum(1 for v in values if isinstance(v, (dt.date, dt.time, dt.timedelta)))
    if has_dates:
        return _NO_VALUE
    if func == "SUM":
        if all(isinstance(x, int) for x in nums):
            return sum(nums)
        return math.fsum(nums)
    if func == "AVERAGE":
        return (math.fsum(nums) / len(nums)) if nums else "#DIV/0!"
    return _NO_VALUE


# --------------------------------------------------------------------------- writing
def write_xlsx(req: "WriteRequest") -> None:
    from openpyxl import Workbook
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    layout = req.layout
    extra = layout.extra or {}
    wb = Workbook(write_only=True)
    if extra.get("date_system") == 1904:
        wb.epoch = CALENDAR_MAC_1904
    rng = req.rng_for("excel-placeholders") if req.rng_for else np.random.default_rng(req.seed)
    ph = Placeholders(rng, bool(getattr(req.settings, "keep_titles", False)))
    written: dict[str, dict[str, Any]] = {}
    active = 0
    for k, table in enumerate(req.tables):
        name = table.name if table.name else f"Sheet{k + 1}"
        ws = wb.create_sheet(title=name)
        tl = table.layout or {}
        ws.sheet_state = tl.get("state") or "visible"
        if extra.get("active_sheet") == table.name:
            active = k

        def storage_rng(col: int, sheet: int = k) -> np.random.Generator:
            # one seeded stream per (sheet, column): which synthetic cells take a minority type/format
            purpose = f"excel-storage:{sheet}:{col}"
            return req.rng_for(purpose) if req.rng_for else np.random.default_rng([req.seed, sheet, col])

        try:
            info = _write_sheet(ws, table, tl, ph, storage_rng)
        except (ValueError, TypeError) as e:
            raise WriteError(f"sheet {name!r} could not be written ({type(e).__name__})",
                             fix="report this as a bug (re-run with -v for details)") from e
        written[name] = info
    if not req.tables:
        wb.create_sheet(title="Sheet1")
    if req.tables and (req.tables[active].layout or {}).get("state", "visible") == "visible":
        wb.active = active
    buf = io.BytesIO()
    wb.save(buf)
    req.path.write_bytes(normalize_package(buf.getvalue()))
    register_written(req.path, written)


def normalize_package(data: bytes) -> bytes:
    """Fixed zip timestamps and neutral core document properties (byte-identical reruns, D26)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = _CORE_XML
            zi = zipfile.ZipInfo(info.filename, date_time=_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, body)
    return out.getvalue()


class _CellMaker:
    """Reusable write-only cells per (column, cell type, number format)."""

    def __init__(self, ws: Any, date_kind: str | None = None) -> None:
        self.ws = ws
        self.date_kind = date_kind
        self.cells: dict[tuple[str, str | None], tuple[Any, str | None]] = {}

    def make(self, v: Any, t: str, fmt: str | None) -> Any:
        from openpyxl.cell import WriteOnlyCell

        hit = self.cells.get((t, fmt))
        if hit is None:
            c = WriteOnlyCell(self.ws)
            safe = safe_number_format(fmt)  # never a source literal (io.numfmt)
            if safe and safe != "General":
                c.number_format = safe
            hit = self.cells[(t, fmt)] = (c, safe)
        return set_value(hit[0], v, t, hit[1], self.date_kind)


class _IsoDates:
    """A worksheet (or its workbook) as openpyxl's cell writer sees it from one cell, with ``iso_dates`` on.

    The writer stores a date cell as ISO 8601 text (``t="d"``) when ``cell.parent.parent.iso_dates`` is set,
    else as a serial number; every other attribute is the real worksheet's / workbook's (styles, number
    formats). Only the date cells :func:`set_value` keeps in a format that is no date format see it (K-14).
    """

    def __init__(self, target: Any, parent: Any = None) -> None:  # noqa: ANN401 - an openpyxl sheet or book
        object.__setattr__(self, "_target", target)
        object.__setattr__(self, "parent", parent)

    iso_dates = True

    def __getattr__(self, name: str) -> Any:  # noqa: ANN401 - delegated
        return getattr(object.__getattribute__(self, "_target"), name)


def _iso_view(ws: Any) -> Any:  # noqa: ANN401 - an openpyxl worksheet
    """The worksheet seen through a workbook with ``iso_dates`` on (one view per worksheet)."""
    view = getattr(ws, "_datanonymous_iso", None)
    if view is None:
        view = _IsoDates(ws, _IsoDates(ws.parent))
        ws._datanonymous_iso = view
    return view


@lru_cache(maxsize=256)
def _date_format(fmt: str | None) -> bool:
    """``fmt`` is a date/time number format (openpyxl reads a serial number in it as a date)."""
    from openpyxl.styles.numbers import is_date_format

    return bool(is_date_format(fmt or "General"))


def set_value(c: Any, v: Any, t: str, fmt: str | None = None, date_kind: str | None = None) -> Any:
    """Assign ``v`` as cell type ``t``. ``fmt``: the (safe) number format the cell must keep.

    K-14: openpyxl stamps its default date format on a date assigned to a cell whose format is no date format
    (``General``). A source date cell can hold such a format only when it was stored as ISO 8601 text
    (``t="d"``; a serial number in ``General`` is a number to every reader), so the cell keeps the source's
    format and is written as ISO text too (a date, not a date-time, when the column holds dates only).
    """
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.utils.exceptions import IllegalCharacterError

    if t == "d" and isinstance(v, (dt.date, dt.time)) and not _date_format(fmt):
        if date_kind == "date" and isinstance(v, dt.datetime) and v.time() == dt.time():
            v = v.date()
        c.value = v
        c.number_format = fmt or "General"
        if not isinstance(c.parent, _IsoDates):  # a reused cell keeps its view
            c.parent = _iso_view(c.parent)
        return c
    try:
        c.value = v
    except IllegalCharacterError:
        c.value = ILLEGAL_CHARACTERS_RE.sub("", v)
    if t == "s":
        c.data_type = "s"
    elif t == "e":
        c.data_type = "e"
    return c


def _new_cell(ws: Any, v: Any, t: str, fmt: str | None, date_kind: str | None = None) -> Any:
    from openpyxl.cell import WriteOnlyCell

    c = WriteOnlyCell(ws)
    fmt = safe_number_format(fmt)  # never a source literal (io.numfmt)
    if fmt and fmt != "General":
        c.number_format = fmt
    return set_value(c, v, t, fmt, date_kind)


def _apply_sheet_layout(ws: Any, tl: dict[str, Any], excluded: set[int],
                        col_map: dict[int, int] | None = None) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    if tl.get("default_width"):
        ws.sheet_format.defaultColWidth = tl["default_width"]
    for lo, hi, width, hidden in _shift_widths(tl.get("widths") or [], col_map or {}):
        letter = get_column_letter(lo)
        ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, min=lo, max=hi, width=width,
                                                       hidden=bool(hidden), customWidth=width is not None)
    if tl.get("freeze_panes"):
        ws.freeze_panes = tl["freeze_panes"]
    if tl.get("autofilter"):
        ws.auto_filter.ref = _shift_ref(tl["autofilter"], excluded)
    for ref in tl.get("merged") or []:
        ws.merged_cells.add(ref)


def _shift_widths(widths: list[list[Any]], col_map: dict[int, int]) -> list[list[Any]]:
    """Column widths/hidden flags moved with their table columns when ``--exclude`` removed some.

    ``col_map``: source column -> output column (-1 = excluded) for the table's
    columns; columns outside the table keep their position (as title, footer
    and side-note cells do). Without exclusions the source ranges are kept.
    """
    if not col_map or all(src == out for src, out in col_map.items()):
        return widths
    cols: dict[int, tuple[Any, bool]] = {}
    for lo, hi, width, hidden in widths:
        for c in range(lo, min(hi, 16384) + 1):
            oc = col_map.get(c, c)
            if oc >= 1 and oc not in cols:
                cols[oc] = (width, bool(hidden))
    out: list[list[Any]] = []
    for c in sorted(cols):
        width, hidden = cols[c]
        if out and out[-1][1] == c - 1 and out[-1][2] == width and out[-1][3] == hidden:
            out[-1][1] = c
        else:
            out.append([c, c, width, hidden])
    return out


def _shift_ref(ref: str, excluded: set[int]) -> str:
    """Shrink a range's last column by the excluded (``--exclude``) columns inside it."""
    if not excluded:
        return ref
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    try:
        c1, r1, c2, r2 = range_boundaries(ref)
    except (ValueError, TypeError):
        return ref
    if c1 is None or c2 is None:
        return ref
    removed = sum(1 for c in excluded if c1 <= c <= c2)
    if not removed:
        return ref
    return f"{get_column_letter(c1)}{r1}:{get_column_letter(max(c1, c2 - removed))}{r2}"


def _set_dimension(ws: Any, max_row: int, max_col: int) -> None:
    """Write the used range up front (write-only sheets have none, and readers then scan the whole sheet)."""
    from openpyxl.utils import get_column_letter

    ref = f"A1:{get_column_letter(max(1, max_col))}{max(1, max_row)}" if max_row and max_col else "A1"
    ws.calculate_dimension = lambda: ref


def _extent(rows: list[tuple[int, dict[int, XCell]]]) -> tuple[int, int]:
    rr = [r for r, cells in rows if cells]
    cc = [c for _, cells in rows for c in cells]
    return (max(rr) if rr else 0), (max(cc) if cc else 0)


def _write_sheet(ws: Any, table: "OutTable", tl: dict[str, Any], ph: Placeholders,
                 storage_rng: Callable[[int], np.random.Generator] | None = None) -> dict[str, Any]:
    if tl.get("notes_only") or tl.get("header_row") is None:
        _apply_sheet_layout(ws, tl, set())
        rows = sorted(tl.get("_title_rows") or [], key=lambda rc: rc[0])
        _set_dimension(ws, *_extent(rows))
        last = 0
        for r, cells in rows:
            for _ in range(last + 1, r):
                ws.append([])
            ws.append(_layout_row(ws, cells, ph, footer=False))
            last = r
        return {"notes_only": True}
    h = int(tl["header_row"])
    first = int(tl["first_col"])
    src_last = int(tl["last_col"])
    n_src = src_last - first + 1
    src_index = list(table.source_index)
    out_last = first + len(src_index) - 1
    # source column -> output column (or -1 when excluded)
    col_map: dict[int, int] = {}
    pos = {j: k for k, j in enumerate(src_index)}
    for j in range(n_src):
        col_map[first + j] = first + pos[j] if j in pos else -1
    excluded = {c for c, o in col_map.items() if o < 0}
    _apply_sheet_layout(ws, tl, excluded, col_map if excluded else None)
    t_rows, t_cols = _extent(tl.get("_title_rows") or [])
    f_rows, f_cols = _extent(tl.get("_footer_rows") or [])
    s_cols = max((c for cells in (tl.get("_side_cells") or {}).values() for c in cells), default=0)
    _set_dimension(ws, max(t_rows, h + table.n_rows, f_rows), max(t_cols, out_last, f_cols, s_cols))
    natives = []
    for j in src_index:
        p = table.profile.columns[j] if table.profile is not None and j < len(table.profile.columns) else None
        natives.append(p.native if p is not None else None)
    # title rows
    last = 0
    title_values: dict[tuple[int, int], Any] = {}
    for r, cells in sorted(tl.get("_title_rows") or [], key=lambda rc: rc[0]):
        if r >= h:
            continue
        for _ in range(last + 1, r):
            ws.append([])
        ws.append(_layout_row(ws, cells, ph, footer=False, record=(r, title_values)))
        last = r
    for _ in range(last + 1, h):
        ws.append([])
    # header row
    hdr_cells = tl.get("_header_cells") or {}
    hrow: list[Any] = [None] * (out_last)
    for k, (name, j) in enumerate(zip(table.header, src_index)):
        if name == "":
            continue
        t, fmt, hkind = (list(hdr_cells.get(first + j) or ["s", "General"]) + [None])[:3]
        v, tt = _header_value(name, t, hkind)
        hrow[first - 1 + k] = _new_cell(ws, v, tt, fmt if tt != "s" else (fmt if fmt == "@" else None), hkind)
    ws.append(hrow)
    # data rows: each cell's type and number format by exact quota of the source column's mix (excel_storage)
    n = table.n_rows
    plans = [plan_column(col, nat, storage_rng(k) if storage_rng else np.random.default_rng(k))
             for k, (col, nat) in enumerate(zip(table.columns, natives))]
    makers = [_CellMaker(ws, (nat or {}).get("date_kind")) for nat in natives]
    base = [None] * (first - 1)
    side = tl.get("_side_cells") or {}
    for pos, i in enumerate(_row_order([p.strings for p in plans], n, side, bool(tl.get("_footer_rows")))):
        row = list(base)
        for k, plan in enumerate(plans):
            v, t, fmt = plan.cell(i)
            row.append(None if v is None else makers[k].make(v, t, fmt))
        if pos in side:
            _put_cells(row, _layout_row(ws, side[pos], ph, footer=False))
        ws.append(row)
    data_end = h + n
    # footer rows
    footer = sorted(tl.get("_footer_rows") or [], key=lambda rc: rc[0])
    formulas = tl.get("_footer_formulas") or {}
    footer_values: dict[tuple[int, int], Any] = {}
    last = data_end
    resolver = _Resolver(h, first, table, plans, col_map, title_values, footer_values)
    for r, cells in footer:
        for _ in range(last + 1, r):
            ws.append([])
        row: list[Any] = []
        for c in sorted(cells):
            x = cells[c]
            oc = col_map.get(c, c)
            if oc < 0:
                continue
            val, t, fmt = _footer_cell(x, formulas.get(r, {}).get(c), resolver, ph, r)
            if val is None:
                continue
            footer_values[(r, oc)] = val
            while len(row) < oc - 1:
                row.append(None)
            row.append(_new_cell(ws, val, t, fmt))
        ws.append(row)
        last = r
    return {"notes_only": False, "header_row": h, "data_end": data_end, "first_col": first, "last_col": out_last}


def _row_order(cols: list[Any], n: int, side: dict[int, Any], has_footer: bool) -> Any:
    """Data rows in written order. Excel stores no empty row, so a trailing run of all-blank synthetic
    rows would be lost to every reader (N rows in != N rows out): without a footer below the data, the
    last non-blank row is moved to the end instead (the source's last data row is never blank)."""
    if has_footer or n == 0 or (n - 1) in side:
        return range(n)

    def blank(i: int) -> bool:
        return all(c[i] is None or c[i] == "" for c in cols)

    if not blank(n - 1):
        return range(n)
    last = next((i for i in range(n - 2, -1, -1) if not blank(i)), None)
    if last is None:
        return range(n)
    return [*range(last), *range(last + 1, n), last]


def _put_cells(row: list[Any], extra: list[Any]) -> None:
    """Overlay the non-empty cells of ``extra`` (a row list from column 1) on ``row`` (side notes)."""
    for k, cell in enumerate(extra):
        if cell is None:
            continue
        while len(row) <= k:
            row.append(None)
        if row[k] is None:
            row[k] = cell


def _header_value(name: str, t: str, kind: str | None) -> tuple[Any, str]:
    """The header cell as stored in the source: numbers, dates and booleans stay typed (pandas labels)."""
    if t == "n" and _NUM_RE.fullmatch(name):
        return (int(name) if _INT_RE.fullmatch(name) else float(name)), "n"
    if t == "d":
        v = parse_date(name, kind)
        if v is not None:
            return v, "d"
    if t == "b" and name in ("TRUE", "FALSE"):
        return name == "TRUE", "b"
    if t == "e" and name.startswith("#"):
        return name, "e"
    return name, "s"


def _layout_row(ws: Any, cells: dict[int, XCell], ph: Placeholders, footer: bool,
                record: tuple[int, dict[tuple[int, int], Any]] | None = None) -> list[Any]:
    row: list[Any] = []
    for c in sorted(cells):
        x = cells[c]
        if x.v is None:
            continue
        if x.t != "e" and isinstance(x.v, str):
            v: Any = ph.text(x.v) if not (footer and _KEEP_LABEL_RE.match(x.v)) else x.v
            t = "s"
        elif x.t == "e":
            v, t = x.v, "e"
        else:
            v, t = ph.value(x), ("b" if isinstance(x.v, bool) else "d" if x.t == "d" else "n")
        while len(row) < c - 1:
            row.append(None)
        row.append(_new_cell(ws, v, t, x.fmt))
        if record is not None:
            record[1][(record[0], c)] = v
    return row


class _Resolver:
    """Written value at (row, column) of the output sheet, for footer formulas."""

    def __init__(self, h: int, first: int, table: "OutTable", plans: list[ColumnPlan], col_map: dict[int, int],
                 titles: dict[tuple[int, int], Any], footers: dict[tuple[int, int], Any]) -> None:
        self.h = h
        self.first = first
        self.table = table
        self.plans = plans  # the written cells (a number stored as text is not summed, as in Excel)
        self.col_map = col_map
        self.titles = titles
        self.footers = footers
        self.header = {first + k: name for k, name in enumerate(table.header)}

    def values(self, c1: int, r1: int, c2: int, r2: int, upto: int) -> list[Any]:
        out: list[Any] = []
        last_row = r2 if r2 else upto - 1
        n = self.table.n_rows
        for c in range(c1, c2 + 1):
            oc = self.col_map.get(c, c)
            if oc < 0:
                continue
            k = oc - self.first
            for r in range(r1, last_row + 1):
                if r < self.h:
                    out.append(self.titles.get((r, c)))
                elif r == self.h:
                    out.append(self.header.get(oc) or None)
                elif r <= self.h + n:
                    if 0 <= k < len(self.plans):
                        out.append(self.plans[k].cell(r - self.h - 1)[0])
                else:
                    out.append(self.footers.get((r, oc)))
        return out


def _footer_cell(x: XCell, formula: str | None, resolver: _Resolver, ph: Placeholders,
                 row: int) -> tuple[Any, str, str]:
    """(value, type, number format) written for one footer cell in ``row``; value None = blank."""
    agg = parse_aggregate(formula) if formula else None
    if agg is not None:
        func, ranges = agg
        vals: list[Any] = []
        for (c1, r1, c2, r2) in ranges:
            vals.extend(resolver.values(c1, r1, c2, min(r2, row - 1) if r2 else 0, upto=row))
        res = evaluate_aggregate(func, vals)
        if res is _NO_VALUE:
            return None, "", x.fmt
        if isinstance(res, str):
            return res, "e", x.fmt
        return res, "n", x.fmt
    if formula or x.t != "s" or not isinstance(x.v, str):
        return None, "", x.fmt  # other static totals and formulas: blank
    return (x.v if _KEEP_LABEL_RE.match(x.v) else ph.text(x.v)), "s", x.fmt
