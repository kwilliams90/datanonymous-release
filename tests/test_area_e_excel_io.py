"""AREA E reader/writer integration: identity round trips, layout diff, .xls, HTML, extras, options."""

from __future__ import annotations

import dataclasses
import datetime as dt
import io
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

from datanonymous.config import Settings
from datanonymous.detect.context import Detection
from datanonymous.errors import ReadError
from datanonymous.generate.rng import stream
from datanonymous.io import excel_read, excelio, registry
from datanonymous.io.registry import OutTable, ReadRequest, WriteRequest
from datanonymous.model import ColumnProfile, ColumnType, TableProfile

FIX = Path(__file__).resolve().parent / "fixtures" / "excel"
XLSX = FIX / "multi_sheet.xlsx"


def read(path: Path, fmt: str = "xlsx", settings: Settings | None = None, purpose: str = "source"):
    return registry.get(fmt).read(ReadRequest(path, path.name, fmt, None, settings or Settings(mode="fidelity"), purpose=purpose))


def out_tables(rr, drop: dict[str | None, set[int]] | None = None, values=None) -> list[OutTable]:
    tables = []
    for t in rr.tables:
        cols = [ColumnProfile(i, c.name, ColumnType.FACTOR, native=c.native) for i, c in enumerate(t.columns)]
        tp = TableProfile(t.name, t.header, t.n_rows, cols, layout=dict(t.layout))
        keep = [i for i in range(len(t.columns)) if i not in (drop or {}).get(t.name, set())]
        vals = [(values or {}).get((t.name, i), t.columns[i].values) for i in keep]
        tables.append(OutTable(t.name, [t.header[i] for i in keep], vals, keep, t.n_rows, dict(t.layout), tp))
    return tables


def write(rr, path: Path, fmt: str = "xlsx", settings: Settings | None = None, tables=None) -> None:
    req = WriteRequest(path, rr.layout, tables if tables is not None else out_tables(rr), settings or Settings(mode="fidelity"),
                       seed=1, rng_for=lambda purpose: stream(1, "test", purpose))
    registry.get(fmt).write(req)


def roundtrip(src: Path, tmp_path: Path, fmt: str = "xlsx", settings: Settings | None = None, name: str = "out.xlsx"):
    rr = read(src, fmt, settings)
    out = tmp_path / name
    write(rr, out, fmt, settings)
    rr2 = read(out, registry.output_handler(fmt).key, settings, purpose="selfcheck")
    return rr, rr2, out


def cells(path: Path, sheet: str, rows: range) -> list[list[tuple]]:
    """(value, type, format) of the non-empty cells (styled empty cells are not written)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    return [[(c.coordinate, c.value, c.data_type, c.number_format) for c in ws[r] if c.value is not None]
            for r in rows]


# --------------------------------------------------------------------------- reading the fixture
def test_read_fixture_tables_layout_and_natives() -> None:
    rr = read(XLSX)
    assert [t.name for t in rr.tables] == ["Enrollment", "Courses", "Lookup", "Empty"]
    enr = rr.tables[0]
    assert enr.layout["header_row"] == 4 and enr.layout["n_title_rows"] == 2 and enr.n_rows == 40
    assert enr.layout["n_footer_rows"] == 2 and enr.layout["freeze_panes"] == "A5"
    assert enr.layout["autofilter"] == "A4:K44" and enr.layout["merged"] == ["A1:K1"]
    assert enr.header[:4] == ["student_id", "last_name", "grade", "enroll_date"]
    col = {c.name: c for c in enr.columns}
    assert col["student_id"].values[0] == "000120" and col["student_id"].native["main"] == "s"
    assert col["enroll_date"].values[0] == "2023-08-28" and col["enroll_date"].native["formats"]["d"] == "mm/dd/yyyy"
    assert col["gpa"].native["formats"]["n"] == "0.00" and col["attendance_rate"].native["formats"]["n"] == "0.0%"
    assert col["zip_code"].native["formats"]["n"] == "00000" and col["zip_code"].values[0] == "7030"
    assert set(col["econ_dis"].values.tolist()) == {"TRUE", "FALSE", ""}
    assert "#N/A" in col["test_score"].values.tolist() and col["test_score"].native["errors"] == ["#N/A"]
    assert col["points"].native["formulas"] == 40 and col["points"].values[0] == "30"  # cached J5*10
    lookup, empty = rr.tables[2], rr.tables[3]
    assert lookup.layout["state"] == "hidden" and empty.n_rows == 0 and empty.columns == []
    pub = rr.layout.to_public()["extra"]
    assert pub["date_system"] == 1900 and pub["formula_cells"] == 44
    assert pub["extras_dropped"] == {"comments": 1, "data validation": 1, "document properties (author, title ...)": 1,
                                     "hyperlinks": 1, "named ranges": 1}
    # private title/footer content never reaches the public layout
    assert "jsmith" not in str(rr.layout.to_public())
    codes = {m.code for m in rr.messages}
    assert {"excel_extras_dropped", "excel_formulas", "excel_hidden_sheets", "excel_empty_sheets"} <= codes
    assert "excel_formulas_uncached" not in codes


def test_identity_roundtrip_keeps_cells_types_formats(tmp_path: Path) -> None:
    rr, rr2, out = roundtrip(XLSX, tmp_path)
    assert excelio.layout_diff(rr.layout, rr2.layout, [t.layout for t in rr.tables], rr2.tables) == []
    for a, b in zip(rr.tables, rr2.tables):
        assert a.header == b.header and a.n_rows == b.n_rows
        for ca, cb in zip(a.columns, b.columns):
            assert ca.values.tolist() == cb.values.tolist(), ca.name
            assert ca.native["main"] == cb.native["main"] and ca.native["formats"] == cb.native["formats"]
    src = cells(XLSX, "Enrollment", range(4, 46))
    got = cells(out, "Enrollment", range(4, 46))
    assert src[:-1] == got[:-1]  # header and data (formulas -> their cached values)
    assert [(c[0], pytest.approx(c[1]) if isinstance(c[1], float) else c[1], c[2], c[3]) for c in src[-1]] == got[-1]
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Enrollment", "Courses", "Lookup", "Empty"]
    assert wb["Lookup"].sheet_state == "hidden" and wb.active.title == "Enrollment"
    ws = wb["Enrollment"]
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref == "A4:K44"
    assert [str(m) for m in ws.merged_cells.ranges] == ["A1:K1"]
    assert ws.column_dimensions["B"].width == 18 and ws.column_dimensions["F"].width == 15.25
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)  # formulas become static values
    assert ws["A1"].value != "Lincoln Unified School District - Enrollment Report"  # title placeholder
    assert ws["A45"].value == "Total"  # total vocabulary kept
    core = zipfile.ZipFile(out).read("docProps/core.xml")
    assert b"creator" not in core and b"title" not in core and b"2000-01-01T00:00:00Z" in core
    for name in ("comments", "vml", "vbaProject"):
        assert not any(name in n for n in zipfile.ZipFile(out).namelist())


def test_writer_is_byte_identical_for_the_same_seed(tmp_path: Path) -> None:
    rr = read(XLSX)
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    write(rr, a)
    write(rr, b)
    assert a.read_bytes() == b.read_bytes()
    excel_read._WRITTEN.clear()


def test_layout_diff_reports_differences(tmp_path: Path) -> None:
    rr, rr2, _ = roundtrip(XLSX, tmp_path)
    src_tables = [t.layout for t in rr.tables]
    lay2 = dataclasses.replace(rr2.layout, extra=dict(rr2.layout.extra))
    lay2.extra["sheets"] = list(reversed(lay2.extra["sheets"]))
    assert "sheet names/order differ from the source" in excelio.layout_diff(rr.layout, lay2, src_tables, rr2.tables)
    lay3 = dataclasses.replace(rr2.layout, extra=dict(rr2.layout.extra))
    lay3.extra["sheets"] = [dict(s, state="visible") for s in lay3.extra["sheets"]]
    assert "sheet visibility differs from the source" in excelio.layout_diff(rr.layout, lay3, src_tables, rr2.tables)
    t0 = rr2.tables[0]
    t0.layout["header_row"] = 5
    gpa = next(c for c in t0.columns if c.name == "gpa")
    n_gpa = sum(gpa.native["storage"]["n"].values())
    gpa.native = dict(gpa.native, formats={"n": "General"}, storage={"n": {"General": n_gpa}})
    sid = next(c for c in t0.columns if c.name == "student_id")
    sid.native = dict(sid.native, main="n")
    issues = excelio.layout_diff(rr.layout, rr2.layout, src_tables, rr2.tables)
    assert any("header row" in i for i in issues)
    # R3-05: the formats of every type are compared by exact quota (not only the dominant one)
    assert any(f"'gpa': number formats of the number cells {{'General': {n_gpa}}}, expected {{'0.00': {n_gpa}}}"
               in i for i in issues), issues
    assert any("'student_id': cell type s -> n" in i for i in issues)


def test_excluded_column_shifts_and_autofilter_shrinks(tmp_path: Path) -> None:
    rr = read(XLSX)
    tables = out_tables(rr, drop={"Enrollment": {1}})  # drop last_name
    out = tmp_path / "x.xlsx"
    write(rr, out, tables=tables)
    rr2 = read(out, purpose="selfcheck")
    assert rr2.tables[0].header[:2] == ["student_id", "grade"]
    assert openpyxl.load_workbook(out)["Enrollment"].auto_filter.ref == "A4:J44"
    # columns are matched by header, so a dropped column is not a type/format mismatch
    assert excelio.layout_diff(rr.layout, rr2.layout, [t.layout for t in rr.tables], rr2.tables) == []


def test_footer_recomputed_from_synthetic_rows(tmp_path: Path) -> None:
    rr = read(XLSX)
    t = rr.tables[0]
    j_credits = t.header.index("credits")
    new = t.columns[j_credits].values.copy()
    new[:] = "2"
    tables = out_tables(rr, values={("Enrollment", j_credits): new})
    out = tmp_path / "f.xlsx"
    write(rr, out, tables=tables)
    ws = openpyxl.load_workbook(out)["Enrollment"]
    assert ws["J45"].value == 80 and ws["B45"].value == 40  # SUM over the synthetic credits, COUNTA
    assert ws["A47"].value != "Source: SIS nightly export. Questions: jsmith@lincoln-usd.example"
    assert "jsmith" not in ws["A47"].value
    excel_read._WRITTEN.clear()


# --------------------------------------------------------------------------- built workbooks
def build(path: Path, fn) -> Path:
    wb = openpyxl.Workbook()
    fn(wb)
    wb.save(path)
    return path


def test_uncached_formulas_warn_and_static_totals_blank(tmp_path: Path) -> None:
    def fill(wb):
        ws = wb.active
        ws.append(["id", "amount", "note"])
        for i in range(6):
            ws.append([f"A{i}", i * 2, "x"])
        ws.append(["Total", "=SUM(B2:B7)", None])
        ws.append([None, None, None])
        ws.append(["Max", 10, None])
    p = build(tmp_path / "f.xlsx", fill)
    rr = read(p)
    msgs = {m.code: m for m in rr.messages}
    assert "excel_formulas_uncached" in msgs and "open the workbook in Excel" in msgs["excel_formulas_uncached"].fix
    t = rr.tables[0]
    assert t.n_rows == 6 and t.layout["n_footer_rows"] == 2  # 'Total' (formula) and 'Max' (static) rows
    out = tmp_path / "o.xlsx"
    write(rr, out)
    ws = openpyxl.load_workbook(out).active
    assert ws["B8"].value == 30 and ws["B8"].data_type == "n"  # SUM recomputed although never cached
    assert ws["A10"].value == "Max" and ws["B10"].value is None  # other static totals are left blank
    excel_read._WRITTEN.clear()


def test_static_total_left_blank(tmp_path: Path) -> None:
    def fill(wb):
        ws = wb.active
        ws.append(["id", "amount"])
        for i in range(5):
            ws.append([f"A{i}", i])
        ws.append(["Total", 10])
    p = build(tmp_path / "s.xlsx", fill)
    rr = read(p)
    assert rr.tables[0].n_rows == 5
    out = tmp_path / "o.xlsx"
    write(rr, out)
    ws = openpyxl.load_workbook(out).active
    assert ws["A7"].value == "Total" and ws["B7"].value is None


def test_date_system_1904_and_time_kinds_roundtrip(tmp_path: Path) -> None:
    def fill(wb):
        wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
        ws = wb.active
        ws.append(["when", "at", "took", "stamp"])
        for i in range(4):
            ws.append([dt.datetime(2020, 1, 1 + i), dt.time(8, i * 10), dt.timedelta(hours=20 + i * 3),
                       dt.datetime(2020, 1, 1, 9, i)])
            ws.cell(i + 2, 1).number_format = "dd/mm/yyyy"
            ws.cell(i + 2, 2).number_format = "h:mm"
            ws.cell(i + 2, 3).number_format = "[h]:mm:ss"
            ws.cell(i + 2, 4).number_format = "yyyy-mm-dd hh:mm"
    p = build(tmp_path / "d.xlsx", fill)
    rr, rr2, out = roundtrip(p, tmp_path)
    t = rr.tables[0]
    assert rr.layout.extra["date_system"] == 1904
    assert [c.native["date_kind"] for c in t.columns] == ["date", "time", "duration", "datetime"]
    assert t.columns[0].values.tolist()[0] == "2020-01-01" and t.columns[2].values.tolist()[3] == "29:00:00"
    assert [c.values.tolist() for c in rr2.tables[0].columns] == [c.values.tolist() for c in t.columns]
    wb = openpyxl.load_workbook(out)
    assert wb.epoch == openpyxl.utils.datetime.CALENDAR_MAC_1904
    assert wb.active["A2"].value == dt.datetime(2020, 1, 1) and wb.active["C2"].number_format == "[h]:mm:ss"


def test_prose_sheet_is_kept_as_notes(tmp_path: Path) -> None:
    def fill(wb):
        ws = wb.active
        ws.title = "Data"
        ws.append(["id", "score"])
        ws.append(["001", 5])
        ws.append(["002", 7])
        notes = wb.create_sheet("ReadMe")
        notes.append(["This workbook lists every student enrolled at Lincoln ES this fall."])
        notes.append(["Questions go to Jane Smith in the district office."])
    p = build(tmp_path / "n.xlsx", fill)
    rr, rr2, out = roundtrip(p, tmp_path)
    assert rr.tables[1].layout["notes_only"] and rr.tables[1].columns == []
    assert any(m.code == "excel_notes_sheet" for m in rr.messages)
    ws = openpyxl.load_workbook(out)["ReadMe"]
    assert ws["A1"].value and "Lincoln" not in ws["A1"].value and "Jane Smith" not in ws["A2"].value
    assert len(ws["A1"].value) == len("This workbook lists every student enrolled at Lincoln ES this fall.")
    assert excelio.layout_diff(rr.layout, rr2.layout, [t.layout for t in rr.tables], rr2.tables) == []


def test_header_that_looks_like_data_fails_with_fix(tmp_path: Path) -> None:
    def fill(wb):
        ws = wb.active
        ws.append(["jdoe@example.org", "555-123-4567"])
        ws.append(["asmith@example.org", "555-987-6543"])
    p = build(tmp_path / "h.xlsx", fill)
    with pytest.raises(ReadError) as e:
        read(p)
    assert "--header-row" in e.value.fix
    assert "jdoe" not in e.value.message


def test_sheet_header_row_and_range_options(tmp_path: Path) -> None:
    def fill(wb):
        ws = wb.active
        ws.title = "Main"
        ws.append(["Report"])
        ws.append(["a", "b", "c"])
        ws.append([1, 2, 3])
        ws.append([4, 5, 6])
        ws.append(["Other table"])
        ws.append(["x", "y"])
        wb.create_sheet("Second").append(["k", "v"])
    p = build(tmp_path / "o.xlsx", fill)
    rr = read(p, settings=Settings(mode="fidelity", sheet="Second"))
    assert [t.name for t in rr.tables] == ["Second"]
    assert any(m.code == "excel_sheets_not_written" for m in rr.messages)
    with pytest.raises(ReadError):
        read(p, settings=Settings(mode="fidelity", sheet="Nope"))
    rr = read(p, settings=Settings(mode="fidelity", sheet="Main", header_row=2))
    assert rr.tables[0].header == ["a", "b", "c"] and rr.tables[0].layout["n_title_rows"] == 1
    rr = read(p, settings=Settings(mode="fidelity", range=["Main!A2:C4"], sheet="Main"))
    t = rr.tables[0]
    assert t.header == ["a", "b", "c"] and t.n_rows == 2
    assert any(m.code == "excel_range_dropped" for m in rr.messages)


def test_hidden_only_selection_is_written_visible(tmp_path: Path) -> None:
    def fill(wb):
        wb.active.append(["a"])
        h = wb.create_sheet("H")
        h.append(["k", "v"])
        h.append([1, 2])
        h.sheet_state = "hidden"
    p = build(tmp_path / "h.xlsx", fill)
    rr, rr2, out = roundtrip(p, tmp_path, settings=Settings(mode="fidelity", sheet="H"))
    assert openpyxl.load_workbook(out)["H"].sheet_state == "visible"
    assert any(m.code == "excel_hidden_selected" for m in rr.messages)


def test_xlsm_macros_and_chart_extras(tmp_path: Path) -> None:
    from openpyxl.chart import BarChart, Reference

    def fill(wb):
        ws = wb.active
        ws.append(["k", "v"])
        for i in range(4):
            ws.append([f"k{i}", i])
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
        ws.add_chart(ch, "D2")
    p = build(tmp_path / "c.xlsx", fill)
    m = tmp_path / "book.xlsm"
    with zipfile.ZipFile(p) as src, zipfile.ZipFile(m, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"\x00fake vba")
    rr = read(m)
    codes = {x.code: x.text for x in rr.messages}
    assert "xlsm_macros" in codes and "charts (1)" in codes["excel_extras_dropped"]
    out = tmp_path / "book.xlsx"
    write(rr, out)
    names = zipfile.ZipFile(out).namelist()
    assert not any("vba" in nm.lower() or "chart" in nm for nm in names)
    excel_read._WRITTEN.clear()


def test_formatting_and_protection_extras_are_listed(tmp_path: Path) -> None:
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill

    def fill(wb):
        ws = wb.active
        ws.append(["k", "v"])
        for i in range(4):
            ws.append([f"k{i}", i])
        ws.conditional_formatting.add("B2:B5", CellIsRule(operator="lessThan", formula=["2"],
                                                          fill=PatternFill("solid", start_color="FF0000")))
        ws.protection.sheet = True
        ws.oddHeader.left.text = "Printed by jsmith"
    p = build(tmp_path / "x.xlsx", fill)
    rr = read(p)
    extras = rr.layout.extra["extras_dropped"]
    assert extras["conditional formatting"] == 1 and extras["sheet protection"] == 1
    assert extras["page header/footer text"] == 1
    out = tmp_path / "o.xlsx"
    write(rr, out)
    assert b"jsmith" not in zipfile.ZipFile(out).read("xl/worksheets/sheet1.xml")
    ws = openpyxl.load_workbook(out).active
    assert not ws.protection.sheet and not ws.conditional_formatting
    excel_read._WRITTEN.clear()


def test_config_file_accepts_range_and_keep_titles(tmp_path: Path) -> None:
    from datanonymous.config import load_config
    from conftest import build_fidelity_settings as build_settings

    cfg = tmp_path / "c.yaml"
    cfg.write_text("range: ['Main!A3:D40']\nkeep_titles: true\nsheet: Main\n", encoding="utf-8")
    s = build_settings({"range": ("Other!B2:C9",)}, load_config(cfg))
    assert s.range == ["Main!A3:D40", "Other!B2:C9"] and s.keep_titles is True and s.sheet == "Main"


def _date_ctx(values: list[str], native: dict):
    """A real detection context (the hook contract): AREA A's detector reads its value counts."""
    import numpy as np

    from datanonymous.config import ColumnOptions
    from datanonymous.detect.context import DetectContext

    return DetectContext(index=0, name="c", values=np.array(values, dtype=object), file_rel="x.xlsx", table="S",
                         format="xlsx", options=ColumnOptions(), settings=Settings(mode="fidelity"), native=native)


def test_detect_native_routes_date_cells_to_area_a(monkeypatch) -> None:
    # integration: the fake context of the foundation-stub era had no values; AREA A's real detector reads
    # ctx.typed_values, so the hook is exercised with a real DetectContext
    from datanonymous.detect import dates

    nat = {"kind": "excel", "main": "d", "date_kind": "date", "temporal": {"kind": "date", "tz": None, "unit": None}}
    assert excelio.detect_native(type("C", (), {"native": None})()) is None
    assert excelio.detect_native(type("C", (), {"native": {"kind": "excel", "main": "n"}})()) is None
    assert excelio.detect_native(_date_ctx(["", ""], nat)) is None  # no values: the date hook has no opinion
    det = excelio.detect_native(_date_ctx(["2024-03-01", "2024-03-04", "", "2024-04-15"], nat))
    assert det.type is ColumnType.DATE and det.reason.startswith("Excel date cells") and det.source == "native"
    monkeypatch.setattr(dates, "detect_dates", lambda ctx: Detection(ColumnType.DATE, params={"format": "%Y-%m-%d"},
                                                                     reason="ISO dates"))
    det = excelio.detect_native(_date_ctx(["2024-03-01"], nat))
    assert det.type is ColumnType.DATE and det.reason.startswith("Excel date cells")


def test_duration_cells_under_24h_stay_durations(tmp_path: Path) -> None:
    """The reader describes date cells with AREA A's native ``temporal`` kind ([h]:mm cells are durations)."""
    from datanonymous.detect import dates

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["c1", "c2"])
    for i in range(30):
        ws.append([dt.timedelta(hours=i % 20, minutes=7 * i % 60), dt.time(7 + i % 8, 15 * (i % 4))])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "[h]:mm"
        row[1].number_format = "h:mm AM/PM"
    wb.save(tmp_path / "d.xlsx")
    rr = read(tmp_path / "d.xlsx")
    c1, c2 = rr.tables[0].columns
    assert c1.native["temporal"]["kind"] == "duration" and c2.native["temporal"]["kind"] == "time"
    assert dates.native_temporal(c1.native)["kind"] == "duration"
    det = excelio.detect_native(_date_ctx(list(c1.values), c1.native))
    assert det.type is ColumnType.DURATION, det
    assert excelio.detect_native(_date_ctx(list(c2.values), c2.native)).type is ColumnType.TIME


def test_inferred_dtypes_reads_at_the_header_row() -> None:
    rr = read(XLSX)
    dts = excelio.inferred_dtypes(XLSX, rr.layout, "Enrollment")
    assert len(dts) == 11 and dts[3].startswith("datetime64") and dts[4] == "float64"
    assert excelio.inferred_dtypes(XLSX, rr.layout, "Empty") is None


# --------------------------------------------------------------------------- .xls (xlrd) and disguises
def test_read_real_xls_with_xlrd(tmp_path: Path) -> None:
    rr = read(FIX / "real_biff.xls", "xls")
    att, codes = rr.tables
    assert att.layout["header_row"] == 3 and att.layout["n_title_rows"] == 1 and att.n_rows == 20
    assert codes.layout["state"] == "hidden"
    col = {c.name: c for c in att.columns}
    assert col["student_id"].values[0] == "0051000" and col["student_id"].native["main"] == "s"
    assert col["rate"].native["formats"]["n"] == "0.00" and col["last_absence"].native["main"] == "d"
    assert col["chronic"].native["main"] == "b" and col["days_absent"].values[0] == "8"
    rr2_out = tmp_path / "real_biff.xlsx"
    write(rr, rr2_out)
    rr2 = read(rr2_out, purpose="selfcheck")
    assert excelio.layout_diff(rr.layout, rr2.layout, [t.layout for t in rr.tables], rr2.tables) == []
    assert rr.layout.format == "xls" and rr2.layout.format == "xlsx"
    assert [c.values.tolist() for c in rr2.tables[0].columns] == [c.values.tolist() for c in att.columns]


def test_html_table_read_and_mirror(tmp_path: Path) -> None:
    src = FIX / "fake_html.xls"
    rr = read(src, "html_xls")
    t = rr.tables[0]
    assert t.name is None and t.header == ["Student_Number", "LastFirst", "Grade_Level", "Entry_Date", "GPA"]
    assert t.n_rows == 30 and t.layout["header_row"] == 2 and t.columns[0].values[0] == "004100"
    assert rr.layout.encoding == "utf-8"
    assert any(m.code == "html_extras_dropped" for m in rr.messages)
    out = tmp_path / "fake_html.xls"
    write(rr, out, "html_xls")
    data = out.read_bytes()
    text = data.decode("utf-8")
    assert data.count(b"\r\n") == data.count(b"\n")  # CRLF kept
    for secret in ("jsmith", "10.1.2.3", "Roster Report", "x:num", 'name="author"'):
        assert secret not in text
    assert "mso-number-format" in text and "<th>Student_Number</th>" in text
    rr2 = read(out, "html_xls", purpose="selfcheck")
    assert [c.values.tolist() for c in rr2.tables[0].columns] == [c.values.tolist() for c in t.columns]
    assert excelio.layout_diff(rr.layout, rr2.layout, [x.layout for x in rr.tables], rr2.tables) == []
    dts = excelio.inferred_dtypes_html(src, rr.layout, None)  # None without lxml/bs4 (default install)
    assert dts is None or len(dts) == 5


def test_html_refuses_spans_in_data_and_nested_tables(tmp_path: Path) -> None:
    p = tmp_path / "s.xls"
    p.write_text("<table><tr><th>a</th><th>b</th></tr><tr><td rowspan=2>1</td><td>2</td></tr>"
                 "<tr><td>3</td></tr></table>", encoding="utf-8")
    with pytest.raises(ReadError):
        read(p, "html_xls")
    p.write_text("<table><tr><td><table><tr><td>x</td></tr></table></td></tr></table>", encoding="utf-8")
    with pytest.raises(ReadError):
        read(p, "html_xls")


def test_html_without_end_tags(tmp_path: Path) -> None:
    p = tmp_path / "u.xls"
    p.write_bytes(b"<table>\n<tr><td>id<td>grade\n<tr><td>01<td>9\n<tr><td>02<td>10\n<tr><td>03<td>9\n</table>\n")
    rr = read(p, "html_xls")
    assert rr.tables[0].header == ["id", "grade"] and rr.tables[0].n_rows == 3
    out = tmp_path / "o.xls"
    write(rr, out, "html_xls")
    rr2 = read(out, "html_xls", purpose="selfcheck")
    assert [c.values.tolist() for c in rr2.tables[0].columns] == [c.values.tolist() for c in rr.tables[0].columns]
    assert b"</td>" not in out.read_bytes()


def test_xlsx_disguised_as_xls_is_read_as_xlsx(tmp_path: Path) -> None:
    from datanonymous.discover import classify

    p = tmp_path / "book.xls"
    shutil.copy(XLSX, p)
    f, _ = classify(p, "book.xls")
    assert f.format == "xlsx" and f.out_rel_path == "book.xls"
    rr = read(p)
    assert len(rr.tables) == 4


def test_broken_workbook_is_a_read_error(tmp_path: Path) -> None:
    p = tmp_path / "b.xlsx"
    p.write_bytes(b"PK\x03\x04 not really a zip")
    with pytest.raises(ReadError):
        read(p)
    q = tmp_path / "q.xlsx"
    with zipfile.ZipFile(q, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(ReadError):
        read(q)


def test_sheet_xml_layout_without_and_with_formulas(tmp_path: Path) -> None:
    from datanonymous.io import excel_book

    def fill(wb):
        ws = wb.active
        ws.append(["a", "b", "c"])
        for i in range(6):
            ws.append([i, i * 2, i * 3])
        ws.column_dimensions["B"].width = 21.5
        ws.column_dimensions["C"].hidden = True
        ws.freeze_panes = "B2"
        ws["A9"] = "note"
        ws["A11"] = "end of report"
        ws.merge_cells("A11:C11")
    p = build(tmp_path / "l.xlsx", fill)
    with zipfile.ZipFile(p) as z:
        sx = excel_book.read_sheet_xml(z, "xl/worksheets/sheet1.xml")
    assert sx.n_formulas == 0 and sx.freeze_panes == "B2" and sx.merged == ["A11:C11"]
    assert [2, 2, 21.5, False] in sx.widths and any(w[0] == 3 and w[3] for w in sx.widths)
    # shared formulas (as Excel writes a row of totals) are translated from their master cell
    with zipfile.ZipFile(p) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
        parts = {n: z.read(n) for n in z.namelist()}
    xml = xml.replace('<c r="A9" t="inlineStr"><is><t>note</t></is></c>',
                      '<c r="A9" t="inlineStr"><is><t>Total</t></is></c>'
                      '<c r="B9"><f t="shared" ref="B9:C9" si="0">SUM(B2:B7)</f><v>30</v></c>'
                      '<c r="C9"><f t="shared" si="0"/><v>45</v></c>')
    parts["xl/worksheets/sheet1.xml"] = xml.encode()
    q = tmp_path / "shared.xlsx"
    with zipfile.ZipFile(q, "w") as z:
        for name, body in parts.items():
            z.writestr(name, body)
    with zipfile.ZipFile(q) as z:
        sx2 = excel_book.read_sheet_xml(z, "xl/worksheets/sheet1.xml")
    assert sx2.n_formulas == 2 and sx2.n_formulas_uncached == 0
    assert dict(sx2.formula_tail)[9] == {2: "=SUM(B2:B7)", 3: "=SUM(C2:C7)"}
    rr = read(q)
    t = rr.tables[0]
    assert t.n_rows == 6 and t.layout["n_footer_rows"] == 2 and t.layout["merged"] == ["A11:C11"]
    assert not any(m.code == "excel_merged_dropped" for m in rr.messages)  # the merge is in the footer
    out = tmp_path / "o.xlsx"
    tables = out_tables(rr, values={(t.name, 2): t.columns[2].values.copy()})
    tables[0].columns[2][:] = "1"
    write(rr, out, tables=tables)
    ws = openpyxl.load_workbook(out).active
    assert ws["B9"].value == 30 and ws["C9"].value == 6 and ws["A9"].value == "Total"
    excel_read._WRITTEN.clear()


def test_merges_inside_data_rows_are_dropped_with_a_warning(tmp_path: Path) -> None:
    def fill(wb):
        ws = wb.active
        ws.append(["Title"])
        ws.merge_cells("A1:C1")
        ws.append(["a", "b", "c"])
        for i in range(5):
            ws.append([f"x{i}", i, i])
        ws.merge_cells("B3:B4")
    p = build(tmp_path / "m.xlsx", fill)
    rr, rr2, out = roundtrip(p, tmp_path)
    assert rr.tables[0].layout["merged"] == ["A1:C1"]
    assert any(m.code == "excel_merged_dropped" for m in rr.messages)
    assert [str(m) for m in openpyxl.load_workbook(out).active.merged_cells.ranges] == ["A1:C1"]


def test_html_with_two_tables(tmp_path: Path) -> None:
    p = tmp_path / "two.xls"
    p.write_text("<html><body><h1>District roster</h1>"
                 "<table><tr><th>id</th><th>grade</th></tr><tr><td>01</td><td>9</td></tr><tr><td>02</td><td>10</td>"
                 "</tr></table><p>Prepared by J. Doe</p>"
                 "<table><tr><th>school</th><th>n</th></tr><tr><td>A</td><td>5</td></tr></table></body></html>",
                 encoding="utf-8")
    rr = read(p, "html_xls")
    assert [t.name for t in rr.tables] == ["Table 1", "Table 2"]
    out = tmp_path / "o.xls"
    write(rr, out, "html_xls")
    text = out.read_text(encoding="utf-8")
    assert "District roster" not in text and "J. Doe" not in text and text.count("<table>") == 2
    rr2 = read(out, "html_xls", purpose="selfcheck")
    assert [c.values.tolist() for t in rr2.tables for c in t.columns] == [
        c.values.tolist() for t in rr.tables for c in t.columns]


def test_large_sheet_streams(tmp_path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Big")
    ws.append(["id", "score", "flag"])
    for i in range(60000):
        ws.append([f"{i:07d}", i % 97, i % 3 == 0])
    p = tmp_path / "big.xlsx"
    wb.save(p)
    rr, rr2, out = roundtrip(p, tmp_path)
    assert rr.tables[0].n_rows == 60000 and rr2.tables[0].n_rows == 60000
    assert rr2.tables[0].columns[1].values[-1] == str(59999 % 97)
    buf = io.BytesIO(out.read_bytes())
    assert zipfile.ZipFile(buf).testzip() is None
