"""D34 hardening (passes 1-3): leaks and false refusals of the screen of header, sheet, file and folder names.

Leaks closed: ``Ms`` titles (``Ms`` is also the institution word of ``Lincoln MS``), stub + ID names without a
separator (Stata ``reshape wide`` ``score482900``), names quoted by read-time errors of files that hold no table
(sheet listings, header-row checks, unsupported columns and Arrow struct fields), name layouts the lists missed
(``Park, Jennifer``, ``Hall, J.``, ``Okafor, Summer``, ``Smith, John Day 2``, ``J Smith``, ``Smith_J``), the
folders of the input path, ``--keep-titles`` title rows, HTML class attributes and typed column references
(``inspect``). False refusals closed: a flagged real ID's other words (``Reading``, ``Gradebook``) never make
clean files FAIL, ordinary numbers equal to an ID by chance (time stamps, report, room and CIP numbers) need
repeated evidence, a lookup table's ``id`` codes are no person IDs, and schools, programs and awards named after
people pass.

Pass 2 leaks closed: a given name and a last initial (``Jayden M.``), user names (``jgarcia``, ``a.okonkwo``),
Spanish, German and role titles (``Sr. Garcia``, ``Herr Schmidt``, ``Counselor Nwosu``, ``Mrs. G``), given names
off Faker's lists (``Adeyemi, Oluwaseun``, ``Hodan Warsame``), a per-student file named by its own real ID in a
large namespace, real IDs and free text in HTML attributes, extensionless ``first.last`` names, Office lock files,
unquoted names in usage errors and the ``--config`` path. Pass 2 false refusals closed: campus, county and house
names among siblings, gradebook assignments naming books, authors, characters and science laws, bare school
namesakes, federal programs, vendors, 10-digit time stamps and finance numbers, 4-digit room numbers and MMDD
dates beside dense student numbers.

Pass 3 leaks closed: full names beside school-document words (``Ana Garcia College Essay``, ``Senior Garcia,
Ana``, ``Ms. Okafor Room``: only a school name form after a name makes it a school), a given name beside a surname
off the lists inside a longer name (``Kaden Yazzie IEP 2024``), pivots by student (one word around a shared text,
popular given names that are also surnames, surname-only sheets), title rows listing given names, two per-student
files named by real IDs in a large namespace, user names (``smithj``, ``JSMITH``, ``Rodriguez.M``), typed paths
in any printed form and YAML error lines, names in Cyrillic, Greek, Arabic, Chinese and Korean, and CSS classes
on no word list. Pass 3 false refusals closed: ``Week 1``..``Week 36`` label columns in the leak scanner, court
cases, book titles, authors, characters and local campus namesakes, macOS AppleDouble files, ZIP-code sheets
beside dense 5-digit IDs, 10-digit export numbers and ``City, State``.

D34 freeze (user, 2026-10-06): the screen is a best-effort layer. The low-precision name rules are dropped (a
given name beside a word off the lists, header-word given names in a pair, single given names or surnames among
siblings, a surname folder with a given-name child, hyphenated romanized given names, title rows listing given
names), so their assertions are gone: those names may pass (README: What D34 does not catch). The reliable checks
stay asserted (real IDs, personal-data shapes, explicit person layouts, known first + last name pairs, the
namesake allowlist, ``--keep-name``, the redaction). Every name here is invented fixture data or a public name.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from conftest import run_cli
from datanonymous.privacy import headers as d34
from datanonymous.privacy import leaks

N = 40


# --------------------------------------------------------------------------- helpers
def write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return path


def numbers(n_cols: int, n: int = N) -> list[list[Any]]:
    return [[(r * 7 + c * 3) % 50 for c in range(n_cols)] for r in range(n)]


def run(src: Path, tmp: Path, *extra: str, name: str = "run"):
    out, rep = tmp / f"{name}-out", tmp / f"{name}-rep"
    res = run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", "7", *extra)
    return res, out, rep


def everything(*dirs: Path, console: str = "") -> str:
    parts = [console]
    for d in dirs:
        if not d.exists():
            continue
        for p in sorted(d.rglob("*")):
            parts.append(p.relative_to(d).as_posix())
            if p.is_file():
                data = p.read_bytes()
                parts.append(data.decode("utf-8", errors="replace"))
                parts.append(data.decode("latin-1"))
    return "\n".join(parts).casefold()


def assert_never_printed(secrets: list[str], *dirs: Path, console: str = "") -> None:
    blob = everything(*dirs, console=console)
    found = sorted(s for s in secrets if s.casefold() in blob)
    assert not found, f"{len(found)} flagged name form(s) reached the out dir, the report or the console"


def report(rep: Path) -> dict[str, Any]:
    return json.loads((rep / "report.json").read_text(encoding="utf-8"))


def statuses(rep: Path) -> dict[str, str]:
    return {f["source"]: f["status"] for f in report(rep)["files"]}


def written(out: Path) -> list[str]:
    return sorted(p.relative_to(out).as_posix() for p in out.rglob("*") if p.is_file())


class NS:
    """A fake ID namespace: raw values (``generate.ids.IdNamespace``)."""

    def __init__(self, values: list[str]) -> None:
        self.raw_values = set(values)


# --------------------------------------------------------------------------- person-name layouts (leaks)
@pytest.mark.parametrize("name", [
    # titles: 'Ms' is also the institution word of 'Lincoln MS'
    "Ms. Garcia", "Ms Okonkwo Period 3", "Ms_Garcia_Period3", "Ms_Okonkwo_grades", "Ms Rodriguez P3", "MS Nguyen",
    # 'Surname, Given' / 'Surname, I.', also with an institution-word surname or a header-word given name
    "Park, Jennifer", "Park, J.", "Hall, J.", "Hall, Ana", "Day, Doris", "Fields, Sarah", "House, Eddie",
    "Okafor, Summer", "Rodriguez, Max", "Smith, John Day 2", "Begay, Kaden", "Ivanova, Olga", "Price, Sarah",
    # a pair before 'Day' / 'Field', bare initials, invisible characters
    "Garcia_Ana_Day_1", "Ana Garcia Field Trip", "J Smith", "Jo​hn Smith",
    "Smith_J_IEP_goals", "GarciaA", "SmithJ", "NguyenT", "PatelR",
    # names of newcomer families on the lists
    "Kaden Begay",
    # a namesake inside a longer person name is still a person
    "Jose Marti Garcia",
])
def test_person_name_layouts_are_flagged(name: str) -> None:
    assert d34.name_kinds(name) == {"name"}


@pytest.mark.parametrize("name", [
    "Lincoln MS", "Jefferson Davis MS", "Martin Luther King Jr Day", "Martin_Luther_King_Jr_Day",
    "Total, Grace", "Score, Max", "Reading, Summer", "Attendance, June", "Retention, May", "Washington, D.C.",
    "Total, N", "Block_A", "Bell_B", "Plan_B", "Group_B", "Grade_A", "Vitamin_D", "Section A", "Lunch_A",
    "Amira_Learning", "Max_Price", "Mark_Total", "Summer_Davis", "X Ray", "Marco Polo",
    "Alexa Skills", "Self-check", "Pre-test Score", "Lee Follow-up",
    "Gradebook_Summary", "Gradebook_Weight", "www.district.org",
])
def test_ordinary_labels_pass(name: str) -> None:
    assert d34.name_kinds(name) == frozenset()


@pytest.mark.parametrize("name", [
    # schools, charters, programs and awards named after people
    "Rosa_Parks", "Sam_Houston", "Barbara_Jordan", "KIPP_Austin", "Green_Dot_Enrollment", "Bishop_Lynch_Roster",
    "Jose_Marti_MAST_Grades", "Escuela_Benito_Juarez", "Ruby_Payne_Framework_Survey", "Cal_Grant_Eligibility",
    "Sanford_Harmony_Lessons", "Jan_Richardson_Guided_Reading_Levels", "Rosa Parks", "Cesar Chavez",
    "Thurgood Marshall", "Cesar_Chavez_Middle", "Booker T Washington", "Stephen F Austin", "John F Kennedy",
    "Cardinal Newman", "Martin Luther King Jr", "Dr Martin Luther King Jr", "IDEA_Rundberg", "AVID_Roster",
])
def test_namesakes_and_programs_pass(name: str) -> None:
    assert d34.name_kinds(name) == frozenset()


# --------------------------------------------------------------------------- real IDs: chance-aware
def test_stub_and_id_without_a_separator_is_a_real_id() -> None:
    sparse = d34.IdKeys([NS([str(482900 + i) for i in range(N)])])
    for name in ("score482900", "X482905", "id482910", "score.482911", "q1_482912", "482913math"):
        assert sparse.match(name), name
    prefixed = d34.IdKeys([NS([f"S{482900 + i}" for i in range(N)])])
    for name in ("S482909math", "mathS482910", "s-482901", "score_482902", "S0482903", "Ｓ482915", "S482917​"):
        assert prefixed.match(name), name
    hits = d34.id_group_flags(sparse, ["item", "score482900", "score482901", "Notes"])
    assert set(hits) == {1, 2} and hits[1] == ("482900",)


def test_ordinary_numbers_equal_to_a_dense_id_by_chance_need_repeated_evidence() -> None:
    dense = d34.IdKeys([NS([str(150000 + i) for i in range(5000)])])  # 5,000 6-digit IDs 150000-154999
    for name in ("Report_153017", "PowerSchool_export_20250115_152233", "Pct_151200_to_152000", "CIP_150301",
                 "Job 154321", "Room_150101"):
        assert not dense.match(name), name
    assert d34.id_group_flags(dense, ["grade_level", "Pct_151200_to_152000", "CIP_150301", "Report_153017"]) == {}
    assert d34.id_group_flags(dense, ["Room_150101", "Room_150102", "Room_150103"]) == {}  # numbered things
    # the pivot: one stub around distinct real IDs (time-like ones such as 150001 = 15:00:01 count only within it)
    assert set(d34.id_group_flags(dense, ["item", "score150061", "score150072", "score150001"])) == {1, 2, 3}
    assert d34.id_group_flags(dense, ["item", "score150001", "score150002"]) == {}
    staff = d34.IdKeys([NS([str(10000 + i) for i in range(300)])])  # 300 5-digit staff IDs 10000-10299
    assert d34.id_group_flags(staff, ["grade_level", "Room_10101", "Room_10250", "Course 10199"]) == {}
    assert d34.natural_number("152233") and d34.natural_number("250115") and not d34.natural_number("482913")


def test_only_digit_bearing_ids_count() -> None:
    codes = d34.IdKeys([NS(["MATH", "READ", "SCIE", "ELA1", "AB12"])])
    for name in ("Math", "MAP_Math_Fall", "Reading", "ELA1", "AB12_score"):
        assert not codes.matches(name), name
    assert d34.person_id_shape("S482913") and d34.person_id_shape("STU-0042") and not d34.person_id_shape("ELA1")


def test_a_flagged_real_id_indexes_its_id_never_its_other_words() -> None:
    screen = d34.NameScreen()
    screen.add("S482901_Reading_Score", "column", frozenset({"id"}), "column 2 of 2", written=True, number=2,
               id_tokens=("S482901",))
    words, seqs, exact = screen.flags[0].terms()
    assert exact == ["S482901", "S482901_Reading_Score"] and words == []
    assert "Reading" not in exact and "Score" not in exact
    idx = leaks.SensitiveIndex.build([], screen)
    allow = leaks.Allow(tokens={"reading", "score", "gradebook"}, phrases=set())
    scanner = leaks.Scanner(idx, allow)
    assert not scanner.hits_in_text("columns Reading and Score of clean.csv")
    assert scanner.hits_in_text("the header S482901 was renamed")


def test_split_name_keeps_long_extensions() -> None:
    assert d34.split_name("Garcia_Ana.parquet") == ("Garcia_Ana", ".parquet")
    assert d34.split_name("Nguyen_Thi.feather") == ("Nguyen_Thi", ".feather")
    assert d34.split_name("Okonkwo_Chinedu.sas7bdat") == ("Okonkwo_Chinedu", ".sas7bdat")
    assert d34.split_name("Smith_John.jsonl.gz") == ("Smith_John", ".jsonl.gz")
    assert d34.split_name("Smith_John_IEP.csv.gz") == ("Smith_John_IEP", ".csv.gz")


def test_quoted_names_and_usage_errors_are_screened() -> None:
    msg = "no sheet named 'Data' (sheets: 'Okonkwo, Chinedu', \"O'Brien, Patrick\")"
    assert d34.quoted_names(msg) == ["Data", "Okonkwo, Chinedu", "O'Brien, Patrick"]
    out = d34.redact_message("--range 'Okonkwo, Chinedu!A1:XX': expected [Sheet!]A5:H200")
    assert "Okonkwo" not in out and "[redacted]" in out and "expected [Sheet!]A5:H200" in out


def test_html_classes_and_css_selectors_holding_names_are_dropped() -> None:
    from datanonymous.io.excel_html import clean_css, safe_class

    assert safe_class("xl65 Abernathy-Zelda") == "xl65" and safe_class("header odd") == "header odd"
    css = clean_css(".xl65 {mso-number-format:\"0\";}\n.Abernathy-Zelda {color:red;}\n#Okonkwo_Chinedu {x:1}\n")
    assert "Abernathy" not in css and "Okonkwo" not in css and "xl65" in css


def test_arrow_types_are_printed_without_struct_field_names() -> None:
    from datanonymous.io.arrowio import type_text

    t = pa.struct([("Okonkwo, Chinedu", pa.int64()), ("math", pa.int64())])
    assert type_text(t) == "struct<2 fields>" and type_text(pa.list_(t)) == "list<struct<2 fields>>"
    assert type_text(pa.int64()) == "int64" and type_text(pa.list_(pa.int64())) == str(pa.list_(pa.int64()))


# --------------------------------------------------------------------------- end to end: false refusals
def test_a_flagged_name_never_cascades_to_clean_files(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "long.csv", ["student_id", "Reading", "Score"],
              [[f"S{482900 + r}", r % 9, (r * 3) % 11] for r in range(N)])
    write_csv(src / "wide.csv", ["item", "S482901_Reading_Score"], numbers(2))
    write_csv(src / "clean.csv", ["grade_level", "Reading", "Score"], numbers(3))
    write_csv(src / "Gradebook_S482913.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "Gradebook_Summary.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "attendance.csv", ["grade_level", "Gradebook_Weight"], numbers(2))
    write_csv(src / "surnames.csv", ["grade_level", "Garcia", "Ana"], numbers(3))  # single words: clean
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    failed = sorted(k for k, v in st.items() if v == "FAIL")
    assert len(failed) == 2 and "wide.csv" in failed, failed
    for clean in ("long.csv", "clean.csv", "Gradebook_Summary.csv", "attendance.csv", "surnames.csv"):
        assert st[clean] != "FAIL", (clean, report(rep)["files"])
    assert not any("leak scan" in r for f in report(rep)["files"] for r in f["reasons"])


def test_ordinary_numbers_in_names_pass_beside_a_dense_id_namespace(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "students.csv", ["student_id", "grade_level"], [[150000 + i, i % 13] for i in range(5000)])
    write_csv(src / "PowerSchool_export_20250115_152233.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "Report_153017.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "attendance.csv", ["grade_level", "Pct_151200_to_152000", "CIP_150301"], numbers(3))
    write_csv(src / "staff.csv", ["staffid", "dept"], [[10000 + i, i % 4] for i in range(300)])
    write_csv(src / "rooms.csv", ["grade_level", "Room_10101", "Room_10250", "Course 10199"], numbers(4))
    write_csv(src / "pivot.csv", ["item", "score150061", "score150072", "score150085"], numbers(4))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    assert {k for k, v in statuses(rep).items() if v == "FAIL"} == {"pivot.csv"}
    reason = next(f for f in report(rep)["files"] if f["source"] == "pivot.csv")["reasons"][0]
    assert "the headers of 3 of 4 columns (columns 2-4; a real ID of this run: 3)" in reason


def test_lookup_table_id_codes_are_no_person_ids(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "subjects.csv", ["id", "description"],
              [[c, d] for c, d in (("MATH", "Mathematics"), ("READ", "Reading"), ("SCIE", "Science"),
                                   ("WRIT", "Writing"), ("HIST", "History"), ("ARTS", "Fine arts"))])
    write_csv(src / "MAP_Math_Fall.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "scores.csv", ["grade_level", "Math"], numbers(2))
    write_csv(src / "Reading" / "fall.csv", ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert "MAP_Math_Fall.csv" in written(out) and "Reading/fall.csv" in written(out)


def test_schools_and_programs_named_after_people_pass(tmp_path: Path) -> None:
    src = tmp_path / "in"
    for folder in ("Rosa_Parks", "KIPP_Austin", "Sam_Houston", "Barbara_Jordan", "Cesar_Chavez"):
        write_csv(src / folder / "enrollment.csv", ["grade_level", "score"], numbers(2))
    for f in ("Green_Dot_Enrollment.csv", "Bishop_Lynch_Roster.csv", "Jose_Marti_MAST_Grades.csv",
              "Escuela_Benito_Juarez.csv", "Cal_Grant_Eligibility.csv", "Sanford_Harmony_Lessons.csv"):
        write_csv(src / f, ["grade_level", "score"], numbers(2))
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    for sh in ("Rosa Parks", "Cesar Chavez", "Thurgood Marshall"):
        wb.create_sheet(sh)
    for ws in wb.worksheets:
        ws.append(["grade_level", "score"])
        for r in numbers(2):
            ws.append(r)
    wb.save(src / "District_Enrollment_by_School.xlsx")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert not any(r.startswith(d34.D34_PREFIX) for f in report(rep)["files"] for r in f["reasons"])
    assert "Rosa_Parks/enrollment.csv" in written(out)


# --------------------------------------------------------------------------- end to end: leaks
def test_ms_titles_fail_headers_sheets_files_and_folders(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "grades.csv", ["student_id", "Ms. Garcia", "Ms Okonkwo Period 3"],
              [[100000 + r, r % 7, r % 5] for r in range(N)])
    write_csv(src / "Ms_Garcia_Period3" / "roster.csv", ["grade_level", "score"], numbers(2))
    wb = openpyxl.Workbook()
    wb.active.title = "Ms. Nguyen"
    wb.active.append(["grade_level", "score"])
    for r in numbers(2):
        wb.active.append(r)
    wb.save(src / "gradebooks.xlsx")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    assert written(out) == ["DATA_DICTIONARY.md", "README.md", "schema.json"]
    assert_never_printed(["Garcia", "Okonkwo", "Nguyen"], out, rep, console=res.output)


@pytest.mark.parametrize("fmt", ["csv", "dta"])
def test_stub_and_id_pivot_without_a_separator_fails(tmp_path: Path, fmt: str) -> None:
    src = tmp_path / "in"
    src.mkdir()
    ids = [482900 + i for i in range(N)]
    long = pd.DataFrame({"student_id": [i for i in ids for _ in range(3)], "item": [j for _ in ids for j in range(3)],
                         "score": [(i + j) % 10 for i in ids for j in range(3)]})
    wide = long.pivot(index="item", columns="student_id", values="score")
    wide.columns = [f"score{c}" for c in wide.columns]  # Stata 'reshape wide score, i(item) j(student_id)'
    wide = wide.reset_index()
    if fmt == "dta":
        long.to_stata(src / "long.dta", write_index=False)
        wide.to_stata(src / "wide.dta", write_index=False)
    else:
        long.to_csv(src / "long.csv", index=False)
        wide.to_csv(src / "wide.csv", index=False)
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    assert statuses(rep)[f"wide.{fmt}"] == "FAIL" and statuses(rep)[f"long.{fmt}"] != "FAIL"
    assert_never_printed(["score482900", "score482939", "482913"], out, rep, console=res.output)


def test_a_real_id_flagged_by_its_siblings_fails_every_file_that_writes_it(tmp_path: Path) -> None:
    """One name, one verdict: a pivot's ``score150061`` is weak alone (5,000 dense IDs) but a real ID by its
    siblings; another file's lone header or file name holding it FAILs too, with the same kind."""
    src = tmp_path / "in"
    write_csv(src / "students.csv", ["student_id", "grade_level"], [[150000 + i, i % 13] for i in range(5000)])
    write_csv(src / "pivot.csv", ["item", "score150061", "score150072", "score150085"], numbers(4))
    write_csv(src / "other.csv", ["grade_level", "score150061"], numbers(2))
    write_csv(src / "score150072.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "Report_153017.csv", ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    failed = {k for k, v in statuses(rep).items() if v == "FAIL"}
    assert len(failed) == 3 and {"pivot.csv", "other.csv"} < failed, failed
    assert any(f.startswith("[redacted-") and f.endswith("].csv") for f in failed)
    assert all("a real ID of this run" in f["reasons"][0] for f in report(rep)["files"] if f["status"] == "FAIL")
    assert_never_printed(["score150061", "150072", "150085"], out, rep, console=res.output)


def test_keep_titles_fails_every_file_with_the_same_title(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    for name in ("a.xlsx", "b.xlsx"):
        wb = openpyxl.Workbook()
        wb.active.append(["Gradebook of Okonkwo, Chinedu"])
        wb.active.append([])
        wb.active.append(["grade_level", "score"])
        for r in numbers(2):
            wb.active.append(r)
        wb.save(src / name)
    res, out, rep = run(src, tmp_path, "--keep-titles")
    assert res.exit_code == 3 and set(statuses(rep).values()) == {"FAIL"}
    assert_never_printed(["Okonkwo", "Chinedu"], out, rep, console=res.output)


def test_read_errors_never_quote_a_person_named_sheet(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    frame = pd.DataFrame(numbers(2), columns=["grade_level", "score"])
    with pd.ExcelWriter(src / "a_book.xlsx") as xw:
        frame.to_excel(xw, sheet_name="Data", index=False)
    with pd.ExcelWriter(src / "b_gradebook.xlsx") as xw:
        frame.to_excel(xw, sheet_name="Okonkwo, Chinedu", index=False)
        frame.to_excel(xw, sheet_name="Abernathy, Zelda", index=False)
    res, out, rep = run(src, tmp_path, "--sheet", "Data")
    assert res.exit_code == 3
    reason = next(f for f in report(rep)["files"] if f["source"] == "b_gradebook.xlsx")["reasons"][0]
    assert reason.startswith("no sheet named 'Data' (the workbook has 2 sheet(s))")
    ins = run_cli("inspect", src, "--sheet", "Data", "--json")
    secrets = ["Okonkwo", "Chinedu", "Abernathy", "Zelda"]
    assert_never_printed(secrets, out, rep, console=res.output + ins.output)
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    wb.active.append(["grade", "score"])
    wb.create_sheet("Abernathy, Zelda").append(["grade", "score"])
    wb.save(src / "a_book.xlsx")
    res = run_cli(src / "a_book.xlsx", "--out-dir", tmp_path / "o2", "--report-dir", tmp_path / "r2",
                  "--header-row", "10")
    assert res.exit_code == 3
    assert_never_printed(secrets, tmp_path / "o2", tmp_path / "r2", console=res.output)


def test_unsupported_columns_and_struct_fields_never_print_their_names(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    pq.write_table(pa.table({"grade_level": [i % 5 for i in range(N)],
                             "Okonkwo, Chinedu": pa.array([b"\x00" for _ in range(N)], type=pa.binary())}),
                   src / "binary.parquet")
    pq.write_table(pa.table({"grade_level": [i % 5 for i in range(N)],
                             "s": pa.array([{"Abernathy, Zelda": i, "math": 1} for i in range(N)])}),
                   src / "nested.parquet")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    secrets = ["Okonkwo", "Chinedu", "Abernathy", "Zelda"]
    ins = run_cli("inspect", src, "--json", "-v")
    assert_never_printed(secrets, out, rep, console=res.output + ins.output)
    # kept with --null, the struct keeps its type: its field names are written, so they FAIL like headers
    res, out, rep = run(src / "nested.parquet", tmp_path, "--null", "s", name="null")
    assert res.exit_code == 3 and "a person's name" in report(rep)["files"][0]["reasons"][0]
    assert_never_printed(secrets, out, rep, console=res.output)


def test_input_path_folders_are_redacted(tmp_path: Path) -> None:
    src = tmp_path / "in" / "Okonkwo_Chinedu"
    write_csv(src / "scores.csv", ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert "scores.csv" in written(out)
    res2, out2, rep2 = run(src / "scores.csv", tmp_path, name="single")
    assert res2.exit_code == 0, res2.output
    assert_never_printed(["Okonkwo", "Chinedu"], out, rep, out2, rep2, console=res.output + res2.output)
    manifest = json.loads((rep / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["input"].endswith("/[redacted-1]")


def test_keep_titles_screens_title_rows_written_verbatim(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws.append(["Gradebook for Okonkwo, Chinedu (teacher)"])
    ws.append(["Prepared by abernathy.zelda@school.example"])
    ws.append([])
    ws.append(["grade_level", "score"])
    for r in numbers(2):
        ws.append(r)
    wb.save(src / "roster.xlsx")
    rows = "".join(f"<tr><td>{a}</td><td>{b}</td></tr>" for a, b in numbers(2))
    (src / "export.xls").write_text("<html><body><p>Teacher: Abernathy, Zelda</p><table>"
                                    "<tr><th>grade_level</th><th>score</th></tr>" + rows + "</table></body></html>",
                                    encoding="utf-8")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output  # placeholders: nothing verbatim
    res, out, rep = run(src, tmp_path, "--keep-titles", name="keep")
    assert res.exit_code == 3
    assert set(statuses(rep).values()) == {"FAIL"}
    assert all("written verbatim with --keep-titles" in f["reasons"][0] for f in report(rep)["files"])
    assert_never_printed(["Okonkwo", "Chinedu", "abernathy", "Zelda"], out, rep, console=res.output)


def test_html_class_attributes_holding_names_are_not_copied(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    rows = "".join(f"<tr><td class=\"xl65\">{a}</td><td>{b}</td></tr>" for a, b in numbers(2))
    (src / "export.xls").write_text(
        "<html><head><style>.Abernathy-Zelda {color:red;}</style></head><body><table><tr><th>grade_level</th>"
        "<th id=\"Okonkwo_Chinedu\" class=\"Abernathy-Zelda\">score</th></tr>" + rows + "</table></body></html>",
        encoding="utf-8")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    text = (out / "export.xls").read_text(encoding="utf-8")
    assert "xl65" in text
    assert_never_printed(["Abernathy", "Zelda", "Okonkwo", "Chinedu"], out, rep, console=res.output)


def test_inspect_redacts_a_typed_near_miss_of_a_flagged_column(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "t.csv", ["grade_level", "Okonkwo, Chinedu", "score"], numbers(3))
    ins = run_cli("inspect", src, "--exclude", "Okonkwo, Chinedu ", "--force-factor", "Okonkwo Chinedu")
    ins_json = run_cli("inspect", src, "--json", "--force-factor", "okonkwo chinedu")
    assert ins.exit_code == 3
    assert_never_printed(["Okonkwo", "Chinedu"], console=ins.output + ins_json.output)


def test_usage_errors_never_print_a_person_named_range(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    pd.DataFrame(numbers(2), columns=["grade_level", "score"]).to_excel(src / "book.xlsx", index=False)
    res = run_cli("inspect", src / "book.xlsx", "--range", "Okonkwo, Chinedu!XYZ")
    assert res.exit_code == 2
    assert "Okonkwo" not in res.output and "Chinedu" not in res.output


def test_a_flagged_parquet_file_name_shows_its_extension(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    pd.DataFrame(numbers(2), columns=["grade_level", "score"]).to_parquet(src / "Garcia_Ana.parquet", index=False)
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    f = report(rep)["files"][0]
    assert f["source"] == "[redacted-1].parquet" and "; .parquet, " in f["reasons"][0]


def test_an_empty_chart_sheet_is_read(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    wb.active.append(["grade_level", "score"])
    for r in numbers(2):
        wb.active.append(r)
    wb.create_chartsheet("Chart")
    wb.save(src / "book.xlsx")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert openpyxl.load_workbook(out / "book.xlsx").sheetnames == ["Data"]


# =========================================================================== pass 2
# --------------------------------------------------------------------------- leaks: layouts the screen missed
@pytest.mark.parametrize("name", [
    # a given name and a last initial (teacher gradebooks, reading logs)
    "Jayden M.", "Jayden_M", "JaydenM", "Jayden.M", "Sofia R.", "Aaliyah R", "Jayden M - Reading Log", "ana_g",
    # user names of staff and students
    "jgarcia", "aokonkwo", "mrodriguez", "agarcia2", "j.garcia", "a.okonkwo", "a.garcia", "anagarcia",
    # Spanish, German and role titles, a title before a bare initial
    "Sr. Garcia", "Señor Garcia", "Profe Lopez", "Maestra Lopez", "Mtra. Lopez", "Herr Schmidt",
    "Frau Müller", "Counselor Nwosu", "Nurse Garcia", "Principal Garcia", "Officer Garcia", "Father Murphy",
    "Brother John", "Rabbi Cohen", "Imam Ali", "Teacher Garcia", "Mrs. G", "Mr. G's Class",
    # given names off Faker's lists: 'Surname, Given', a distinctive surname, newcomer communities
    "Adeyemi, Oluwaseun", "OKAFOR, CHIDI", "Yamamoto, Takeshi", "Kim, Seo-yeon", "Oluwaseun Adeyemi",
    "Takeshi Yamamoto", "Hodan Warsame", "Dmitri Volkov", "Raj Kumar", "Warsame, Hodan", "Okafor; Summer",
    # comma layouts with a place-word surname, look-alike letters, labelled numbers
    "St. Clair, Ana", "King, Ana", "Church, J.", "Rivers, Ana", "Gаrcia, Ana",
])
def test_pass2_person_name_layouts_are_flagged(name: str) -> None:
    assert d34.name_kinds(name) == {"name"}


@pytest.mark.parametrize("name, kind", [
    ("Ph 5551234567", "phone"), ("Call 312-555-0199", "phone"), ("Contact 3125550199", "phone"),
    ("SSN 123456789", "ssn"),
])
def test_pass2_labelled_numbers_are_personal_data(name: str, kind: str) -> None:
    assert kind in d34.name_kinds(name)


@pytest.mark.parametrize("name", [
    # letter labels, Roman numerals, period labels
    "Vitamin D", "Plan B", "Elizabeth I", "Richard III", "Henry VIII", "Section A", "Day B", "Garcia P1", "Lunch_A",
    # ordinary words shaped like a user name
    "chart", "layers", "fallen", "space", "comma", "decimal_comma", "stata_prop", "players", "aware", "legal",
    # role words before header words, companies and credentials after 'Surname,', places
    "Counselor Name", "Nurse Visits", "Principal Amount", "Father Name", "Sr Project", "Coach Rating",
    "Dean List", "Teacher Email", "Smith, Inc.", "Garcia, PhD", "Jackson, Mississippi", "Houston, TX",
])
def test_pass2_ordinary_labels_pass(name: str) -> None:
    assert d34.name_kinds(name) == frozenset()


def test_pass2_usernames_beside_digits_and_in_title_text() -> None:
    assert d34.name_kinds("agarcia2_gradebook") == {"name"}
    assert d34.person_name_hit("Prepared by agarcia on 10/03/2025") is not None


# --------------------------------------------------------------------------- false refusals: public names
@pytest.mark.parametrize("name", [
    # gradebook assignments: books, authors, characters, historical figures, science laws and models
    "Junie B. Jones Quiz", "Percy Jackson Ch 1-5", "Harry Potter", "Charles's Law Lab", "Lewis Dot Structures",
    "Julius Caesar Act 1", "Lincoln-Douglas Debate", "Lewis & Clark", "John Brown's Raid", "Charles Darwin",
    "Holden Caulfield", "Atticus Finch Essay", "Stuart Little", "Amelia Bedelia", "Ms. Frizzle", "Mrs. Frisby",
    "Edgar Allan Poe Unit", "John Dalton", "Neil Armstrong Moon", "Isaac Newton Laws", "Niels Bohr", "Greg Heffley",
    "Ming Dynasty", "Hardy-Weinberg Principle", "Darwin's Finches",
    # bare school namesakes
    "Harriet Tubman", "Sally Ride", "John Marshall", "Walt Whitman", "Mary McLeod Bethune", "Phillis Wheatley",
    "Amelia Earhart", "Paul Robeson", "Sonia Sotomayor", "Robert E Lee", "Stonewall Jackson", "Jefferson Davis",
    "Mirabeau Lamar", "Lamar Consolidated", "Francis Scott Key", "Isaac Newton",
    # federal programs and funds named after people
    "Carl D. Perkins CTE", "Carl D Perkins Basic Grant", "Richard B. Russell NSLA",
    "Jacob K. Javits Gifted and Talented", "Stewart B. McKinney Homeless", "Ronald E. McNair Scholars",
    "Augustus F. Hawkins", "Harry S Truman Scholarship",
    # vendors and brands (a brand that is a common person's name only beside a business word or a number)
    "Ben E. Keith", "Tyler Traversa", "Dell Chromebooks", "Dell Latitude", "Dell Laptops", "Sherwin Williams",
    "John Deere", "Lamb Weston", "Tony's Pizza", "General Mills", "Little Debbie", "Mrs Baird's Bread",
    "Sara Lee Bread", "Edward Jones 403b", "Raymond James Financial", "Benjamin Moore Paint",
    "Wendy's Fundraiser", "Ruiz Foods", "Johnson Controls",
    # 10-digit time stamps and finance numbers are no phone numbers
    "roster_1696342522", "students_1700000000", "export_2023101512", "PO 4500012345", "Invoice 3105551234",
    "Run_6175550123", "ISBN 0134685997",
])
def test_pass2_public_names_and_numbers_pass(name: str) -> None:
    assert d34.name_kinds(name) == frozenset()


def test_pass2_common_person_names_stay_flagged_beside_namesakes() -> None:
    # a namesake entry keeps its initial; a brand that is a common name needs a business word; a longer name; a
    # Roman numeral that is also a last initial ('Henry V.')
    for name in ("Robert Lee", "Edward Jones", "Jose Marti Garcia", "Ana Cesar Chavez", "Henry V.", "James I."):
        assert d34.name_kinds(name) == {"name"}, name


def test_pass2_campus_county_and_house_lists_pass_among_siblings() -> None:
    """Single words among siblings are never screened as a pivot by student (D34 freeze): campus, county and house
    lists pass, and so does a list of given names (the accepted residual risk)."""
    aisd = ["grade", "Akins", "Anderson", "Austin", "Bowie", "Crockett", "Eastside", "Lanier", "Lamar", "LBJ",
            "McCallum", "Navarro", "Northeast", "Travis"]
    screen = d34.NameScreen()
    for siblings in (aisd, ["Brandon", "Burgess", "Jones", "Kimberly", "Scott", "Smith", "Woodson"],
                     ["Anthony", "Hale", "Harrison", "Kenny", "Marcy", "Seward", "Sheridan"],
                     ["Bell", "Carson", "Dorsey", "Fremont", "Garfield", "Taft", "Venice"],
                     ["Travis", "Harris", "Collin", "Denton"], ["Apollo", "Athena", "Hermes"],
                     ["Hexter", "Kiest", "Truett", "Marsalis", "Junkins", "Zilker", "Casis"]):
        assert not any(k for k, _, _ in screen.group_kinds([(t, "") for t in siblings])), siblings


def test_pass2_real_ids_weigh_dense_ranges_dates_and_the_files_own_ids() -> None:
    dense = d34.IdKeys([NS([str(x) for x in range(1001, 1401)])])  # sequential 4-digit student numbers
    # D37: numbers that are all real IDs are a pivot whatever the density (room numbers among them too)
    assert set(d34.id_group_flags(dense, ["1375", "1380", "1392"])) == {0, 1, 2}
    assert d34.id_group_flags(dense, ["Room 1375", "Room 1380", "Room 1392"]) == {}  # a numbered word: rooms
    assert d34.id_group_flags(dense, ["Attendance_1001", "Attendance_1002", "Attendance_1003"]) == {}  # MMDD
    sparse = d34.IdKeys([NS([str(1000 + 173 * i) for i in range(50)])])
    assert set(d34.id_group_flags(sparse, ["item", "score1173", "score1346", "score1519"])) == {1, 2, 3}
    big = d34.IdKeys([NS([str(482913 + 11 * i) for i in range(2000)])])  # many 6-digit IDs: one alone is weak
    x = str(482913 + 11 * 1234)
    # K-01: after a person or document label (IEP_496487) a real ID is strong whatever the density; the file's own
    # IDs are read here after a word that is no label
    assert big.match(f"IEP_{x}")
    assert not big.match(f"Goals_{x}")
    assert big.match(f"Goals_{x}", big.local_keys([x]))  # the file holds that one student: strong
    assert not big.match(f"Goals_{x}", big.local_keys([str(482913 + 11 * i) for i in range(2000)]))
    assert d34.id_group_flags(big, [f"Goals_{x}", "notes"], [big.local_keys([x]), None]) == {0: (x,)}


def test_pass2_split_name_keeps_unknown_suffixes_in_the_name() -> None:
    assert d34.split_name("Ana.Garcia") == ("Ana.Garcia", "")
    assert d34.split_name("notes.Okonkwo") == ("notes.Okonkwo", "")
    assert d34.split_name("report.2023.csv") == ("report.2023", ".csv")
    assert d34.split_name("scan.PDF") == ("scan", ".PDF")
    assert d34.lock_file_rest("~$rcia_Ana_IEP.docx") == "rcia_Ana_IEP"
    assert d34.lock_file_rest(".~lock.Okonkwo_Chinedu_IEP.xlsx#") == "Okonkwo_Chinedu_IEP"


def test_pass2_html_attributes_never_carry_ids_or_names() -> None:
    from datanonymous.io.excel_html import _known_fonts, _safe_attr, clean_css, safe_class

    assert safe_class("xl65 stu482913") == "xl65" and safe_class("s0 x12 sid-0042 header") == "s0 x12 header"
    assert safe_class("s482913") == "" and safe_class("Zelmira_Abernathy odd") == "odd"
    assert not _safe_attr("lang", "Quixote") and _safe_attr("lang", "en-US")
    assert not _safe_attr("dir", "Quixote") and _safe_attr("dir", "rtl")
    assert not _safe_attr("xmlns:x", "urn:Quixote") and _safe_attr("xmlns:x", "urn:schemas-microsoft-com:office:excel")
    assert not _safe_attr("width", "Philomena Quixote") and _safe_attr("width", "64")
    assert _known_fonts('"Philomena Quixote Script", Calibri, sans-serif') == "Calibri,sans-serif"
    css = clean_css('.xl65 {font-family:"Philomena Quixote Script"; color:red}')
    assert "Quixote" not in css and "color:red" in css


# --------------------------------------------------------------------------- end to end (pass 2)
def test_pass2_given_name_and_initial_columns_and_sheets_fail(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "gradebook.csv", ["Assignment", "Jayden M.", "Aaliyah R.", "Brayden S.", "Ximena T.", "Kaden B."],
              [[f"HW{r}"] + [(r * 7 + c) % 40 + 60 for c in range(5)] for r in range(N)])
    wb = openpyxl.Workbook()
    wb.active.title = "Jayden M."
    for name in ("Sr. Garcia", "Herr Schmidt", "Profe Lopez", "Counselor Nwosu", "Mrs. G"):
        wb.create_sheet(name)
    for ws in wb.worksheets:
        ws.append(["grade_level", "score"])
        for r in numbers(2):
            ws.append(r)
    wb.save(src / "classes.xlsx")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    assert set(statuses(rep).values()) == {"FAIL"}
    assert written(out) == ["DATA_DICTIONARY.md", "README.md", "schema.json"]
    assert_never_printed(["Jayden M", "Aaliyah R", "Ximena T", "Kaden B", "Garcia", "Schmidt", "Lopez", "Nwosu",
                          "Mrs. G"], out, rep, console=res.output)


def test_pass2_username_folders_fail_and_their_files_keep_their_names(tmp_path: Path) -> None:
    src = tmp_path / "in"
    for user in ("jgarcia", "aokonkwo", "tnguyen", "mrodriguez", "j.garcia"):
        write_csv(src / user / "gradebook.csv", ["student_id", "score"],
                  [[100000 + r, 50 + r] for r in range(N)])
    write_csv(src / "players" / "stats.csv", ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    st = statuses(rep)
    assert st["players/stats.csv"] != "FAIL"
    assert sum(1 for k, v in st.items() if v == "FAIL" and k.endswith("]/gradebook.csv")) == 5
    assert_never_printed(["jgarcia", "aokonkwo", "tnguyen", "mrodriguez", "j.garcia"], out, rep, console=res.output)


def test_pass2_names_off_the_lists_fail(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    wb = openpyxl.Workbook()
    wb.active.title = "Adeyemi, Oluwaseun"
    for name in ("Warsame, Hodan", "Volkov, Dmitri", "Yamamoto, Takeshi"):
        wb.create_sheet(name)
    for ws in wb.worksheets:
        ws.append(["grade_level", "minutes_read"])
        for r in numbers(2):
            ws.append(r)
    wb.save(src / "reading_logs.xlsx")
    write_csv(src / "gradebook.csv", ["Assignment", "Oluwaseun Adeyemi", "Hodan Warsame", "Dmitri Volkov"],
              [[f"HW{r}"] + [(r * 3 + c) % 40 + 60 for c in range(3)] for r in range(N)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3 and set(statuses(rep).values()) == {"FAIL"}
    assert_never_printed(["Adeyemi", "Warsame", "Volkov", "Yamamoto", "Oluwaseun", "Hodan"], out, rep,
                         console=res.output)


def test_pass2_a_per_student_file_named_by_its_own_id_fails_in_a_large_namespace(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = [482913 + 11 * i for i in range(2000)]
    write_csv(src / "roster.csv", ["student_id", "grade_level"], [[x, 9 + x % 4] for x in ids])
    x = ids[1234]
    write_csv(src / f"IEP_{x}.csv", ["student_id", "goal", "score"], [[x, f"G{k}", k] for k in range(N)])
    write_csv(src / f"Roster_{ids[7]}.csv", ["student_id", "grade_level"], [[y, 9 + y % 4] for y in ids])
    ids5 = [10000 + 3 * i for i in range(300)]
    write_csv(src / "five" / "roster.csv", ["student_id", "grade_level"], [[y, 9 + y % 4] for y in ids5])
    write_csv(src / "five" / f"{ids5[41]}_transcript.csv", ["student_id", "course", "mark"],
              [[ids5[41], f"C{k}", "ABCDF"[k % 5]] for k in range(N)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    failed = sorted(k for k, v in statuses(rep).items() if v == "FAIL")  # the full roster's name is weak
    assert len(failed) == 2 and failed[0].startswith("[redacted-") and failed[1].startswith("five/[redacted-"), failed
    assert_never_printed([str(x), str(ids5[41])], out, rep, console=res.output)


def test_pass2_ordinary_numbered_files_and_sheets_pass_beside_dense_ids(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "roster.csv", ["student_number", "grade_level"], [[x, 9 + x % 4] for x in range(1001, 1401)])
    for day in range(1, 6):
        write_csv(src / "daily" / f"Attendance_10{day:02d}.csv", ["student_number", "status"],
                  [[x, "PAT"[x % 3]] for x in range(1001, 1061)])
    wb = openpyxl.Workbook()
    wb.active.title = "1375"
    for room in ("1380", "1392"):
        wb.create_sheet(room)
    for ws in wb.worksheets:
        ws.append(["seat", "period"])
        for r in range(N):
            ws.append([r + 1, r % 7 + 1])
    # D37: sheets named by numbers that are all real IDs FAIL whatever the density, unless the workbook's name says
    # what they are (rooms)
    wb.save(src / "seating_charts_by_room.xlsx")
    for name in ("roster_1696342522.csv", "export_2023101512.csv"):
        write_csv(src / name, ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert "daily/Attendance_1003.csv" in written(out) and "roster_1696342522.csv" in written(out)
    assert "seating_charts_by_room.xlsx" in written(out)
    wb.save(src / "seating_charts.xlsx")
    res2, out2, rep2 = run(src, tmp_path, name="unnamed")
    assert res2.exit_code == 3 and statuses(rep2)["seating_charts.xlsx"] == "FAIL", res2.output
    assert {k for k, v in statuses(rep2).items() if v == "FAIL"} == {"seating_charts.xlsx"}


def test_pass2_k12_names_pass_end_to_end(tmp_path: Path) -> None:
    src = tmp_path / "in"
    campuses = ["Akins", "Anderson", "Austin", "Bowie", "Crockett", "Eastside", "Lanier", "Lamar", "LBJ",
                "McCallum", "Navarro", "Northeast", "Travis"]
    write_csv(src / "enrollment_by_campus.csv", ["grade_level"] + campuses,
              [[g] + [300 + 7 * i + g for i in range(len(campuses))] for g in range(9, 13) for _ in range(10)])
    for c in ("Crockett", "Lamar", "Travis", "Anderson", "Harriet Tubman", "Sally Ride", "John Marshall"):
        write_csv(src / "campuses" / c / "attendance.csv", ["grade_level", "score"], numbers(2))
    heads = ["Student", "Section", "Junie B. Jones Quiz (40213)", "Charles's Law Lab (40214)",
             "Lewis Dot Structures (40215)", "Julius Caesar Act 1 (40216)", "Lincoln-Douglas Debate (40217)",
             "Percy Jackson Ch 1-5 (40218)", "Current Score"]
    write_csv(src / "2025-10-03T1415_Grades-ENGL_101.csv", heads,
              [[f"Student{r % 7}", f"P{r % 3}"] + [(r + c) % 10 for c in range(7)] for r in range(N)])
    budget = ["fund", "Carl D. Perkins CTE", "Richard B. Russell NSLA", "Sherwin Williams", "Ben E. Keith",
              "Dell Chromebooks"]
    write_csv(src / "budget.csv", budget, [[f"F{r % 9}"] + numbers(5)[r] for r in range(N)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert not any(r.startswith(d34.D34_PREFIX) for f in report(rep)["files"] for r in f["reasons"])
    assert "campuses/Crockett/attendance.csv" in written(out)


def test_pass2_html_class_ids_and_free_attribute_text_are_never_written(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    ids = [482913 + 7 * i for i in range(N)]
    rows = "\n".join(f'<tr class="stu{x}"><td class="xl65 sid{x}">{x}</td><td class=xl66>{50 + i}</td></tr>'
                     for i, x in enumerate(ids))
    (src / "roster.xls").write_text(
        '<html xmlns:o="urn:Philomena-Quixote"><head><style>.xl65 {font-family:"Philomena Quixote Script", Calibri;}'
        '</style></head><body lang="Quixote" dir="Abernathy"><table><tr><th>student_id</th><th>score</th></tr>'
        + rows + "</table></body></html>", encoding="utf-8")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    text = (out / "roster.xls").read_text(encoding="utf-8")
    assert "xl65" in text and "Calibri" in text
    assert_never_printed(["Quixote", "Philomena", "Abernathy"] + [f"{x}" for x in ids[:5]], out, rep,
                         console=res.output)


def test_pass2_lock_files_and_extensionless_names_are_skipped_and_redacted(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "Garcia_Ana_IEP.csv", ["grade_level", "score"], numbers(2))
    (src / "~$rcia_Ana_IEP.xlsx").write_bytes(b"\x0bA user" + b"\x00" * 40)
    (src / ".~lock.Okonkwo_Chinedu_IEP.xlsx#").write_text("A user,HOST,user,03.10.2025 14:15;\n", encoding="utf-8")
    (src / "~$Okafor Summer.docx").write_bytes(b"\x0bA user" + b"\x00" * 40)
    (src / "Kaden.Begay").write_text("hello\n", encoding="utf-8")
    (src / "Maria.Lopez").write_text("hello\n", encoding="utf-8")
    write_csv(src / "scores.csv", ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3  # only Garcia_Ana_IEP.csv FAILs: the lock files are skipped, never read
    st = statuses(rep)
    assert len(st) == 2 and list(st.values()).count("FAIL") == 1 and st["scores.csv"] != "FAIL", st
    skipped = report(rep)["skipped"]
    assert len(skipped) == 5 and all(s["path"].startswith("[redacted-") for s in skipped), skipped
    assert all("lock file" in s["reason"] or "no known extension" in s["reason"] for s in skipped)
    assert_never_printed(["rcia_Ana", "Okonkwo", "Begay", "Lopez", "Garcia", "Okafor"], out, rep,
                         console=res.output)


def test_pass2_usage_errors_and_the_config_path_never_print_a_name(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "s.csv", ["student_id", "score"], [[100000 + r, r] for r in range(N)])
    outs = []
    for extra in (["--force-type", "Garcia_Ana=bogus"], ["--pii-col", "Mrs. Nguyen=foo"],
                  ["--header-row", "Garcia_Ana"], ["--sheet", "Okafor, Summer", "--header-row", "x"]):
        res = run_cli(src, "--out-dir", tmp_path / "o", *extra)
        assert res.exit_code == 2, (extra, res.output)
        outs.append(res.output)
    cfg_dir = tmp_path / "Okafor, Summer"
    cfg_dir.mkdir()
    (cfg_dir / "cfg.yaml").write_text("id_col: [student_id]\n", encoding="utf-8")
    res, out, rep = run(src, tmp_path, "--config", str(cfg_dir / "cfg.yaml"), "--include", "*Okafor*s.csv")
    assert_never_printed(["Garcia_Ana", "Garcia Ana", "Nguyen", "Okafor", "Summer"], out, rep,
                         console=res.output + "".join(outs))


def test_pass2_keep_titles_screens_user_names_initials_and_spanish_titles(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    for i, line in enumerate(("Period 3 Roster - Teacher: Jayden M.", "Prepared by agarcia on 10/03/2025",
                              "Questions? Contact Sr. Garcia")):
        wb = openpyxl.Workbook()
        wb.active.append([line])
        wb.active.append([])
        wb.active.append(["grade_level", "score"])
        for r in numbers(2):
            wb.active.append(r)
        wb.save(src / f"roster{i}.xlsx")
    res, out, rep = run(src, tmp_path, "--keep-titles")
    assert res.exit_code == 3 and set(statuses(rep).values()) == {"FAIL"}
    assert_never_printed(["Jayden M", "agarcia", "Sr. Garcia"], out, rep, console=res.output)


# =========================================================================== pass 3
# --------------------------------------------------------------------------- leaks: school words beside full names
@pytest.mark.parametrize("name", [
    # an institution, business or eponym word right after a full name: the person's document, no place
    "Ana Garcia College Essay", "Garcia_Ana_Scholarship", "Chinedu Okafor College Applications",
    "Ana Garcia Senior Project", "Liam Johnson Senior Exit", "Mr. Garcia Senior Seminar", "Ms. Okafor Room",
    "Counselor: Ms. Okafor Room 214", "Ana Garcia Medication Administration Record", "Ana Garcia Dental Screening",
    "Garcia Ana Insurance Card", "Ana Garcia Trial Placement", "Ana Garcia High Risk", "Ana Garcia School Records",
    # an institution word right before one
    "Scholarship - Ana Garcia", "Senior Garcia, Ana", "Kindergarten Chinedu Okafor", "Room Ms Garcia",
    "Fort Worth Ana Garcia",
])
def test_pass3_full_names_beside_school_words_are_flagged(name: str) -> None:
    assert d34.name_kinds(name) == {"name"}


@pytest.mark.parametrize("name", [
    # school name forms after a full name, a single name before any institution word
    "Jesse H Jones High School", "Lily B Clayton Elementary Attendance", "Sam Rayburn ECHS",
    "Dr. Phillips High School", "Jack Yates Senior High School", "Lincoln High", "Smith Hall", "Kennedy Center",
    "Lincoln Campus", "Patel Brothers", "Garcia LLC",
    # city and state, court cases, books, authors and characters, local campus namesakes with their initials
    "Allen, Texas", "Henderson, Nevada", "Sherman Texas", "Morse v. Frederick", "Roe v. Wade",
    "Brown v. Board of Education", "Dr. Jekyll and Mr. Hyde", "Miss Nelson Is Missing", "Mrs. Wishy-Washy",
    "Mr. Tiger Goes Wild", "Mr Poppers Penguins Quiz", "Mr. Lincoln's Way", "Jan Brett Author Study",
    "Andrew Clements", "Jack Prelutsky", "Doreen Cronin", "Judy Moody", "Esperanza Ortega", "Brian Robeson",
    "Because of Winn-Dixie", "Charles H Milby", "Michael E DeBakey", "H Grady Spruce", "Ross S Sterling",
    "Justin F Kimball", "Irma Lerma Rangel", "Calvin-Benson Cycle", "Colonial Era", "John Brown's Raid",
    "Jimmy Dean Sausage", "Charles Schwab Brokerage",
    # export and job numbers are no phone numbers; labels in lower case
    "Extract_5839201746", "Survey_Responses_5839201746", "Qualtrics_SV_5839201746", "Student ID 5125550199",
    "geo_long", "lat_lon_independent", "holds a link", "History of the Americas from the Colonial Era to Today",
    # ordinary headers in other scripts
    "成绩", "学号", "周一", "高一", "姓名", "Оценка", "Фамилия", "الاسم", "رقم الطالب", "이름", "주소지", "학생번호",
])
def test_pass3_school_name_forms_and_public_names_pass(name: str) -> None:
    assert d34.name_kinds(name) == frozenset()


@pytest.mark.parametrize("name", [
    # a known given name and surname (newcomer families' names on the lists) inside a longer name
    "Kaden Yazzie IEP 2024", "Olivia Thibodeaux - Case File", "504 Plan Sofia Bhattacharya", "Aiden Quarshie Goals",
    "Roster for Olivia Thibodeaux", "Sofia M Bhattacharya", "kaden_yazzie",
    # user names: surname and initial, upper case, glued titles; a surname, a period and an initial
    "Gradebook_JSMITH", "smithj_period3", "rodriguezm", "nguyent2", "MRODRIGUEZ", "Roster_AGARCIA", "msgarcia",
    "mrsgarcia", "Rodriguez.M", "garcia.a",
    # other name shapes
    "Max T.", "Muhammad Abdullah", "Le Minh", "Minh Le", "Lee Minho", "Oscar Mayer",
    # other scripts: names on the lists (transliterated or read by script)
    "Иван Петров", "张伟", "王梓涵", "محمد علي", "عبد الله حسن", "김민준", "김 민준",
])
def test_pass3_more_person_name_shapes_are_flagged(name: str) -> None:
    assert d34.name_kinds(name) == {"name"}


def test_pass3_user_name_lookalikes_and_labels_with_phone_words() -> None:
    for word in ("smithy", "moored", "younger", "decimal_comma", "players"):
        assert d34.name_kinds(word) == frozenset(), word
    assert d34.name_kinds("Parent_5125550199") == {"phone"} and d34.name_kinds("Ph 5551234567") == {"phone"}
    assert "phone" in d34.name_kinds("Garcia_5125550199")


# --------------------------------------------------------------------------- title rows
def test_pass3_title_rows_naming_a_person_are_flagged() -> None:
    screen = d34.NameScreen()
    for text in ("IEP Progress Report Kaden Yazzie", "Kaden Yazzie, Grade 3",
                 "Student: Ana Garcia Senior", "Student Name Olivia Thibodeaux", "Teacher Kaden Yazzie Room 12"):
        assert d34._verbatim_kinds(screen, text) == {"name"}, text
    for text in ("Fall 2024 Reading Benchmark - Grade 3", "Lincoln High School - Period 3 Roster",
                 "Students in Grades 3, 4 and 5", "Prepared by the Office of Assessment"):
        assert d34._verbatim_kinds(screen, text) == frozenset(), text


# --------------------------------------------------------------------------- real IDs: repeated evidence, context
def test_pass3_two_per_student_files_in_a_large_namespace_are_repeated_evidence() -> None:
    import random

    rng = random.Random(5)
    ids = rng.sample(range(100000, 1000000), 50000)  # p = 0.056 per six-digit number
    keys = d34.IdKeys([NS([str(x) for x in ids])])
    x, y = str(ids[17]), str(ids[4242])
    assert keys.match(f"IEP_{x}")  # K-01: after a person or document label a real ID is strong on its own
    assert not keys.match(f"Goals_{x}")  # one alone is weak
    assert set(d34.id_group_flags(keys, [f"Goals_{x}", f"Goals_{y}"])) == {0, 1}
    others = [str(v) for v in range(100000, 1000000, 7919) if v not in set(ids)][:8]
    survey = [f"Survey_{x}", f"Survey_{y}"] + [f"Survey_{v}" for v in others]
    assert d34.id_group_flags(keys, survey) == {}  # two of ten numbers: chance
    dense = d34.IdKeys([NS([str(v) for v in range(10001, 38001)])])
    zips = ["30303", "30305", "30308", "30310", "30311", "30312"]
    assert d34.id_group_flags(dense, zips, context=d34.numbered_context("enrollment_by_zip")) == {}
    assert d34.numbered_context("Rooms") and d34.numbered_context("by_campus")
    assert not d34.numbered_context("Period 3") and not d34.numbered_context("iep")


def test_pass3_numbered_labels_of_a_column_are_not_indexed() -> None:
    assert leaks.numbered_labels([f"Week {k}" for k in range(1, 37)] + ["Week of Ana"], "Week") == {
        f"Week {k}" for k in range(1, 37)}
    assert leaks.numbered_labels(["Question 12", "Q12", "Garcia 12"], "Item Question") == {"Question 12"}


def test_pass3_html_classes_are_an_allowlist() -> None:
    from datanonymous.io.excel_html import clean_css, safe_class

    assert safe_class("quixote") == "" and safe_class("xl65 quixote") == "xl65"
    assert safe_class("header odd total-row text-center") == "header odd total-row text-center"
    css = clean_css("td.quixote { color: red; } .header {font-weight:bold}")
    assert "quixote" not in css and ".header" in css


def test_pass3_usage_messages_screen_typed_path_parts_in_any_form() -> None:
    typed = ["tests/x/Garcia Ana/roster.csv", "--out-dir", "o"]
    msg = "input not found: tests\\x\\Garcia Ana\\roster.csv"
    assert "Garcia" not in d34.redact_message(msg, typed=typed)
    msg = "--out-dir C:\\data\\Okonkwo Chinedu\\o is inside the input folder C:\\data\\Okonkwo Chinedu"
    assert "Okonkwo" not in d34.redact_message(msg)


# --------------------------------------------------------------------------- end to end (pass 3)
def data_written(out: Path) -> list[str]:
    """The synthetic files of the out dir (not its README, data dictionary and schema)."""
    return [p for p in written(out) if p not in ("README.md", "DATA_DICTIONARY.md", "schema.json")]


def _book(path: Path, sheets: list[str], header: list[str] | None = None, title: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    for k, name in enumerate(sheets):
        ws = wb.active if k == 0 else wb.create_sheet()
        ws.title = name
        if title is not None:
            ws.append([title])
            ws.append([])
        ws.append(header or ["grade_level", "score"])
        for r in numbers(len(header or ["grade_level", "score"])):
            ws.append(r)
    wb.save(path)


def test_pass3_full_names_beside_school_words_fail_files_folders_sheets_and_headers(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "counseling" / "Ana Garcia Senior Project" / "rubric.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "counseling" / "Garcia_Ana_Scholarship.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "nurse" / "Senior Garcia, Ana.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "nurse" / "Kindergarten Chinedu Okafor.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "nurse" / "Ana Garcia Medication Administration Record.csv", ["grade_level", "dose"], numbers(2))
    _book(src / "counseling" / "exit_interviews.xlsx", ["Liam Johnson Senior Exit", "Summary"])
    write_csv(src / "counseling" / "seminar.csv", ["grade_level", "Mr. Garcia Senior Seminar"], numbers(2))
    write_csv(src / "campus" / "Jesse H Jones High School" / "attendance.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "campus" / "Senior Exit Survey.csv", ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    st = statuses(rep)
    assert sum(1 for v in st.values() if v == "FAIL") == 7, st
    assert "campus/Jesse H Jones High School/attendance.csv" in written(out)
    assert "campus/Senior Exit Survey.csv" in written(out)
    assert_never_printed(["Ana Garcia", "Garcia_Ana", "Garcia, Ana", "Chinedu Okafor", "Liam Johnson",
                          "Mr. Garcia"], out, rep, console=res.output)


def test_pass3_known_pairs_in_longer_names_fail(tmp_path: Path) -> None:
    src = tmp_path / "in"
    _book(src / "sped" / "Kaden Yazzie IEP 2024.xlsx", ["Goals"])
    write_csv(src / "sped" / "Olivia Thibodeaux - Case File" / "progress.csv", ["grade_level", "score"], numbers(2))
    _book(src / "sped" / "caseload.xlsx", ["Aiden Quarshie Goals", "Summary"])
    write_csv(src / "sped" / "goals_wide.csv", ["goal", "baseline", "Zyaire Okafor IEP"], numbers(3))
    roster = src / "roster.csv"
    roster.write_text("Roster for Maria Lopez,,\nstudent_id,score,grade_level\n"
                      + "".join(f"{400000 + r},{r},9\n" for r in range(N)), encoding="utf-8")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3 and set(statuses(rep).values()) == {"FAIL"}, statuses(rep)
    assert_never_printed(["Yazzie", "Thibodeaux", "Quarshie", "Zyaire", "Maria Lopez"], out, rep,
                         console=res.output)


def test_pass3_real_ids_of_per_student_files_and_zip_sheets(tmp_path: Path) -> None:
    import random

    src = tmp_path / "in"
    rng = random.Random(9)
    ids = rng.sample(range(100000, 1000000), 50000)
    write_csv(src / "big" / "roster.csv", ["student_id", "grade_level"], [[x, 9 + x % 4] for x in ids])
    x, y = ids[11], ids[2222]
    for v in (x, y):
        write_csv(src / "big" / "iep" / f"IEP_{v}.csv", ["goal", "baseline", "current"],
                  [[f"G{k}", k, k + 5] for k in range(N)])
    write_csv(src / "dense" / "students.csv", ["student_id", "zip", "grade_level"],
              [[v, 30301 + v % 24, 9 + v % 4] for v in range(10001, 38001)])
    _book(src / "dense" / "enrollment_by_zip.xlsx", ["30303", "30305", "30308", "30310", "30311", "30312"])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    failed = sorted(k for k, v in statuses(rep).items() if v == "FAIL")
    assert failed == ["big/iep/[redacted-1].csv", "big/iep/[redacted-2].csv"], failed
    assert "dense/enrollment_by_zip.xlsx" in written(out)
    assert_never_printed([str(x), str(y)], out, rep, console=res.output)


def test_pass3_numbered_label_columns_and_macos_files_pass(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "reading_minutes.csv", ["student_id", "Week", "minutes"],
              [[300000 + s, f"Week {w}", (s * w) % 200] for s in range(8) for w in range(1, 37)])
    write_csv(src / "items.csv", ["student_id", "Question", "correct"],
              [[300000 + s, f"Question {q}", (s + q) % 2] for s in range(6) for q in range(1, 41)])
    (src / "._reading_minutes.csv").write_bytes(b"\x00\x05\x16\x07\x00\x02\x00\x00" + b"\x00" * 200)
    (src / "__MACOSX").mkdir()
    (src / "__MACOSX" / "._items.csv").write_bytes(b"\x00\x05\x16\x07\x00\x02\x00\x00" + b"\x00" * 200)
    (src / "._Garcia_Ana_IEP.csv").write_bytes(b"\x00\x05\x16\x07\x00\x02\x00\x00" + b"\x00" * 200)
    for name in ("Extract_5839201746.csv", "Qualtrics_SV_5839201746.csv"):
        write_csv(src / name, ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    skipped = report(rep)["skipped"]
    assert len(skipped) == 3 and all("AppleDouble" in s["reason"] for s in skipped), skipped
    assert sum(s["path"].startswith("[redacted-") for s in skipped) == 1
    assert_never_printed(["Garcia_Ana"], out, rep, console=res.output)


def test_pass3_other_scripts_fail_and_never_print(tmp_path: Path) -> None:
    src = tmp_path / "in"
    write_csv(src / "Иван Петров" / "scores.csv", ["grade_level", "score"], numbers(2))
    write_csv(src / "wide.csv", ["Item", "Points", "张伟", "محمد علي"], [[f"I{r}", 10, r % 2, r % 3] for r in range(N)])
    write_csv(src / "clean.csv", ["成绩", "学号", "이름"], numbers(3))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    assert data_written(out) == ["clean.csv"], written(out)
    assert_never_printed(["Иван", "Петров", "张伟", "محمد"], out, rep, console=res.output)


def test_pass3_keep_titles_and_html_classes(tmp_path: Path) -> None:
    src = tmp_path / "in"
    _book(src / "progress.xlsx", ["Progress"], ["goal", "baseline", "current"],
          title="IEP Progress Report Kaden Yazzie")
    _book(src / "roster.xlsx", ["Roster"], title="Class list - Teacher: Ana Garcia")
    _book(src / "plain.xlsx", ["Data"], title="Fall 2024 Reading Benchmark - Grade 3")
    rows = "\n".join(f'<tr><td class="quixote">{r}</td><td class="xl65">{50 + r}</td></tr>' for r in range(N))
    (src / "gradebook.xls").write_text(
        "<html><head><style>td.quixote { color: red; } .xl65 {color: blue}</style></head><body><table>"
        "<tr><th>item</th><th>score</th></tr>" + rows + "</table></body></html>", encoding="utf-8")
    res, out, rep = run(src, tmp_path, "--keep-titles")
    assert res.exit_code == 3
    assert data_written(out) == ["gradebook.xls", "plain.xlsx"], written(out)
    assert "xl65" in (out / "gradebook.xls").read_text(encoding="utf-8")
    assert_never_printed(["Yazzie", "Ana Garcia", "quixote"], out, rep, console=res.output)


def test_pass3_usage_and_yaml_errors_never_print_a_name(tmp_path: Path) -> None:
    src = tmp_path / "Okonkwo Chinedu"
    write_csv(src / "s.csv", ["student_id", "score"], [[100000 + r, r] for r in range(N)])
    outs = []
    for args in ([(tmp_path / "Garcia Ana" / "roster.csv").as_posix(), "--out-dir", tmp_path / "o"],
                 ["./" + (tmp_path / "Garcia Ana").as_posix() + "/", "--out-dir", tmp_path / "o"],
                 [src.as_posix(), "--out-dir", (src / "o").as_posix()],
                 [src, "--out-dir", tmp_path / "o", "--config", (tmp_path / "Lopez Maria.yaml").as_posix()]):
        res = run_cli(*args)
        assert res.exit_code == 2, res.output
        outs.append(res.output)
    bad = tmp_path / "bad.yaml"
    bad.write_text('force_factor:\n  - "Okonkwo, Chinedu"\n  - [Lopez Maria\n', encoding="utf-8")
    bad2 = tmp_path / "bad2.yaml"
    bad2.write_text("id_col: Okonkwo Chinedu: 3\n", encoding="utf-8")
    for cfg in (bad, bad2):
        res = run_cli(src, "--out-dir", tmp_path / "o", "--config", cfg)
        assert res.exit_code == 2 and "not valid YAML" in res.output and "line " in res.output, res.output
        outs.append(res.output)
    blob = "\n".join(outs)
    assert not any(s in blob for s in ("Garcia", "Okonkwo", "Lopez", "Maria")), blob


def test_pass3_many_numbered_headers_equal_to_ids_by_chance_are_no_evidence() -> None:
    import random

    keys = d34.IdKeys([NS([str(x) for x in random.Random(1).sample(range(100000, 1000000), 50000)])])
    items = [f"Q{k}_{random.Random(k).randint(100000, 999999)}" for k in range(2000)]
    assert d34.id_group_flags(keys, items) == {}  # about 112 of them equal a real ID by chance
    real = [x for x in sorted(keys._index()[0]) if not d34.natural_number(x)][100:130]
    assert len(d34.id_group_flags(keys, ["Assignment"] + [f"score_{x}" for x in real])) == 30


# =========================================================================== D34 freeze (best-effort layer)
# --------------------------------------------------------------------------- false refusals of the last attack pass
@pytest.mark.parametrize("name", [
    # campuses named after people: a middle initial, more name words before the school form, AskTED upper case,
    # New York school numbers
    "Victor H Hexter ES", "Maria T. Quenby Elementary School", "Edward R Murrow HS", "Louis D Brandeis HS",
    "Gus Garcia YMLA", "Gus Garcia Young Men's Leadership Academy", "Bertha Sadler Means YWLA",
    "Kay Bailey Hutchison Elementary", "John Paul Stevens HS", "Dr. Wright L. Lassiter Jr. Early College High School",
    "Tom W Field Jr ES", "Barbara M Manns Education Center", "JACK YATES H S", "VICTOR H HEXTER EL",
    "TOM C CLARK H S", "Victor_H_Hexter_EL_Attendance", "PS 123 Mahalia Jackson", "MS 51 William Alexander",
    "Gus Garcia", "Norman-Sims", "O Henry",
    # authors, characters and historical figures in gradebook headers
    "Gary Soto Seventh Grade Quiz", "Maniac Magee Ch 1-10", "Thank You Mr Falker Response", "Harrison Bergeron Quiz",
    "Charlie Gordon Progress Report", "Because of Mr Terupt Book Talk", "Sandra Cisneros Eleven",
    "Jason Reynolds Ghost Ch 1-3", "Bernardo de Galvez DBQ", "Andrew Johnson Impeachment Quiz",
    "Abigail Adams Letter", "Benedict Arnold Debate", "James Armistead Spy Report", "James Fannin Goliad",
    "Gail Borden Inventions", "Jim Hogg Reforms", "Bessie Coleman Biography", "Kurt Vonnegut Harrison Bergeron",
    "Jess Aarons Leslie Burke", "Harper Lee", "Amy Tan Two Kinds", "Mrs Luella Bates Washington Jones",
    # places named like people: everyday words are never one word of a name pair
    "Oak Hill", "Oak_Hill_Enrollment", "Tom Bean", "Tom_Bean_ISD_TAPR", "Barton Hills", "Pleasant Hill",
    "Wells Branch", "Forest Oak", "Forest Hill",
    # vendors and brands, court cases, cities and states
    "Sara Lee", "Charles Schwab", "Ernst Young", "Dell Technologies", "Mr Peanut", "Mrs Doubtfire",
    "Brown v Board", "Plessy v Ferguson", "Columbus, Georgia", "Augusta, Georgia", "Charlotte, North Carolina",
    # 10 digits: no valid North American number, or after a numbered thing
    "Students_5129110199", "Survey_2925550199", "Room 5125550199", "Extract_5839201746", "Student ID 5125550199",
])
def test_freeze_false_refusals_of_the_last_attack_pass(name: str) -> None:
    assert d34.name_kinds(name) == frozenset()


@pytest.mark.parametrize("name, kinds", [
    # the reliable checks stay: known pairs (also with an everyday middle name), layouts, titles, user names
    ("Ana Garcia", {"name"}), ("Garcia, Ana", {"name"}), ("Ana Rose Garcia", {"name"}), ("Mr. Quarshie", {"name"}),
    ("Ms Okafor Period 3", {"name"}), ("Jayden M.", {"name"}), ("jgarcia", {"name"}), ("Victor H Hexter", {"name"}),
    ("Fort Worth Ana Garcia", {"name"}), ("Kaden Begay", {"name"}), ("Le Minh", {"name"}),
    # URL-decoded names (percent escapes, '+' as a space)
    ("Kaden%20Yazzie%20IEP", {"name"}), ("Ana+Garcia+504", {"name"}), ("frontoffice%40school.org", {"email"}),
    # phone shapes
    ("Garcia_5125550199", {"phone"}), ("Parent_5125550199", {"phone"}), ("Call 312-555-0199", {"phone"}),
])
def test_freeze_reliable_checks_still_flag(name: str, kinds: set[str]) -> None:
    assert d34.name_kinds(name) == kinds


def test_freeze_dropped_low_precision_rules() -> None:
    """The rules the freeze dropped (README: What D34 does not catch): a given name beside a word off the lists, a
    header or everyday word in a pair, a hyphenated romanized given name, single given names or surnames among
    siblings, title rows listing given names. These may pass: the accepted residual risk."""
    for name in ("Quixote_Philomena", "Kaden Vrexley IEP 2024", "Kim Seo-yeon", "Summer Davis", "Joy Taylor",
                 "Mark Johnson", "Ana Price", "Ana Hall"):
        assert d34.name_kinds(name) == frozenset(), name
    screen = d34.NameScreen()
    for siblings in (["Aaliyah", "Brayden", "Ximena", "Nevaeh"], ["Okafor", "Nguyen", "Kowalski", "Hernandez",
                                                                 "Yamamoto"]):
        assert not any(k for k, _, _ in screen.group_kinds([(t, "") for t in siblings])), siblings
    assert d34._verbatim_kinds(screen, "Class list Liam Noah Emma") == frozenset()


def test_freeze_nanp_numbers() -> None:
    assert d34.nanp_number("5125550199") and d34.nanp_number("3125550187")
    for d in ("1125550199", "5121550199", "2115550199", "5129110199", "2925550199", "512555019"):
        assert not d34.nanp_number(d), d


def test_freeze_html_selectors_classes_and_attributes_are_allowlists() -> None:
    from datanonymous.io.excel_html import _safe_attr, clean_css, safe_class, safe_doctype, safe_selector

    for sel in ("td", ".xl65", "td.xl65, th", "@page", "<!--table", "a:link", "tr:nth-child(2n+1)", ".style0"):
        assert safe_selector(sel), sel
    for sel in ("Thibodeaux", "td.quixote", "#Okonkwo_Chinedu", 'td[title="Ximena Villarreal"]', ":not(Quixote)",
                "@media Okafor"):
        assert not safe_selector(sel), sel
    assert safe_class("xl65 font5 style0 header num quixote Okonkwo") == "xl65 font5 style0 header num"
    css = clean_css('<!--table {mso-displayed-decimal-separator:"\\.";}\n.style0 {mso-style-name:"Okonkwo Chinedu";'
                    ' mso-style-id:0;}\n.style1 {mso-style-name:Normal;}\ntd {font:10.0pt "Philomena Quixote", Arial;'
                    ' --student:"Kaden Yazzie"; list-style-type:"Zamarripa"; counter-reset:Bhattacharya;}\n'
                    'Thibodeaux {color:blue;}\n.xl65 {mso-number-format:"\\@"; color:red;}\n-->')
    for secret in ("Okonkwo", "Quixote", "Yazzie", "Zamarripa", "Bhattacharya", "Thibodeaux"):
        assert secret not in css, secret
    for kept in ('mso-displayed-decimal-separator:"\\."', "mso-style-name:Normal", "font:10.0pt Arial",
                 'mso-number-format:"\\@"', "color:red", "<!--table", "-->"):
        assert kept in css, kept
    assert _safe_attr("bgcolor", "#FFCC00") and not _safe_attr("bgcolor", "Ada")
    assert _safe_attr("align", "center") and not _safe_attr("align", "Garcia")
    assert _safe_attr("charset", "utf-8") and not _safe_attr("charset", "Okonkwo")
    assert safe_doctype('DOCTYPE html PUBLIC "-//W3C//DTD HTML 4.0 Transitional//EN"').endswith("Transitional//EN\"")
    assert safe_doctype('DOCTYPE html PUBLIC "-//Philomena Quixote//DTD//EN" "garcia_ana.dtd"') == "DOCTYPE html"


# --------------------------------------------------------------------------- end to end (freeze)
def test_freeze_campuses_authors_and_places_pass_end_to_end(tmp_path: Path) -> None:
    src = tmp_path / "in"
    _book(src / "Dallas_ES_Attendance_2024.xlsx", ["Victor H Hexter ES", "Edwin J Kiest ES", "George W Truett ES",
                                                  "Thomas L Marsalis ES", "Jerry R Junkins ES"])
    for c in ("Gus Garcia YMLA", "Bertha Sadler Means YWLA", "Louis D Brandeis HS", "John Paul Stevens HS",
              "Oliver Wendell Holmes HS", "Edward R Murrow HS", "Pleasant Hill", "Tom Bean ISD"):
        write_csv(src / "campuses" / c / "attendance.csv", ["grade_level", "score"], numbers(2))
    heads = ["Student ID", "Gary Soto Seventh Grade Quiz", "Maniac Magee Ch 1-10", "Thank You Mr Falker Response",
             "Harrison Bergeron Quiz", "Charlie Gordon Progress Report", "Because of Mr Terupt Book Talk",
             "Sandra Cisneros Eleven", "Jason Reynolds Ghost Ch 1-3"]
    write_csv(src / "Gradebook_ELA7_P3.csv", heads, [[482900 + r] + [(r + c) % 50 + 50 for c in range(8)]
                                                     for r in range(N)])
    _book(src / "AISD_ES_Enrollment.xlsx", ["Oak Hill", "Barton Hills", "Wells Branch", "Zilker", "Casis"])
    write_csv(src / "Oak_Hill_Enrollment.csv", ["grade_level", "score"], numbers(2))
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert not any(r.startswith(d34.D34_PREFIX) for f in report(rep)["files"] for r in f["reasons"])
    assert "campuses/Gus Garcia YMLA/attendance.csv" in written(out)


def test_freeze_url_encoded_file_names_are_screened_decoded(tmp_path: Path) -> None:
    src = tmp_path / "in"
    for name in ("Kaden%20Yazzie%20IEP.csv", "Ana+Garcia+504.csv", "attendance%20summary.csv"):
        write_csv(src / "iep" / name, ["student_id", "score"], [[700000 + r, r % 50] for r in range(N)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3
    st = statuses(rep)
    assert list(st.values()).count("FAIL") == 2, st
    assert st["iep/attendance%20summary.csv"] != "FAIL"
    assert_never_printed(["Yazzie", "Ana+Garcia", "Ana Garcia"], out, rep, console=res.output)


def test_freeze_a_missing_sheet_never_lists_the_sheets(tmp_path: Path) -> None:
    src = tmp_path / "in"
    _book(src / "reading_logs.xlsx", ["Jayden", "Aaliyah", "Brayden", "Ximena"])
    _book(src / "summary.xlsx", ["Summary"])
    res, out, rep = run(src, tmp_path, "--sheet", "Summary")
    assert res.exit_code == 3
    reason = next(f for f in report(rep)["files"] if f["source"] == "reading_logs.xlsx")["reasons"][0]
    assert reason == "no sheet named 'Summary' (the workbook has 4 sheet(s))", reason
    ins = run_cli("inspect", src, "--sheet", "Summary")
    assert_never_printed(["Jayden", "Aaliyah", "Brayden", "Ximena"], out, rep, console=res.output + ins.output)


def test_freeze_html_doctype_processing_instructions_and_css_names_are_not_copied(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    rows = "".join(f"<tr><td class=\"xl65\" style='font:9pt \"Warsame Hodan\"'>{482900 + r}</td><td>{r % 50}</td>"
                   f"</tr>\n" for r in range(N))
    (src / "grades.xls").write_text(
        '<?xml version="1.0" encoding="utf-8"?>\n<?xml-stylesheet href="Okonkwo_Chinedu_style.xsl"?>\n'
        '<!DOCTYPE html PUBLIC "-//Philomena Quixote//DTD for Ana Garcia//EN" "garcia_ana.dtd">\n'
        '<html><head><meta charset="utf-8"><style>\n.style0 {mso-style-name:"Okonkwo Chinedu highlight";}\n'
        'td {font:10.0pt "Philomena Quixote Script", Arial; --student:"Kaden Yazzie";}\n'
        'td[title="Ximena Villarreal"] {color:red;}\nThibodeaux {color:blue;}\n'
        '.xl65 {mso-number-format:"0"; list-style-type:"Josue Zamarripa"; counter-reset:Bhattacharya;}\n'
        '</style></head><body>\n<table border="1">\n<tr><th>student_id</th><th>score</th></tr>\n' + rows
        + "</table></body></html>\n", encoding="utf-8")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    text = (out / "grades.xls").read_text(encoding="utf-8")
    assert text.startswith('<?xml version="1.0" encoding="utf-8"?>') and "<!DOCTYPE html>" in text
    assert "xl65" in text and "mso-number-format" in text and "Arial" in text
    assert_never_printed(["Okonkwo", "Chinedu", "Quixote", "Philomena", "Ana Garcia", "garcia_ana", "Yazzie",
                          "Villarreal", "Thibodeaux", "Zamarripa", "Bhattacharya", "Warsame"], out, rep,
                         console=res.output)
