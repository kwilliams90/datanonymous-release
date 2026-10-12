"""AREA E end to end: the CLI on the Excel fixtures (native typing, layout, privacy, determinism, options)."""

from __future__ import annotations

import json
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

from conftest import run_cli

FIX = Path(__file__).resolve().parent / "fixtures" / "excel"
DATA_FILES = ("multi_sheet.xlsx", "fake_html.xls", "fake_tsv.xls", "real_biff.xls")
#: strings that only exist in titles, notes, comments, links, properties or HTML chrome of the fixtures
SECRETS = ("jsmith", "Jane Q. Registrar", "Confidential roster", "Pemberton-Hale", "lincoln-usd",
           "Lincoln Unified", "10.1.2.3", "Roster Report", "CONFIDENTIAL", "Attendance summary")


def fixture_input(tmp_path: Path) -> Path:
    inp = tmp_path / "input"
    inp.mkdir()
    for name in DATA_FILES:
        shutil.copy(FIX / name, inp / name)
    return inp


@pytest.fixture(scope="module")
def excel_run(tmp_path_factory: pytest.TempPathFactory):
    base = tmp_path_factory.mktemp("excel_run")
    inp = fixture_input(base)
    out, rep = base / "out", base / "rep"
    res = run_cli(inp, "--out-dir", out, "--report-dir", rep, "--seed", "7")
    assert res.exit_code == 0, res.output
    return inp, out, rep


def xlsx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist())


def string_cells(path: Path) -> set[str]:
    wb = openpyxl.load_workbook(path, data_only=True)
    return {c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)}


def column_values(path: Path, sheet: str, header_row: int, name: str) -> list:
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    hdr = [c.value for c in ws[header_row]]
    j = hdr.index(name)
    return [row[j] for row in ws.iter_rows(min_row=header_row + 1, values_only=True)]


def test_outputs_names_and_statuses(excel_run) -> None:
    _, out, rep = excel_run
    names = sorted(p.name for p in out.iterdir())
    assert names == sorted(["multi_sheet.xlsx", "fake_html.xls", "fake_tsv.xls", "real_biff.xlsx", "README.md",
                            "DATA_DICTIONARY.md", "schema.json"])
    data = json.loads((rep / "report.json").read_text(encoding="utf-8"))
    assert data["totals"]["FAIL"] == 0 and data["totals"]["files"] == 4
    report = (rep / "files" / "multi_sheet.xlsx.md").read_text(encoding="utf-8")
    assert "Excel extras not copied" in report and "hidden sheet(s) processed" in report
    xls_report = (rep / "files" / "real_biff.xls.md").read_text(encoding="utf-8")
    assert "written as .xlsx" in xls_report
    for f in ("multi_sheet.xlsx", "fake_html.xls", "real_biff.xls"):
        assert "## Self-check" not in (rep / "files" / f"{f}.md").read_text(encoding="utf-8"), f


def test_xlsx_native_typing_and_layout(excel_run) -> None:
    inp, out, _ = excel_run
    p = out / "multi_sheet.xlsx"
    wb = openpyxl.load_workbook(p, data_only=True)
    assert [(ws.title, ws.sheet_state) for ws in wb.worksheets] == [
        ("Enrollment", "visible"), ("Courses", "visible"), ("Lookup", "hidden"), ("Empty", "visible")]
    ws = wb["Enrollment"]
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref == "A4:K44"
    assert [str(m) for m in ws.merged_cells.ranges] == ["A1:K1"] and ws["A1"].data_type == "s"
    assert [c.value for c in ws[4]] == ["student_id", "last_name", "grade", "enroll_date", "gpa", "attendance_rate",
                                        "zip_code", "econ_dis", "test_score", "credits", "points"]
    data_rows = list(ws.iter_rows(min_row=5, max_row=44))
    for row in data_rows:
        sid, last, grade, date, gpa, rate, zip_, econ, score, credits, points = row
        # text ID cells stay text (the width follows the run-level student_id namespace, shared with real_biff.xls)
        assert sid.data_type == "s" and re.fullmatch(r"\d{6,7}", sid.value)
        assert grade.data_type == "s"
        assert date.is_date and date.number_format == "mm/dd/yyyy"
        assert gpa.value is None or (gpa.data_type == "n" and gpa.number_format == "0.00")
        assert rate.data_type == "n" and rate.number_format == "0.0%"
        assert zip_.data_type == "n" and zip_.number_format == "00000"
        assert econ.value is None or econ.data_type == "b"
        assert score.data_type == "e" and score.value == "#N/A" or score.data_type == "n"
        assert credits.number_format == "0.0" and points.data_type == "n"
    src_rows = list(openpyxl.load_workbook(inp / "multi_sheet.xlsx", data_only=True)["Enrollment"].iter_rows(
        min_row=5, max_row=44, values_only=True))
    for j in (4, 7, 8):  # blanks and #N/A error cells at their exact counts
        assert (sum(1 for r in data_rows if r[j].value in (None, "#N/A"))
                == sum(1 for r in src_rows if r[j] in (None, "#N/A"))), j
    assert sum(1 for r in data_rows if r[8].value == "#N/A") == 5
    assert ws["A45"].value == "Total" and ws["B45"].value == 40
    assert ws["J45"].value == pytest.approx(sum(r[9].value for r in data_rows))
    gpas = [r[4].value for r in data_rows if r[4].value is not None]
    assert ws["E45"].value == pytest.approx(sum(gpas) / len(gpas))
    assert ws["A47"].value and ws["A47"].value != "Source: SIS nightly export. Questions: jsmith@lincoln-usd.example"
    assert ws.max_row == 47
    courses = wb["Courses"]
    assert all(c.data_type == "s" and len(c.value) == 5 for c in courses["A"][1:])  # '00120' stays text
    assert all(c.is_date and c.number_format == "yyyy-mm-dd" for c in courses["D"][1:])


def test_pandas_reads_the_same_dtypes(excel_run) -> None:
    import pandas as pd

    inp, out, _ = excel_run
    a = pd.read_excel(inp / "multi_sheet.xlsx", sheet_name="Enrollment", header=3, nrows=40)
    b = pd.read_excel(out / "multi_sheet.xlsx", sheet_name="Enrollment", header=3, nrows=40)
    assert list(a.columns) == list(b.columns) and a.dtypes.astype(str).tolist() == b.dtypes.astype(str).tolist()
    c = pd.read_excel(out / "multi_sheet.xlsx", sheet_name=None)
    assert list(c) == ["Enrollment", "Courses", "Lookup", "Empty"]


def test_no_real_sensitive_value_in_outputs_or_report(excel_run) -> None:
    inp, out, rep = excel_run
    real_ids = {v for v in column_values(inp / "multi_sheet.xlsx", "Enrollment", 4, "student_id")[:40]}
    real_names = {v for v in column_values(inp / "multi_sheet.xlsx", "Enrollment", 4, "last_name")[:40]}
    syn_ids = set(column_values(out / "multi_sheet.xlsx", "Enrollment", 4, "student_id")[:40])
    assert len(syn_ids) == 40 and not (syn_ids & real_ids)
    cells = string_cells(out / "multi_sheet.xlsx")
    assert not (cells & real_names) and not (cells & real_ids)
    text = xlsx_text(out / "multi_sheet.xlsx") + xlsx_text(out / "real_biff.xlsx")
    html_out = (out / "fake_html.xls").read_text(encoding="utf-8")
    html_in = (inp / "fake_html.xls").read_text(encoding="utf-8")
    real_html_names = set(re.findall(r"<td>([A-Z][a-z]+, [A-Z][a-z]+)</td>", html_in))
    real_html_ids = set(re.findall(r">(\d{6})</td>", html_in))
    assert real_html_names and not any(f">{n}<" in html_out for n in real_html_names)
    assert not any(f">{i}<" in html_out for i in real_html_ids)
    report_text = "\n".join(p.read_text(encoding="utf-8") for p in rep.rglob("*") if p.is_file())
    docs_text = "\n".join((out / d).read_text(encoding="utf-8") for d in ("README.md", "DATA_DICTIONARY.md",
                                                                           "schema.json"))
    for secret in SECRETS:
        assert secret not in text, secret
        assert secret not in html_out, secret
        assert secret not in report_text, secret
        assert secret not in docs_text, secret
    for name in real_names | real_html_names:
        assert name not in report_text and name not in docs_text, name
    for sid in real_ids | real_html_ids:
        assert not re.search(rf"(?<!\d){sid}(?!\d)", report_text), sid


def test_disguised_tsv_and_html_are_mirrored(excel_run) -> None:
    inp, out, _ = excel_run
    tsv = (out / "fake_tsv.xls").read_bytes()
    assert tsv.count(b"\r\n") == 26 and tsv.count(b"\t") == (inp / "fake_tsv.xls").read_bytes().count(b"\t")
    tsv.decode("cp1252")
    assert tsv.splitlines()[0] == b"Student_Number\tLast_Name\tGrade\tSchool"
    html_out = (out / "fake_html.xls").read_bytes()
    assert html_out.startswith(b"<html ") and html_out.count(b"<tr>") == 32 and b"\r\n" in html_out


def test_same_seed_same_bytes(excel_run, tmp_path: Path) -> None:
    inp, out, _ = excel_run
    out2 = tmp_path / "out2"
    res = run_cli(inp, "--out-dir", out2, "--report-dir", tmp_path / "rep2", "--seed", "7")
    assert res.exit_code == 0, res.output
    for name in ("multi_sheet.xlsx", "fake_html.xls", "fake_tsv.xls", "real_biff.xlsx"):
        assert (out / name).read_bytes() == (out2 / name).read_bytes(), name


def test_keep_titles_copies_titles_verbatim(tmp_path: Path) -> None:
    inp = tmp_path / "in"
    inp.mkdir()
    shutil.copy(FIX / "multi_sheet.xlsx", inp / "multi_sheet.xlsx")
    # the fixture's footer quotes an e-mail address: written verbatim it is personal data in the output (D34)
    res = run_cli(inp / "multi_sheet.xlsx", "--out-dir", tmp_path / "o1", "--seed", "3", "--keep-titles")
    assert res.exit_code == 3 and "a footer row of sheet 1 of 4 (written verbatim with --keep-titles)" in res.output
    assert "jsmith@" not in res.output and not (tmp_path / "o1" / "multi_sheet.xlsx").exists()
    assert "a title/notes row of sheet 1 of 4" in res.output  # 'Printed by jsmith': a user name (D34)
    wb = openpyxl.load_workbook(inp / "multi_sheet.xlsx")
    ws = wb["Enrollment"]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "@" in c.value:
                c.value = "Source: SIS nightly export."
            elif isinstance(c.value, str) and "jsmith" in c.value:
                c.value = c.value.replace("jsmith", "the registrar")
    wb.save(inp / "multi_sheet.xlsx")
    out = tmp_path / "out"
    res = run_cli(inp / "multi_sheet.xlsx", "--out-dir", out, "--seed", "3", "--keep-titles")
    assert res.exit_code == 0, res.output
    ws = openpyxl.load_workbook(out / "multi_sheet.xlsx")["Enrollment"]
    assert ws["A1"].value == "Lincoln Unified School District - Enrollment Report"
    assert ws["A2"].value == "Printed by the registrar on 10/01/2024 08:15 AM"


def test_inspect_xlsx_writes_nothing(tmp_path: Path) -> None:
    inp = fixture_input(tmp_path)
    before = sorted(p.name for p in tmp_path.rglob("*"))
    res = run_cli("inspect", inp / "multi_sheet.xlsx")
    assert res.exit_code == 0, res.output
    assert "IDENTIFIER" in res.output and "Enrollment" in res.output
    assert sorted(p.name for p in tmp_path.rglob("*")) == before


def test_sheet_header_row_and_range_flags(tmp_path: Path) -> None:
    inp = fixture_input(tmp_path)
    out = tmp_path / "o1"
    res = run_cli(inp / "multi_sheet.xlsx", "--out-dir", out, "--seed", "1", "--sheet", "Courses")
    assert res.exit_code == 0, res.output
    assert openpyxl.load_workbook(out / "multi_sheet.xlsx").sheetnames == ["Courses"]
    out2 = tmp_path / "o2"
    res = run_cli(inp / "multi_sheet.xlsx", "--out-dir", out2, "--seed", "1", "--range", "Enrollment!A4:K44",
                  "--sheet", "Enrollment")
    assert res.exit_code == 0, res.output
    ws = openpyxl.load_workbook(out2 / "multi_sheet.xlsx")["Enrollment"]
    assert ws.max_row == 44 and ws["A4"].value == "student_id"  # footer rows below the range are not written
    out3 = tmp_path / "o3"
    res = run_cli(inp / "multi_sheet.xlsx", "--out-dir", out3, "--seed", "1", "--sheet", "Missing")
    assert res.exit_code == 3 and "no sheet named" in res.output
    res = run_cli(inp / "multi_sheet.xlsx", "--out-dir", tmp_path / "o4", "--range", "Enrollment!zz")
    assert res.exit_code == 2


def test_ids_share_the_run_registry_across_sheets(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    st = wb.active
    st.title = "Students"
    st.append(["student_id", "grade"])
    ids = [f"{7000 + i * 3:06d}" for i in range(30)]
    for i, sid in enumerate(ids):
        st.append([sid, ["09", "10", "11", "12"][i % 4]])
        st.cell(i + 2, 1).data_type = "s"
    sc = wb.create_sheet("Scores")
    sc.append(["student_id", "score"])
    for i in range(60):
        sc.append([ids[(i * 7) % 20], 50 + i % 40])
        sc.cell(i + 2, 1).data_type = "s"
    inp = tmp_path / "in"
    inp.mkdir()
    wb.save(inp / "book.xlsx")
    out = tmp_path / "out"
    res = run_cli(inp / "book.xlsx", "--out-dir", out, "--seed", "5")
    assert res.exit_code == 0, res.output
    o = openpyxl.load_workbook(out / "book.xlsx")
    syn_students = [r[0] for r in o["Students"].iter_rows(min_row=2, values_only=True)]
    syn_scores = [r[0] for r in o["Scores"].iter_rows(min_row=2, values_only=True)]
    assert set(syn_scores) <= set(syn_students) and len(set(syn_scores)) == 20  # joins survive across sheets
    assert not set(syn_students) & set(ids)
    assert all(isinstance(v, str) and len(v) == 6 for v in syn_students)


def test_xlsm_written_as_xlsx_without_macros(tmp_path: Path) -> None:
    inp = tmp_path / "in"
    inp.mkdir()
    with zipfile.ZipFile(FIX / "multi_sheet.xlsx") as src, zipfile.ZipFile(inp / "book.xlsm", "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"\x00not real vba")
    out, rep = tmp_path / "out", tmp_path / "rep"
    res = run_cli(inp, "--out-dir", out, "--report-dir", rep, "--seed", "2")
    assert res.exit_code == 0, res.output
    assert (out / "book.xlsx").is_file() and not (out / "book.xlsm").exists()
    assert not any("vba" in n.lower() for n in zipfile.ZipFile(out / "book.xlsx").namelist())
    assert "macros (VBA) never copied" in (rep / "files" / "book.xlsm.md").read_text(encoding="utf-8")
