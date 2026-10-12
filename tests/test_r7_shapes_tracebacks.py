"""Round 7 (shapes-tracebacks cluster): names and values share one shape test, and no diagnostic prints a value.

* R7-03: a header, sheet, file or folder name is screened with THE copy gate's shape test of values
  (``detect.text.unit_shapes``: also a bare valid North American number, R6-01), so a JSON key, a file, sheet or
  folder named ``3125550187`` is a phone number (D34: the file FAILs, the name is never printed). The D34
  exceptions (a time stamp, an export/job/record number, no valid number) only ever excuse a bare run of digits,
  never another shape; a phone shape no excused run accounts for stays a phone number (fail-safe).
* R7-04: nothing the tool prints about an exception holds its text: ``-v``/``-vv`` tracebacks (internal errors while
  reading, profiling, generating, leak-scanning, and the CLI's top-level handlers) show the exception classes and
  the code file, line and function of each frame only, never a message, argument, note, chained exception's text or
  local variable; a library's warning is printed with ``-v`` only, as its category and code location; a tool error
  never quotes a foreign exception's message.

Every value here is invented fixture data.
"""

from __future__ import annotations

import ast
import csv
import io
import json
import random
import warnings
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest

from conftest import run_cli
from datanonymous import cli, diagnostics, pipeline
from datanonymous.detect import text as dtext
from datanonymous.io import csvio
from datanonymous.privacy import headers as d34
from datanonymous.privacy import leaks

SRC = Path(__file__).resolve().parents[1] / "src" / "datanonymous"
N = 30
PHONE = "3125550187"
SECRET = "Student takes lithium nightly and needs monitoring"
SECRET_WORD = "lithium"


# --------------------------------------------------------------------------- helpers
def write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return path


def roster(path: Path) -> Path:
    return write_csv(path, ["student_id", "score"], [[f"S{1000 + r}", (r * 7) % 50] for r in range(N)])


def run(src: Path, tmp: Path, *extra: str, name: str = "run"):
    out, rep = tmp / f"{name}-out", tmp / f"{name}-rep"
    res = run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", "7", *extra)
    return res, out, rep


def everything(*dirs: Path, console: str = "") -> str:
    parts = [console]
    for d in dirs:
        if not d.exists():
            continue
        for p in sorted(d.rglob("*")):
            parts.append(p.relative_to(d).as_posix())
            if p.is_file():
                parts.append(p.read_bytes().decode("utf-8", errors="replace"))
    return "\n".join(parts)


def report(rep: Path) -> dict[str, Any]:
    return json.loads((rep / "report.json").read_text(encoding="utf-8"))


def statuses(rep: Path) -> dict[str, str]:
    return {f["source"]: f["status"] for f in report(rep)["files"]}


def written(out: Path) -> list[str]:
    return sorted(p.relative_to(out).as_posix() for p in out.rglob("*") if p.is_file())


# =========================================================================== R7-03: one shape test
@pytest.mark.parametrize("name", [
    PHONE, "13125550187", "+13125550187", f"{PHONE}.0", f"{PHONE} 5125550199", "312_555_0187", "312-555-0187",
    "(312) 555-0187", f"Parent_{PHONE}", f"score {PHONE}", f"Q17_{PHONE}",
])
def test_a_bare_valid_number_in_a_name_is_a_phone_number(name: str) -> None:
    assert d34.name_kinds(name) == {"phone"}, name


@pytest.mark.parametrize("name", [
    # the D34 exceptions stay: time stamps, export/job/record numbers, numbers that are no valid phone number
    "2023101512", "export_2023101512", "roster_1696342522", "1696342522", "Extract_5839201746",
    "Qualtrics_SV_5839201746", "PO 4500012345", "Student ID 5125550199", "1234567890", "5551234567",
    "Students_5129110199", "ISBN 0134685997", "S3125550187", "3125550187abc", "512555019",
])
def test_the_d34_exceptions_still_excuse_bare_digit_runs(name: str) -> None:
    assert "phone" not in d34.name_kinds(name), name


def test_names_and_values_share_the_nanp_test() -> None:
    assert d34.nanp_number is dtext.nanp_digits


_SHAPES = {"email", "phone", "ssn", "address"}
#: Names the D34 exceptions excuse although a value would be a phone number (a time stamp, a record number).
_EXCUSED = {"2023101512", "Student ID 5125550199", "Invoice 3105551234", "Run 6175550123"}


@pytest.mark.parametrize("text", [
    PHONE, "13125550187", "+13125550187", f"{PHONE}.0", f"{PHONE}, 5125550199", "Call 312-555-0199",
    "+44 20 7946 0958", "123-45-6789", "123 45 6789", "ana.garcia@school.org", "Front Office;ana@x.edu",
    "12 Oak Street", "lives at 4410 Elm Ave", "Call mom 3125550199", "Ph 5551234567", *sorted(_EXCUSED),
])
def test_a_name_never_misses_a_shape_its_value_would_have(text: str) -> None:
    """The header screen starts from the value test (``unit_shapes``) and only the documented D34 exceptions narrow
    it: no personal-data shape a value would have is lost in a name (R7-03: no pre-classification gap)."""
    want = set(dtext.unit_shapes(text)) & _SHAPES
    got = set(d34.name_kinds(text))
    if text in _EXCUSED:
        assert want == {"phone"} and "phone" not in got
    else:
        assert want <= got, (text, want, got)


@pytest.mark.parametrize("text, kind", [
    (f"Office_{PHONE}", "phone"), (f"{PHONE}_5125550199", "phone"), ("312_555_0187", "phone"),
    ("123_45_6789", "ssn"), ("12_Oak_Street", "address"),
])
def test_an_underscore_is_a_space_for_values_too(text: str, kind: str) -> None:
    """One shape test: what a name's screen sees through an underscore, a value's does too, so a shared value such
    as a front-office number written ``Office_3125550187`` is never kept for its count (D35 never-kept rule)."""
    assert kind in dtext.unit_shapes(text) and kind in d34.name_kinds(text)
    assert not dtext.unit_kept(text, holders=500, min_cell=10)
    assert dtext.unit_kept("Office_Main", holders=500, min_cell=10)


def test_every_bare_ten_digit_name_is_a_phone_exactly_when_valid_and_no_time_stamp() -> None:
    rng = random.Random(7)
    samples = {"".join(rng.choice("0123456789") for _ in range(10)) for _ in range(3000)}
    samples |= {"2125550187", "9175550187", "2023101512", "2024013123", "2024013124"}
    for d in sorted(samples):
        stamp = d34.date_like(d[:8]) and int(d[8:]) <= 23
        assert ("phone" in d34.name_kinds(d)) == (dtext.nanp_digits(d) and not stamp), d
        # a value is the stricter of the two: every bare number a name flags, a value flags too
        if "phone" in d34.name_kinds(d):
            assert "phone" in dtext.unit_shapes(d), d


def _book(path: Path, sheets: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for s in sheets:
        ws = wb.create_sheet(s)
        ws.append(["student_id", "score"])
        for r in range(N):
            ws.append([f"S{2000 + r}", (r * 3) % 40])
    wb.save(path)
    return path


def test_bare_phone_numbers_in_keys_files_sheets_folders_and_fields_fail_and_are_never_printed(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    # the codex reproduction: a JSON key that is a phone number
    (src / "keys.json").write_text(json.dumps([{PHONE: r % 9, "score": r} for r in range(N)]), encoding="utf-8")
    roster(src / f"{PHONE}.csv")  # a file named by a phone number
    roster(src / "5125550199" / "roster.csv")  # a folder named by one
    _book(src / "book.xlsx", ["Summary", "2125550143"])  # a sheet named by one
    pd.DataFrame({"week": range(N), "9175550123": range(N)}).to_parquet(src / "wide.parquet")  # a field
    roster(src / "clean.csv")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    assert sorted(k for k, v in st.items() if v != "FAIL") == ["clean.csv"], st
    assert written(out) == ["DATA_DICTIONARY.md", "README.md", "clean.csv", "schema.json"]
    blob = everything(out, rep, console=res.output)
    for secret in (PHONE, "5125550199", "2125550143", "9175550123"):
        assert secret not in blob, "a phone-number name reached the out dir, the report or the console"
    reasons = " ".join(r for f in report(rep)["files"] for r in f["reasons"])
    assert reasons.count("a phone number") >= 5


def test_keep_name_keeps_a_bare_number_that_is_no_phone(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    (src / "keys.json").write_text(json.dumps([{PHONE: r % 9, "score": r} for r in range(N)]), encoding="utf-8")
    res, out, rep = run(src, tmp_path, "--keep-name", PHONE)
    assert res.exit_code == 0, res.output
    assert PHONE in (out / "keys.json").read_text(encoding="utf-8")


# =========================================================================== R7-04: no value in a diagnostic
def _assert_clean(output: str) -> None:
    assert SECRET_WORD not in output.casefold(), "a planted source value reached the console"


def _xlsx_with_text_in_a_number_cell(path: Path) -> Path:
    """The codex reproduction: a cell tagged as a number that holds text (openpyxl raises ValueError quoting it)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["student_id", "score"])
    ws.append(["S1001", 123456789])
    buf, bad = io.BytesIO(), io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(bad, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v>123456789</v>", f"<v>{SECRET}</v>".encode())
            zout.writestr(item, data)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bad.getvalue())
    return path


@pytest.mark.parametrize("verbosity", [(), ("-v",), ("-vv",)])
def test_a_read_error_quoting_a_cell_never_prints_it(tmp_path: Path, verbosity: tuple[str, ...]) -> None:
    src = _xlsx_with_text_in_a_number_cell(tmp_path / "in" / "input.xlsx")
    res, out, rep = run(src, tmp_path, *verbosity)
    assert res.exit_code == 3, res.output
    _assert_clean(everything(out, rep, console=res.output))
    if verbosity:  # the traceback still locates the error: classes and code locations
        assert "Traceback" in res.output and "ValueError" in res.output
        assert "openpyxl/worksheet/_reader.py:" in res.output and "datanonymous/pipeline.py:" in res.output


class _Planted(Exception):
    pass


def _raise_planted(*_args: Any, **_kwargs: Any) -> Any:
    """An internal error whose every part quotes a source value: message, arguments, note, cause, context, local."""
    cell = SECRET  # a local variable holding the value (a traceback with locals would print it)
    try:
        try:
            raise KeyError(cell)
        except KeyError as inner:
            raise ValueError(f"could not convert string to float: {cell!r}") from inner
    except ValueError:
        err = _Planted(f"bad cell {cell!r}", cell, {"value": cell})
        err.add_note(f"while reading {cell}")
        raise err


def _inputs(tmp_path: Path) -> Path:
    src = tmp_path / "in"
    write_csv(src / "notes.csv", ["student_id", "note"], [[f"S{1000 + r}", SECRET if r == 0 else "ok"]
                                                         for r in range(N)])
    roster(src / "clean.csv")
    return src


@pytest.mark.parametrize("stage", ["profile", "generate"])
@pytest.mark.parametrize("verbosity", [("-v",), ("-vv",)])
def test_an_internal_error_while_profiling_or_generating_never_prints_its_text(
        tmp_path: Path, monkeypatch: pytest.MonkeyPatch, stage: str, verbosity: tuple[str, ...]) -> None:
    src = _inputs(tmp_path)
    target = "profile_table" if stage == "profile" else "generate_table"
    orig = getattr(pipeline, target)

    def boom(tp: Any, *args: Any, **kwargs: Any) -> Any:
        rel = args[0] if stage == "profile" else args[0].file_rel
        if str(rel).endswith("notes.csv"):
            _raise_planted()
        return orig(tp, *args, **kwargs)

    monkeypatch.setattr(pipeline, target, boom)
    res, out, rep = run(src, tmp_path, *verbosity)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    assert st["notes.csv"] == "FAIL" and st["clean.csv"] != "FAIL"
    _assert_clean(everything(out, rep, console=res.output))
    assert "_Planted" in res.output and "KeyError" in res.output and "ValueError" in res.output
    assert "direct cause" in res.output and "During handling" in res.output
    assert "test_r7_shapes_tracebacks.py:" in res.output and "in _raise_planted" in res.output


@pytest.mark.parametrize("verbosity", [("-v",), ("-vv",)])
def test_inspect_never_prints_the_text_of_an_internal_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                                                            verbosity: tuple[str, ...]) -> None:
    src = _inputs(tmp_path)
    monkeypatch.setattr(pipeline, "profile_table", _raise_planted)
    res = run_cli("inspect", src, *verbosity)
    assert res.exit_code == 3, res.output
    _assert_clean(res.output)
    assert "_Planted" in res.output


def test_leak_scan_errors_never_print_their_text(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    src = _inputs(tmp_path)
    monkeypatch.setattr(leaks, "_scan_data_file", _raise_planted)  # one file's scan (per-file isolation)
    res, out, rep = run(src, tmp_path, "-v", name="file")
    assert res.exit_code == 3, res.output
    _assert_clean(everything(out, rep, console=res.output))
    monkeypatch.undo()
    monkeypatch.setattr(leaks, "_scan_outputs", _raise_planted)  # the whole scan (every output unverified)
    res, out, rep = run(src, tmp_path, "-v", name="run")
    assert res.exit_code == 3, res.output
    _assert_clean(everything(out, rep, console=res.output))
    assert "_Planted" in res.output


@pytest.mark.parametrize("verbosity", [(), ("-v",), ("-vv",)])
def test_the_cli_top_level_handlers_never_print_the_text(tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                                                         verbosity: tuple[str, ...]) -> None:
    src = _inputs(tmp_path)
    monkeypatch.setattr(pipeline, "execute", _raise_planted)  # inside the command
    res, out, rep = run(src, tmp_path, *verbosity, name="cmd")
    assert res.exit_code == 1
    _assert_clean(res.output)
    assert "unexpected error: _Planted" in res.output
    assert ("Traceback" in res.output) == bool(verbosity)
    monkeypatch.setattr(cli, "_run_guarded", _raise_planted)  # outside it (click's own main)
    res, out, rep = run(src, tmp_path, *verbosity, name="top")
    assert res.exit_code == 1
    _assert_clean(res.output)
    assert ("Traceback" in res.output) == bool(verbosity)


def test_a_non_numeric_exit_is_never_printed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def leave(*_a: Any, **_k: Any) -> None:
        raise SystemExit(SECRET)  # Python prints a text exit code to stderr

    monkeypatch.setattr(cli, "_run_guarded", leave)
    res, _, _ = run(_inputs(tmp_path), tmp_path)
    assert res.exit_code == 1
    _assert_clean(res.output)


@pytest.mark.parametrize("verbosity", [(), ("-v",)])
def test_library_warnings_never_print_their_message(tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                                                    verbosity: tuple[str, ...]) -> None:
    src = _inputs(tmp_path)
    orig = pipeline.profile_table

    def warn(tp: Any, *args: Any, **kwargs: Any) -> Any:
        warnings.warn(f"tzname {SECRET} identified but not understood", UserWarning, stacklevel=1)
        return orig(tp, *args, **kwargs)

    monkeypatch.setattr(pipeline, "profile_table", warn)
    with warnings.catch_warnings():
        warnings.simplefilter("always")
        res, out, rep = run(src, tmp_path, *verbosity)
    assert res.exit_code == 0, res.output
    _assert_clean(everything(out, rep, console=res.output))
    assert ("warning: UserWarning at" in res.output) == bool(verbosity)
    assert "test_r7_shapes_tracebacks.py:" in res.output or not verbosity


_PROCESS = """
import sys, warnings
from datanonymous import pipeline
secret = sys.argv.pop()
orig = pipeline.profile_table
def noisy(tp, *a, **k):
    warnings.warn("tzname " + secret + " identified but not understood", UserWarning)
    return orig(tp, *a, **k)
pipeline.profile_table = noisy
from datanonymous.cli import main
sys.argv[0] = "datanonymous"
main()
"""


@pytest.mark.parametrize("verbosity", [(), ("-v",)])
def test_a_real_process_prints_no_warning_message_and_no_traceback_message(tmp_path: Path,
                                                                           verbosity: tuple[str, ...]) -> None:
    """Python itself prints warnings and tracebacks to stderr outside pytest's capture: in a real process, neither
    the planted warning text nor the openpyxl error quoting a cell reaches stdout or stderr."""
    import os
    import subprocess
    import sys

    src = tmp_path / "in"
    _xlsx_with_text_in_a_number_cell(src / "input.xlsx")
    roster(src / "clean.csv")
    env = dict(os.environ, PYTHONPATH=str(SRC.parent), PYTHONWARNINGS="always")
    args = [sys.executable, "-c", _PROCESS, "--mode", "fidelity", str(src), "--out-dir", str(tmp_path / "out"), "--report-dir",
            str(tmp_path / "rep"), "--seed", "7", *verbosity, SECRET]
    proc = subprocess.run(args, capture_output=True, env=env, timeout=300)
    console = proc.stdout.decode("utf-8", "replace") + proc.stderr.decode("utf-8", "replace")
    assert proc.returncode == 3, console
    _assert_clean(console)
    assert ("warning: UserWarning at" in console) == bool(verbosity)
    assert ("Traceback" in console) == bool(verbosity)


def test_a_csv_parser_error_is_reported_in_the_tools_words(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    class Broken:
        line_num = 3

        def __iter__(self) -> Any:
            raise csv.Error(f"field holds {SECRET!r}")

    real = csv.reader
    # This case exercises the stdlib table parser's error wording, so explicitly
    # decline the Arrow fast path before injecting csv.Error.
    monkeypatch.setattr(csvio, "_arrow_records", lambda *args, **kwargs: None)
    # the table parser only (the delimiter sniffer reads with the real one)
    monkeypatch.setattr(csvio.csv, "reader", lambda *a, **k: Broken() if "skipinitialspace" in k else real(*a, **k))
    res, out, rep = run(_inputs(tmp_path) / "notes.csv", tmp_path, "-v")
    assert res.exit_code == 3
    _assert_clean(everything(out, rep, console=res.output))
    assert "CSV parse error on line 3" in res.output


@pytest.mark.parametrize("kind", ["ArrowInvalid", "ArrowNotImplementedError", "UnicodeError", "RuntimeError"])
def test_an_arrow_parser_exception_never_prints_its_source_value(
        tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str) -> None:
    import pyarrow as pa
    import pyarrow.csv as pacsv

    exception = {"UnicodeError": UnicodeError, "RuntimeError": RuntimeError}.get(kind)
    if exception is None:
        exception = getattr(pa, kind)
    calls = []

    def broken(*args: Any, **kwargs: Any) -> Any:
        calls.append(True)
        raise exception(f"field holds {SECRET!r}")

    monkeypatch.setattr(pacsv, "read_csv", broken)
    res, out, rep = run(_inputs(tmp_path) / "notes.csv", tmp_path, "-v")
    assert calls, "the injected Arrow parser was never reached"
    _assert_clean(everything(out, rep, console=res.output))
    if kind == "RuntimeError":  # unexpected failures retain safe class/location diagnostics
        assert res.exit_code == 3
        assert "RuntimeError" in res.output and "Traceback" in res.output
        assert not (out / "notes.csv").exists()
    else:  # ordinary Arrow parse failures silently retry the authoritative stdlib parser
        assert res.exit_code == 0, res.output
        assert (out / "notes.csv").is_file()


# --------------------------------------------------------------------------- the diagnostics themselves
def test_safe_traceback_holds_classes_and_code_locations_only() -> None:
    try:
        _raise_planted()
    except _Planted as e:
        text = diagnostics.safe_traceback(e)
        place = diagnostics.error_place(e)
    assert SECRET_WORD not in text and "bad cell" not in text and "while reading" not in text
    assert text.count("Traceback") == 3 and "KeyError" in text and "ValueError" in text and "_Planted" in text
    assert "tests/test_r7_shapes_tracebacks.py:" in text or "test_r7_shapes_tracebacks.py:" in text
    assert place.startswith("_Planted at test_r7_shapes_tracebacks.py:")


def test_safe_traceback_of_groups_cycles_and_unusual_exceptions() -> None:
    a = ValueError(SECRET)
    b = TypeError(SECRET)
    a.__context__ = b
    b.__context__ = a  # a cycle
    group = ExceptionGroup(SECRET, [a, KeyError(SECRET), ExceptionGroup(SECRET, [OSError(SECRET)])])
    text = diagnostics.safe_traceback(group)
    assert SECRET_WORD not in text
    for name in ("ExceptionGroup", "ValueError", "TypeError", "KeyError", "OSError"):
        assert name in text, name

    class Loud(Exception):
        def __str__(self) -> str:
            raise RuntimeError(SECRET)

    assert SECRET_WORD not in diagnostics.safe_traceback(Loud())
    assert diagnostics.code_file("C:/data/" + SECRET + ".csv") == "<code>"

    def deep(n: int) -> None:
        if n == 0:
            raise ValueError(SECRET)
        deep(n - 1)

    try:
        deep(200)
    except ValueError as e:
        text = diagnostics.safe_traceback(e)
    assert SECRET_WORD not in text and "repeated 199 more time(s)" in text and len(text.splitlines()) < 10
    assert SECRET_WORD not in diagnostics.warning_text(UserWarning, SECRET, 3)


def test_errors_python_ignores_never_print_their_text() -> None:
    """An error in a finalizer or a thread is printed by Python with its message: within a command it goes to the
    ``-v`` channel as classes and code locations only (``sys.unraisablehook``, ``threading.excepthook``)."""
    import gc
    import sys
    import threading

    class Finalizer:
        def __del__(self) -> None:
            raise ValueError(SECRET)

    def worker() -> None:
        raise KeyError(SECRET)

    lines: list[str] = []
    before = (sys.unraisablehook, threading.excepthook)
    with warnings.catch_warnings():
        warnings.simplefilter("always")
        with diagnostics.quiet_diagnostics(lines.append):
            obj = Finalizer()
            del obj
            gc.collect()
            t = threading.Thread(target=worker)
            t.start()
            t.join()
            warnings.warn(SECRET, RuntimeWarning)
    assert (sys.unraisablehook, threading.excepthook) == before
    text = "\n".join(lines)
    assert SECRET_WORD not in text
    assert "ValueError" in text and "KeyError" in text and "RuntimeWarning" in text


#: ``except ... as e`` handlers that may print the exception's text: they quote what the user typed (a config value,
#: a --date-format), which the CLI screens (D34), never a source value.
_TYPED_TEXT_ONLY = {("config.py", "validate"), ("dates.py", "_explicit_format")}


def test_no_tool_message_quotes_a_caught_exception() -> None:
    """Structural guard (R7-04): a caught exception is printed as its class and attributes the tool chose
    (``type(e).__name__``, ``e.reason``, ``e.lineno``), never as itself (``{e}``, ``str(e)``, ``repr(e)``,
    ``e.args``, ``traceback.format_exception``)."""
    found = []
    for path in sorted(SRC.rglob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        funcs = {id(n): f.name for f in ast.walk(tree) if isinstance(f, (ast.FunctionDef, ast.AsyncFunctionDef))
                 for n in ast.walk(f)}  # node -> its innermost function (ast.walk visits outer functions first)
        for node in ast.walk(tree):
            if not isinstance(node, ast.ExceptHandler) or not node.name:
                continue
            name = node.name
            for sub in ast.walk(ast.Module(body=node.body, type_ignores=[])):
                bad = (isinstance(sub, ast.FormattedValue) and isinstance(sub.value, ast.Name)
                       and sub.value.id == name)
                bad |= (isinstance(sub, ast.Call) and isinstance(sub.func, ast.Name)
                        and sub.func.id in ("str", "repr", "format") and any(
                            isinstance(a, ast.Name) and a.id == name for a in sub.args))
                bad |= (isinstance(sub, ast.Attribute) and sub.attr in ("args", "__notes__")
                        and isinstance(sub.value, ast.Name) and sub.value.id == name)
                if bad and (path.name, funcs.get(id(node), "")) not in _TYPED_TEXT_ONLY:
                    found.append(f"{path.relative_to(SRC).as_posix()}:{sub.lineno}")
    assert not found, found
    for path in sorted(SRC.rglob("*.py")):
        src = path.read_text(encoding="utf-8")
        for banned in ("format_exception(", "format_exc(", "print_exc(", "print_exception(", "exc_info=True"):
            assert banned not in src, (path.name, banned)
