"""Build the AREA E Excel fixtures deterministically (run: python tests/fixtures/excel/make_fixtures.py).

Outputs (committed, byte-exact via .gitattributes):

* ``multi_sheet.xlsx`` - "Enrollment": merged title row, a printed-by title row
  with a date, a blank row, header in row 4, 40 data rows (text IDs with
  leading zeros, a 30-name surname column, text grades, mm/dd/yyyy dates, 0.00
  GPAs, 0.0% rates, 00000 ZIPs, booleans with blanks, scores with #N/A error
  cells, credits, a formula column with cached values), a footer "Total" row
  with COUNTA/AVERAGE/SUM formulas (cached values), a blank row and a note
  row; freeze panes, autofilter, widths, a comment, a hyperlink, a data
  validation, a named range and document properties (extras that must not be
  copied). "Courses": a plain table. "Lookup": a hidden sheet. "Empty": an
  empty sheet.
* ``fake_html.xls`` - an SIS-style HTML table saved as .xls (title row with
  colspan, CSS with a page-header text, a comment, x:num attributes).
* ``fake_tsv.xls`` - tab-delimited text saved as .xls (CRLF, cp1252).
* ``real_biff.xls`` - a minimal BIFF8 workbook in an OLE2 container (built
  here by hand: no .xls writer is installed) for the xlrd path.
* ``adversarial.xlsx`` / ``banner_html.xls`` - layouts that broke the first
  reader (``test_area_e_hardening.py``): a first name "Max" in the last data
  row, per-row SUM formulas, side notes beside the table, formulas filled down
  past the data returning "", titles above a one-column table, a notes-only
  sheet; an HTML banner table, CSS comments/urls and a script inside the
  data table.

Every name/value is invented. The workbook's zip timestamps and document
dates are fixed so reruns are byte-identical.
"""

from __future__ import annotations

import datetime as dt
import io
import random
import re
import struct
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent

SURNAMES = ["Abernathy", "Brightwater", "Castellano", "Delacroix", "Eastwood", "Fairbanks", "Gallagher",
            "Holloway", "Ingersoll", "Jablonski", "Kowalczyk", "Lindqvist", "Montgomery", "Northcott",
            "Okonkwo", "Pemberton", "Quintero", "Rasmussen", "Sorensen", "Thistlewood", "Underhill", "Valdivia",
            "Whitcombe", "Xiomarra", "Yarborough", "Zielinski", "Ashcombe", "Blackwood", "Crestfield", "Dunmore"]
DEPTS = ["Math", "Science", "English", "History"]


def _student_rows(rng: random.Random) -> list[list]:
    rows = []
    base_dates = [dt.datetime(2023, 8, 14) + dt.timedelta(days=d) for d in (0, 1, 2, 3, 4, 7, 8, 9, 10, 11, 14, 21)]
    for i in range(40):
        sid = f"{100120 + i * 7:06d}"[-6:]
        sid = "00" + sid[2:]
        name = SURNAMES[i % len(SURNAMES)] if i < 30 else SURNAMES[rng.randrange(len(SURNAMES))]
        grade = rng.choice(["KG", "01", "02", "03", "04", "05"])
        date = rng.choice(base_dates)
        gpa = None if i % 9 == 4 else round(rng.uniform(1.2, 4.0), 2)
        rate = round(rng.uniform(0.82, 1.0), 4)
        zip_code = rng.choice([2134, 2138, 75001, 75002, 7030, 10001])
        econ = None if i % 5 == 3 else (rng.random() < 0.6)
        score = "#N/A" if i % 8 == 6 else rng.randrange(52, 100)
        credits = rng.choice([0.5, 1.0, 1.5, 2.0, 3.0])
        rows.append([sid, name, grade, date, gpa, rate, zip_code, econ, score, credits])
    return rows


def build_xlsx(path: Path) -> None:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    rng = random.Random(20241001)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enrollment"
    ws["A1"] = "Lincoln Unified School District - Enrollment Report"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:K1")
    ws["A2"] = "Printed by jsmith on 10/01/2024 08:15 AM"
    header = ["student_id", "last_name", "grade", "enroll_date", "gpa", "attendance_rate", "zip_code",
              "econ_dis", "test_score", "credits", "points"]
    for j, h in enumerate(header, 1):
        ws.cell(4, j, h)
    rows = _student_rows(rng)
    cached: dict[str, str] = {}
    for i, row in enumerate(rows):
        r = 5 + i
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j)
            if v is None:
                continue
            c.value = v
            if j == 1 or j == 3:
                c.data_type = "s"
        ws.cell(r, 4).number_format = "mm/dd/yyyy"
        ws.cell(r, 5).number_format = "0.00"
        ws.cell(r, 6).number_format = "0.0%"
        ws.cell(r, 7).number_format = "00000"
        ws.cell(r, 10).number_format = "0.0"
        ws.cell(r, 11).value = f"=J{r}*10"
        cached[f"K{r}"] = _num(row[9] * 10)
    last = 4 + len(rows)
    total = last + 1
    ws.cell(total, 1, "Total")
    ws.cell(total, 2, f"=COUNTA(B5:B{last})")
    cached[f"B{total}"] = str(len(rows))
    ws.cell(total, 5, f"=AVERAGE(E5:E{last})")
    gpas = [r[4] for r in rows if r[4] is not None]
    cached[f"E{total}"] = _num(sum(gpas) / len(gpas))
    ws.cell(total, 5).number_format = "0.00"
    ws.cell(total, 10, f"=SUM(J5:J{last})")
    cached[f"J{total}"] = _num(sum(r[9] for r in rows))
    ws.cell(total, 11, f"=SUM(K5:K{last})")
    cached[f"K{total}"] = _num(sum(r[9] * 10 for r in rows))
    ws.cell(total + 2, 1, "Source: SIS nightly export. Questions: jsmith@lincoln-usd.example")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{last}"
    for letter, width in (("A", 12.5), ("B", 18), ("D", 13), ("F", 15.25)):
        ws.column_dimensions[letter].width = width
    ws["B7"].comment = Comment("Check spelling with Mrs. Pemberton-Hale", "jsmith")
    ws["A6"].hyperlink = "https://sis.lincoln-usd.example/student/000127"
    dv = DataValidation(type="list", formula1='"KG,01,02,03,04,05"')
    ws.add_data_validation(dv)
    dv.add(f"C5:C{last}")
    wb.defined_names["StudentList"] = DefinedName("StudentList", attr_text=f"Enrollment!$A$5:$A${last}")

    cs = wb.create_sheet("Courses")
    cs.append(["course_code", "department", "credits", "start_date"])
    for k in range(12):
        cs.append([f"00{120 + k * 3}", DEPTS[k % 4], [0.5, 1.0][k % 2], dt.datetime(2023, 8, 14 + k % 5)])
        cs.cell(k + 2, 1).data_type = "s"
        cs.cell(k + 2, 3).number_format = "0.0"
        cs.cell(k + 2, 4).number_format = "yyyy-mm-dd"
    cs.column_dimensions["B"].width = 14

    lk = wb.create_sheet("Lookup")
    lk.append(["code", "label"])
    for k, lab in enumerate(["Active", "Withdrawn", "Graduated", "Transferred", "Inactive"], 1):
        lk.append([k, lab])
    lk.sheet_state = "hidden"
    wb.create_sheet("Empty")

    wb.properties.creator = "Jane Q. Registrar"
    wb.properties.title = "Confidential roster"
    fixed = dt.datetime(2024, 10, 1, 8, 15)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(_finish_zip(buf.getvalue(), cached, fixed))


def _num(x: float) -> str:
    return repr(float(x)) if not float(x).is_integer() else str(int(x))


def _finish_zip(data: bytes, cached: dict[str, str], fixed: dt.datetime) -> bytes:
    """Insert cached formula values (as Excel would) and fix timestamps for byte-identical rebuilds."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    stamp = fixed.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                def fill(m: re.Match) -> bytes:
                    ref = m.group(1).decode()
                    val = cached.get(ref)
                    if val is None:
                        return m.group(0)
                    return m.group(0).replace(b"<v />", b"<v>" + val.encode() + b"</v>").replace(
                        b"<v/>", b"<v>" + val.encode() + b"</v>")
                body = re.sub(rb'<c r="([A-Z]+\d+)"[^>]*><f>[^<]*</f><v ?/></c>', fill, body)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", stamp, body)
            zi = zipfile.ZipInfo(info.filename, date_time=(2024, 10, 1, 8, 15, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, body)
    return out.getvalue()


def build_html(path: Path) -> None:
    rng = random.Random(77)
    lines = [
        '<html xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:x="urn:schemas-microsoft-com:office:excel">',
        "<head>",
        '<meta http-equiv="Content-Type" content="text/html; charset=utf-8">',
        '<meta name="author" content="jsmith">',
        "<title>Roster export for jsmith</title>",
        "<style>",
        'td { mso-number-format:"\\@"; }',
        '.num { mso-number-format:"0.00"; }',
        '@page { mso-header-data:"&LPrinted by jsmith"; }',
        "</style>",
        "<!-- exported by SIS for user jsmith from 10.1.2.3 -->",
        "</head>",
        "<body>",
        '<table border="1">',
        '<tr><td colspan="5">Roster Report - Printed 10/01/2024</td></tr>',
        "<tr><th>Student_Number</th><th>LastFirst</th><th>Grade_Level</th><th>Entry_Date</th><th>GPA</th></tr>",
    ]
    firsts = ["Ada", "Bram", "Cleo", "Dov", "Esme", "Finn", "Gita", "Hugo", "Iris", "Jude"]
    for i in range(30):
        sid = f"{4100 + i * 13:06d}"
        name = f"{SURNAMES[i % len(SURNAMES)]}, {firsts[i % len(firsts)]}"
        grade = str(rng.choice([9, 10, 11, 12]))
        date = f"0{rng.randint(1, 9)}/{rng.randint(10, 28)}/2024"
        gpa = f"{rng.uniform(1.5, 4.0):.2f}"
        lines.append(f'<tr><td style="mso-number-format:\'\\@\'">{sid}</td><td>{name}</td><td>{grade}</td>'
                     f'<td>{date}</td><td class="num" x:num="{gpa}">{gpa}</td></tr>')
    lines += ["</table>", "</body>", "</html>"]
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))


def build_tsv(path: Path) -> None:
    rng = random.Random(5)
    out = ["Student_Number\tLast_Name\tGrade\tSchool"]
    schools = ["Lincoln ES", "Garc\u00eda MS", "Roosevelt HS"]
    for i in range(25):
        out.append(f"{3000 + i * 11:06d}\t{SURNAMES[(i * 7) % len(SURNAMES)]}\t{rng.choice(['9', '10', '11', '12'])}\t"
                   f"{schools[i % 3]}")
    path.write_bytes(("\r\n".join(out) + "\r\n").encode("cp1252"))


# --------------------------------------------------------------------------- minimal BIFF8 / OLE2
def _rec(rtype: int, data: bytes) -> bytes:
    return struct.pack("<HH", rtype, len(data)) + data


def _ustr16(s: str) -> bytes:
    return struct.pack("<HB", len(s), 1) + s.encode("utf-16-le")


def _ustr8(s: str) -> bytes:
    return struct.pack("<BB", len(s), 1) + s.encode("utf-16-le")


def _biff_workbook(sheets: list[tuple[str, int, list[list]]]) -> bytes:
    """sheets: (name, visibility, rows); cell values str / float / bool / ('date', serial)."""
    glob = [
        _rec(0x0809, struct.pack("<HHHHII", 0x0600, 0x0005, 0x0DBB, 0x07CC, 0, 0x06)),
        _rec(0x0042, struct.pack("<H", 1200)),
        _rec(0x0022, struct.pack("<H", 0)),
    ]
    font = struct.pack("<HHHHHBBBB", 200, 0, 0x7FFF, 400, 0, 0, 0, 0, 0) + _ustr8("Arial")
    for _ in range(5):
        glob.append(_rec(0x0031, font))
    glob.append(_rec(0x041E, struct.pack("<H", 164) + _ustr16("mm/dd/yyyy")))
    glob.append(_rec(0x041E, struct.pack("<H", 165) + _ustr16("0.00")))

    def xf(fmt: int, style: bool) -> bytes:
        type_par = 0xFFF5 if style else 0x0001
        return _rec(0x00E0, struct.pack("<HHHBBBBIiH", 0, fmt, type_par, 0x20, 0, 0, 0 if style else 0x04,
                                         0, 0, 0x20C0))

    for _ in range(16):
        glob.append(xf(0, True))  # style XFs 0..15
    glob += [xf(0, False), xf(164, False), xf(165, False)]  # cell XFs 16 general, 17 date, 18 0.00
    head = b"".join(glob)
    bounds_len = sum(len(_rec(0x0085, struct.pack("<IBB", 0, 0, 0) + _ustr8(n))) for n, _, _ in sheets)
    eof = _rec(0x000A, b"")
    offset = len(head) + bounds_len + len(eof)
    bodies = []
    bounds = []
    for name, vis, rows in sheets:
        body = [_rec(0x0809, struct.pack("<HHHHII", 0x0600, 0x0010, 0x0DBB, 0x07CC, 0, 0x06))]
        ncols = max((len(r) for r in rows), default=0)
        body.append(_rec(0x0200, struct.pack("<IIHHH", 0, len(rows), 0, ncols, 0)))
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                if v is None:
                    continue
                if isinstance(v, bool):
                    body.append(_rec(0x0205, struct.pack("<HHHBB", r, c, 16, int(v), 0)))
                elif isinstance(v, tuple) and v[0] == "date":
                    body.append(_rec(0x0203, struct.pack("<HHHd", r, c, 17, float(v[1]))))
                elif isinstance(v, tuple) and v[0] == "dec":
                    body.append(_rec(0x0203, struct.pack("<HHHd", r, c, 18, float(v[1]))))
                elif isinstance(v, (int, float)):
                    body.append(_rec(0x0203, struct.pack("<HHHd", r, c, 16, float(v))))
                else:
                    body.append(_rec(0x0204, struct.pack("<HHH", r, c, 16) + _ustr16(str(v))))
        body.append(eof)
        data = b"".join(body)
        bounds.append(_rec(0x0085, struct.pack("<IBB", offset, vis, 0) + _ustr8(name)))
        bodies.append(data)
        offset += len(data)
    stream = head + b"".join(bounds) + eof + b"".join(bodies)
    if len(stream) < 4096:
        stream += b"\x00" * (4096 - len(stream))
    return stream


def _ole2(stream: bytes, name: str = "Workbook") -> bytes:
    sec = 512
    if len(stream) % sec:
        stream += b"\x00" * (sec - len(stream) % sec)
    n_data = len(stream) // sec
    endc, free, fatsect, nostream = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    fat = [fatsect, endc] + [3 + k if k < n_data - 1 else endc for k in range(n_data)]
    assert len(fat) <= 128
    fat += [free] * (128 - len(fat))
    header = (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16
              + struct.pack("<HHHHH", 0x003E, 0x0003, 0xFFFE, 9, 6) + b"\x00" * 6
              + struct.pack("<IIIIIIIII", 0, 1, 1, 0, 4096, endc, 0, endc, 0)
              + struct.pack("<I", 0) + struct.pack("<I", free) * 108)
    assert len(header) == 512

    def entry(nm: str, typ: int, child: int, start: int, size: int) -> bytes:
        raw = nm.encode("utf-16-le") + b"\x00\x00" if nm else b""
        return (raw.ljust(64, b"\x00") + struct.pack("<HBB", len(raw), typ, 1)
                + struct.pack("<III", nostream, nostream, child) + b"\x00" * 16 + struct.pack("<I", 0)
                + struct.pack("<QQ", 0, 0) + struct.pack("<IQ", start, size))

    dirs = (entry("Root Entry", 5, 1, endc, 0) + entry(name, 2, nostream, 2, len(stream))
            + entry("", 0, nostream, 0, 0) * 2)
    assert len(dirs) == 512
    return header + struct.pack("<128I", *fat) + dirs + stream


def build_xls(path: Path) -> None:
    rows: list[list] = [["Attendance summary - CONFIDENTIAL"], [],
                        ["student_id", "school", "days_absent", "rate", "last_absence", "chronic"]]
    rng = random.Random(9)
    for i in range(20):
        serial = 45170 + rng.randrange(0, 120)  # 2023-2024 dates
        rows.append([f"{51000 + i * 3:07d}", rng.choice(["Lincoln ES", "Roosevelt HS"]), rng.randrange(0, 25),
                     ("dec", round(rng.uniform(0.7, 1.0), 3)), ("date", serial), rng.random() < 0.3])
    hidden = [["code", "meaning"], [1, "Excused"], [2, "Unexcused"]]
    path.write_bytes(_ole2(_biff_workbook([("Attendance", 0, rows), ("Codes", 1, hidden)])))


# --------------------------------------------------------------------------- adversarial inputs (hardening)
#: invented strings that exist only in titles, notes and banners of the adversarial fixtures: never in any output
ADVERSARIAL_SECRETS = ("Quillfeather", "Wilhelmina", "555-0142", "Okafor", "nurse.okafor", "jdoe", "Hartfield")
FIRST_NAMES = ["Anselm", "Briony", "Caspian", "Delphine", "Evander", "Fenella", "Gideon", "Honora", "Ignatius",
               "Juniper", "Kestrel", "Leopold", "Marisol", "Nicodemus", "Ottoline", "Peregrine", "Quentin",
               "Rosalind", "Silvanus", "Tamsin", "Ulysses", "Verity", "Wystan", "Xanthe", "Yolanda", "Zebedee",
               "Aurelio", "Bettany", "Cyprian"]


def build_adversarial(path: Path) -> None:
    """``adversarial.xlsx``: layouts that broke the first Excel reader (see ``test_area_e_hardening``).

    * "Roster": a title row, a person whose first name is "Max" in the last data
      row (a summary word, but data), per-row ``=SUM(D5:E5)`` formulas in every
      row (data, not a footer), and side notes two columns right of the table
      naming a person and a phone number.
    * "FillDown": a formula filled down far below the data that returns ""
      (pandas reads 25 rows, not 120).
    * "Immunization": two title rows above a single-column table of numbers.
    * "Info": a single-column notes block (title + printed-by line), no table.
    """
    import openpyxl

    rng = random.Random(4242)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws["A1"] = "Roster for Ms. Quillfeather"
    header = ["student_id", "first_name", "grade", "q1", "q2", "total"]
    for j, h in enumerate(header, 1):
        ws.cell(3, j, h)
    cached: dict[str, dict[str, str]] = {"sheet1": {}, "sheet2": {}}
    names = FIRST_NAMES + ["Max"]
    for i, name in enumerate(names):
        r = 4 + i
        q1, q2 = rng.randrange(0, 11), rng.randrange(0, 11)
        ws.cell(r, 1, f"{310 + i * 7:06d}").data_type = "s"
        ws.cell(r, 2, name)
        ws.cell(r, 3, rng.choice(["09", "10"])).data_type = "s"
        ws.cell(r, 4, q1)
        ws.cell(r, 5, q2)
        ws.cell(r, 6, f"=SUM(D{r}:E{r})")
        cached["sheet1"][f"F{r}"] = str(q1 + q2)
    ws["H5"] = "Note: Wilhelmina Quillfeather withdrew 9/14"
    ws["H6"] = "Call 555-0142 before 10/03/2024"

    fd = wb.create_sheet("FillDown")
    fd.append(["id", "score", "pass"])
    for r in range(2, 121):
        if r <= 26:
            fd.cell(r, 1, f"{r * 13:05d}").data_type = "s"
            score = rng.randrange(40, 100)
            fd.cell(r, 2, score)
            cached["sheet2"][f"C{r}"] = "Y" if score >= 60 else "N"
        else:
            cached["sheet2"][f"C{r}"] = ""
        fd.cell(r, 3, f'=IF(B{r}="","",IF(B{r}>=60,"Y","N"))')

    im = wb.create_sheet("Immunization")
    im["A1"] = "Immunization follow-up list"
    im["A2"] = "Printed 10/01/2024 by nurse.okafor"
    im["A3"] = "student_number"
    for k in range(18):
        im.cell(4 + k, 1, 52000 + rng.randrange(0, 900))

    info = wb.create_sheet("Info")
    info["A1"] = "Class roster - Mr. Okafor"
    info["A2"] = "Printed 10/01/2024 by jdoe"

    fixed = dt.datetime(2024, 10, 1, 8, 15)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(_cache_formulas(buf.getvalue(), cached, fixed))


def _cache_formulas(data: bytes, cached: dict[str, dict[str, str]], fixed: dt.datetime) -> bytes:
    """Insert cached formula results per sheet part (text results get t="str"); fixed zip timestamps."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    stamp = fixed.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            m = re.fullmatch(r"xl/worksheets/(sheet\d+)\.xml", info.filename)
            if m and m.group(1) in cached:
                values = cached[m.group(1)]

                def fill(mm: re.Match) -> bytes:
                    val = values.get(mm.group(1).decode())
                    if val is None:
                        return mm.group(0)
                    attrs, formula = mm.group(2), mm.group(3)
                    is_num = re.fullmatch(r"-?\d+(\.\d+)?", val) is not None
                    t = b"" if is_num else b' t="str"'
                    return b'<c r="%s"%s%s><f>%s</f><v>%s</v></c>' % (mm.group(1), attrs, t, formula,
                                                                     val.encode())
                body = re.sub(rb'<c r="([A-Z]+\d+)"([^>]*)><f>([^<]*)</f><v ?/></c>', fill, body)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", stamp, body)
            zi = zipfile.ZipInfo(info.filename, date_time=(2024, 10, 1, 8, 15, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, body)
    return out.getvalue()


def build_banner_html(path: Path) -> None:
    """``banner_html.xls``: an SIS HTML export with a one-column banner table above the data table, CSS
    comments/content/url naming people, and a script inside the data table."""
    rng = random.Random(31)
    rows = []
    for i in range(24):
        rows.append(f'<tr><td>{2000 + i * 17:06d}</td><td title="Hartfield">{rng.choice(["A", "B", "C"])}</td>'
                    f"<td>{rng.randrange(10, 99)}</td></tr>")
    lines = [
        "<html>",
        "<head>",
        '<meta http-equiv="Content-Type" content="text/html; charset=utf-8">',
        "<style>/* exported for Okafor */ td { color: black; content: \"jdoe\"; }"
        " .x { background: url(file:///C:/Users/jdoe/logo.png); mso-number-format:\"0\\.00\"; }</style>",
        "</head>",
        "<body>",
        "<table><tr><td><b>Class roster - Mr. Okafor</b></td></tr><tr><td>Printed 10/01/2024 by jdoe</td></tr></table>",
        "<table>",
        "<tr><th>id</th><th>grp</th><th>score</th></tr>",
        *rows,
        "<script>var exportedBy = 'Hartfield';</script>",
        "</table>",
        "</body>",
        "</html>",
    ]
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))


def main() -> None:
    build_xlsx(HERE / "multi_sheet.xlsx")
    build_html(HERE / "fake_html.xls")
    build_tsv(HERE / "fake_tsv.xls")
    build_xls(HERE / "real_biff.xls")
    build_adversarial(HERE / "adversarial.xlsx")
    build_banner_html(HERE / "banner_html.xls")


if __name__ == "__main__":
    main()
