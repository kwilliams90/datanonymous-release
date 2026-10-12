"""Cross-area regressions found when the seven feature areas were merged (integrate/areas).

Each test pins a contract between two areas that their own suites could not see:
G's leak scan vs B's decimal-comma text, G's schema vs the A/B/C doc hooks, G's fix
rendering vs E's quoted ``--range`` values, D's ID-candidate warning vs C's generators,
and E's Excel date cells vs A's native-storage contract.
"""

from __future__ import annotations

import datetime as dt
import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from conftest import REPO, read_dicts, run_cli
from datanonymous import pipeline
from datanonymous.config import Settings
from datanonymous.privacy import leaks
from datanonymous.report.report import fix_commands, fix_markdown

NUMERIC = REPO / "tests" / "fixtures" / "numeric"
DATES = REPO / "tests" / "fixtures" / "dates" / "dates.csv"


# --------------------------------------------------------------------------- B x G: decimal-comma text vs the leak scan
@pytest.fixture(scope="module")
def comma_run(tmp_path_factory: pytest.TempPathFactory):
    base = tmp_path_factory.mktemp("comma")
    (base / "in").mkdir()
    shutil.copyfile(NUMERIC / "decimal_comma.csv", base / "in" / "decimal_comma.csv")
    run = pipeline.execute(Settings(mode="fidelity", seed=7), base / "in", base / "out", base / "rep")
    return run, base


def test_decimal_comma_text_passes_the_leak_scan(comma_run) -> None:
    run, _ = comma_run
    assert pipeline.exit_code(run) == 0
    assert run.files[0].written and not run.files[0].failed


def test_numeric_label_matches_only_its_own_notation(comma_run) -> None:
    run, base = comma_run
    real = [line.split(";") for line in (NUMERIC / "decimal_comma.csv").read_bytes().decode("cp1252").split("\r\n")]
    i = real[0].index("faktor")
    value = next(r[i] for r in real[1:] if r[i].count(",") == 1)  # e.g. '0,31'
    scanner = leaks.Scanner(leaks.SensitiveIndex.build(run.files), leaks.Allow())
    assert scanner.hits_in_text(f"the real value {value} quoted") != []  # the real text is still a leak
    assert scanner.hits_in_text(f"a share of {value.replace(',', '.')}") == []  # a statistic is not


# --------------------------------------------------------------------------- A/B/C x G: schema.json types/formats
def test_schema_has_no_hook_statistics_and_honours_r_class(tmp_path: Path) -> None:
    res = run_cli(DATES, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "5")
    assert res.exit_code == 0, res.output
    schema = json.loads((tmp_path / "o" / "schema.json").read_text(encoding="utf-8"))
    cols = {c["name"]: c for c in schema["files"][0]["tables"][0]["columns"]}
    assert cols["updated_epoch"]["r_class"] == "POSIXct"  # AREA A's r_class override (epoch seconds are instants)
    for c in cols.values():
        assert not {"note", "char_length_p50", "n_items", "bounding_box"} & set(c), c["name"]
        assert not isinstance(c.get("formats"), dict), c["name"]  # mixed formats without their row shares
    assert all("parse" in c for c in cols.values() if c["readr"] == "c" and c["type"] != "SCHOOL_YEAR"
               and c["type"].startswith(("DATE", "TIME", "DURATION", "YEAR_MONTH")))


def test_text_hooks_stay_out_of_the_schema(tmp_path: Path) -> None:
    src = REPO / "tests" / "fixtures" / "text" / "free_text.csv"
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "3")
    assert res.exit_code == 0, res.output
    schema = json.loads((tmp_path / "o" / "schema.json").read_text(encoding="utf-8"))
    cols = {c["name"]: c for c in schema["files"][0]["tables"][0]["columns"]}
    assert cols["first_name"]["name_part"] == "first" and "note" not in cols["first_name"]
    assert "bounding_box" not in cols["home_lat"] and "char_length_p50" not in cols["comments"]
    dd = (tmp_path / "o" / "DATA_DICTIONARY.md").read_text(encoding="utf-8")
    assert "`first_name`" in dd


# --------------------------------------------------------------------------- E x G: quoted --range values in a fix
def test_fix_with_a_quoted_range_value_is_one_command() -> None:
    fix = '--range "Roster!A3:H33" to read them as data columns'
    assert fix_markdown(fix) == '`--range "Roster!A3:H33"` to read them as data columns'
    assert fix_commands('--header-row N or --range "Sheet 1!A5:H200" (or --sheet NAME to read one sheet)') == [
        "--header-row N", '--range "Sheet 1!A5:H200"', "--sheet NAME"]
    assert fix_commands("--id-check luhn") == ["--id-check luhn"]


# --------------------------------------------------------------------------- C x D: the ID-candidate warning
def test_id_candidate_warning_says_which_values_are_written(tmp_path: Path) -> None:
    src = REPO / "tests" / "fixtures" / "ids" / "ids.csv"
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "3")
    assert res.exit_code == 0, res.output
    data = json.loads((tmp_path / "r" / "files" / "ids.csv.json").read_text(encoding="utf-8"))
    msg = next(m for m in data["tables"][0]["messages"] if m["code"] == "id_candidates")
    written, _, replaced = msg["text"].partition("their values are replaced")
    assert "row_num" in written and "md5" not in written  # sequential integers are written; hashes are masked
    assert "md5" in replaced and "record_uuid" in replaced
    real = read_dicts(src)
    syn = read_dicts(tmp_path / "o" / "ids.csv")
    assert not {r["md5"] for r in real} & {s["md5"] for s in syn}


# --------------------------------------------------------------------------- A x E: Excel date cells keep their kind
def test_excel_duration_cells_end_to_end(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["rid", "elapsed_x", "clock_x"])
    for i in range(40):
        ws.append([i + 1, dt.timedelta(hours=i % 19, minutes=(11 * i) % 60), dt.time(7 + i % 9, 15 * (i % 4))])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "[h]:mm"
        row[2].number_format = "h:mm"
    (tmp_path / "in").mkdir()
    wb.save(tmp_path / "in" / "t.xlsx")
    res = run_cli(tmp_path / "in", "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "2")
    assert res.exit_code == 0, res.output
    data = json.loads((tmp_path / "r" / "files" / "t.xlsx.json").read_text(encoding="utf-8"))
    types = {c["name"]: c["type"] for c in data["tables"][0]["columns"]}
    assert types["elapsed_x"] == "DURATION" and types["clock_x"] == "TIME"
    out = openpyxl.load_workbook(tmp_path / "o" / "t.xlsx").active
    assert all(isinstance(r[1].value, dt.timedelta) for r in out.iter_rows(min_row=2))
