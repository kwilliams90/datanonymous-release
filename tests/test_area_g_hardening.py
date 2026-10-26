"""AREA G hardening: adversarial fixtures and regressions found by the adversarial review.

* ``tests/fixtures/privacy/adversarial/``: UTF-8 BOM + CRLF roster with 10-digit
  state IDs next to 10-digit phones, ``Last, First`` names with typographic
  apostrophes and hyphens, short medical alerts, an all-missing column; and a
  one-row file. Every out-dir, report and console text is grepped for every
  real value (assertions name kinds and files, never values).
* Data scan: a synthetic ID (or number, date, reproduced label) that equals a
  real phone/email of another column by chance is not a leak; a real value in a
  text column still is. Stale outputs are scanned column by column.
* Report/docs scan: ``O'Brien`` = ``O’Brien``, a part of a hyphenated name,
  short free-text values.
* Report and docs polish: section order, link targets, the fidelity "not
  compared" note, README statements about kept row numbers and panels.
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import zipfile
from pathlib import Path

import numpy as np
import pytest

from conftest import run_cli
from datanonymous import pipeline
from datanonymous.config import Settings
from datanonymous.model import ColumnModel, ColumnProfile, ColumnType
from datanonymous.privacy import leaks
from datanonymous.privacy.leaks import DataScanner, Scanner, SensitiveIndex, build_allow, cross_column_values
from datanonymous.report import docs, fidelity
from datanonymous.report import report as report_mod
from test_area_g_grep import FIXTURES, SCHOOL, read_texts, sensitive

ADV = FIXTURES / "adversarial"
WORD = re.compile(r"[^\W_]+(?:['\-][^\W_]+)*")
DIGIT_RUN = re.compile(r"\+?\(?\d(?:[ ().\-]{0,2}\d){6,14}")


def adv_values() -> dict[str, list[str]]:
    return json.loads((FIXTURES / "adversarial_values.json").read_text(encoding="utf-8"))


def _fold(s: str) -> str:
    return s.replace("’", "'").replace("‘", "'").casefold()


def _name_tokens(values: list[str]) -> set[str]:
    out: set[str] = set()
    for v in values:
        for t in WORD.findall(_fold(v)):
            out.add(t)
            out.update(p for p in re.split(r"['\-]", t) if len(p) >= 3)
    return out


def adv_hits(text: str) -> list[str]:
    """Kinds of real adversarial values found in ``text`` (never the values)."""
    s = adv_values()
    found = []
    if any(re.search(rf"(?<![0-9]){v}(?![0-9])", text) for v in s["state_ids"] + s["ids"]):
        found.append("ids")
    toks: set[str] = set()
    for t in WORD.findall(_fold(text)):
        toks.add(t)
        toks.update(re.split(r"['\-]", t))
    if toks & _name_tokens(s["names"]):
        found.append("names")
    digits = {"".join(ch for ch in m.group() if ch.isdigit()) for m in DIGIT_RUN.finditer(text)}
    if digits & {"".join(ch for ch in v if ch.isdigit()) for v in s["phones"]}:
        found.append("phones")
    folded = " ".join(_fold(text).split())
    if any(a.casefold() in folded for a in s["alerts"]):
        found.append("alerts")
    if any(e.casefold() in folded or e.split("@")[0].casefold() in folded for e in s["emails"]):
        found.append("emails")
    return found


def _run(tmp_path: Path, src: Path, *extra: str, name: str = "run"):
    out, rep = tmp_path / f"{name}_out", tmp_path / f"{name}_rep"
    res = run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", "4", *extra)
    return res, out, rep


def _report(rep: Path) -> dict:
    return json.loads((rep / "report.json").read_text(encoding="utf-8"))


def _assert_no_adv_value(res, *dirs: Path) -> None:
    texts = {"console": res.output, **read_texts(*dirs)}
    dirty = {k: adv_hits(v) for k, v in texts.items() if adv_hits(v)}
    assert not dirty, dirty  # {file: [kinds]}, never the values


# --------------------------------------------------------------------------- the adversarial fixtures
def test_adversarial_grep_finds_every_kind_in_the_sources() -> None:
    roster = (ADV / "roster.csv").read_bytes().decode("utf-8-sig")
    assert set(adv_hits(roster)) == {"ids", "names", "phones", "alerts"}
    assert set(adv_hits((ADV / "one_row.csv").read_text(encoding="utf-8"))) == {"ids", "names", "emails"}
    # an ASCII apostrophe finds the curly-apostrophe name; a hyphen part finds the hyphenated name
    assert adv_hits("D'ALESSANDROVA") == ["names"] and adv_hits("met Nkemelu") == ["names"]


def test_adversarial_run_is_clean_complete_and_reproducible(tmp_path: Path) -> None:
    res, out, rep = _run(tmp_path, ADV)
    assert res.exit_code == 0, res.output
    _assert_no_adv_value(res, out, rep)
    data = _report(rep)
    for f in data["files"]:
        assert f["written"] and f["rows"] == f["rows_out"], f["source"]
        assert f["checks"]["rows"] == "PASS" and f["checks"]["header"] == "PASS" and f["checks"]["leak_scan"] == "PASS"
    raw = (out / "roster.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and raw.count(b"\r\n") == 81 and raw.count(b"\n") == 81  # BOM, CRLF
    with open(out / "roster.csv", newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == ["state_id", "student_name", "nickname", "guardian_phone", "medical_alert", "program", "notes",
                       "score"]
    assert len(rows) == 81 and all(r[6] == "" for r in rows[1:])  # the all-missing column stays all-missing
    assert len((out / "one_row.csv").read_text(encoding="utf-8").splitlines()) == 2
    res2, out2, _ = _run(tmp_path, ADV, name="again")
    assert res2.exit_code == 0
    files = {p.relative_to(out).as_posix(): p.read_bytes() for p in out.rglob("*") if p.is_file()}
    files2 = {p.relative_to(out2).as_posix(): p.read_bytes() for p in out2.rglob("*") if p.is_file()}
    assert files == files2


def test_adversarial_inspect_and_json_have_no_real_value(tmp_path: Path) -> None:
    res = run_cli("inspect", ADV)
    assert res.exit_code == 0
    _assert_no_adv_value(res)
    res, out, rep = _run(tmp_path, ADV, "--json")
    assert res.exit_code == 0
    _assert_no_adv_value(res)


# --------------------------------------------------------------------------- data scan: chance coincidences
def _write_ids_and_phones(path: Path, ids: list[str], phones: list[str]) -> None:
    lines = ["state_id,guardian_phone,program"]
    lines += [f"{i},{p},{'ELL' if k % 3 else 'GT'}" for k, (i, p) in enumerate(zip(ids, phones))]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def test_synthetic_id_equal_to_a_real_phone_is_not_a_leak(tmp_path: Path) -> None:
    """A synthetic 10-digit state ID equal to a real guardian phone's digits is a coincidence, not a copy.

    Regression: the run-wide phone-digit rule used to FAIL the file (exit 3), which at district scale
    (10-digit state IDs and phones) happens by chance in a large share of runs.
    """
    rng = np.random.default_rng(3)
    ids = [str(v) for v in sorted(rng.choice(np.arange(10**9, 10**9 + 10**7), 200, replace=False))]
    phones = [f"({a}) {b}-{c}" for a, b, c in zip(rng.integers(201, 989, 200), rng.integers(201, 989, 200),
                                                   rng.integers(1000, 9999, 200))]
    src = tmp_path / "in"
    src.mkdir()
    _write_ids_and_phones(src / "s.csv", ids, phones)
    res, out, _ = _run(tmp_path, src, name="first")
    assert res.exit_code == 0, res.output
    with open(out / "s.csv", newline="", encoding="utf-8") as fh:
        syn_ids = [r[0] for r in list(csv.reader(fh))[1:]]
    planted = [f"({s[:3]}) {s[3:6]}-{s[6:]}" for s in syn_ids[:5]]
    _write_ids_and_phones(src / "s.csv", ids, planted + phones[5:])
    res, out, rep = _run(tmp_path, src, name="second")
    with open(out / "s.csv", newline="", encoding="utf-8") as fh:
        again = [r[0] for r in list(csv.reader(fh))[1:]]
    assert again[:5] == syn_ids[:5]  # precondition: the IDs do not depend on the phone column
    assert res.exit_code == 0, res.output
    assert _report(rep)["files"][0]["checks"]["leak_scan"] == "PASS"


def _profile(name: str, typ: ColumnType, subtype: str | None = None, **params) -> ColumnProfile:  # noqa: ANN003
    return ColumnProfile(index=0, name=name, type=typ, subtype=subtype, params=dict(params))


def test_cross_column_values_skip_generated_shapes_and_reproduced_levels() -> None:
    from datanonymous.detect.text import COPY_GATE_KEY

    gated = {COPY_GATE_KEY: "screened"}  # real values a detection writes are excused once the copy gate cleared them
    vals = ["3125550199", "jane@x.org"]
    assert cross_column_values(_profile("sid", ColumnType.IDENTIFIER, "digits"), vals) == []
    assert cross_column_values(_profile("code", ColumnType.DIGIT_CODE), vals) == []
    assert cross_column_values(_profile("amount", ColumnType.FLOAT, **gated), vals) == []
    assert cross_column_values(_profile("n", ColumnType.INTEGER, "continuous", **gated), vals) == []
    assert cross_column_values(_profile("d", ColumnType.DATE, **gated), ["2024-03-15 10:30"]) == []
    factor = _profile("school_phone", ColumnType.FACTOR, **gated)
    factor.model = ColumnModel(nonmissing_counts={"3125550199": 4, "x": 1})
    assert cross_column_values(factor, ["3125550199", "jane@x.org"]) == ["jane@x.org"]  # a non-level is checked
    # R4: values that never passed the copy gate are never excused (one source of truth with the gate)
    assert cross_column_values(_profile("amount", ColumnType.FLOAT), vals) == vals
    factor.params.pop(COPY_GATE_KEY)
    assert cross_column_values(factor, ["3125550199", "jane@x.org"]) == ["3125550199", "jane@x.org"]
    assert cross_column_values(_profile("note", ColumnType.FREE_TEXT), vals) == vals
    assert cross_column_values(_profile("label", ColumnType.HIGH_CARD_CATEGORICAL), vals) == vals


@pytest.fixture(scope="module")
def school_run(tmp_path_factory: pytest.TempPathFactory):
    base = tmp_path_factory.mktemp("hardening_school")
    run = pipeline.execute(Settings(mode="fidelity", seed=5), SCHOOL, base / "out", base / "rep")
    assert pipeline.exit_code(run) == 0
    return run


def test_a_real_phone_in_a_text_column_is_still_caught(school_run) -> None:
    fr = next(f for f in school_run.files if f.rel_path == "people.csv")
    tp = fr.tables[0]
    data = DataScanner(SensitiveIndex.build(school_run.files), build_allow(school_run))
    phone = sensitive()["phones"][0]
    digits = "".join(ch for ch in phone if ch.isdigit())
    sid = next(p for p in tp.columns if p.name == "student_id")
    comment = next(p for p in tp.columns if p.name == "counselor_comment")
    assert not data.scan_column(fr, 0, sid, np.array([digits[:6]], dtype=object))
    hits = data.scan_column(fr, 0, comment, np.array([f"call {digits}"], dtype=object))
    assert any(h.rule == "phone digits" for h in hits)


def test_stale_output_is_scanned_column_by_column(tmp_path: Path) -> None:
    """An output left from an earlier run (its source FAILs now) gets the data rules, not raw text patterns."""
    rng = np.random.default_rng(8)
    ids = [str(v) for v in sorted(rng.choice(np.arange(10**9, 10**9 + 10**7), 60, replace=False))]
    phones = [f"({a}) {b}-{c}" for a, b, c in zip(rng.integers(201, 989, 60), rng.integers(201, 989, 60),
                                                   rng.integers(1000, 9999, 60))]
    src = tmp_path / "in"
    src.mkdir()
    _write_ids_and_phones(src / "s.csv", ids, phones)
    out, rep = tmp_path / "o", tmp_path / "r"
    assert run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", "2").exit_code == 0
    with open(out / "s.csv", newline="", encoding="utf-8") as fh:
        syn_ids = [r[0] for r in list(csv.reader(fh))[1:]]
    # the real phones now equal the old synthetic IDs, and the gate makes the file FAIL (old output stays)
    _write_ids_and_phones(src / "s.csv", ids, [f"({s[:3]}) {s[3:6]}-{s[6:]}" for s in syn_ids])
    res = run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", "2", "--force", "--max-row-matches", "0")
    data = _report(rep)
    assert data["files"][0]["status"] == "FAIL" and (out / "s.csv").exists()
    assert not [e for e in data["errors"] if "s.csv" in e["text"] and e["code"] == "leak_found"], res.output
    # a real phone planted in the stale file's phone column is still a run-level error
    text = (out / "s.csv").read_text(encoding="utf-8").splitlines()
    cells = text[1].split(",")
    cells[1] = "(" + syn_ids[3][:3] + ") " + syn_ids[3][3:6] + "-" + syn_ids[3][6:]
    text[1] = ",".join(cells)
    (out / "s.csv").write_text("\n".join(text) + "\n", encoding="utf-8")
    res = run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", "2", "--force", "--max-row-matches", "0")
    assert res.exit_code == 3
    errs = [e for e in _report(rep)["errors"] if "s.csv" in e["text"] and e["code"] == "leak_found"]
    assert errs and errs[0]["fix"] == "delete s.csv from the out dir"


def test_xlsx_view_skips_cell_values_but_keeps_other_text(tmp_path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Roster printed for Ysolde Quennell"])  # a title row (not a cell of the scanned table)
    ws.append(["state_id"])
    ws.append([3125550199])
    ws.append(["3125550198"])
    wb.properties.creator = "creator-xyz"
    path = tmp_path / "w.xlsx"
    wb.save(path)
    view = leaks.xlsx_text_view(path, covered={"3125550199", "3125550198", "state_id"})
    assert "Ysolde Quennell" in view and "creator-xyz" in view
    assert "3125550199" not in view and "3125550198" not in view
    full = leaks.xlsx_text_view(path)
    assert "3125550199" in full
    with zipfile.ZipFile(path) as zf:  # a real xlsx
        assert any(n.startswith("xl/worksheets/") for n in zf.namelist())


# --------------------------------------------------------------------------- report/docs scan: variants
@pytest.fixture(scope="module")
def adv_run(tmp_path_factory: pytest.TempPathFactory):
    base = tmp_path_factory.mktemp("hardening_adv")
    run = pipeline.execute(Settings(mode="fidelity", seed=5), ADV, base / "out", base / "rep")
    assert pipeline.exit_code(run) == 0
    return run


@pytest.fixture(scope="module")
def adv_scanner(adv_run) -> Scanner:
    return Scanner(SensitiveIndex.build(adv_run.files), build_allow(adv_run))


@pytest.mark.parametrize("text", [
    "value D'Alessandrova looks odd",  # ASCII apostrophe; the real name has a typographic one
    "value d’alessandrova looks odd",
    "student Nkemelu was absent",  # one part of a hyphenated name
    "the Haverford's file",  # a part inside a possessive
    "Kasprowicz—Severin",  # an em dash between two names
    "alert: SEVERE PEANUT ALLERGY",  # a short free-text value (3 words)
    "note type 1 diabetes here",
], ids=["ascii-apostrophe", "curly-apostrophe", "hyphen-part", "possessive-part", "em-dash", "short-alert",
        "alert-with-digit"])
def test_report_scan_catches_variants(adv_scanner: Scanner, text: str) -> None:
    hits = adv_scanner.hits_in_text(text)
    assert hits
    red = adv_scanner.redact_text(text)
    assert "[redacted]" in red and not adv_scanner.hits_in_text(red)


def test_report_scan_ignores_tool_text_and_labels(adv_run, adv_scanner: Scanner) -> None:
    payload = report_mod.build_report(adv_run, Settings(mode="fidelity", seed=5))
    for rel, data in payload.items():
        assert not adv_scanner.hits_in_bytes(rel, data), rel
    # tool literals of several words are excused as sequences, labels of a factor too
    assert not adv_scanner.hits_in_text("no ID column matched; each row is its own entity")
    assert not adv_scanner.hits_in_text("program ELL and SPED")


def test_short_alert_quoted_in_a_message_fails_the_file(tmp_path: Path, monkeypatch) -> None:
    from datanonymous.model import Message

    alert = adv_values()["alerts"][0]
    orig = pipeline.profile_file

    def profile_file(f, settings, run):  # noqa: ANN001
        fr = orig(f, settings, run)
        if fr is not None and fr.rel_path == "roster.csv":
            fr.tables[0].messages.append(Message("warning", "buggy", f"odd value {alert.upper()}",
                                                 column="medical_alert"))
        return fr

    monkeypatch.setattr(pipeline, "profile_file", profile_file)
    res, out, rep = _run(tmp_path, ADV)
    assert res.exit_code == 3
    entry = next(f for f in _report(rep)["files"] if f["source"] == "roster.csv")
    assert entry["status"] == "FAIL" and entry["checks"]["leak_scan"] == "FAIL"
    assert not (out / "roster.csv").exists() and (out / "one_row.csv").exists()
    _assert_no_adv_value(res, out, rep)


# --------------------------------------------------------------------------- report polish
def test_report_sections_follow_the_spec_order(tmp_path: Path) -> None:
    src = tmp_path / "in"
    (src / "a#b").mkdir(parents=True)
    shutil.copy(SCHOOL / "staff.csv", src / "a#b" / "staff (1).csv")
    (src / "notes.md").write_text("# notes\n", encoding="utf-8")
    (src / "broken.csv").write_bytes(b"\n\nx,y\n1,2\n")  # a blank header line FAILs the file
    res, out, rep = _run(tmp_path, src)
    assert res.exit_code == 3
    md = (rep / "REPORT.md").read_text(encoding="utf-8")
    order = [md.index(h) for h in ("## What was generated", "## What needs attention", "## Appendix", "### Verdict", "### Files",
                                  "### Skipped inputs (not copied)")]
    assert order == sorted(order)
    assert "(files/a%23b/staff%20%281%29.csv.md)" in md  # '#', spaces and parentheses are encoded


def test_fidelity_says_when_level_shares_were_not_compared() -> None:
    p = _profile("x", ColumnType.INTEGER, "discrete")
    p.model = ColumnModel(nonmissing_counts={str(v): 1 for v in range(300)})
    real = {"n": 300, "missing": 0, "missing_rate": 0.0, "distinct": 300, "pmf_omitted": "every value is unique"}
    out = fidelity.compare_column(p, real, dict(real))
    check = next(c for c in out["checks"] if c["check"] == "level_shares")
    assert check["compared"] is False and "not compared" in check["reason"] and out["verdict"] == "ok"
    code = _profile("site_code", ColumnType.DIGIT_CODE, levels_mode=True)
    code.model = ColumnModel(nonmissing_counts={"0012": 5, "0040": 7})
    real = {"n": 12, "missing": 0, "missing_rate": 0.0, "distinct": 2}
    out = fidelity.compare_column(code, real, dict(real))
    assert "invented" not in out["reason"] and "kept as levels" in out["reason"]
    empty = _profile("e", ColumnType.EMPTY)
    out = fidelity.compare_column(empty, {"n": 3, "missing": 3, "missing_rate": 1.0},
                                  {"n": 3, "missing": 3, "missing_rate": 1.0})
    assert [c["check"] for c in out["checks"]] == ["missing_rate"]


def test_readme_names_kept_row_numbers_and_detected_panels(tmp_path: Path) -> None:
    from conftest import EXAMPLES

    run = pipeline.execute(Settings(mode="fidelity", seed=1), EXAMPLES / "student-scores.csv", tmp_path / "o", tmp_path / "r")
    readme = (tmp_path / "o" / "README.md").read_text(encoding="utf-8")
    assert "row-number ID columns that run 1..n keep those numbers" in readme
    assert "`student-scores.csv :: id`" in readme
    fr = run.files[0]
    fr.tables[0].structure.panels.append({"columns": ["id"], "kind": "panel"})
    docs.rewrite_docs(run, 10)
    readme = (tmp_path / "o" / "README.md").read_text(encoding="utf-8")
    assert "1 panel / slowly-changing pattern(s) detected" in readme
    schema = json.loads((tmp_path / "o" / "schema.json").read_text(encoding="utf-8"))
    col = schema["files"][0]["tables"][0]["columns"][0]
    assert col["name"] == "id" and col["row_numbers_kept"] is True


def test_readme_has_no_row_number_claim_without_kept_ids(tmp_path: Path) -> None:
    pipeline.execute(Settings(mode="fidelity", seed=1), SCHOOL, tmp_path / "o", tmp_path / "r")
    readme = (tmp_path / "o" / "README.md").read_text(encoding="utf-8")
    assert "row-number id" not in readme.lower()
    assert "Replacement IDs keep their shape" in readme and "avoid detected source IDs" in readme


def test_a_broken_leak_scan_never_leaves_unverified_outputs(tmp_path: Path, monkeypatch) -> None:
    def boom(files):  # noqa: ANN001
        raise RuntimeError("scanner bug")

    monkeypatch.setattr(leaks.SensitiveIndex, "build", staticmethod(boom))
    res, out, rep = _run(tmp_path, ADV)
    assert res.exit_code == 3
    assert not (out / "roster.csv").exists() and not (out / "one_row.csv").exists()
    assert (out / "README.md").exists()  # docs rewritten without the removed files
    data = _report(rep)
    assert data["verdict"]["overall"] == "FAIL"
    assert all(f["status"] == "FAIL" and f["checks"]["leak_scan"] == "FAIL" for f in data["files"])
    assert any(e["code"] == "leak_scan_incomplete" for e in data["errors"])
    assert "scanner bug" not in (rep / "REPORT.md").read_text(encoding="utf-8")  # exception text never reported


# --------------------------------------------------------------------------- docs and warnings details
def test_tsv_snippet_escapes_the_tab_and_null_columns_show_empty_cells(tmp_path: Path) -> None:
    src = tmp_path / "in"
    src.mkdir()
    (src / "t.tsv").write_bytes(b"id\tgrade\tscore\n1\tA\t5\n2\tB\t6\n3\tA\t7\n4\tB\t8\n")
    res, out, rep = _run(tmp_path, src, "--null", "score")
    assert res.exit_code == 0, res.output
    dd = (out / "DATA_DICTIONARY.md").read_text(encoding="utf-8")
    assert 'delim = "' + "\\" + 't"' in dd and "\t" not in dd  # an escaped tab, never a literal one
    assert "<null>" not in dd  # a --null column of a text file holds empty cells
    assert 'na = c(""))' in dd
    schema = json.loads((out / "schema.json").read_text(encoding="utf-8"))
    score = next(c for c in schema["files"][0]["tables"][0]["columns"] if c["name"] == "score")
    assert score["na_tokens"] == [""]


def test_dtype_drift_is_a_listed_warning_also_with_exclude(tmp_path: Path) -> None:
    from conftest import EXAMPLES

    res = run_cli(EXAMPLES / "students_data.csv", "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r",
                  "--seed", "1", "--exclude", "Course", "--null", "Age")
    assert res.exit_code == 0, res.output
    data = _report(tmp_path / "r")
    assert data["files"][0]["checks"]["dtypes"] == "WARN"
    drift = [w for w in data["warnings"] if w["code"] == "dtype_mismatch"]
    assert drift and "'Age'" in drift[0]["text"]
    assert data["files"][0]["warnings"] == len([w for w in data["warnings"] if w["file"] == "students_data.csv"])


def test_large_literal_word_lists_are_not_tool_text() -> None:
    """A future name/state list in the package must not excuse real names in the report scan."""
    names = ", ".join(f'"Name{i}"' for i in range(60))
    source = f'"""Doc words."""\nNAMES = ({names},)\nSMALL = ("Alpha", "Beta")\nMSG = "written by the tool"\n'
    texts = leaks.literal_texts(source)
    assert "Name7" not in texts and "Doc words." not in texts
    assert {"Alpha", "Beta", "written by the tool"} <= set(texts)
    assert ("no", "id", "column", "matched") in leaks.tool_phrases()  # a sequence inside one tool literal


def test_same_column_phone_check_ignores_the_country_code() -> None:
    real = leaks._RealSets.of(leaks.PHONE, ["1-312-555-0199", "(415) 555-0123"])
    assert leaks._same_column_hits(leaks.PHONE, real, ["(312) 555-0199"]) == 1
    assert leaks._same_column_hits(leaks.PHONE, real, ["+1 415 555 0123"]) == 1
    assert leaks._same_column_hits(leaks.PHONE, real, ["(415) 555-0124"]) == 0


def test_a_decimal_statistic_is_not_read_as_a_phone(adv_scanner: Scanner) -> None:
    phone = adv_values()["phones"][0]
    digits = "".join(ch for ch in phone if ch.isdigit())
    assert adv_scanner.hits_in_text(f"call {phone}")  # the phone itself is caught
    assert not adv_scanner.hits_in_text(f"mean {digits[:4]}.{digits[4:]}")  # a mean with these digits is not
    assert adv_scanner.hits_in_text(f"call {digits[:3]}.{digits[3:6]}.{digits[6:]}")  # dotted phone still is


def _dense_ids(src: Path) -> None:
    src.mkdir()
    ids = list(range(1000, 9990)) + list(range(10000, 10050))  # 4-digit IDs almost all taken
    rows = ["student_id,grade"] + [f"{i},{('09', '10', '11')[i % 3]}" for i in ids]
    (src / "d.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")


def test_a_count_equal_to_a_dense_numeric_id_is_not_a_leak(tmp_path: Path) -> None:
    """Regression: '8980 synthetic ID(s) have a different number of digits' FAILed the file when 8980 is a real ID."""
    _dense_ids(tmp_path / "in")
    res, out, rep = _run(tmp_path, tmp_path / "in")
    assert res.exit_code == 0, res.output
    data = _report(rep)
    assert any(w["code"] == "id_width_changed" for w in data["warnings"])
    assert data["files"][0]["checks"]["leak_scan"] == "PASS"


def test_a_dense_numeric_id_quoted_in_another_message_is_still_caught(tmp_path: Path, monkeypatch) -> None:
    from datanonymous.model import Message

    _dense_ids(tmp_path / "in")
    orig = pipeline.profile_file

    def profile_file(f, settings, run):  # noqa: ANN001
        fr = orig(f, settings, run)
        if fr is not None:
            fr.tables[0].messages.append(Message("warning", "buggy", "duplicate key 4321 in two rows",
                                                 column="student_id"))
        return fr

    monkeypatch.setattr(pipeline, "profile_file", profile_file)
    res, out, rep = _run(tmp_path, tmp_path / "in")
    assert res.exit_code == 3
    assert "4321" not in (rep / "REPORT.md").read_text(encoding="utf-8")
