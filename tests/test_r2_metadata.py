"""Round-2 review fixes, metadata cluster.

* R2-02: real text survived in format metadata. Every Excel number format the tool reads or writes
  (``.xlsx``/``.xls`` cells, title/header/footer cells, the ``mso-number-format`` CSS of HTML "Excel" files)
  is rebuilt with its literal text masked (``io.numfmt``: quoted strings, ``\\x`` escapes, unquoted
  non-code letters, ``[$text-LCID]`` currency text; units, currencies and symbols kept), and the same holds
  for the ``!c`` literals of Stata ``%t`` display formats. The leak scan also reads number formats.
* R2-03: unused Arrow dictionary entries and Stata value labels were copied although no row holds them
  (so nothing ever checked them). A dictionary / label table now holds the source labels the synthetic
  rows hold, plus unused labels that are words of the tool's own vocabularies (grade levels, ordinal
  scales, boolean words ...); a Stata table whose labels could not be shown as values is not written.
  The self-check flags any other unused entry, and the leak scan reads label text no cell shows.
* R2-12: numbers wider than float64 (int64/uint64 from 2**53, decimal text with 16+ significant digits,
  formatted numbers, numbers mixed with tokens, within-ID constants) are generated on exact scaled integers
  with their exact min and max (D14), and the extremes self-check compares them exactly.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import random
import zipfile
from decimal import Decimal
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.feather as feather
import pyarrow.parquet as pq
import pytest

from conftest import run_cli
from datanonymous.generate import numeric as G
from datanonymous.generate.table import _ensure_level_extremes
from datanonymous.io import arrowio, excel_html, excel_read, statio
from datanonymous.io.numfmt import mask_literal, safe_css_number_format, safe_number_format
from datanonymous.model import ColumnModel, ColumnProfile, ColumnType
from datanonymous.pipeline import _extremes_missing
from datanonymous.privacy import leaks

NAME = "Zelmira Quenby"
EMAIL = "zelmira.quenby@private.example.edu"
PHONE = "312-555-0199"


def _run(inp: Path, tmp: Path, *flags: str, seed: str = "123", tag: str = "") -> tuple[int, Path, Path]:
    out, rep = tmp / f"out{tag}", tmp / f"rep{tag}"
    res = run_cli(inp, "--out-dir", out, "--report-dir", rep, "--seed", seed, *flags)
    return res.exit_code, out, rep


def _report(rep: Path) -> dict:
    return json.loads((rep / "report.json").read_text(encoding="utf-8"))


def _file_report(rep: Path, name: str) -> dict:
    return json.loads((rep / "files" / f"{name}.json").read_text(encoding="utf-8"))


def _statuses(rep: Path) -> dict[str, str]:
    return {f["source"]: f["status"] for f in _report(rep)["files"]}


def _blob(*dirs: Path) -> bytes:
    """Every file under ``dirs``, with zip members (xlsx parts) decompressed."""
    chunks = []
    for d in dirs:
        for p in sorted(d.rglob("*")):
            if not p.is_file():
                continue
            data = p.read_bytes()
            chunks.append(data)
            if data.startswith(b"PK\x03\x04"):
                with zipfile.ZipFile(p) as z:
                    chunks.extend(z.read(n) for n in z.namelist())
    return b"\n".join(chunks)


def _codes(rep: Path, name: str) -> list[str]:
    f = _file_report(rep, name)
    return [m["code"] for m in f["messages"]] + [m["code"] for t in f["tables"] for c in t["columns"]
                                                 for m in c["messages"]]


# =========================================================================== R2-02 number formats
KEPT_FORMATS = [
    "General", "0", "0.00", "#,##0", "#,##0.00", "0%", "0.00%", "0.00E+00", "##0.0E+0", "# ?/?", "# ??/??",
    "mm-dd-yy", "d-mmm-yy", "mmm-yy", "h:mm AM/PM", "h:mm:ss a/p", "m/d/yy h:mm", "[h]:mm:ss", "[ss].00",
    "mmss.0", "@", "yyyy-mm-dd", "dd/mm/yyyy hh:mm:ss.000", "dddd, mmmm dd, yyyy",
    "#,##0 ;[Red](#,##0)", "#,##0.00;[Red]-#,##0.00", "[Color10]0", "[=1]\"Yes\";[=0]\"No\"",
    '_(* #,##0_);_(* \\(#,##0\\);_(* "-"_);_(@_)', '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)',
    '"$"#,##0.00', "[$-409]mmmm d, yyyy", "[$€-407] #,##0.00", "[$USD-409]#,##0", "[$-x-sysdate]dddd, mmmm dd",
    "[$-F800]dddd\\,\\ mmmm\\ dd\\,\\ yyyy", "[<=9999999]###-####;(###) ###-####", "000-00-0000", "00000-0000",
    '0.0" kg"', '#,##0 "USD"', '0" pts"', '0.0 "m²"', '0.00\\ \\k\\g', 'yyyy"年"m"月"d"日"', '"Yes";"Yes";"No"',
    "0;-0;;@", "[DBNum1]0", "[$-en-US]dddd, mmmm d", "[$-zh-Hans-CN]yyyy", "[$-1010409]0", "yyyymmdd", "HH:MM:SS",
    '#,##0.00 "Kč"', '[$R$-416] #,##0.00', '#,##0.00 "лв."', '0" credits"', "0.0 °C",
]


@pytest.mark.parametrize("fmt", KEPT_FORMATS)
def test_structural_formats_are_unchanged(fmt: str) -> None:
    assert safe_number_format(fmt) == fmt


@pytest.mark.parametrize("fmt,expected", [
    ('"Zelmira Quenby "0', '"Xxxxxxx Xxxxxx "0'),                    # quoted literal (the reported case)
    ("\\Z\\e\\l\\m\\i\\r\\a 0", "\\X\\x\\x\\x\\x\\x\\x 0"),           # escapes, judged as one word
    ("Zelmira 0", '"Xxxxxxx" 0'),                                     # unquoted non-code letters: quoted
    ('"Mrs. "@', '"Xxx. "@'),
    ("[$Zelmira Quenby-409]0", "[$Xxxxxxx Xxxxxx-409]0"),            # currency-tag text
    ("[Zelmira]0", "0"),                                              # unknown bracket dropped
    ("[$Abe-Bee]0", "[$Xxx-Xxx]0"),                                   # hex letters are not an LCID
    ("Bess 0", '"Xxxx" 0'),                                           # mixed-case code letters: a word
    (f'"Call {PHONE} "0', '"Xxxx 000-000-0000 "0'),                  # digits zeroed
    (f'"{EMAIL} "0', '"xxxxxxx.xxxxxx@xxxxxxx.xxxxxxx.xxx "0'),
    ('"张伟 "0', '"xx "0'),
    ('0 "Zelmira', '0 "Xxxxxxx"'),                                   # unterminated quote
    ("_Z0*Q", "_X0*X"),                                               # spacing / fill characters
    ('General;"Quenby"General', 'General;"Xxxxxx"General'),
])
def test_format_literals_are_masked(fmt: str, expected: str) -> None:
    out = safe_number_format(fmt)
    assert out == expected
    assert safe_number_format(out) == out  # idempotent: the self-check compares rebuilt formats
    assert "Zelmira" not in out and "Quenby" not in out and "555" not in out


def test_mask_literal_keeps_units_symbols_and_length() -> None:
    assert mask_literal(" kg") == " kg" and mask_literal("€ $ ° ²") == "€ $ ° ²"
    assert mask_literal("Total: 12") == "Xxxxx: 00"
    s = "Zelmira Quenby (ID 004512)"
    assert len(mask_literal(s)) == len(s)


def test_css_mso_number_format_is_masked() -> None:
    v = '"\\0022Zelmira Quenby\\0022 0"'
    out = safe_css_number_format(v)
    assert "Zelmira" not in out and "Quenby" not in out
    assert safe_css_number_format(out) == out
    for keep in ('"\\@"', "Standard", '"Short Date"', '"0\\.00"', "'\\#\\,\\#\\#0\\.00'", '"0" !important'):
        assert safe_css_number_format(keep) == keep
    css = 'td { mso-number-format:"\\0022Dr. Zelmira Quenby\\0022 0"; color: red; }'
    cleaned = excel_html.clean_css(css)
    assert "Zelmira" not in cleaned and "color: red" in cleaned and "mso-number-format" in cleaned


def _grades_book(path: Path, fmt: str = '"Zelmira Quenby "0') -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Grades"
    ws.append(["score", "full_name"])
    for i in range(30):
        ws.append([i, NAME])
        ws.cell(i + 2, 1).number_format = fmt
    wb.save(path)
    return path


def test_xlsx_number_format_literal_with_a_detected_name(tmp_path: Path) -> None:
    """The reported case: the name of a PERSON_NAME column is also quoted in a number format."""
    src = _grades_book(tmp_path / "in" / "grades.xlsx")
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0 and _statuses(rep)["grades.xlsx"] in ("PASS", "WARN")
    blob = _blob(out, rep)
    assert b"Zelmira" not in blob and b"Quenby" not in blob
    with zipfile.ZipFile(out / "grades.xlsx") as z:
        styles = z.read("xl/styles.xml").decode("utf-8")
    assert "&quot;Xxxxxxx Xxxxxx &quot;0" in styles  # the formatting (a literal prefix) is kept, masked
    f = _file_report(rep, "grades.xlsx")
    assert f["checks"]["layout_selfcheck"]["result"] == "PASS"
    extras = [m["text"] for m in f["messages"] if m["code"] == "excel_extras_dropped"]
    assert extras and "number-format literal text (masked) (1)" in extras[0]


def test_xlsx_title_header_footer_formats_and_units(tmp_path: Path) -> None:
    """Formats of title, header and footer cells are rebuilt too; unit/currency formats are kept."""
    src = tmp_path / "in" / "book.xlsx"
    src.parent.mkdir(parents=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 2024
    ws["A1"].number_format = '"Prepared by Zelmira Quenby for "0'
    ws.append([])
    ws.append([1, "weight", "amount", "category"])  # a number header cell, formatted
    ws["A3"].number_format = '"Quenby "0'
    for i in range(30):
        ws.append([i, 50 + i / 10, 1000 + i, "A" if i % 2 else "B"])
        ws.cell(i + 4, 2).number_format = '0.0" kg"'
        ws.cell(i + 4, 3).number_format = '[$€-407] #,##0.00'
    ws.append(["Total", None, "=SUM(C4:C33)"])
    ws.cell(34, 3).number_format = '"Signed: Zelmira "#,##0'
    wb.save(src)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0, _report(rep)
    blob = _blob(out, rep)
    assert b"Zelmira" not in blob and b"Quenby" not in blob and b"Prepared" not in blob
    wb2 = openpyxl.load_workbook(out / "book.xlsx")
    ws2 = wb2["Data"]
    assert ws2["B4"].number_format == '0.0" kg"' and ws2["C4"].number_format == '[$€-407] #,##0.00'
    assert ws2["A1"].number_format == '"Xxxxxxxx xx Xxxxxxx Xxxxxx xxx "0'
    assert ws2["A3"].number_format == '"Xxxxxx "0'
    assert _file_report(rep, "book.xlsx")["checks"]["layout_selfcheck"]["result"] == "PASS"


def test_xls_formats_are_masked_on_read() -> None:
    class _Fmt:
        format_str = '"Zelmira Quenby "0.00'

    class _Xf:
        format_key = 7

    class _Book:
        xf_list = [_Xf()]
        format_map = {7: _Fmt()}

    masked = excel_read.Counter()
    cache: dict[int, str] = {}
    assert excel_read._xls_fmt(_Book(), 0, cache, masked) == '"Xxxxxxx Xxxxxx "0.00'
    assert masked[excel_read.MASKED_FORMAT_EXTRA] == 1


def test_html_xls_mso_number_format_is_masked(tmp_path: Path) -> None:
    src = tmp_path / "in" / "roster.xls"
    src.parent.mkdir(parents=True)
    lines = ['<html xmlns:o="urn:schemas-microsoft-com:office:office" '
             'xmlns:x="urn:schemas-microsoft-com:office:excel">', "<head>",
             '<meta http-equiv="Content-Type" content="text/html; charset=utf-8">', "<style>",
             'td { mso-number-format:"\\@"; }', '.n { mso-number-format:"\\0022Zelmira Quenby\\0022 0"; }',
             "</style>", "</head>", "<body>", '<table border="1">', "<tr><th>score</th><th>grp</th></tr>"]
    for i in range(30):
        lines.append(f'<tr><td class="n" style="mso-number-format:\'\\0022Quenby\\0022 0\'">{i}</td>'
                     f"<td>{'A' if i % 2 else 'B'}</td></tr>")
    lines += ["</table>", "</body>", "</html>"]
    src.write_bytes("\r\n".join(lines).encode("utf-8"))
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0, _report(rep)
    text = (out / "roster.xls").read_text(encoding="utf-8")
    assert "Zelmira" not in text and "Quenby" not in text
    assert 'mso-number-format:"\\@"' in text and text.count("mso-number-format") >= 3


def test_leak_scan_reads_number_formats(tmp_path: Path) -> None:
    """Defence in depth: number formats (an attribute of styles.xml) are part of the scanned metadata."""
    p = tmp_path / "f.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["A1"].number_format = f'"{EMAIL}"0'  # written directly by openpyxl, not by the tool
    wb.save(p)
    assert EMAIL in leaks.xlsx_text_view(p)


@pytest.mark.parametrize("fmt,expected", [
    ("%tdCCYY-NN-DD!_!Z!e!l!m!i!r!a", "%tdCCYY-NN-DD!_!X!x!x!x!x!x!x"),
    ("%tcHH:MM!_!J!S", "%tcHH:MM!_!X!X"),
    ("%td!_!k!g", "%td!_!k!g"),
    ("%9.0g", "%9.0g"), ("%-20s", "%-20s"), ("%td", "%td"), ("%10.2fc", "%10.2fc"),
])
def test_stata_display_format_literals(fmt: str, expected: str) -> None:
    assert statio.safe_stata_format(fmt) == expected
    assert statio.safe_stata_format(expected) == expected


def test_stata_display_format_literal_is_not_copied(tmp_path: Path) -> None:
    src = tmp_path / "in" / "d.dta"
    src.parent.mkdir(parents=True)
    pd.DataFrame({"visit": [dt.datetime(2024, 1, 1) + dt.timedelta(days=i) for i in range(30)],
                  "score": list(range(30))}).to_stata(src, write_index=False, version=118,
                                                     convert_dates={"visit": "td"})
    statio.patch_formats(src, 118, ["%tdCCYY-NN-DD!_!Z!e!l!m!i!r!a!_!Q!u!e!n!b!y", "%9.0g"])
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0, _report(rep)
    data = (out / "d.dta").read_bytes()
    assert b"!Z!e!l" not in data and b"!Q!u!e" not in data
    with pd.io.stata.StataReader(out / "d.dta") as r:
        r.read()
        assert r._fmtlist[0] == "%tdCCYY-NN-DD!_!X!x!x!x!x!x!x!_!X!x!x!x!x!x"


# =========================================================================== R2-03 unused labels
def _dict_table(secret: str, ordered: bool = False, extra: dict | None = None) -> pa.Table:
    arr = pa.DictionaryArray.from_arrays(pa.array([0, 1] * 15, type=pa.int8()), pa.array(["A", "B", secret]),
                                         ordered=ordered)
    return pa.table({"category": arr, **(extra or {})})


@pytest.mark.parametrize("fmt", ["parquet", "feather", "dta"])
def test_unused_sensitive_label_is_not_written(tmp_path: Path, fmt: str) -> None:
    """The reported case: an unused label equal to a value of an email column."""
    src = tmp_path / "in" / f"data.{fmt}"
    src.parent.mkdir(parents=True)
    if fmt == "dta":
        pd.DataFrame({"category": [1, 2] * 15, "email": [EMAIL] * 30}).to_stata(
            src, write_index=False, version=118, value_labels={"category": {1: "A", 2: "B", 3: EMAIL}})
    else:
        t = _dict_table(EMAIL, extra={"email": [EMAIL] * 30})
        (pq.write_table if fmt == "parquet" else feather.write_feather)(t, src)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0 and _statuses(rep)[src.name] in ("PASS", "WARN")
    assert EMAIL.encode() not in _blob(out, rep)
    out_file = out / src.name
    if fmt == "dta":
        with pd.io.stata.StataReader(out_file) as r:
            r.read()
            assert r.value_labels()["category"] == {1: "A", 2: "B"}
        assert "value_labels_unused" in _codes(rep, src.name)
    else:
        t = (pq.read_table if fmt == "parquet" else feather.read_table)(out_file)
        assert t.column("category").combine_chunks().dictionary.to_pylist() == ["A", "B"]
        assert "dictionary_unused_levels" in _codes(rep, src.name)
    assert _file_report(rep, src.name)["checks"]["layout_selfcheck"]["result"] == "PASS"


def test_unused_ordered_level_never_reaches_data_docs_or_report(tmp_path: Path) -> None:
    """An ordered dictionary (ORDINAL) used to quote every level in its params, docs (R level order) and
    report; an unused real name there is not detectable by any scan, so it must never be copied."""
    src = tmp_path / "in" / "levels.parquet"
    src.parent.mkdir(parents=True)
    arr = pa.DictionaryArray.from_arrays(pa.array([0, 2] * 15, type=pa.int8()), pa.array(["low", NAME, "high"]),
                                         ordered=True)
    pq.write_table(pa.table({"level": arr, "score": list(range(30))}), src)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0, _report(rep)
    assert b"Zelmira" not in _blob(out, rep)
    d = pq.read_table(out / "levels.parquet").column("level").combine_chunks()
    assert d.dictionary.to_pylist() == ["low", "high"] and d.type.ordered
    col = _file_report(rep, "levels.parquet")["tables"][0]["columns"][0]
    assert col["type"] == "ORDINAL" and col["params"]["levels"] == ["low", "high"]


def test_dictionary_variants_chunked_forced_and_pandas_categories(tmp_path: Path) -> None:
    inp = tmp_path / "in"
    inp.mkdir()
    a1 = pa.DictionaryArray.from_arrays(pa.array([0, 1] * 8, type=pa.int8()), pa.array(["A", "B", EMAIL]))
    a2 = pa.DictionaryArray.from_arrays(pa.array([0, 1] * 7, type=pa.int8()), pa.array(["B", "A"]))
    pq.write_table(pa.Table.from_batches([pa.record_batch({"category": a1, "email": [EMAIL] * 16}),
                                          pa.record_batch({"category": a2, "email": [EMAIL] * 14})]),
                   inp / "chunked.parquet")
    pq.write_table(_dict_table(EMAIL, extra={"email": [EMAIL] * 30}), inp / "forced.parquet")
    pd.DataFrame({"category": pd.Categorical(["A", "B"] * 15, categories=["A", "B", NAME]),
                  "score": list(range(30))}).to_parquet(inp / "pandas_cat.parquet", index=False)
    arr = pa.DictionaryArray.from_arrays(pa.array([0, 1] * 15, type=pa.int8()), pa.array([1, 2, 3125550199]))
    feather.write_feather(pa.table({"code": arr, "phone": [PHONE] * 30}), inp / "int_dict.feather")
    code, out, rep = _run(inp, tmp_path, "--force-factor", "forced.parquet::category")
    assert code == 0, _report(rep)
    assert set(_statuses(rep).values()) <= {"PASS", "WARN"}
    blob = _blob(out, rep)
    assert EMAIL.encode() not in blob and b"Zelmira" not in blob
    for name in ("chunked.parquet", "forced.parquet", "pandas_cat.parquet"):
        assert pq.read_table(out / name).column("category").combine_chunks().dictionary.to_pylist() == ["A", "B"]
    md = json.loads(pq.read_schema(out / "pandas_cat.parquet").metadata[b"pandas"])
    assert md["columns"][0]["metadata"]["num_categories"] == 2  # follows the written dictionary
    assert pd.read_parquet(out / "pandas_cat.parquet")["category"].cat.categories.tolist() == ["A", "B"]
    assert feather.read_table(out / "int_dict.feather").column("code").combine_chunks().dictionary.to_pylist() \
        == [1, 2]


def test_tool_vocabulary_labels_stay_when_unused(tmp_path: Path) -> None:
    """A grade dictionary keeps a grade level no row holds; a Likert label too; anything else goes."""
    inp = tmp_path / "in"
    inp.mkdir()
    grades = pa.DictionaryArray.from_arrays(pa.array([0, 1, 2] * 10, type=pa.int8()),
                                            pa.array(["KG", "01", "02", "03", NAME]), ordered=True)
    pq.write_table(pa.table({"grade": grades, "score": list(range(30))}), inp / "g.parquet")
    pd.DataFrame({"agree": [1, 2] * 15, "score": list(range(30))}).to_stata(
        inp / "s.dta", write_index=False, version=118,
        value_labels={"agree": {1: "Disagree", 2: "Agree", 3: "Strongly agree", 9: NAME}})
    code, out, rep = _run(inp, tmp_path)
    assert code == 0, _report(rep)
    assert b"Zelmira" not in _blob(out, rep)
    assert pq.read_table(out / "g.parquet").column("grade").combine_chunks().dictionary.to_pylist() == \
        ["KG", "01", "02", "03"]
    with pd.io.stata.StataReader(out / "s.dta") as r:
        r.read()
        assert r.value_labels()["agree"] == {1: "Disagree", 2: "Agree", 3: "Strongly agree"}
    for name in ("g.parquet", "s.dta"):
        assert _file_report(rep, name)["checks"]["layout_selfcheck"]["result"] == "PASS"


@pytest.mark.parametrize("labels", [{1: "A", 2: "A", 3: EMAIL}, {1: "A", 2: " ", 9: EMAIL}])
def test_stata_labels_never_shown_as_values_are_not_written(tmp_path: Path, labels: dict) -> None:
    """Duplicate or blank labels are shown as codes: no label text was checked, so no table is written."""
    src = tmp_path / "in" / "d.dta"
    src.parent.mkdir(parents=True)
    pd.DataFrame({"category": [1, 2, 3] * 10, "email": [EMAIL] * 30}).to_stata(
        src, write_index=False, version=118, value_labels={"category": labels})
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0 and EMAIL.encode() not in _blob(out, rep)
    with pd.io.stata.StataReader(out / "d.dta") as r:
        df = r.read(convert_categoricals=False)
        assert "category" not in r.value_labels()
    assert sorted(set(df["category"].tolist())) == [1, 2, 3]
    assert "value_labels_not_written" in _codes(rep, "d.dta")


def test_vocabulary_label() -> None:
    # D35: the standard education vocabulary counts too ('Refused', 'Title I', 'Free/Reduced' quote nobody)
    for s in ("KG", "06", "12", "Strongly agree", "Proficient", "Yes", "N/A", "2023-2024", " ", "sophomore", "Refused",
              "Title I", "Free/Reduced"):
        assert arrowio.vocabulary_label(s), s
    for s in (NAME, EMAIL, PHONE, "Adams MS", "Zeta Cohort", "3125550199", None):
        assert not arrowio.vocabulary_label(s), s


def test_self_check_and_leak_scan_catch_a_written_unused_label(tmp_path: Path, monkeypatch) -> None:
    """Regression guard: if a writer copied an unused label again, the self-check FAILs the file (strict) and
    the leak scan reads the label text no cell shows (--no-strict)."""
    src = tmp_path / "in" / "data.feather"
    src.parent.mkdir(parents=True)
    feather.write_feather(_dict_table(EMAIL, extra={"email": [EMAIL] * 30}), src)

    def old_writer(values, field, native, profile):  # the old behaviour: the whole source dictionary
        pos = {s: i for i, s in enumerate(native["dictionary_values"]) if s is not None}
        idx = [None if v in (None, "") else pos[v] for v in values.tolist()]
        return pa.DictionaryArray.from_arrays(pa.array(idx, type=field.type.index_type), native["dictionary"])

    monkeypatch.setattr(arrowio, "_dictionary_array", old_writer)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 3
    f = _report(rep)["failures"][0]
    assert "self-check" in f["reasons"][0] and "dictionary entr" in f["reasons"][0]
    code, out, rep = _run(src.parent, tmp_path, "--no-strict", tag="2")
    assert code == 3 and not (out / "data.feather").exists()
    assert "leak scan" in _report(rep)["failures"][0]["reasons"][0]
    assert EMAIL.encode() not in _blob(out)


def test_hidden_label_text() -> None:
    from datanonymous.model import FileLayout, RawColumn, RawTable, ReadResult

    d = RawColumn("c", np.array(["A", None], dtype=object), {"kind": "dictionary",
                                                             "dictionary_values": ["A", "B", EMAIL]})
    s = RawColumn("s", np.array(["Yes", "1"], dtype=object), {"family": "stata", "labels": {0: "Yes", 3: PHONE},
                                                              "var_label": "about Zelmira"})
    rr = ReadResult(FileLayout("dta", extra={"_data_label": "by Quenby"}), [RawTable(None, ["c", "s"], [d, s], 2)],
                    [])
    text = leaks.hidden_label_text(rr)
    assert EMAIL in text and PHONE in text and "B" in text.split("\n")
    assert "about Zelmira" in text and "by Quenby" in text and "Yes" not in text.split("\n")


# =========================================================================== R2-12 wide numbers
B53 = 2 ** 53


def _write_csv(path: Path, header: list[str], rows: list[list]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return path


def _col(path: Path, col: str) -> list:
    if path.suffix == ".csv":
        return [r[col] for r in csv.DictReader(path.open(encoding="utf-8"))]
    if path.suffix == ".json":
        return [r[col] for r in json.loads(path.read_text(encoding="utf-8"))]
    return (pq.read_table if path.suffix == ".parquet" else feather.read_table)(path).column(col).to_pylist()


def _assert_exact_extremes(real: list, got: list, num=lambda s: Decimal(str(s))) -> None:
    rv = [num(x) for x in real if x not in (None, "") and not str(x).startswith("<")]
    gv = [num(x) for x in got if x not in (None, "") and not str(x).startswith("<")]
    assert len(got) == len(real)  # N in = N out
    assert min(rv) in gv and max(rv) in gv  # D14: the exact min and max
    assert min(gv) >= min(rv) and max(gv) <= max(rv)  # never outside the real range


@pytest.mark.parametrize("kind", ["csv", "json", "parquet"])
def test_sparse_wide_integers_keep_the_exact_max(tmp_path: Path, kind: str) -> None:
    """The reported case: 62 values 2**53 + 101*i; the odd maximum was rounded away by float64."""
    vals = [B53 + 101 * i for i in range(62)]
    src = tmp_path / "in" / f"data.{kind}"
    if kind == "csv":
        _write_csv(src, ["amount"], [[v] for v in vals])
    elif kind == "json":
        src.parent.mkdir(parents=True)
        src.write_text(json.dumps([{"amount": v} for v in vals]), encoding="utf-8")
    else:
        src.parent.mkdir(parents=True)
        pq.write_table(pa.table({"amount": pa.array(vals, pa.int64())}), src)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0, _report(rep)
    got = _col(out / src.name, "amount")
    _assert_exact_extremes(vals, got)
    assert "extremes_not_emitted" not in _codes(rep, src.name)


def _wide_cases(inp: Path) -> dict[str, tuple[str, list, object]]:
    """name -> (column, real values, parser) for the wide-number variants."""
    cases: dict[str, tuple[str, list, object]] = {}
    vals = [-B53 - 101 * i - 7 for i in range(62)]
    _write_csv(inp / "neg.csv", ["amount"], [[v] for v in vals])
    cases["neg.csv"] = ("amount", vals, Decimal)
    vals = [2 ** 64 - 1 - 1009 * i for i in range(62)]  # used to FAIL: rounded values overflowed uint64
    pq.write_table(pa.table({"amount": pa.array(vals, pa.uint64())}), inp / "uint64.parquet")
    cases["uint64.parquet"] = ("amount", vals, Decimal)
    vals = [2 ** 63 - 1 - 997 * i for i in range(62)]  # used to FAIL: rounded values overflowed int64
    feather.write_feather(pa.table({"amount": pa.array(vals, pa.int64())}), inp / "int64max.feather")
    cases["int64max.feather"] = ("amount", vals, Decimal)
    vals = [f"{123456789012345 + 7 * i}.{(i * 37) % 1000:03d}" for i in range(62)]  # 18 significant digits
    _write_csv(inp / "float18.csv", ["amount"], [[v] for v in vals])
    cases["float18.csv"] = ("amount", vals, Decimal)
    r = random.Random(5)
    vals = [repr(r.random()) for _ in range(62)]  # float64 repr text: 16-17 significant digits
    _write_csv(inp / "repr17.csv", ["share"], [[v] for v in vals])
    cases["repr17.csv"] = ("share", vals, Decimal)
    vals = [f"{B53 + 1 + 101 * i}.0" for i in range(62)]
    _write_csv(inp / "floattext.csv", ["amount"], [[v] for v in vals])
    cases["floattext.csv"] = ("amount", vals, Decimal)
    vals = [f"{B53 + 1 + 101 * i}.5" for i in range(62)]  # float64 read these as integers
    _write_csv(inp / "half.csv", ["amount"], [[v] for v in vals])
    cases["half.csv"] = ("amount", vals, Decimal)
    vals = [f"${B53 + 1 + 1013 * i:,}" for i in range(62)]
    _write_csv(inp / "formatted.csv", ["amount"], [[v] for v in vals])
    cases["formatted.csv"] = ("amount", vals, lambda s: Decimal(str(s).replace("$", "").replace(",", "")))
    vals = [str(B53 + 1 + 1013 * i) for i in range(60)] + ["<5"] * 12
    _write_csv(inp / "tokens.csv", ["amount"], [[v] for v in vals])
    cases["tokens.csv"] = ("amount", vals, Decimal)
    vals = [-(B53 + 3 + 101 * i) for i in range(62)]
    (inp / "neg.json").write_text(json.dumps([{"amount": v} for v in vals]), encoding="utf-8")
    cases["neg.json"] = ("amount", vals, Decimal)
    return cases


def test_wide_number_variants_keep_exact_extremes(tmp_path: Path) -> None:
    inp = tmp_path / "in"
    inp.mkdir()
    cases = _wide_cases(inp)
    code, out, rep = _run(inp, tmp_path)
    assert code == 0, _report(rep)["failures"]
    st = _statuses(rep)
    for name, (col, vals, num) in cases.items():
        assert st[name] in ("PASS", "WARN"), name
        _assert_exact_extremes(vals, _col(out / name, col), num)
        assert "extremes_not_emitted" not in _codes(rep, name), name
    types = {c["name"]: c["type"] for c in _file_report(rep, "half.csv")["tables"][0]["columns"]}
    assert types["amount"] == "FLOAT"  # not misread as integer-valued by float64


def test_wide_within_id_constant_keeps_row_level_extremes(tmp_path: Path) -> None:
    rows = []
    for i in range(150):
        v = B53 + 1 + 2 * 1013 * i
        rows += [[f"S{i:03d}", v], [f"S{i:03d}", v]]
    rows.append(["S000", B53 + 1 + 2 * 1013 * 150 + 2])  # the max: a minority row of one ID
    src = _write_csv(tmp_path / "in" / "e.csv", ["student_id", "amount"], rows)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0, _report(rep)
    cols = {c["name"]: c for c in _file_report(rep, "e.csv")["tables"][0]["columns"]}
    assert cols["amount"]["role"] == "entity" and cols["amount"]["flags"]["extreme_value_adjusted"]
    _assert_exact_extremes([r[1] for r in rows], _col(out / "e.csv", "amount"))
    assert "extremes_not_emitted" not in _codes(rep, "e.csv")


def _profile(t: ColumnType, raws: list[str], subtype: str | None = None, **params) -> ColumnProfile:
    p = ColumnProfile(index=0, name="x", type=t, subtype=subtype, params=dict(params))
    counts: dict[str, int] = {}
    for r in raws:
        counts[r] = counts.get(r, 0) + 1
    p.model = ColumnModel(nonmissing_counts=counts,
                          sorted_numeric=np.sort(np.array([float(r) for r in raws], dtype=float)))
    return p


def test_extremes_self_check_compares_wide_values_exactly() -> None:
    raws = [str(B53 + 101 * i) for i in range(62)]
    p = _profile(ColumnType.INTEGER, raws, "continuous", decimals=0)
    ok = np.array(raws, dtype=object)
    rounded = np.array(raws[:-1] + [str(B53 + 101 * 61 - 1)], dtype=object)  # the float64-rounded max
    assert float(rounded[-1]) == float(raws[-1])  # float64 cannot tell them apart
    assert _extremes_missing(ok, p) is False
    assert _extremes_missing(rounded, p) is True
    small = _profile(ColumnType.INTEGER, [str(i) for i in range(62)], "continuous", decimals=0)
    assert G.exact_extremes_missing(small, np.array(["0", "61"], dtype=object), lambda r: False) is None


def test_discrete_level_extremes_are_ordered_exactly() -> None:
    levels = [str(B53 + k) for k in (1, 3, 5, 7, 9)]  # float64: 2**53+5 == 2**53+4 == 2**53+3 ...
    values = np.array([levels[0]] * 6 + [levels[2]] * 6, dtype=object)
    assert _ensure_level_extremes(values, levels, np.random.default_rng(1))
    assert levels[-1] in values.tolist() and levels[0] in values.tolist()


def test_generator_and_renderer_are_exact_for_each_type() -> None:
    rng = np.random.default_rng(3)
    p = _profile(ColumnType.INTEGER, [str(B53 + 101 * i) for i in range(62)], "continuous", decimals=0)
    y, adj = G.generate_continuous_wide(p, p.model, 62, rng, 0)
    assert all(isinstance(v, Decimal) for v in y) and max(y) == B53 + 101 * 61
    assert G.render_numeric(p, y)[int(np.argmax([int(v) for v in y]))] == str(B53 + 101 * 61)
    small = _profile(ColumnType.INTEGER, [str(i) for i in range(62)], "continuous", decimals=0)
    assert G.generate_continuous_wide(small, small.model, 62, rng, 0) is None  # float generator, unchanged
    sci = _profile(ColumnType.FLOAT, ["1.5e-300", "2.5e300"], scientific={"mantissa_decimals": 1})
    assert G.generate_continuous_wide(sci, sci.model, 4, rng, 0) is None
    assert G.format_number(Decimal("-9007199254740993.50"), {"decimals": 2, "group_sep": ",", "positive": "$#",
                                                               "negative": "($#)"}) == "($9,007,199,254,740,993.50)"


def test_wide_numbers_are_deterministic(tmp_path: Path) -> None:
    inp = tmp_path / "in"
    inp.mkdir()
    _wide_cases(inp)
    _grades_book(inp / "grades.xlsx")
    feather.write_feather(_dict_table(EMAIL, extra={"email": [EMAIL] * 30}), inp / "dict.feather")
    c1, out1, _ = _run(inp, tmp_path, seed="9", tag="a")
    c2, out2, _ = _run(inp, tmp_path, seed="9", tag="b")
    assert c1 == c2 == 0
    names = sorted(p.relative_to(out1).as_posix() for p in out1.rglob("*") if p.is_file())
    assert names == sorted(p.relative_to(out2).as_posix() for p in out2.rglob("*") if p.is_file())
    for n in names:
        assert (out1 / n).read_bytes() == (out2 / n).read_bytes(), n


def test_wide_integer_statistics_and_exposure_are_exact() -> None:
    """Report statistics of wide integers: exact min/max and pmf (float64 merged neighbours), and the D14
    exposure count of rows at an exact extreme."""
    from datanonymous.profile.stats import column_stats
    from datanonymous.report.report import _rows_at_extreme

    raws = [str(B53 + 3)] + [str(B53 + 4)] * 9 + [str(B53 + 5)] * 2  # float64: 2**53+3 == 2**53+4 == 2**53+5
    p = _profile(ColumnType.INTEGER, raws, "discrete", decimals=0)
    st = column_stats(np.array(raws, dtype=object), p, 10)
    assert st["min"] == B53 + 3 and st["max"] == B53 + 5
    assert st["pmf"] == {str(B53 + 3): 1, str(B53 + 4): 9, str(B53 + 5): 2}
    assert _rows_at_extreme(p, "min") == 1 and _rows_at_extreme(p, "max") == 2
    small = _profile(ColumnType.INTEGER, ["1", "2", "2"], "discrete", decimals=0)
    assert column_stats(np.array(["1", "2", "2"], dtype=object), small, 10)["pmf"] == {"1": 1, "2": 2}
