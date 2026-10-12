"""Round 8, cluster registry-titles: the run's real-text registry knows every real text, wherever it is held.

The bug class: a real value reached the output because one planning step did not consult the run-wide registry of
real texts (``privacy.texts``), or the registry did not know a place where real text lives.

* R8-03: ID mappings were built before the registry existed. E-mail and username IDs avoided only the addresses of
  EMAIL columns, so a synthetic ID took the username of an excluded or nulled column's address, or of an address
  inside a pseudonymised label; the scan did not check ID columns against the registry. Now ``plan_ids`` reads the
  registry (every real e-mail address of the run, wherever it is held) before any ID is drawn, and the scan checks
  every ID column against the registry's e-mail addresses and usernames (both layers).
* R8-04: source text outside the tables (Excel title, notes, footer and side-note text; HTML text outside the table)
  was not in the registry nor in the scan's index, so Faker could recreate a name that only a title held. Now it
  is registered (placeholders: real text the run does not write; ``--keep-titles``: written verbatim, so kept),
  the name words of a longer title are linked pair by pair in the full-name index, and the e-mail addresses,
  phone numbers and SSNs inside it are indexed and avoided. The audit adds the text no output holds: HTML
  comments and dropped attributes, free-form Arrow metadata, Stata variable and data labels.
* R8-06: title placeholders masked an e-mail address character by character, giving an address on an arbitrary
  domain. Now every placeholder goes through the shared embedded-PII rules (``generate.embedded.SpanSynth``:
  reserved domains, fresh phone bases) and the generators' last check, and never equals a real text, phone or SSN.

Fixtures are built from the tool's own output where a hot case needs it: a first run gives the value the tool
draws (an ID's local part, a Faker name); a second run plants exactly that value as real data in the new place.
Assertion messages never print a real value beyond the invented fixture data.
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook

from conftest import read_dicts, run_cli
from datanonymous import pipeline
from datanonymous.config import Settings
from datanonymous.generate import pii
from datanonymous.generate import text as gtext
from datanonymous.generate.ids import RealUsers
from datanonymous.io import excel_write
from datanonymous.model import ColumnProfile, ColumnType
from datanonymous.privacy import leaks
from datanonymous.privacy import texts as rt

SEED = "77"
RESERVED = ("example.org", "example.com", "example.net")
EMAIL_RE = re.compile(r"[\w.%+'\-]+@[\w\-]+(?:\.[\w\-]+)+")


def write_csv(path: Path, columns: dict[str, list[Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(list(columns))
        w.writerows(zip(*columns.values()))
    return path


def run(src: Path, tmp: Path, *extra: str, name: str = "run"):
    out, rep = tmp / f"{name}-out", tmp / f"{name}-rep"
    res = run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", SEED, *extra)
    return res, out, rep


def file_page(rep: Path, file: str) -> dict[str, Any]:
    return json.loads((rep / "files" / f"{file}.json").read_text(encoding="utf-8"))


def leak_scan(rep: Path, file: str) -> str:
    return file_page(rep, file)["checks"]["leak_scan"]["result"]


def everything(*dirs: Path, console: str = "") -> str:
    return "\n".join([p.read_bytes().decode("utf-8", errors="replace") for d in dirs if d.exists()
                      for p in sorted(d.rglob("*")) if p.is_file()] + [console]).casefold()


def folded(values: list[str]) -> set[str]:
    return {" ".join(str(v).casefold().split()) for v in values if v and str(v).strip()}


# =========================================================================== R8-03: e-mail IDs and usernames
def email_key(n: int = 20, **cols: list[Any]) -> dict[str, list[Any]]:
    return {"Student_ID": ["pupil@district.org"] * n, "Score": [str(i) for i in range(n)], **cols}


def ids_of(out: Path, file: str = "email_key.csv") -> list[str]:
    path = out / file
    return [r["Student_ID"] for r in read_dicts(path)] if path.exists() else []


def harvest_local(tmp_path: Path) -> str:
    """The local part the ID registry draws for the one student when nothing else is in the run."""
    src = write_csv(tmp_path / "harvest" / "email_key.csv", email_key())
    res, out, _ = run(src, tmp_path, name="harvest")
    assert res.exit_code == 0, res.output
    ids = set(ids_of(out))
    assert len(ids) == 1 and next(iter(ids)).endswith(RESERVED)
    return next(iter(ids)).partition("@")[0].casefold()


def planted(tmp_path: Path, local: str, how: str) -> tuple[Path, tuple[str, ...]]:
    """``email_key.csv`` with a real address whose local part is ``local`` held in another column (``how``)."""
    addr = f"{local}@district.org"
    n = 20
    if how in ("exclude", "null"):
        cols, flags = email_key(Guardian_Email=[addr] * n), (f"--{how}", "Guardian_Email")
    elif how == "label":  # one student's room is an address: D35 pseudonymises it (never kept)
        cols, flags = email_key(Room=[addr] + ["Rm 42"] * (n - 1)), ()
    elif how == "note":  # an address inside free text
        cols, flags = email_key(Notes=[f"Call home or write to {addr} about the trip number {i}" for i in range(n)]), ()
    else:
        raise AssertionError(how)
    return write_csv(tmp_path / how / "email_key.csv", cols), flags


@pytest.mark.parametrize("how", ["exclude", "null", "label", "note"])
def test_email_ids_never_take_a_real_username_held_anywhere(tmp_path: Path, how: str) -> None:
    local = harvest_local(tmp_path)
    src, flags = planted(tmp_path, local, how)
    res, out, rep = run(src, tmp_path, *flags, name="planted")
    assert res.exit_code == 0, res.output
    ids = ids_of(out)
    assert len(ids) == 20 and len(set(ids)) == 1, "N in = N out, one synthetic ID per real ID"
    assert all(i.endswith(RESERVED) for i in ids)
    assert local not in {i.partition("@")[0].casefold() for i in ids}, "a real username was given to a student"
    assert leak_scan(rep, "email_key.csv") == "PASS"


def test_email_ids_avoid_an_address_in_an_excel_title(tmp_path: Path) -> None:
    local = harvest_local(tmp_path)
    p = tmp_path / "xl" / "email_key.xlsx"
    p.parent.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.append([f"Questions: {local}@district.org"])
    ws.merge_cells("A1:B1")
    ws.append(["Student_ID", "Score"])
    for i in range(20):
        ws.append(["pupil@district.org", i])
    wb.save(p)
    res, out, rep = run(p, tmp_path, name="xl")
    assert res.exit_code == 0, res.output
    rows = list(load_workbook(out / "email_key.xlsx").active.values)
    ids = {r[0] for r in rows[2:]}
    assert len(ids) == 1 and local not in {str(i).partition("@")[0].casefold() for i in ids}
    assert local not in str(rows[0][0]).casefold()
    assert leak_scan(rep, "email_key.xlsx") == "PASS"


def no_id_avoidance(monkeypatch: pytest.MonkeyPatch) -> None:
    """ID planning without the registry (what the code did before R8-03, minus the EMAIL columns): only the scan
    stands. (R9-03 made the ID registry refuse every registry text too, its e-mail local parts included: off here.)"""
    monkeypatch.setattr(pipeline.real_texts, "email_addresses", lambda texts: set())
    monkeypatch.setattr(RealUsers, "text_taken", lambda self, value: False)


def test_the_id_fixture_is_hot_without_the_fix(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    local = harvest_local(tmp_path)
    src, flags = planted(tmp_path, local, "exclude")
    no_id_avoidance(monkeypatch)
    monkeypatch.setattr(leaks, "scan_table", lambda tp, st: [])
    monkeypatch.setattr(leaks, "scan_outputs", lambda run, payload: [])
    res, out, _ = run(src, tmp_path, *flags, name="hot")
    assert res.exit_code == 0, res.output
    assert {i.partition("@")[0].casefold() for i in ids_of(out)} == {local}


@pytest.mark.parametrize("layer", ["before writing", "written file"])
def test_scanner_fails_an_id_with_a_real_username(tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                                                  layer: str) -> None:
    local = harvest_local(tmp_path)
    src, flags = planted(tmp_path, local, "exclude")
    no_id_avoidance(monkeypatch)
    if layer == "written file":
        monkeypatch.setattr(leaks, "run_texts", lambda tp: None)  # layer 1 sees no registry: the data scan must
    res, out, rep = run(src, tmp_path, *flags, name="leaky")
    assert res.exit_code == 3, res.output
    assert not (out / "email_key.csv").exists()
    page = json.dumps(file_page(rep, "email_key.csv"))
    assert ("username of, a real e-mail address held in the run" in page if layer == "before writing"
            else "real e-mail username" in page)
    blob = everything(out, rep, console=res.output)
    assert f"{local}@" not in blob, "a real address or username was printed"


def test_user_hits_unit() -> None:
    texts = rt.RealTexts(frozenset({"robin@district.org", "write to ann.lee@school.net today"}), frozenset(),
                         (), frozenset(rt.local_parts(["robin@district.org", "write to ann.lee@school.net today"])))
    assert rt.email_addresses(texts) == {"robin@district.org", "ann.lee@school.net"}
    assert texts.user_hits(["Robin@example.org", "ANN.LEE@example.com", "julie@example.org"]) == {
        "Robin@example.org", "ANN.LEE@example.com"}
    # a username ID (letters of a real local part, read without the kept prefix) and an address with that username
    assert texts.user_hits(["robin42", "annlee7", "S104019", "zz@example.org"], prefix="") == {"robin42", "annlee7"}
    assert texts.user_hits(["Srobin1"], prefix="s") == {"Srobin1"}
    assert texts.user_hits(["Srobin1"], prefix="") == set()
    id_col = ColumnProfile(index=0, name="Student_ID", type=ColumnType.IDENTIFIER)
    assert leaks.id_user_hits(texts, id_col, ["robin@example.org"]) == {"robin@example.org"}
    id_col.params["_identity_mapped"] = True  # kept row numbers 1..n are no synthetic IDs
    assert leaks.id_user_hits(texts, id_col, ["robin@example.org"]) == set()
    label = ColumnProfile(index=1, name="Notes", type=ColumnType.FREE_TEXT)
    assert leaks.id_user_hits(texts, label, ["robin@example.org"]) == set()  # texts: checked as texts


def test_id_generation_avoids_everything_the_scan_checks() -> None:
    """One source of truth: every local part and username the scan checks ID columns against is one the ID registry
    avoids (whatever the real address's domain looks like), so the scan never refuses what generation could not
    know."""
    from datanonymous.generate.ids import RealUsers

    keys = ["robin@district.org", "write to ann.lee@school.net today", "x.y.z@host.c0", "José.Núñez@colegio.mx",
            "o'neil@b.io", "kept.one@school.org"]
    texts = rt.RealTexts(frozenset(rt.key(k) for k in keys), frozenset({"kept.one@school.org"}), (),
                         frozenset(rt.local_parts(rt.key(k) for k in keys[:-1])))
    users = RealUsers.of([], sorted(rt.email_addresses(texts)))
    assert texts.locals <= users.locals
    assert texts.user_stems <= users.stems
    assert "kept.one" in users.locals  # an address written verbatim somewhere is avoided too (superset)


# =========================================================================== R8-04: text outside the tables
NAMES = ("John Smith", "Mary Jones")


def roster_xlsx(path: Path, title: str | None, names: tuple[str, ...] = NAMES, *, notes: list[str] | None = None,
                side: str | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    if title is not None:
        ws.append([title])
        ws.merge_cells("A1:C1")
    ws.append(["Student_ID", "Student_Name", "Score"])
    for i in range(60):
        ws.append([f"S{123000 + i}", names[i % len(names)], i])
    if side is not None:  # a headerless column beside the table, separated by an empty column: a side note
        start = 3 if title is not None else 2
        ws.cell(row=start, column=5, value=side)
    if notes:
        ns = wb.create_sheet("Notes")
        for line in notes:
            ns.append([line])
    wb.save(path)
    return path


def names_of(out: Path, file: str = "roster.xlsx") -> list[str]:
    rows = list(load_workbook(out / file, data_only=True)["Data"].values)
    head = next(i for i, r in enumerate(rows) if r and r[0] == "Student_ID")
    return [r[1] for r in rows[head + 1:] if r and r[0]]


def harvest_name(tmp_path: Path) -> str:
    src = roster_xlsx(tmp_path / "harvest" / "roster.xlsx", "Teacher report")
    res, out, _ = run(src, tmp_path, name="harvest")
    assert res.exit_code == 0, res.output
    fake = names_of(out)
    assert len(fake) == 60 and len(set(fake)) == 2
    return fake[0]


def pairs(values: list[str]) -> set[frozenset[str]]:
    return {frozenset(v.casefold().split()) for v in values}


@pytest.mark.parametrize("where", ["title", "long title", "side note", "notes sheet"])
def test_a_name_held_only_outside_the_table_is_never_generated(tmp_path: Path, where: str) -> None:
    name = harvest_name(tmp_path)
    p = tmp_path / where.replace(" ", "_") / "roster.xlsx"
    if where == "title":
        src = roster_xlsx(p, name)
    elif where == "long title":
        src = roster_xlsx(p, f"Class roster for Ms. {name}, printed 10/01/2024 (period 3)")
    elif where == "side note":
        src = roster_xlsx(p, "Teacher report", side=f"Call {name} first")
    else:
        src = roster_xlsx(p, "Teacher report", notes=["Notes for the spring term", f"Field trip lead: {name}",
                                                      "Bring the signed forms back by Friday"])
    res, out, rep = run(src, tmp_path, name="planted")
    assert res.exit_code == 0, res.output
    got = names_of(out)
    assert len(got) == 60
    assert name.casefold() not in folded(got), "a real name from the text outside the table was generated"
    assert frozenset(name.casefold().split()) not in pairs(got), "a real person's two names were generated"
    assert leak_scan(rep, "roster.xlsx") == "PASS"


def test_a_name_in_an_html_xls_title_is_never_generated(tmp_path: Path) -> None:
    def html(title: str) -> str:
        rows = "".join(f"<tr><td>S{5000 + i}</td><td>{NAMES[i % 2]}</td><td>{i}</td></tr>" for i in range(60))
        return (f"<html><body><h2>{title}</h2><table><tr><td colspan=\"3\">{title}</td></tr>"
                f"<tr><th>Student_ID</th><th>Student_Name</th><th>Score</th></tr>{rows}</table></body></html>")

    def names(out: Path) -> list[str]:
        text = (out / "roster.xls").read_text(encoding="utf-8")
        return re.findall(r"<td>S\d+</td><td>([^<]*)</td>", text)

    p = tmp_path / "h1" / "roster.xls"
    p.parent.mkdir(parents=True)
    p.write_text(html("Teacher report"), encoding="utf-8")
    res, out, _ = run(p, tmp_path, name="harvest")
    assert res.exit_code == 0, res.output
    fake = names(out)
    assert len(fake) == 60
    name = fake[0]
    p = tmp_path / "h2" / "roster.xls"
    p.parent.mkdir(parents=True)
    p.write_text(html(name), encoding="utf-8")
    res, out, rep = run(p, tmp_path, name="planted")
    assert res.exit_code == 0, res.output
    got = names(out)
    assert len(got) == 60 and name.casefold() not in folded(got)
    assert name.casefold() not in (out / "roster.xls").read_text(encoding="utf-8").casefold()
    assert leak_scan(rep, "roster.xls") == "PASS"


def no_name_avoidance(monkeypatch: pytest.MonkeyPatch) -> None:
    """Generation without the run-wide registry (what the code did before R7/R8): only the scan stands."""
    monkeypatch.setattr(gtext, "real_text", lambda text: False)
    plan = pii.plan_names
    monkeypatch.setattr(pipeline.pii, "plan_names", lambda files, extra=(), texts=None: plan(files, extra, None))


def test_the_title_fixture_is_hot_without_the_fix(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    name = harvest_name(tmp_path)
    src = roster_xlsx(tmp_path / "hot" / "roster.xlsx", name)
    no_name_avoidance(monkeypatch)
    monkeypatch.setattr(leaks, "scan_table", lambda tp, st: [])
    monkeypatch.setattr(leaks, "scan_outputs", lambda run, payload: [])
    res, out, _ = run(src, tmp_path, name="hot")
    assert res.exit_code == 0, res.output
    assert name.casefold() in folded(names_of(out))


@pytest.mark.parametrize("layer", ["before writing", "written file"])
def test_scanner_fails_a_copy_of_title_text(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, layer: str) -> None:
    name = harvest_name(tmp_path)
    src = roster_xlsx(tmp_path / "leaky" / "roster.xlsx", name)
    no_name_avoidance(monkeypatch)
    if layer == "written file":
        monkeypatch.setattr(leaks, "run_texts", lambda tp: None)
    res, out, rep = run(src, tmp_path, name="leaky")
    assert res.exit_code == 3, res.output
    assert not (out / "roster.xlsx").exists()
    page = json.dumps(file_page(rep, "roster.xlsx"))
    assert ("real text value held in the run" in page if layer == "before writing" else "rule: real text" in page)
    if layer == "written file":
        assert "title/notes/metadata text" in page and "--null roster.xlsx::(text" not in page
    assert name.casefold() not in everything(out, rep, console=res.output), "a real value was printed"


def test_title_contacts_are_indexed_and_avoided(tmp_path: Path) -> None:
    """The e-mail addresses, phone numbers and SSNs of the text outside the tables are real data: the data scan
    knows them (every output text and the file metadata view), and every generator avoids them."""
    src = roster_xlsx(tmp_path / "in" / "roster.xlsx", "Office: (312) 867-5309, Ana Diaz <ana.diaz@gmail.com>, "
                                                       "SSN 123-45-6789")
    fr = profiled(src)
    texts = leaks.real_texts_of([fr])
    assert texts.outside and "ana.diaz" in texts.locals
    assert "ana.diaz@gmail.com" in rt.email_addresses(texts)
    keys = leaks.plan_digit_keys([fr])
    assert not keys.phones  # no registry planned yet: only the columns
    for tp in fr.tables:
        tp.private.extra[rt.RUN_KEY] = texts
    keys = leaks.plan_digit_keys([fr])
    assert "3128675309" in keys.phones and "123456789" in keys.ssns
    index = leaks.SensitiveIndex.build([fr])
    assert index.emails.get("ana.diaz@gmail.com")
    assert index.phone_digits.get("3128675309") and index.ssn_digits.get("123456789")
    o = next(iter(index.emails.get("ana.diaz@gmail.com")))
    assert index.origins[o].kind == leaks.LAYOUT and index.origins[o].column == rt.LAYOUT_WHERE


def profiled(path: Path):
    """Pass 1 of a run over one file (read, profile, the run's D35 decision)."""
    from datanonymous.profile.table import finish_copy_gate

    fr = pipeline.profile_file(_discovered(path), Settings(mode="fidelity", seed=1), _run())
    assert fr is not None and not fr.failed
    finish_copy_gate(fr.tables, 10)
    return fr


def _discovered(path: Path):
    from datanonymous.model import DiscoveredFile

    fmt = {".xlsx": "xlsx", ".xls": "html_xls", ".parquet": "parquet", ".dta": "dta"}[path.suffix]
    return DiscoveredFile(path=path, rel_path=path.name, out_rel_path=path.name, format=fmt)


def _run():
    from datanonymous.model import RunResult

    return RunResult(mode="inspect", input=Path("."), out_dir=None, report_dir=None, seed=1, seed_was_random=False)


def test_keep_titles_text_is_kept_not_refused(tmp_path: Path) -> None:
    """--keep-titles writes the title verbatim (D34 screens it): it is kept, so no false refusal, and the
    placeholder path never changes what --keep-titles writes."""
    src = roster_xlsx(tmp_path / "in" / "roster.xlsx", "Grade 5 Science - Spring Roster")
    res, out, rep = run(src, tmp_path, "--keep-titles", name="keep")
    assert res.exit_code == 0, res.output
    assert load_workbook(out / "roster.xlsx")["Data"]["A1"].value == "Grade 5 Science - Spring Roster"
    assert leak_scan(rep, "roster.xlsx") == "PASS"
    fr = profiled(src)
    kept = leaks.real_texts_of([fr], keep_titles=True)
    assert not kept.holds("Grade 5 Science - Spring Roster") and rt.key("Grade 5 Science - Spring Roster") in kept.kept
    plain = leaks.real_texts_of([fr])
    assert plain.holds("grade 5 science - spring roster")
    assert [s.kind for s, _ in plain.sources if s.column == rt.LAYOUT_WHERE] == [rt.LAYOUT]


def test_a_total_label_in_a_footer_is_kept() -> None:
    from types import SimpleNamespace

    from datanonymous.io.excel_read import XCell
    from datanonymous.model import FileResult, TableProfile

    tp = TableProfile(name="Data", header=["a"], n_rows=1, columns=[],
                      layout={"_title_rows": [(1, {1: XCell("Roster for Quenby Zane", "s", "General")})],
                              "_footer_rows": [(9, {1: XCell("Total", "s", "General"),
                                                    2: XCell("Checked by Quenby Zane", "s", "General")})]})
    fr = FileResult("f.xlsx", "f.xlsx", "xlsx", tables=[tp])
    fr.layout = SimpleNamespace(extra={})
    texts = leaks.real_texts_of([fr])
    assert texts.holds("roster for quenby zane") and texts.holds("Checked by Quenby Zane")
    assert not texts.holds("Total") and "total" in texts.kept
    index = pii.plan_names([fr], (), texts)
    assert "quenby" in index.get("zane", set())


def test_long_layout_text_links_adjacent_name_words_only() -> None:
    long = "class roster for ms. tara walsh, printed on monday for the office"
    texts = rt.RealTexts(frozenset({long}), frozenset(), ((rt.Source("f.xlsx", None, rt.LAYOUT_WHERE, rt.LAYOUT),
                                                            frozenset({long})),))
    index = pii.plan_names([], (), texts)
    assert "tara" in index.get("walsh", set()) and "walsh" in index.get("tara", set())
    assert "office" not in index.get("tara", set())  # only neighbours: never every pair of a long text
    other = rt.RealTexts(frozenset({long}), frozenset(), ((rt.Source("f.csv", None, "Notes", rt.TEXT),
                                                             frozenset({long})),))
    assert "tara" not in pii.plan_names([], (), other).get("walsh", set())  # free text: whole values only (R7)


# =========================================================================== R8-06: placeholders, D24 rules
def title_out(tmp_path: Path, title: str, name: str = "t", fmt: str = "xlsx") -> str:
    if fmt == "xlsx":
        src = roster_xlsx(tmp_path / name / "roster.xlsx", title)
        res, out, rep = run(src, tmp_path, name=name)
        assert res.exit_code == 0, res.output
        assert leak_scan(rep, "roster.xlsx") == "PASS"
        return str(load_workbook(out / "roster.xlsx")["Data"]["A1"].value)
    p = tmp_path / name / "roster.xls"
    p.parent.mkdir(parents=True)
    rows = "".join(f"<tr><td>S{5000 + i}</td><td>{NAMES[i % 2]}</td><td>{i}</td></tr>" for i in range(60))
    p.write_text(f"<html><body><p>{title}</p><table><tr><th>Student_ID</th><th>Student_Name</th><th>Score</th></tr>"
                 f"{rows}</table></body></html>", encoding="utf-8")
    res, out, rep = run(p, tmp_path, name=name)
    assert res.exit_code == 0, res.output
    assert leak_scan(rep, "roster.xls") == "PASS"
    return re.search(r"<p>(.*?)</p>", (out / "roster.xls").read_text(encoding="utf-8"), re.S).group(1)


@pytest.mark.parametrize("fmt", ["xlsx", "html_xls"])
def test_a_title_email_gets_a_reserved_domain(tmp_path: Path, fmt: str) -> None:
    got = title_out(tmp_path, "Contact: john.smith@gmail.com", fmt=fmt)
    emails = EMAIL_RE.findall(got)
    assert len(emails) == 1 and emails[0].casefold().endswith(RESERVED), "D24: reserved domains only"
    assert emails[0].casefold().partition("@")[0] != "john.smith"
    assert got.startswith(tuple("ABCDEFGHIJKLMNOPQRSTUVWXYZ")) and got[7:9] == ": "  # the rest keeps its shape
    assert "gmail" not in got.casefold()


def test_title_urls_phones_and_ssns_follow_the_pii_rules(tmp_path: Path) -> None:
    got = title_out(tmp_path, "See https://www.lincolnhigh.org/roster or call (312) 867-5309; SSN 123-45-6789")
    url = re.search(r"https://www\.([\w.\-]+)/", got)
    assert url and url.group(1).casefold().endswith(RESERVED)
    phone = re.search(r"\((\d{3})\) (\d{3})-(\d{4})", got)
    assert phone and phone.group(2)[0] in "23456789", "D46 phone rule: ordinary NXX exchange"
    ssn = re.search(r"(\d{3})-(\d{2})-(\d{4})", got.replace(phone.group(0), ""))
    assert ssn and ssn.group(1).startswith("9"), "SSN rule: area 9xx (never issued)"
    assert "867-5309" not in got and "123-45-6789" not in got and "lincolnhigh" not in got.casefold()


def test_a_plain_title_keeps_the_same_placeholder(tmp_path: Path) -> None:
    """No random number is drawn for a text without a personal-data shape: the placeholder is the same-shape mask
    it always was (the same seed gives the same file, D26)."""
    got = title_out(tmp_path, "Teacher report 2024-10-01", name="a")
    again = title_out(tmp_path, "Teacher report 2024-10-01", name="b")
    assert got == again
    assert re.fullmatch(r"[A-Z][a-z]{6} [a-z]{6} \d{4}-\d{2}-\d{2}", got) and got != "Teacher report 2024-10-01"
    rng = np.random.default_rng(5)
    ph = excel_write.Placeholders(rng, keep=False)
    old = excel_write.Placeholders(np.random.default_rng(5), keep=False)
    assert ph.text("Teacher report") == old._shape("Teacher report")


def test_placeholders_avoid_the_registry() -> None:
    """A placeholder that equals a real text of the run is drawn again (the guard is active while a file is
    written)."""
    first = excel_write.Placeholders(np.random.default_rng(3), keep=False).text("Ab")
    texts = rt.RealTexts(frozenset({first.casefold()}), frozenset(), ())
    with rt.guard(texts):
        again = excel_write.Placeholders(np.random.default_rng(3), keep=False).text("Ab")
    assert again != first and re.fullmatch(r"[A-Z][a-z]", again)


# =========================================================================== audit: metadata no output holds
def test_free_form_arrow_metadata_is_registered(tmp_path: Path) -> None:
    import pyarrow as pa
    import pyarrow.parquet as pq

    names = ["Ana Garcia", "Liam Brown", "Noah Jones", "Mia Lopez"] * 5
    table = pa.table({"student_id": [f"S{1000 + i}" for i in range(20)], "Full_Name": names,
                      "score": list(range(20))})
    a = tmp_path / "a" / "grades.parquet"
    a.parent.mkdir(parents=True)
    pq.write_table(table, a)
    res, out, _ = run(a, tmp_path, name="harvest")
    assert res.exit_code == 0, res.output
    fake = pq.read_table(out / "grades.parquet").column("Full_Name").to_pylist()[0]
    meta = {b"owner": fake.encode(), b"notes": json.dumps({"reviewed_by": "Quenby Zane"}).encode()}
    b = tmp_path / "b" / "grades.parquet"
    b.parent.mkdir(parents=True)
    pq.write_table(table.replace_schema_metadata({**(table.schema.metadata or {}), **meta}), b)
    fr = profiled(b)
    texts = leaks.real_texts_of([fr])
    assert texts.holds(fake) and texts.holds("Quenby Zane") and not texts.holds("Full_Name")
    res, out, rep = run(b, tmp_path, name="planted")
    assert res.exit_code == 0, res.output
    got = pq.read_table(out / "grades.parquet")
    assert fake.casefold() not in folded(got.column("Full_Name").to_pylist())
    assert not (got.schema.metadata or {}).get(b"owner")  # free-form metadata is still never copied
    assert leak_scan(rep, "grades.parquet") == "PASS"


def test_stata_labels_and_html_extras_are_registered(tmp_path: Path) -> None:
    import pandas as pd

    df = pd.DataFrame({"student_id": [f"S{1000 + i}" for i in range(12)], "score": list(range(12))})
    p = tmp_path / "grades.dta"
    df.to_stata(p, write_index=False, variable_labels={"score": "Graded by Quenby Zane"},
                data_label="Prepared for Tamsin Vale")
    fr = profiled(p)
    texts = leaks.real_texts_of([fr])
    assert texts.holds("Graded by Quenby Zane") and texts.holds("prepared for tamsin vale")
    index = pii.plan_names([fr], (), texts)
    assert "quenby" in index.get("zane", set()) and "tamsin" in index.get("vale", set())
    h = tmp_path / "export.xls"
    rows = "".join(f"<tr><td>S{5000 + i}</td><td>{i}</td></tr>" for i in range(12))
    h.write_text("<html><head><meta name=\"author\" content=\"Quill Marrow\">"
                 "<!-- exported for robin.hale@district.org --></head><body>"
                 f"<table><tr><th>Student_ID</th><th>Score</th></tr>{rows}</table></body></html>", encoding="utf-8")
    fr = profiled(h)
    texts = leaks.real_texts_of([fr])
    assert texts.holds("Quill Marrow") and "robin.hale" in texts.locals
    assert "robin.hale@district.org" in rt.email_addresses(texts)


def test_html_css_and_class_text_the_writer_drops_is_registered() -> None:
    from datanonymous.io import excel_html

    events = [("start", "style", []),
              ("data", "/* exported for Okafor */ td.Thibodeaux { color: red } .xl65 { mso-number-format:\"0\\.00\"; }"
                       " @page { mso-header-data:\"&LPrinted by Quenby Zane\"; }"
                       " .x { background: url(C:/jdoe/a.png) }"),
              ("end", "style"),
              ("start", "td", [("class", "xl65 quixote"), ("style", "content: 'Tamsin Vale'; color: red"),
                               ("x:num", "3125550199"), ("title", "Hartfield")]),
              ("data", "7"), ("end", "td"),
              ("comment", " for robin.hale@district.org ")]
    got = excel_html.dropped_texts(events)
    for text in ("exported for Okafor", "td.Thibodeaux", "Thibodeaux", "&LPrinted by Quenby Zane", "C:/jdoe/a.png",
                 "quixote", "Tamsin Vale", "Hartfield", " for robin.hale@district.org "):
        assert text in got, text
    assert "xl65" not in got and "3125550199" not in got and "7" not in got  # kept formatting; cell data
