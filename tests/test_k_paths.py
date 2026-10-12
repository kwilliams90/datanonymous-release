"""K paths cluster: the D34 header and path screen (``privacy.headers``), structural fixes of the known issues K-01,
K-10, K-12, K-16, K-18 and K-19. D34 stays frozen as a best-effort layer: no name list or name heuristic is added.

* K-01 (G1): a real student ID in a name passed the real-ID screen at district scale (one 6-digit number among 6,000
  IDs is weak evidence), and a pivot with one bare-number column per student passed whenever the IDs were a dense
  range of sequential numbers (each number alone is chance there). Now a number right after a person or document
  label (``Student``, ``Stu``, ``Pupil``, ``Learner``, ``IEP``, ``504``, ``Eval``, ``Evaluation``, ``ID``) that equals a
  real ID is strong evidence on its own, and three or more sibling names that are numbers alone and ALL real IDs of
  ONE namespace are a pivot unless chance explains it (the namespace's density over every number of that length;
  dates, numbered lists and a numbered parent are chance). D37 (2026-10-11) replaced that chance threshold: two or
  more such real IDs are a pivot whatever the density, and B4 is gone (``test_d37_id_headers.py``).
* K-10 (G3): ``_xHHHH_`` (OOXML) escapes, HTML character references and digits glued between two words hid a known
  first + last name from the person-name check; they are decoded before screening and in the redaction and leak
  terms of a flagged name.
* K-12 (G3): ``SS#``/``Soc Sec``/``W2`` labels before nine digits are SSNs, and ``NXX-XXXX`` after a phone or
  family word is a phone number (the dotted 3-2-4 SSN is the shared shape test of values since R9-02).
* K-16 (G5): step 3 (one name, one verdict) and the leak scanner's backstop share one matcher: the scanner's token
  sequences (underscores, hyphens, camelCase), so a sibling holding a flagged name gets the D34 verdict and its
  ``--keep-name`` fix instead of a run-level "please report this as a bug" ERROR.
* K-18 (G5): step 3 never copies an ID flag onto a number alone that its own parent names (``enrollment_by_zip``),
  and a zero-padded exact ID is weighed by the share of zero-padded numbers of its width that are real IDs.
* K-19 (G5): a 3-3-4 run joined by spaces or underscores is excused like a bare run (``Lexile_800_900_1000``).

Hardening after an attack on the K-01 fix (P-01 .. P-10, last section): a label after the number (``612587_IEP``), the
``Student Number`` family of labels, D10 person-ID header names and one document word (``IEP Meeting 612587``); pivots
whose IDs are wrapped (``score1007``, ``_1007``, ``1007_score``), hold a withdrawn student or span per-school rosters;
a labelled parent (``iep/1053.csv``); a singular place word of a parent (``Lincoln_Middle_School_Gradebook``) no longer
excuses its children; a labelled ID that reads as a year-month or a date; and a flagged dense pivot no longer FAILs the
statistics, counts and dated sheets of other files.

Hardening after the second attack (A1 .. A8, last section): a numbered word's excuse gives way to sibling evidence
(``Progress_Report_612714`` x25); a zero-padded ID is strong unless the run uses zero-padded numbers of its width as
a code that is no ID (county-district numbers); a numbered parent (``SOC_Projects``) no longer hides a pivot of sparse
IDs; a flagged person name is repeated in either order (``Mark_Patel`` beside ``Patel, Mark``); the folders of the input
path are screened before step 3; real IDs the run's names hold too often for chance count across folders (per-student
documents one per folder); clock times on five-minute marks are no pivot (lunch periods ``1235,1305,1330``).
"""

from __future__ import annotations

import csv
import json
import random
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from conftest import run_cli
from datanonymous.privacy import headers as d34

SEED = "3"


class NS:
    def __init__(self, values: list[str]) -> None:
        self.raw_values = set(values)
        self.representatives = set(values)


def write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return path


def write_book(path: Path, sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def run(src: Path, tmp: Path, *extra: str, name: str = "run"):
    out, rep = tmp / f"{name}-out", tmp / f"{name}-rep"
    res = run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", SEED, "--force", *extra)
    return res, out, rep


def everything(*dirs: Path, console: str = "") -> str:
    return "\n".join([str(p) for d in dirs if d.exists() for p in sorted(d.rglob("*"))]
                     + [p.read_bytes().decode("utf-8", errors="replace") for d in dirs if d.exists()
                        for p in sorted(d.rglob("*")) if p.is_file()] + [console])


def statuses(rep: Path) -> dict[str, str]:
    data = json.loads((rep / "report.json").read_text(encoding="utf-8"))
    return {f["source"]: f["status"] for f in data["files"]}


def roster(src: Path, ids: list[str], name: str = "roster.csv") -> None:
    rng = random.Random(4)
    write_csv(src / name, ["student_id", "grade", "score"], [[i, rng.choice(["6", "7", "8"]), rng.randint(0, 100)]
                                                             for i in ids])


def district_ids(n: int = 6000, seed: int = 7) -> list[str]:
    return [str(x) for x in random.Random(seed).sample(range(100000, 1000000), n)]


# =========================================================================== K-01: person labels and pivots
@pytest.mark.parametrize("stub", ["Student_{}_IEP_Goals", "Student {}", "IEP_{}", "Evaluation_{}", "Eval {}",
                                  "504_{}", "ID {}", "ID#{}", "Student ID {}", "StudentID{}", "Stu_ID_{}", "STUID {}",
                                  "Pupil {}", "Learner_{}", "SID {}", "Teacher_ID_{}", "Student ({})", "IEP [{}]",
                                  "IEP Progress Report - Student {}"])
def test_a_real_id_after_a_person_label_is_strong_whatever_the_density(stub: str) -> None:
    ids = district_ids()
    keys = d34.IdKeys([NS(ids)])
    x = ids[11]
    assert not keys.match(f"Goals_{x}")  # 6,000 of 900,000 six-digit numbers: one alone is chance
    assert keys.match(stub.format(x)), stub
    assert "id" in d34.name_kinds(stub.format(x), keys), stub


@pytest.mark.parametrize("stub", ["Q17_{}", "Assignment_{}", "Course ID {}", "Item_ID_{}", "CourseID {}",
                                  "Room {}", "Section_ID_{}", "Score {}", "Course (ID {})", "Homework 1 ({})"])
def test_a_number_after_another_word_is_still_weighed_by_chance(stub: str) -> None:
    ids = district_ids()
    keys = d34.IdKeys([NS(ids)])
    assert not keys.match(stub.format(ids[11])), stub


@pytest.mark.parametrize("name", ["Student Count 1200", "Students_1200", "Course ID 1050", "Item_ID_1014", "Room 1014",
                                  "HW1014", "Q1014", "Lexile 1014", "Section 504 Plan", "IEP Goals 2024",
                                  "Eval Window 2024-25", "IEP_Due_1015", "StudentCount_1200", "Bus 1104",
                                  "Evaluation Rubric v1200", "ID_Badge_Batch_1200", "Section ID 1101"])
def test_ordinary_school_names_stay_clean_beside_dense_ids(name: str) -> None:
    dense = d34.IdKeys([NS([str(x) for x in range(1001, 1401)])])
    assert d34.name_kinds(name, dense) == frozenset(), name


def test_a_label_never_makes_a_year_or_a_number_that_is_no_real_id_one() -> None:
    keys = d34.IdKeys([NS([str(x) for x in range(1001, 1401)])])
    assert keys.match("Student 1014") and keys.match("IEP_1201")  # dense sequential IDs, a label: strong
    assert not keys.match("Student 2023") and not keys.match("Student 1500") and not keys.match("Attendance_1014")
    assert not keys.match("Section 504 Plan 1014")  # the word right before the number is no label


def _pivot_ids() -> tuple[d34.IdKeys, list[str]]:
    keys = d34.IdKeys([NS([str(x) for x in range(1001, 1401)])])  # 400 sequential 4-digit student numbers
    picks = sorted(random.Random(5).sample(range(1001, 1401), 25))
    return keys, [str(p) for p in picks]


def test_a_dense_pivot_of_bare_student_numbers_is_flagged() -> None:
    keys, picks = _pivot_ids()
    texts = ["Assignment"] + picks
    flags = d34.id_group_flags(keys, texts)
    assert set(flags) == set(range(1, 26))
    assert flags[1] == (picks[0],)
    # every number alone is chance in that dense range: only the set of them is evidence
    assert not any(keys.match(p) for p in picks)


@pytest.mark.parametrize("numbers", [
    [str(x) for x in range(1001, 1026)],  # assignment numbers counting up: a numbered list
    [str(x) for x in range(1010, 1400, 10)],  # a numbered list in steps of ten
    ["1001", "1002", "1005", "1006", "1007", "1008", "1009", "1012", "1013", "1014"],  # daily sheets (MMDD dates)
])
def test_numbered_things_in_a_dense_id_range_are_no_pivot(numbers: list[str]) -> None:
    keys, _ = _pivot_ids()
    assert d34.id_group_flags(keys, ["Assignment"] + numbers) == {}, numbers


@pytest.mark.parametrize("numbers", [["1375", "1380", "1392"], ["1135", "1162", "1247", "1388"]])
def test_d37_numbers_that_are_all_real_ids_are_a_pivot_whatever_the_density(numbers: list[str]) -> None:
    keys, _ = _pivot_ids()  # room numbers inside the range were chance (0.044 ** 3) before D37: now real IDs
    assert set(d34.id_group_flags(keys, ["Assignment"] + numbers)) == set(range(1, len(numbers) + 1)), numbers


def test_a_pivot_needs_nearly_every_number_to_be_a_real_id() -> None:
    keys, picks = _pivot_ids()
    # P-04: a withdrawn student's column (no ID of this run) or a year beside 25 real IDs leaves a pivot
    assert set(d34.id_group_flags(keys, ["Assignment"] + picks + ["1412"])) == set(range(1, 26))
    assert set(d34.id_group_flags(keys, ["Assignment"] + picks + ["2024"])) == set(range(1, 26))
    assert set(d34.id_group_flags(keys, ["Assignment"] + picks + ["1412", "1433"])) == set(range(1, 26))
    # D37: a few withdrawn students (at most a quarter of the set) never excuse the rest
    assert set(d34.id_group_flags(keys, ["Assignment"] + picks + ["1412", "1433", "1450"])) == set(range(1, 26))
    late = [str(x) for x in range(1412, 1700, 37)][:8]
    assert set(d34.id_group_flags(keys, ["Assignment"] + picks + late)) == set(range(1, 26))  # 25 of 33
    assert d34.id_group_flags(keys, ["Assignment"] + picks + late + ["1777"]) == {}  # 25 of 34: numbered things
    # more numbers that are no real ID: numbered things that partly fall in the ID range
    mixed = picks[:12] + [str(x) for x in (1450, 1475, 1502, 1533, 1561, 1598, 1630, 1666)]
    assert d34.id_group_flags(keys, ["Item"] + mixed) == {}
    assert d34.id_group_flags(keys, ["date", "1235", "1305", "1330", "1350", "1415", "1440"]) == {}  # HHMM times
    assert d34.id_group_flags(keys, picks, context=d34.numbered_context("rooms")) == {}  # the parent says rooms


def test_a_pivot_over_ids_that_look_like_years_is_still_a_pivot() -> None:
    wide = d34.IdKeys([NS([str(x) for x in range(1001, 3001)])])  # 2,000 sequential IDs, 1900-2100 among them
    nums = sorted(str(x) for x in random.Random(1).sample(range(1001, 3001), 20))
    assert any(d34.date_like(n) for n in nums)
    assert set(d34.pivot_flags(wide, nums)) == set(range(20))  # a year-like ID counts as chance, it never blocks
    assert set(d34.pivot_flags(wide, nums[:8])) == set(range(8))  # D37: a few of them too, whatever the density
    assert d34.pivot_flags(wide, ["2015", "2019", "2023", "2024"]) == {}  # years alone: dates


def test_a_dense_pivot_fails_and_a_numbered_gradebook_is_written(tmp_path: Path) -> None:
    src = tmp_path / "in"
    keys, picks = _pivot_ids()
    roster(src, [str(x) for x in range(1001, 1401)])
    rng = random.Random(2)
    write_csv(src / "gradebook_pivot.csv", ["Assignment", *picks],
              [[f"HW{k}"] + [rng.randint(50, 100) for _ in picks] for k in range(1, 13)])
    write_csv(src / "items.csv", ["Student", *[str(x) for x in range(1101, 1121)]],
              [[f"S{k}"] + [rng.randint(0, 4) for _ in range(20)] for k in range(30)])  # item numbers 1101-1120
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    assert st["gradebook_pivot.csv"] == "FAIL" and st["items.csv"] != "FAIL", st
    assert "a real ID of this run: 25" in res.output
    header = (out / "items.csv").read_text(encoding="utf-8").splitlines()[0]
    assert header.startswith("Student,1101,1102")
    assert not (out / "gradebook_pivot.csv").exists()
    blob = everything(out, rep, console=res.output)
    assert ",".join(picks[:3]) not in blob


def test_per_student_names_at_district_scale_fail_and_are_never_printed(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = district_ids()
    roster(src, ids)
    a, b, c, d = ids[11], ids[12], ids[13], ids[14]
    write_csv(src / "iep" / f"Student_{a}_IEP_Goals.csv", ["goal", "progress"], [[f"Goal {k}", 50] for k in range(4)])
    write_csv(src / f"IEP_{b}.csv", ["goal", "status"], [[f"Goal {k}", "met"] for k in range(4)])
    write_csv(src / "scores.csv", ["item", f"Student {c}"], [[f"Q{k}", 70] for k in range(6)])
    (src / f"Evaluation_{d}.pdf").write_bytes(b"%PDF-1.4\n")
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    assert st["scores.csv"] == "FAIL" and st["roster.csv"] != "FAIL", st
    assert sum(1 for s in st.values() if s == "FAIL") == 3, st
    blob = everything(out, rep, console=res.output)
    for x in (a, b, c, d):
        assert x not in blob, x


def test_keep_titles_screens_a_title_with_a_labelled_id(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = district_ids()
    roster(src, ids)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Progress"
    ws.append([f"IEP Progress Report - Student {ids[20]}"])
    ws.append([])
    ws.append(["goal", "baseline", "current"])
    for k in range(20):
        ws.append([f"Goal {k % 4}", 10 + k, 20 + k])
    wb.save(src / "progress.xlsx")
    res, out, rep = run(src, tmp_path, "--keep-titles")
    assert statuses(rep)["progress.xlsx"] == "FAIL", res.output
    assert ids[20] not in everything(out, rep, console=res.output)


# =========================================================================== K-10: escaped and glued names
@pytest.mark.parametrize("name", ["Ana_x0020_Garcia", "Kaden_x0020_Yazzie_x0020_IEP", "Ana&nbsp;Garcia",
                                  "Ana&#32;Garcia", "Ana&#x20;Garcia", "Ana1Garcia", "Ana2Garcia", "ana1garcia",
                                  "Garcia_x002C__x0020_Ana", "Ana%5Fx0020%5FGarcia"])
def test_escaped_and_glued_names_are_screened(name: str) -> None:
    assert "name" in d34.name_kinds(name), name


@pytest.mark.parametrize("name", ["Q1Math", "P1Garcia", "Grade3Reading", "Algebra1Honors", "HW_x0020_Score",
                                  "Fees&amp;Costs", "Pass&notes", "AT&T_Roster", "ELA2Writing"])
def test_escapes_and_glued_digits_make_no_name_of_labels(name: str) -> None:
    assert "name" not in d34.name_kinds(name), name


def test_an_escaped_flagged_name_is_redacted_and_indexed_in_its_decoded_forms() -> None:
    screen = d34.NameScreen()
    for text in ("Ana_x0020_Garcia", "Liam&nbsp;Johnson", "Sofia1Martinez"):
        screen.add(text, "column", frozenset({"name"}), "column 2 of 4", written=True, number=2)
    for printed, n in (("Ana Garcia", 1), ("Ana_x0020_Garcia", 1), ("Liam Johnson", 2), ("Liam\xa0Johnson", 2),
                       ("Sofia Martinez", 3), ("Sofia1Martinez", 3)):
        assert screen.redact_text(f"see {printed} here") == f"see [redacted-{n}] here", printed
    words, seqs, exact = screen.flags[0].terms()
    assert ("ana", "garcia") in seqs and "Ana Garcia" in exact
    assert screen.flags[2].words == ("sofia", "martinez")
    plain = d34.NameScreen()
    plain.add("Student907550Goals", "file", frozenset({"id"}), "file 1 of 1", written=True, id_tokens=("907550",))
    words, seqs, exact = plain.flags[0].terms()  # glued digits are a separator of person names only
    assert ("student", "goals") not in seqs and "Student Goals" not in exact
    assert plain.redact_text("Student Goals") == "Student Goals"


def test_escaped_names_in_a_gradebook_and_a_file_name_fail_and_never_print(tmp_path: Path) -> None:
    src = tmp_path / "in"
    names = ["Ana_x0020_Garcia", "Liam_x0020_Johnson", "Sofia_x0020_Martinez"]
    write_csv(src / "gradebook.csv", ["Assignment", *names], [[f"HW{k}", 80, 90, 70] for k in range(30)])
    write_csv(src / "iep" / "Kaden_x0020_Yazzie_x0020_IEP.csv", ["goal", "progress"], [["Goal 1", 50]] * 4)
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    assert set(statuses(rep).values()) == {"FAIL"}
    blob = everything(out, rep, console=res.output)
    for s in ("Ana_x0020_Garcia", "Ana Garcia", "Liam Johnson", "Kaden_x0020_Yazzie", "Kaden Yazzie"):
        assert s not in blob, s


# =========================================================================== K-12: SSN labels and local phones
@pytest.mark.parametrize("name", ["SS# 123456789", "SS#123456789", "SS #123456789", "SS No 123456789",
                                  "Soc Sec 123456789", "Social Security # 123456789", "W2_123456789",
                                  "W-2 123456789", "SSN 123456789", "123.45.6789", "SS# 123.45.6789"])
def test_ssn_labels_and_the_dotted_layout_are_ssns(name: str) -> None:
    assert "ssn" in d34.name_kinds(name), name


@pytest.mark.parametrize("name", ["Call 486-1234", "Ph 486-1234", "Cell_486_1234", "Mom 486 1234",
                                  "ParentPhone 486.1234", "Home Phone 555-0142"])
def test_a_local_number_after_a_phone_word_is_a_phone(name: str) -> None:
    assert "phone" in d34.name_kinds(name), name


@pytest.mark.parametrize("name", ["486-1234", "Room 486-1234", "Call 911-1234", "Lexile 200-1000", "Grade 9 1234",
                                  "W2 Wages 2023", "SS 2023-2024", "Week2_123456789x", "HW2 123456789",
                                  "Score 100-1500", "Call 123-4567", "Work 250-1000", "Text Level 200-1000",
                                  "Home Room 204-1000", "Phone Ext 1234", "SS 123456789"])
def test_no_ssn_or_phone_without_its_label_or_shape(name: str) -> None:
    assert not ({"ssn", "phone"} & d34.name_kinds(name)), name


def test_ssn_labels_and_local_phones_fail_files_folders_sheets_and_headers(tmp_path: Path) -> None:
    src = tmp_path / "in"
    rows = [[k, 80] for k in range(12)]
    write_csv(src / "SS# 345678912" / "grades.csv", ["row", "score"], rows)
    write_csv(src / "W2_234567891.csv", ["row", "score"], rows)
    write_csv(src / "Call 555-0142.csv", ["row", "score"], rows)
    write_csv(src / "roster.csv", ["row", "SS# 123456789", "W2_123456789", "Call 486-1234", "score"],
              [[k, 1, 2, 3, 80] for k in range(12)])
    write_book(src / "book.xlsx", {"Soc Sec 567891234": (["row", "score"], rows)})
    write_csv(src / "plain.csv", ["row", "score"], rows)
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    assert st.pop("plain.csv") != "FAIL" and set(st.values()) == {"FAIL"}, st
    blob = everything(out, rep, console=res.output)
    for s in ("345678912", "234567891", "555-0142", "123456789", "486-1234", "567891234"):
        assert s not in blob, s


# =========================================================================== K-16: one matcher for step 3 and the scanner
def test_repeats_match_the_scanner_tokens_underscores_hyphens_and_camel_case() -> None:
    screen = d34.NameScreen()
    screen.add("Jack Yates", "sheet", frozenset({"name"}), "the name of sheet 1 of 3", written=True)
    for text in ("Jack_Yates_HS_Enrollment", "Jack-Yates-HS", "JackYatesHS", "jack.yates", "Jack Yates",
                 "Yates_Jack"):  # A4: the same person's name words in the other order
        assert screen.repeats(text) == {"Jack Yates"}, text
    for text in ("Jackson_Yates", "Jack_Yatesville", "Westbury", "Jack_HS_Yates"):
        assert screen.repeats(text) == set(), text
    ids = d34.NameScreen()
    ids.add("S482901_Reading_Score", "column", frozenset({"id"}), "column 2 of 2", written=True, number=2,
            id_tokens=("S482901",))
    assert ids.repeats("S482901_Math") == {"S482901_Reading_Score"} and ids.repeats("Reading_Score") == set()
    assert ids.repeats("XS482901") == set()


def test_a_sibling_holding_a_flagged_sheet_name_gets_the_d34_verdict(tmp_path: Path) -> None:
    src = tmp_path / "in"
    rows = [[f"2024-09-{d:02d}", 300 + d, 12] for d in range(1, 7)]
    write_book(src / "attendance_by_campus.xlsx", {n: (["date", "present", "absent"], rows)
                                                   for n in ("Jack Yates", "Westbury", "Madison")})
    write_csv(src / "Jack_Yates_HS_Enrollment.csv", ["grade", "count"], [["9", 300], ["10", 280], ["11", 260]])
    write_csv(src / "JackYatesHS-Staffing.csv", ["role", "count"], [["teacher", 60], ["aide", 12], ["nurse", 2]])
    rng = random.Random(1)
    write_csv(src / "students.csv", ["student_id", "campus", "grade"],
              [[100000 + k, rng.choice(["Yates", "Westbury", "Madison"]), rng.choice(["9", "10"])] for k in range(300)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    assert "report this as a bug" not in res.output and "ERROR" not in res.output, res.output
    st = statuses(rep)
    assert sum(1 for s in st.values() if s == "FAIL") == 3, st
    assert st["students.csv"] != "FAIL"
    assert "D34:" in res.output and "--keep-name TEXT" in res.output
    assert "Jack Yates" not in everything(out, rep, console=res.output)
    res2, out2, rep2 = run(src, tmp_path, "--keep-name", "Jack Yates", name="kept")
    assert res2.exit_code == 0, res2.output
    assert "FAIL" not in set(statuses(rep2).values())


# =========================================================================== K-18: parent context and zero padding
def test_step_3_never_copies_an_id_flag_onto_a_number_its_parent_names(tmp_path: Path) -> None:
    src = tmp_path / "in"
    zips = ["30303", "30305", "30308", "30310", "30311", "30312"]
    rows = [["9", 120], ["10", 110], ["11", 100], ["12", 95]]
    write_book(src / "Atlanta_Enrollment.xlsx", {z: (["grade", "count"], rows) for z in zips})
    write_book(src / "enrollment_by_zip.xlsx", {z: (["grade", "count"], rows) for z in zips})
    rng = random.Random(3)
    write_csv(src / "students.csv", ["student_id", "grade", "score"],
              [[10001 + k, rng.choice(["9", "10"]), rng.randint(0, 100)] for k in range(28000)])
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert st["Atlanta_Enrollment.xlsx"] == "FAIL", st  # no context word: a sheet per real ID (by design)
    assert st["enrollment_by_zip.xlsx"] != "FAIL", res.output
    assert "leak scan" not in res.output and "report this as a bug" not in res.output, res.output
    assert openpyxl.load_workbook(out / "enrollment_by_zip.xlsx").sheetnames == zips


def test_a_zero_padded_exact_id_is_weighed_by_density() -> None:
    dense = d34.IdKeys([NS([f"{x:06d}" for x in range(1, 30001)])])  # 30,000 zero-padded IDs 000001-030000
    # A2: weighed by density only where the run's names use zero-padded numbers of that width as an ordinary code
    # (a county-district number that is no real ID: 057905 lies above the IDs)
    dense.note_run_names([("057905_Dallas_ISD_TAPR", None, False), ("015915_Northside_ISD_TAPR", None, False)])
    assert dense.padded_codes == {6}
    for name in ("015970_Northside_ISD_TAPR", "015915_Northside_ISD_TAPR", "Region_020980"):
        assert not dense.match(name), name
    assert dense.match("IEP_015970")  # a label still makes it strong
    sparse = d34.IdKeys([NS([f"{x:07d}" for x in random.Random(2).sample(range(1, 10_000_000), 2000)])])
    padded = sorted(v for v in sparse._index()[0] if v[0] == "0" and v[1] != "0")[0]
    assert sparse.match(f"Goals_{padded}")  # ~200 padded 7-digit IDs among a million padded numbers: strong
    assert d34.IdKeys([NS(["0482913"])]).padded_chance("0482913") == 1e-6


def test_county_district_numbers_beside_dense_padded_ids_are_written(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = [f"{x:06d}" for x in random.Random(8).sample(range(1, 10000), 3000)]
    cdns = [x for x in ids if not d34.natural_number(x)][:3]
    rng = random.Random(9)
    write_csv(src / "students.csv", ["student_id", "gender"], [[x, rng.choice("FM")] for x in ids])
    # A2: a folder of TAPR files holds county-district numbers that are no real ID too (057905: Dallas ISD), so its
    # zero-padded numbers are an ordinary code there (the triage repro: 015901 is no ID, 015907 and 015915 are)
    files = [f"{cdn}_{district}_ISD_TAPR.csv" for cdn, district in zip(cdns, ("Northside", "San_Antonio", "Judson"))]
    files.append("057905_Dallas_ISD_TAPR.csv")
    for f in files:
        write_csv(src / f, ["indicator", "value"], [[f"I{k}", 50 + k] for k in range(10)])
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    for f in files:
        assert st[f] != "FAIL", res.output


# =========================================================================== K-19: 3-3-4 number lists
@pytest.mark.parametrize("name", ["Lexile_800_900_1000", "Cut_100_200_3000", "Cut 100 200 3000",
                                  "RIT_190_200_2100", "Item_311_555_0199", "Bands 150 300 4500"])
def test_spaced_3_3_4_lists_are_excused_like_bare_runs(name: str) -> None:
    assert "phone" not in d34.name_kinds(name), name


@pytest.mark.parametrize("name", ["312_555_0199", "Parent 312 555 0199", "Score 312 555 0199", "Mom_100_200_3000",
                                  "312-555-0199", "Cut 312.555.0199", "020 123 4567", "Tel_020_123_4567",
                                  "1 312 555 0199"])
def test_phone_numbers_in_a_3_3_4_layout_stay_phones(name: str) -> None:
    assert "phone" in d34.name_kinds(name), name


# =========================================================================== hardening: the attack on the K-01 fix
SPARSE_X = "439563"  # one of district_ids(): no time stamp or short date (natural_number)


def _sparse() -> tuple[d34.IdKeys, list[str]]:
    ids = district_ids()
    assert SPARSE_X in ids and not d34.natural_number(SPARSE_X)
    return d34.IdKeys([NS(ids)]), ids


@pytest.mark.parametrize("stub", ["{}_IEP", "{}_IEP_Goals", "{} Student", "{}_Evaluation", "{}_504_Plan",
                                  "{}-Student-Notes", "{}StudentGoals", "2024_{}_IEP"])
def test_p01_a_real_id_right_before_a_person_label_is_strong(stub: str) -> None:
    keys, _ = _sparse()
    assert keys.match(stub.format(SPARSE_X)), stub
    dense = d34.IdKeys([NS([str(x) for x in range(1001, 1401)])])
    assert dense.match(stub.format("1053")), stub  # whatever the density


@pytest.mark.parametrize("stub", ["Course_{}_Student_Roster", "Section {} Students", "{} Students", "Room_{}_Student",
                                  "Bus {} Students", "Item {} IEP", "{}_Goals"])
def test_p01_a_number_a_numbered_word_claims_or_before_a_plural_stays_weighed_by_chance(stub: str) -> None:
    keys, _ = _sparse()
    assert not keys.match(stub.format(SPARSE_X)), stub


@pytest.mark.parametrize("stub", ["Student Number {}", "Student No {}", "Student No. {}", "Stu Num {}",
                                  "Student_Nbr_{}", "Pupil Number {}", "Learner Number {}", "Student ID Number {}",
                                  "StudentNumber{}", "STUDENTNUMBER {}", "Student_Number_{}", "ID Number {}",
                                  "Student, {}",
                                  "Student/{}", "Student–{}", "IEP—{}", "Student: {}", "Perm ID {}",
                                  "State ID {}", "Local ID {}", "SSID {}", "LASID {}", "IEP Meeting {}", "504 Plan {}",
                                  "IEP Report {}", "Evaluation Report {}", "IEP_Goals_{}", "Eval 2023 {}"])
def test_p02_student_number_labels_and_document_labels_make_a_real_id_strong(stub: str) -> None:
    keys, _ = _sparse()
    assert keys.match(stub.format(SPARSE_X)), stub
    assert "id" in d34.name_kinds(stub.format(SPARSE_X), keys), stub


@pytest.mark.parametrize("stub", ["Course No {}", "Room No {}", "Number {}", "No {}", "Section ID Number {}",
                                  "Course Number {}", "Item_No_{}", "Report {}", "Student Count {}"])
def test_p02_id_words_after_another_noun_stay_weighed_by_chance(stub: str) -> None:
    keys, _ = _sparse()
    assert not keys.match(stub.format(SPARSE_X)), stub


def test_p02_a_dated_document_after_a_label_word_is_no_id() -> None:
    dense = d34.IdKeys([NS([str(x) for x in range(1001, 1401)])])
    for name in ("IEP_Meeting_1015", "504 Plan 1015", "IEP Due 1015", "Section 504 Plan 1014", "IEP_2024-25_Caseload"):
        assert not dense.match(name), name
    assert dense.match("IEP_Meeting_1053")  # no date: the student's


def test_p01_p02_labelled_ids_in_files_headers_sheets_and_folders_fail(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = district_ids()
    roster(src, ids)
    a, b, c, d, e, f, g = (ids[k] for k in (11, 12, 13, 21, 22, 23, 24))
    write_csv(src / "iep" / f"{a}_IEP_Goals.csv", ["goal", "progress"], [[f"Goal {k}", 50] for k in range(4)])
    write_csv(src / "scores.csv", ["item", f"{b} Student", f"Student_No_{e}"], [[f"Q{k}", 70, 60] for k in range(6)])
    write_csv(src / "evals" / f"{c}_Evaluation.csv", ["area", "score"], [["Reading", 80], ["Math", 70]])
    write_csv(src / f"Student Number {d}.csv", ["goal", "progress"], [[f"Goal {k}", 50] for k in range(4)])
    write_book(src / "progress.xlsx", {f"Student No. {f}": (["goal", "baseline"], [[f"G{k}", k] for k in range(6)])})
    write_csv(src / f"Pupil Number {g}" / "goals.csv", ["goal", "progress"], [[f"Goal {k}", 50] for k in range(4)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    assert st.pop("roster.csv") != "FAIL", st
    assert set(st.values()) == {"FAIL"} and len(st) == 6, st
    blob = everything(out, rep, console=res.output)
    for x in (a, b, c, d, e, f, g):
        assert x not in blob, x


@pytest.mark.parametrize("make", ["score#", "_#", "#_score", "bare"])
def test_p03_wrapped_ids_of_a_dense_pivot_are_a_pivot(make: str) -> None:
    keys, picks = _pivot_ids()
    texts = ["Assignment"] + [make.replace("#", p) if make != "bare" else p for p in picks]
    assert set(d34.id_group_flags(keys, texts)) == set(range(1, 26)), make
    for numbered in ("Q#", "Item_#", "Room #", "Lexile_#"):  # numbered things: a list of items, rooms, levels
        assert d34.id_group_flags(keys, ["Student"] + [numbered.replace("#", p) for p in picks]) == {}, numbered


def test_p03_stata_and_suffix_pivots_fail(tmp_path: Path) -> None:
    pd = pytest.importorskip("pandas")
    src = tmp_path / "in"
    _, picks = _pivot_ids()
    roster(src, [str(x) for x in range(1001, 1401)])
    rng = random.Random(2)
    rows = [[f"HW{k}"] + [rng.randint(50, 100) for _ in picks] for k in range(1, 13)]
    pd.DataFrame(rows, columns=["assignment", *[f"_{p}" for p in picks]]).to_stata(
        src / "gradebook_wide.dta", write_index=False, version=118)
    write_csv(src / "gradebook_stub.csv", ["assignment", *[f"score{p}" for p in picks]], rows)
    write_csv(src / "gradebook_suffix.csv", ["Assignment", *[f"{p}_score" for p in picks]], rows)
    write_csv(src / "items.csv", ["Student", *[f"Q{p}" for p in picks]], [[f"S{k}"] + [1] * 25 for k in range(30)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    for f in ("gradebook_wide.dta", "gradebook_stub.csv", "gradebook_suffix.csv"):
        assert st[f] == "FAIL", (f, st)
    assert st["items.csv"] != "FAIL" and st["roster.csv"] != "FAIL", st
    blob = everything(out, rep, console=res.output)
    assert f"score{picks[0]}" not in blob and f"_{picks[1]}" not in blob and f"{picks[2]}_score" not in blob


def test_p04_a_withdrawn_student_and_per_school_rosters_leave_a_pivot(tmp_path: Path) -> None:
    keys, picks = _pivot_ids()
    two = d34.IdKeys([NS([str(x) for x in range(1001, 1201)]), NS([str(x) for x in range(1201, 1401)])])
    split = [p for p in picks if int(p) < 1201][:12] + [p for p in picks if int(p) >= 1201][:12]
    assert set(d34.pivot_flags(two, split)) == set(range(len(split)))  # the run's IDs together
    src = tmp_path / "in"
    roster(src, [str(x) for x in range(1001, 1201)], name="lincoln_es_roster.csv")
    roster(src, [str(x) for x in range(1201, 1401)], name="washington_es_roster.csv")
    rng = random.Random(4)
    write_csv(src / "district_reading_intervention.csv", ["Assignment", *split],
              [[f"HW{k}"] + [rng.randint(50, 100) for _ in split] for k in range(10)])
    write_csv(src / "gradebook_pivot.csv", ["Assignment", *picks[:20], "1412"],
              [[f"HW{k}"] + [rng.randint(50, 100) for _ in range(21)] for k in range(10)])
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert st["district_reading_intervention.csv"] == "FAIL" and st["gradebook_pivot.csv"] == "FAIL", st
    assert ",".join(split[:3]) not in everything(out, rep, console=res.output)


def test_p05_a_small_caseload_under_a_labelled_folder_or_workbook_fails(tmp_path: Path) -> None:
    dense = d34.IdKeys([NS([str(x) for x in range(1001, 1401)])])
    assert set(d34.id_group_flags(dense, ["1007", "1015", "1053"], label=True)) == {0, 1, 2}  # 1007: no date here
    assert set(d34.id_group_flags(dense, ["1053", "1081"], label=True)) == {0, 1}
    assert set(d34.id_group_flags(dense, ["1053", "1081"])) == {0, 1}  # D37: without the label too (a pair)
    assert d34.id_group_flags(dense, ["1053"]) == {}  # one number alone: chance without the label
    assert d34.id_group_flags(dense, ["1053"], label=True) == {0: ("1053",)}
    assert d34.id_group_flags(dense, ["1001", "1002", "1005", "1006", "1007"], label=True) == {}  # daily sheets
    assert d34.id_group_flags(dense, ["1001", "1002", "1003"], label=True) == {}  # a numbered list
    assert d34.id_group_flags(dense, ["1053", "2024"], label=True) == {0: ("1053",)}  # a year never
    for parent in ("iep", "IEPs", "students", "iep_progress", "Evaluations", "Westside_Campus_IEP_Progress"):
        assert d34.person_context(parent), parent
    for parent in ("Students_by_School", "rooms", "attendance_oct", "Gradebook"):
        assert not d34.person_context(parent), parent
    src = tmp_path / "in"
    roster(src, [str(x) for x in range(1001, 1401)])
    for p in ("1007", "1053", "1081"):
        write_csv(src / "iep" / f"{p}.csv", ["goal", "progress"], [[f"G{k}", k] for k in range(5)])
    write_book(src / "iep_progress.xlsx", {p: (["goal", "progress"], [[f"G{k}", k] for k in range(5)])
                                           for p in ("1127", "1195", "1239")})
    days = [f"10{d:02d}" for d in range(1, 32) if d % 7 not in (5, 6)]
    write_book(src / "student_attendance_oct.xlsx", {d: (["period", "present"], [[k, 380 + k] for k in range(1, 8)])
                                                     for d in days})
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert st["iep_progress.xlsx"] == "FAIL" and st["student_attendance_oct.xlsx"] != "FAIL", st
    assert sum(1 for k, v in st.items() if k.startswith("iep/") and v == "FAIL") == 3, st
    blob = everything(out, rep, console=res.output)
    assert "iep/1053" not in blob and "iep/1007" not in blob and "'1195'" not in blob
    assert openpyxl.load_workbook(out / "student_attendance_oct.xlsx").sheetnames == days


@pytest.mark.parametrize("parent", ["Lincoln_Middle_School_Gradebook", "Springfield_School_District",
                                    "Westside_Campus_IEP_Progress", "Algebra_I_Section_B_Grades", "Room_Schedules"])
def test_p06_a_place_the_parent_itself_names_says_nothing_of_its_numbers(parent: str) -> None:
    assert not d34.numbered_context(parent), parent


@pytest.mark.parametrize("parent", ["enrollment_by_zip", "Rooms", "by_campus", "Bus_Routes", "Campuses",
                                    "ZIP_Enrollment", "School_Codes", "enrollment_per_school", "Course_Sections"])
def test_p06_a_collection_or_a_code_system_says_what_its_numbers_are(parent: str) -> None:
    assert d34.numbered_context(parent), parent


def test_p06_school_campus_and_district_parents_no_longer_excuse_real_ids(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = district_ids()
    roster(src, ids)
    _, picks = _pivot_ids()
    rng = random.Random(3)
    sub = tmp_path / "dense"
    roster(sub, [str(x) for x in range(1001, 1401)])
    write_csv(sub / "Lincoln_Middle_School_Gradebook.csv", ["Assignment", *picks],
              [[f"HW{k}"] + [rng.randint(50, 100) for _ in picks] for k in range(12)])
    write_book(src / "Westside_Campus_IEP_Progress.xlsx",
               {ids[k]: (["goal", "baseline"], [[f"G{g}", g] for g in range(5)]) for k in (31, 32, 33)})
    for k in (41, 42, 43):
        write_csv(src / "Springfield_School_District" / f"{ids[k]}.csv", ["goal", "progress"], [["G", 50]] * 4)
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert st.pop("roster.csv") != "FAIL" and set(st.values()) == {"FAIL"} and len(st) == 4, st
    blob = everything(out, rep, console=res.output)
    for k in (31, 32, 33, 41, 42, 43):
        assert ids[k] not in blob, k
    res2, out2, rep2 = run(sub, tmp_path, name="dense")
    assert statuses(rep2)["Lincoln_Middle_School_Gradebook.csv"] == "FAIL", res2.output


def test_p07_a_labelled_id_that_reads_as_a_year_month_or_a_date_is_strong(tmp_path: Path) -> None:
    seq = d34.IdKeys([NS([str(x) for x in range(200001, 206001)])])  # a district numbering from 200000
    assert all(d34.date_like(x) for x in ("200105", "201907", "203112", "200203"))  # YYYYMM, a school year
    for name in ("Student_200105_IEP_Goals", "Student_201907", "IEP_203112", "IEP 200105"):
        assert seq.match(name), name
    # B1 (second hardening): a school year is never the label's ID (2002-03), and a date before a label is weighed by
    # how often a date of its layout is a real ID: every month of 2000-2059 is one in this district
    assert not seq.match("200203_IEP") and not seq.match("IEP_200203") and not seq.match("201907_IEP")
    assert not seq.match("Report_201907") and not seq.match("201907")  # a year-month without a label: no ID
    wide = d34.IdKeys([NS([str(x) for x in range(1001, 3001)])])
    assert not wide.match("Student 2023") and not wide.match("IEP_2024_Goals")  # a 4-digit year never is
    src = tmp_path / "in"
    eight = [f"{y}{s:04d}" for y in (2019, 2020, 2021, 2022, 2023) for s in range(1, 1201)]
    roster(src, eight)
    for x in ("20210915", "20221104"):
        write_csv(src / f"IEP_{x}.csv", ["goal", "progress"], [[f"Goal {g}", 50] for g in range(4)])
    write_csv(src / "IEP_20231215.csv", ["goal", "progress"], [["Goal", 50]] * 4)  # a date that is no real ID
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert st["IEP_20231215.csv"] != "FAIL" and sum(1 for v in st.values() if v == "FAIL") == 2, st
    blob = everything(out, rep, console=res.output)
    assert "20210915" not in blob and "20221104" not in blob


def test_p10_a_dense_pivot_leaves_statistics_and_dated_sheets_of_other_files_alone(tmp_path: Path) -> None:
    src = tmp_path / "in"
    _, picks = _pivot_ids()
    assert {"1007", "1015", "1333", "1375"} <= set(picks)
    roster(src, [str(x) for x in range(1001, 1401)])
    rng = random.Random(11)
    write_csv(src / "gradebook_pivot.csv", ["Assignment", *picks],
              [[f"HW{k}"] + [rng.randint(50, 100) for _ in picks] for k in range(1, 13)])
    days = [f"10{d:02d}" for d in range(1, 32) if d not in (5, 6, 12, 13, 19, 20, 26, 27)]
    write_book(src / "attendance_oct.xlsx",
               {d: (["period", "present", "absent"], [[p, 380 + p, 20 - p] for p in range(1, 8)]) for d in days})
    vals = [rng.randint(300, 1250) for _ in range(1331)] + [210, 1375]  # 1,333 rows; exact max 1375
    write_csv(src / "lexile.csv", ["test_window", "lexile"], [[rng.choice(["Fall", "Spring"]), v] for v in vals])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    assert "report this as a bug" not in res.output and "leak scan" not in res.output, res.output
    st = statuses(rep)
    assert st.pop("gradebook_pivot.csv") == "FAIL" and "FAIL" not in set(st.values()), st
    lexile = (rep / "files" / "lexile.csv.md").read_text(encoding="utf-8")
    assert "max 1375" in lexile and "1333 in -> 1333 out" in lexile and "[redacted" not in lexile
    docs = (out / "DATA_DICTIONARY.md").read_text(encoding="utf-8")
    assert "sheet_name='1007'" in docs and "[redacted" not in docs
    assert openpyxl.load_workbook(out / "attendance_oct.xlsx").sheetnames == days
    assert ",".join(picks[:3]) not in everything(out, rep, console=res.output)


def test_p10_an_id_only_flag_is_repeated_by_its_id_weighed_against_the_flagged_ids() -> None:
    keys, picks = _pivot_ids()
    screen = d34.NameScreen(keys)
    for k, p in enumerate(picks):
        screen.add(p, "column", frozenset({"id"}), f"column {k + 2} of 26", written=True, number=k + 2, id_tokens=(p,))
    for text in ("1375", "1007", "1375_notes", "Attendance_1053"):  # 25 of 9,000 four-digit numbers: chance
        assert screen.repeats(text) == set(), text
    assert screen.repeats("Student 1375") == {"1375"}  # a label: strong
    words, seqs, exact = screen.flags[0].terms()
    assert not words and not exact  # a real ID: the scanner's ID rules judge it, never as a header text
    assert screen.redact_text("max 1375, 1333 rows") == "max 1375, 1333 rows"
    assert screen.redact_text("file iep/1375.csv, sheet '1007'") == "file iep/[redacted-22].csv, sheet '[redacted-1]'"
    sparse = d34.NameScreen()
    sparse.add(f"Student {SPARSE_X}", "column", frozenset({"id"}), "column 2 of 2", written=True, number=2,
               id_tokens=(SPARSE_X,))
    assert sparse.repeats(f"{SPARSE_X}_Notes") == {f"Student {SPARSE_X}"}  # one of a few flagged IDs: strong
    assert sparse.repeats(f"Room {SPARSE_X}") == set() and sparse.repeats("Notes") == set()


# =========================================================================== hardening 2: the second attack (A1 .. A8)
PDF = b"%PDF-1.4\n%fake\n"


def _sparse_ids(n: int, start: int = 100) -> list[str]:
    """Sparse six-digit IDs of district_ids() that read as no time or short date (no per-number excuse)."""
    return [x for x in district_ids()[start:] if not d34.natural_number(x)][:n]


@pytest.mark.parametrize("stub", ["Progress_Report_{}", "Consent_Form_{}", "STAAR_Report_{}",
                                  "Lunch_Account_{}", "Case_{}_Notes"])
def test_a1_numbered_words_give_way_to_sibling_evidence(stub: str) -> None:
    keys, _ = _sparse()
    xs = _sparse_ids(5)
    assert not any(keys.match(stub.format(x)) for x in xs)  # each alone: a report, form, account or case number
    assert set(d34.id_group_flags(keys, [stub.format(x) for x in xs[:3]] + ["Notes"])) == {0, 1, 2}, stub
    assert d34.id_group_flags(keys, [stub.format(x) for x in xs[:2]]) == {}  # two: the word still explains them


def test_a1_numbered_things_that_are_chance_stay_clean() -> None:
    dense = d34.IdKeys([NS([str(x) for x in range(1001, 1401)])])
    rooms = ["1014", "1022", "1105", "1210", "1302", "1311", "1325", "1340", "1362", "1388"]
    assert d34.id_group_flags(dense, [f"Room_{r}_Inventory" for r in rooms]) == {}  # sequential IDs: chance
    dense5 = d34.IdKeys([NS([str(x) for x in range(20001, 60001)])])  # 40,000 five-digit IDs: 0.44 of them
    courses = ["23101", "23205", "31104", "41102", "52101", "33310", "45210", "27105"]
    assert d34.id_group_flags(dense5, [f"Course_{c}" for c in courses]) == {}  # 0.44 ** 8: chance at 1e-6
    keys, _ = _sparse()
    rng = random.Random(6)
    reports = [str(rng.randint(100000, 999999)) for _ in range(300)] + _sparse_ids(3, 700)  # 3 of 303 real IDs
    assert d34.id_group_flags(keys, [f"Report_{r}" for r in reports]) == {}


def test_a1_per_student_documents_after_numbered_words_fail_and_never_print(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = district_ids()
    roster(src, ids)
    xs = _sparse_ids(17, 900)
    (src / "progress_reports").mkdir(parents=True)
    for x in xs[:6]:
        (src / "progress_reports" / f"Progress_Report_{x}.pdf").write_bytes(PDF)
    for x in xs[6:11]:
        write_csv(src / "lunch" / f"Lunch_Account_{x}.csv", ["date", "item", "amount"],
                  [[f"2024-09-{d:02d}", "Lunch", 2.85] for d in range(3, 13)])
    for x in xs[11:17]:
        write_csv(src / "case_notes" / f"Case_{x}_Notes.csv", ["date", "contact_type", "minutes"],
                  [[f"2024-10-{d:02d}", "Phone", 15] for d in range(1, 6)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    assert st.pop("roster.csv") != "FAIL" and set(st.values()) == {"FAIL"} and len(st) == 11, st
    blob = everything(out, rep, console=res.output)
    for x in xs:
        assert x not in blob, x


def test_a2_a_zero_padded_id_is_strong_unless_the_run_uses_such_numbers_as_codes() -> None:
    ids = [f"{v:06d}" for v in random.Random(11).sample(range(1, 70000), 25000)]  # 0.25 of all padded numbers
    keys = d34.IdKeys([NS(ids)])
    x = next(i for i in ids if len(i.lstrip("0")) == 5 and not d34.natural_number(i))
    for name in (f"Goals_{x}", x, f"Case_{x}", f"Progress_{x}"):
        assert keys.match(name), name
    stamp = next(f"students_20241015_{i}" for i in ids if d34.natural_number(i))
    assert not keys.match(stamp)  # the time of a date-time stamp
    code = next(f"{c:06d}" for c in range(99213, 99999) if f"{c:06d}" not in set(ids))
    keys.note_run_names([(f"{code}_Bexar_ISD_TAPR", None, False)])  # a county-district number that is no real ID
    assert keys.padded_codes == {6} and not keys.match(f"Goals_{x}")
    # siblings that hold no such code keep the padding the IDs' (a folder of per-student goal exports)
    goals = [f"Goals_{i}" for i in ids if not d34.natural_number(i)][:3]
    assert set(d34.id_group_flags(keys, goals)) == {0, 1, 2}
    assert d34.id_group_flags(keys, goals[:1] + [f"{code}_Bexar_ISD_TAPR"]) == {}


def test_a2_per_student_padded_files_fail_beside_a_folder_of_county_district_numbers(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = [f"{v:06d}" for v in random.Random(11).sample(range(1, 70000), 25000)]
    known = set(ids)
    roster(src, ids)
    plain = [i for i in ids if not d34.natural_number(i) and len(i.lstrip("0")) == 5][:8]
    for x in plain[:3]:
        write_csv(src / "goal_exports" / f"Goals_{x}.csv", ["goal", "progress"], [[f"Goal {k}", 50] for k in range(4)])
    write_csv(src / "caseload" / plain[3] / "goals.csv", ["goal", "progress"], [[f"Goal {k}", 50] for k in range(4)])
    write_csv(src / "progress" / f"Progress_{plain[4]}.csv", ["goal", "baseline"], [[f"Goal {k}", k] for k in range(5)])
    (src / "iep_pdfs").mkdir()
    (src / "iep_pdfs" / f"Goals_{plain[5]}.pdf").write_bytes(PDF)
    code = next(f"{c:06d}" for c in range(99213, 99999) if f"{c:06d}" not in known)  # no real ID
    tapr = [f"{c}_{n}_ISD_TAPR.csv" for c, n in zip([code, *plain[6:8]], ("Bexar", "Northside", "Judson"))]
    for f in tapr:
        write_csv(src / "tapr" / f, ["indicator", "value"], [[f"I{k}", 50 + k] for k in range(10)])
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    for f in tapr:
        assert st[f"tapr/{f}"] != "FAIL", (f, st)
    assert sum(1 for v in st.values() if v == "FAIL") == 5, st
    blob = everything(out, rep, console=res.output)
    for x in plain[:6]:
        assert x not in blob, x


def test_a3_a_parent_that_names_codes_never_hides_a_pivot_of_sparse_ids() -> None:
    keys, _ = _sparse()
    cls = sorted(_sparse_ids(26, 900))
    assert set(d34.pivot_flags(keys, cls, context=True)) == set(range(26))
    assert set(d34.id_group_flags(keys, cls[:3], context=d34.numbered_context("SOC_Projects"))) == {0, 1, 2}
    dense = d34.IdKeys([NS([str(x) for x in range(10001, 38001)])])  # ZIP codes inside sequential IDs: chance
    zips = ["30303", "30305", "30308", "30310", "30311", "30312", "30313", "30314", "30315", "30316", "30317",
            "30318", "30319", "30324", "30326", "30327", "30328", "30329", "30331", "30336", "30337"]
    assert d34.pivot_flags(dense, zips, context=True) == {}
    # D37 (B4 dropped): without the parent, ZIP codes that are all real IDs are a pivot, whatever their endings;
    # numbers spread over every ending too, and both are chance under the parent
    assert set(d34.pivot_flags(dense, zips)) == set(range(len(zips)))
    spread = sorted(str(x) for x in random.Random(3).sample(range(30300, 31300), 21))
    assert d34.pivot_flags(dense, spread, context=True) == {}
    assert set(d34.pivot_flags(dense, spread)) == set(range(len(spread)))
    assert 0.9 < dense.local_share("30303") <= 1.0 and keys.local_share(cls[0]) < 0.02


def test_a3_social_studies_files_named_soc_fail_like_any_pivot(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = district_ids()
    roster(src, ids)
    cls = sorted(_sparse_ids(26, 900))
    rng = random.Random(3)
    write_csv(src / "Soc_Studies_Gr8_Gradebook.csv", ["Assignment", *cls],
              [[f"HW{j}"] + [rng.randint(60, 100) for _ in cls] for j in range(1, 9)])
    projects = _sparse_ids(3, 950)
    for x in projects:
        write_csv(src / "SOC_Projects" / f"{x}.csv", ["rubric_row", "points"], [["Thesis", 4], ["Evidence", 3]])
    res, out, rep = run(src, tmp_path, "--header-row", "1")
    st = statuses(rep)
    assert st.pop("roster.csv") != "FAIL" and set(st.values()) == {"FAIL"} and len(st) == 4, st
    blob = everything(out, rep, console=res.output)
    for x in cls[:3] + projects:
        assert x not in blob, x


def test_a4_a_flagged_name_is_repeated_in_either_order() -> None:
    screen = d34.NameScreen()
    screen.add("Patel, Mark", "folder", frozenset({"name"}), "folder 1 of 2", written=True)
    screen.add("Garcia Lopez, Maria", "folder", frozenset({"name"}), "folder 2 of 2", written=True)
    assert screen.repeats("Mark_Patel_Walkthrough") == {"Patel, Mark"}
    assert screen.repeats("MarkPatel-Observation") == {"Patel, Mark"}
    assert screen.repeats("Maria_Garcia_Lopez_IEP") == {"Garcia Lopez, Maria"}
    for text in ("Mark_Twain_Unit", "Patel_Brothers", "Mark_Smith_Patel", "Walkthrough"):
        assert screen.repeats(text) == set(), text
    words, seqs, exact = screen.flags[0].terms()
    assert ("patel", "mark") in seqs and ("mark", "patel") in seqs


def test_a4_given_first_files_of_surname_first_folders_fail(tmp_path: Path) -> None:
    src = tmp_path / "in"
    teachers = [("Patel", "Mark"), ("Smith", "Grace"), ("Nguyen", "Rose"), ("Davis", "Summer")]
    for last, first in teachers:
        write_csv(src / "Teachers" / f"{last}, {first}" / "gradebook.csv", ["assignment", "avg"],
                  [[f"HW{k}", 80 + k] for k in range(1, 9)])
        write_csv(src / "Observations" / f"{first}_{last}_Walkthrough.csv", ["date", "domain", "score"],
                  [[f"2024-10-{d:02d}", "D2", 3] for d in range(1, 6)])
    write_csv(src / "Observations" / "district_rubric.csv", ["domain", "weight"], [["D1", 1], ["D2", 2], ["D3", 1]])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3 and "report this as a bug" not in res.output, res.output
    st = statuses(rep)
    assert st.pop("Observations/district_rubric.csv") != "FAIL" and set(st.values()) == {"FAIL"}, st
    assert len(st) == 8, st
    blob = everything(out, rep, console=res.output)
    for last, first in teachers:
        assert f"{first}_{last}" not in blob and f"{last}, {first}" not in blob, (first, last)


def test_a5_files_named_like_the_input_folder_get_the_d34_verdict(tmp_path: Path) -> None:
    root = tmp_path / "in" / "Davis, Summer"
    for stem in ("Davis_Summer_Reading_Log", "Summer_Davis_Reading_Log"):
        write_csv(root / f"{stem}.csv", ["week", "minutes", "books"], [[k, 100 + 5 * k, k % 3] for k in range(1, 13)])
    write_csv(root / "period3_grades.csv", ["assignment", "avg"], [[f"HW{k}", 80 + k] for k in range(1, 9)])
    res, out, rep = run(root, tmp_path)
    assert res.exit_code == 3, res.output
    assert "report this as a bug" not in res.output and "leak scan" not in res.output, res.output
    assert "D34:" in res.output and "--keep-name TEXT" in res.output
    st = statuses(rep)
    assert st.pop("period3_grades.csv") != "FAIL" and set(st.values()) == {"FAIL"} and len(st) == 2, st
    blob = everything(out, rep, console=res.output)
    for s in ("Davis_Summer", "Summer_Davis", "Davis, Summer"):
        assert s not in blob, s


def test_a8_per_student_documents_one_per_folder_count_together() -> None:
    keys, _ = _sparse()
    xs = _sparse_ids(8, 600)
    docs = [f"Transcript_{xs[0]}", f"Child_{xs[1]}", f"Scholar_{xs[2]}_Goals", f"BIP_{xs[3]}", f"FBA_{xs[4]}",
            f"Cumulative_Folder_{xs[5]}", f"LPAC_{xs[6]}", f"Immunization_Record_{xs[7]}"]
    assert not any(keys.match(d) for d in docs)  # each alone: one in 150 is a real ID
    keys.note_run_names([(d, None, False) for d in docs] + [("roster", None, False), ("student_id", None, False)])
    assert keys.evident == frozenset(xs) and all(keys.match(d) for d in docs)
    lone = d34.IdKeys([NS(district_ids())])
    lone.note_run_names([(docs[0], None, False), ("Q17_Responses", None, False)])
    assert not lone.evident and not lone.match(docs[0])  # one document alone stays chance
    rng = random.Random(8)
    many = [(f"Survey_{rng.randint(100000, 999999)}", None, False) for _ in range(500)]
    noisy = d34.IdKeys([NS(district_ids())])
    noisy.note_run_names(many + [(d, None, False) for d in docs[:3]])
    assert not noisy.evident  # three real IDs among 500 survey numbers: chance
    stamps = d34.IdKeys([NS(district_ids())])
    natural = [x for x in district_ids() if d34.natural_number(x)][:4]
    stamps.note_run_names([(f"export_20241015_{x}", None, False) for x in natural])
    assert not stamps.evident  # times of date-time stamps: no evidence


def test_a8_lone_documents_in_their_own_folders_never_print_their_ids(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = district_ids()
    roster(src, ids)
    xs = _sparse_ids(8, 600)
    lone = {"transcripts": f"Transcript_{xs[0]}.pdf", "early_childhood": f"Child_{xs[1]}.pdf",
            "charter": f"Scholar_{xs[2]}_Goals.pdf", "behavior": f"BIP_{xs[3]}.pdf", "fba": f"FBA_{xs[4]}.csv",
            "records": f"Cumulative_Folder_{xs[5]}.pdf", "ell": f"LPAC_{xs[6]}.pdf",
            "health": f"Immunization_Record_{xs[7]}.csv"}
    for folder, name in lone.items():
        if name.endswith(".csv"):
            write_csv(src / folder / name, ["date", "note"], [[f"2024-10-{d:02d}", "seen"] for d in range(1, 6)])
        else:
            (src / folder).mkdir(parents=True)
            (src / folder / name).write_bytes(PDF)
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert st.pop("roster.csv") != "FAIL" and set(st.values()) == {"FAIL"} and len(st) == 2, st
    blob = everything(out, rep, console=res.output)
    for x in xs:
        assert x not in blob, x


def test_h2_unlabelled_ids_that_read_as_dates_among_flagged_siblings_fail(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = [str(x) for x in range(200001, 206001)]  # a district numbering from 200000: some IDs read as YYYYMM
    roster(src, ids)
    picks = ["205903", "205204", "204209", "204395", "202563", "201777", "203381", "200498", "201864", "202247"]
    assert sum(1 for p in picks if d34.date_like(p)) == 3
    for p in picks:
        write_csv(src / "goal_exports" / f"Goals_{p}.csv", ["goal", "progress"], [[f"Goal {k}", 50] for k in range(4)])
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert sum(1 for v in st.values() if v == "FAIL") == 10, st
    blob = everything(out, rep, console=res.output)
    for p in picks:
        assert p not in blob, p


def test_clock_times_on_five_minute_marks_are_no_pivot_and_no_label() -> None:
    charter = d34.IdKeys([NS([str(x) for x in range(1001, 1601)])])  # a 600-student charter numbering 1001-1600
    assert d34.id_group_flags(charter, ["date", "1235", "1305", "1330", "1350", "1415", "1440"]) == {}  # lunch periods
    assert d34.id_group_flags(charter, ["IEP Meeting 1330", "IEP Meeting 1415", "IEP Meeting 1505"]) == {}
    assert charter.match("IEP Meeting 1053") and charter.match("IEP_1330")  # no clock time; a label right before
    times = ["date", "1235", "1305", "1330", "1350", "1415", "1447"]  # 14:47 is no five-minute mark: real IDs
    assert set(d34.id_group_flags(charter, times)) == set(range(1, 7))


# =========================================================================== hardening 3: the third attack (B1 .. B8)
def _dense(lo: int, hi: int) -> d34.IdKeys:
    return d34.IdKeys([NS([str(x) for x in range(lo, hi + 1)])])


@pytest.mark.parametrize("name", ["2425_Student_Roster", "IEP_Caseload_2425", "504_Plans_2425",
                                  "2324_IEP_Goals_Summary", "IEP_2425", "Student_2425", "1015_IEP_Notes",
                                  "1330_IEP_Meeting", "2024_1015_IEP", "IEP_Meeting_1330", "IEP_Due_1015",
                                  "Student_Council_1015"])
def test_b1_school_years_dates_and_times_beside_a_label_are_no_ids_in_a_dense_range(name: str) -> None:
    dense = _dense(1001, 4500)  # 3,500 student numbers: every school-year code 1011 ... 4445 is one
    assert not dense.match(name), name
    assert "id" not in d34.name_kinds(name, dense), name


def test_b1_labels_still_make_ids_that_read_as_no_date_strong() -> None:
    dense = _dense(1001, 4500)
    for name in ("Student_1053", "IEP_Meeting_1053", "1053_IEP_Notes", "IEP_1015", "IEP_1330"):
        assert dense.match(name), name  # no date, or a label right before it (K-01)
    seq = d34.IdKeys([NS([str(x) for x in range(200001, 206001)])])
    assert seq.match("Student_200105") and not seq.match("IEP_202425") and not seq.match("202425_Student_Roster")
    eight = d34.IdKeys([NS([f"{y}{s:04d}" for y in (2021, 2022, 2023, 2024) for s in range(1, 2001)])])
    assert eight.date_share("20241015") > 0.4  # IDs numbered by entry year: most dates of a year are real IDs
    for name in ("20241015_Student_Attendance", "Student_Attendance_20241015", "IEP_Meeting_20241015"):
        assert not eight.match(name), name
    assert eight.match("IEP_20241015")  # right after the label (P-07)
    keys, _ = _sparse()
    assert d34.name_date("160816") and keys.match("160816_IEP_Goals")  # one in 150 dates is a real ID: strong
    assert keys.date_share("160816") < 0.01


def test_b1_siblings_that_all_read_as_dates_are_dated_files() -> None:
    keys, ids = _sparse()
    dated = "110523"  # 2023-11-05 (MMDDYY) is one of the 6,000 IDs
    assert dated in ids and keys.match(f"{dated}_Student_Attendance")
    days = ["110521", "110522", "110123", "110223", dated, "110623", "110723"]
    names = [f"{d}_Student_Attendance" for d in days]
    assert d34.id_group_flags(keys, names) == {}  # a folder of daily exports
    x = _sparse_ids(1)[0]  # a sibling that is no date: they are per-student files
    assert set(d34.id_group_flags(keys, [f"{dated}_Student_Attendance", f"{x}_Student_Attendance"])) == {0, 1}
    dense = _dense(1001, 4500)
    assert set(d34.id_group_flags(dense, ["IEP_2425", "IEP_1053"])) == {0, 1}  # a school year beside a student
    assert d34.id_group_flags(dense, ["IEP_Caseload_2324", "IEP_Caseload_2425"]) == {}


def test_b1_school_year_and_daily_files_beside_student_numbers_are_written(tmp_path: Path) -> None:
    src = tmp_path / "in"
    rng = random.Random(15)
    roster(src, [str(x) for x in range(1001, 4501)])
    write_csv(src / "2425_Student_Roster.csv", ["Grade", "Count"], [[g, rng.randint(20, 30)] for g in range(6, 9)])
    write_book(src / "sped" / "IEP_Caseload_2425.xlsx", {"Caseload": (["Disability", "Count"],
                                                                     [["SLD", 4], ["OHI", 3], ["AU", 2]])})
    write_csv(src / "sped" / "504_Plans_2425.csv", ["Accommodation", "Count"], [["Extra time", 12], ["Oral", 5]])
    write_csv(src / "sped" / "2324_IEP_Goals_Summary.csv", ["Area", "Goals"], [["Reading", 30], ["Math", 22]])
    for d in ("1101", "1104", "1105", "1106", "1107"):  # MMDD daily files beside a label: dates
        write_csv(src / "attendance" / f"{d}_Student_Attendance.csv", ["Period", "Present"],
                  [[p, rng.randint(300, 400)] for p in range(1, 8)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert "FAIL" not in set(statuses(rep).values())
    assert (out / "2425_Student_Roster.csv").exists() and (out / "sped" / "2324_IEP_Goals_Summary.csv").exists()


def test_b1_a_folder_of_daily_exports_is_written_and_a_lone_labelled_id_still_fails(tmp_path: Path) -> None:
    src = tmp_path / "in"
    roster(src, district_ids())
    rng = random.Random(2)
    for d in ("110123", "110223", "110323", "110423", "110523", "110623", "110723"):  # 110523 is a real ID
        write_csv(src / "attendance" / f"{d}_Student_Attendance.csv", ["Period", "Present", "Absent"],
                  [[p, rng.randint(300, 400), rng.randint(5, 30)] for p in range(1, 8)])
    write_csv(src / "iep" / "160816_IEP_Goals.csv", ["goal", "progress"], [[f"Goal {k}", 50] for k in range(4)])
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert sum(1 for k, v in st.items() if k.startswith("attendance/") and v != "FAIL") == 7, st
    assert sum(1 for v in st.values() if v == "FAIL") == 1, st
    assert "160816" not in everything(out, rep, console=res.output)


@pytest.mark.parametrize("stub", ["Student_Schedule_{}", "Student_Report_Card_{}", "Student_Info_{}", "Student_File_{}",
                                  "Student_Transcript_{}", "Student_Progress_Report_{}", "IEP_Progress_Report_{}",
                                  "IEP_Goals_Report_{}", "IEP_Annual_Review_Notes_{}", "IEP_2024-25_{}",
                                  "Evaluation_SY2425_{}", "IEP_Gr3_{}", "IEP_Amendment_2_{}", "STAAR_Student_Report_{}",
                                  "StudentReportCard_{}", "{}_Report_Card_Student"])
def test_b2_a_label_elsewhere_in_the_name_holds_a_real_id_at_district_scale(stub: str) -> None:
    keys, _ = _sparse()  # 6,000 of 900,000 six-digit numbers: 1 in 150, at most 1 in 100 with a label in the name
    assert keys.match(stub.format(SPARSE_X)), stub
    big = d34.IdKeys([NS([str(x) for x in random.Random(2).sample(range(100000, 1000000), 50000)] + [SPARSE_X])])
    assert not big.match(stub.format(SPARSE_X)), stub  # 1 in 18: the residual of a lone document


@pytest.mark.parametrize("stub", ["Student Count {}", "Course_{}_Student_Roster", "Room_{}_Student", "Item {} IEP",
                                  "Students_Report_{}", "Section {} Students", "Goals_{}", "Report_{}"])
def test_b2_a_numbered_thing_or_a_plural_keeps_its_number(stub: str) -> None:
    keys, _ = _sparse()
    assert not keys.match(stub.format(SPARSE_X)), stub


def test_b2_ids_that_read_as_a_time_with_seconds_or_no_calendar_day_stay_the_labels() -> None:
    keys, ids = _sparse()
    for x in ("150631", "175954"):  # 15:06:31 (and June 31), 17:59:54: no date a person writes beside a label
        assert x in ids and d34.natural_number(x) and not d34.name_date(x)
        for name in (f"IEP_Meeting_{x}", f"504_Plan_{x}", f"{x}_IEP", f"Student_Schedule_{x}"):
            assert keys.match(name), name
    assert d34.name_date("101524") and d34.name_date("241015") and d34.name_date("15062024")


def test_b2_labelled_single_documents_fail_and_never_print_their_ids(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = district_ids()
    roster(src, ids)
    plain = [x for x in ids if not d34.natural_number(x)]
    a, b = plain[0], plain[1]
    write_csv(src / "counselor" / f"Student_Schedule_{a}.csv", ["Period", "Course"],
              [[p, "ENG 9"] for p in range(1, 8)])
    (src / "report_cards").mkdir(parents=True)
    (src / "report_cards" / f"Student_Report_Card_{b}.pdf").write_bytes(PDF)
    write_csv(src / "sped" / "IEP_Meeting_150631.csv", ["Goal", "Target"], [[f"G{k}", 50 + k] for k in range(5)])
    (src / "sped" / "504_Plan_175954.pdf").write_bytes(PDF)
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3, res.output
    st = statuses(rep)
    assert st.pop("roster.csv") != "FAIL" and set(st.values()) == {"FAIL"} and len(st) == 2, st
    blob = everything(out, rep, console=res.output)
    for x in (a, b, "150631", "175954"):
        assert x not in blob, x


def test_b3_a_flagged_name_with_middle_names_or_initials_is_repeated_as_given_name_and_surname() -> None:
    screen = d34.NameScreen()
    for text in ("Davis, Summer Joy", "Mark A. Patel", "Nguyen, Rose Thi", "Garcia Lopez, Maria"):
        screen.add(text, "folder", frozenset({"name"}), "folder 1 of 4", written=True)
    assert screen.flags[1].words == ("mark", "patel")
    for text, flagged in (("Summer_Davis_Walkthrough", "Davis, Summer Joy"), ("Davis_Summer_IEP", "Davis, Summer Joy"),
                          ("Mark_Patel_Walkthrough", "Mark A. Patel"), ("PatelMark", "Mark A. Patel"),
                          ("Rose_Nguyen_Walkthrough", "Nguyen, Rose Thi"), ("Maria_Garcia_IEP", "Garcia Lopez, Maria")):
        assert screen.repeats(text) == {flagged}, text
    for text in ("Summer_Joy_Recital", "Mark_Twain_Unit", "Rose_Thi", "Joy_Davis_Award_Night_Patel"):
        assert screen.repeats(text) == set(), text
    plain = d34.NameScreen()
    plain.add("Mr. A. Patel", "folder", frozenset({"name"}), "folder 1 of 1", written=True)
    assert plain.flags[0].words == ("patel",)  # a title is no given name


def test_b3_walkthrough_files_of_teachers_with_middle_names_fail(tmp_path: Path) -> None:
    src = tmp_path / "in"
    rows = [[f"HW{k}", 80 + k] for k in range(1, 9)]
    for folder in ("Davis, Summer Joy", "Mark A. Patel", "Nguyen, Rose Thi"):
        write_csv(src / "Teachers" / folder / "gradebook.csv", ["assignment", "avg"], rows)
    for f in ("Summer_Davis_Walkthrough.csv", "Mark_Patel_Walkthrough.csv", "Rose_Nguyen_Walkthrough.csv"):
        write_csv(src / "Observations" / f, ["domain", "score"], [["D1", 3], ["D2", 4], ["D3", 2]])
    write_csv(src / "Observations" / "district_rubric.csv", ["domain", "weight"], [["D1", 1], ["D2", 2], ["D3", 1]])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 3 and "report this as a bug" not in res.output, res.output
    st = statuses(rep)
    assert st.pop("Observations/district_rubric.csv") != "FAIL" and set(st.values()) == {"FAIL"} and len(st) == 6, st
    blob = everything(out, rep, console=res.output)
    for s in ("Summer_Davis", "Mark_Patel", "Rose_Nguyen", "Summer Joy", "Mark A. Patel"):
        assert s not in blob, s


ROOMS = ["1104", "1107", "1112", "1203", "1208", "1215", "1222", "2104", "2110", "2118", "2203", "2210"]
COURSES = ["1101", "1102", "1105", "1121", "1201", "1203", "1221", "1301", "1310", "1402", "1501", "1503", "1610",
           "1702", "1801", "2101", "2203", "2301"]


def test_b4_dropped_rooms_and_courses_that_are_all_real_ids_are_a_pivot() -> None:
    students = _dense(1001, 2500)  # 1,500 student numbers: one in six 4-digit numbers
    # D37: B4 (numbers ending too low for a sample of the IDs) no longer excuses: they are real IDs all the same
    assert set(d34.pivot_flags(students, ROOMS)) == set(range(12))
    assert set(d34.pivot_flags(students, COURSES)) == set(range(18))
    assert set(d34.id_group_flags(students, ["Grade", *COURSES])) == set(range(1, 19))
    staff = _dense(1001, 1300)
    assert d34.pivot_flags(staff, ["1104", "1107", "1112", "1203", "1208", "1215"]) == {}  # all read as MMDD dates
    assert set(d34.pivot_flags(staff, ["1104", "1107", "1112", "1203", "1208", "1299"])) == set(range(6))
    rng = random.Random(4)
    for _ in range(20):  # a class of 12 students is a sample of the IDs
        cls = sorted(str(x) for x in rng.sample(range(1001, 2501), 12))
        assert set(d34.pivot_flags(students, cls)) == set(range(12)), cls
    # what still excuses numbered things: one constant step, a numbered word before them or a numbered parent
    assert d34.pivot_flags(students, ["1100", "1200", "1300", "1400"]) == {}
    assert d34.pivot_flags(students, [f"Room {r}" for r in ROOMS]) == {}
    assert d34.pivot_flags(students, ROOMS, context=d34.numbered_context("Rooms")) == {}


def test_b4_dropped_room_rosters_and_course_tallies_fail_unless_named_or_kept(tmp_path: Path) -> None:
    src = tmp_path / "in"
    ids = [str(x) for x in range(1001, 2501)]
    roster(src, ids)
    rng = random.Random(5)
    pool = ids[:]
    rng.shuffle(pool)
    sheets = {r: (["Student_ID", "Grade", "Subject"], [[s, "7", "Reading"] for s in pool[k * 20:(k + 1) * 20]])
              for k, r in enumerate(ROOMS)}
    write_book(src / "STAAR_Spring2025_Rosters.xlsx", sheets)
    write_book(src / "STAAR_Spring2025_Rosters_by_Room.xlsx", sheets)  # the workbook says what the numbers are
    tally = [[g] + [rng.randint(0, 60) for _ in COURSES] for g in ("9", "10", "11", "12")]
    write_csv(src / "Course_Request_Tally.csv", ["Grade", *COURSES], tally)
    write_csv(src / "Course_Requests_by_Course.csv", ["Grade", *COURSES], tally)
    res, out, rep = run(src, tmp_path)
    st = statuses(rep)
    assert st["STAAR_Spring2025_Rosters.xlsx"] == "FAIL" and st["Course_Request_Tally.csv"] == "FAIL", st
    assert st["STAAR_Spring2025_Rosters_by_Room.xlsx"] != "FAIL" and st["Course_Requests_by_Course.csv"] != "FAIL", st
    assert openpyxl.load_workbook(out / "STAAR_Spring2025_Rosters_by_Room.xlsx").sheetnames == ROOMS
    # --keep-name keeps a false alarm, one per flagged name
    kept = [a for x in sorted(set(ROOMS) | set(COURSES)) for a in ("--keep-name", x)]
    res2, out2, rep2 = run(src, tmp_path, *kept, name="kept")
    assert res2.exit_code == 0, res2.output
    assert "FAIL" not in set(statuses(rep2).values())
    assert openpyxl.load_workbook(out2 / "STAAR_Spring2025_Rosters.xlsx").sheetnames == ROOMS
    header = (out2 / "Course_Request_Tally.csv").read_text(encoding="utf-8").splitlines()[0]
    assert header == ",".join(["Grade", *COURSES])


def test_b5_a_pivot_whose_headers_carry_initials_is_a_pivot(tmp_path: Path) -> None:
    students = _dense(1001, 2500)
    rng = random.Random(12)
    picks = sorted(rng.sample([str(x) for x in range(1001, 2501)], 24))
    heads = [f"{chr(65 + rng.randint(0, 25))}{chr(65 + rng.randint(0, 25))} ({p})" for p in picks]
    assert set(d34.pivot_flags(students, ["Assignment", *heads])) == set(range(1, 25))
    assert set(d34.pivot_flags(students, ["Assignment", *heads[2:4]])) == {1, 2}  # D37: a pair is enough
    assert d34.pivot_flags(students, ["Assignment", *heads[:2]]) == {}  # 1004 and 1023 read as MMDD dates
    # D37: B4 no longer excuses course numbers that are all real student numbers; a short subject code reads like
    # initials (a pair is a set), whole words like surnames (three real IDs are a set); a numbered word excuses
    assert set(d34.pivot_flags(students, ["Name", "ELA (1101)", "SS (2110)"])) == {1, 2}
    assert d34.pivot_flags(students, ["Name", "Algebra (1101)", "Biology (2104)"]) == {}
    assert set(d34.pivot_flags(students, ["Name", "Algebra (1101)", "Biology (2104)", "Chemistry (2110)"])) == {1, 2, 3}
    assert d34.pivot_flags(students, ["Name", "Course 1101", "Course 1203", "Course 2104", "Course 2110"]) == {}
    src = tmp_path / "in"
    roster(src, [str(x) for x in range(1001, 2501)])
    write_csv(src / "Gradebook_P2.csv", ["Assignment", *heads], [[f"HW{k}"] + [80] * 24 for k in range(1, 9)])
    res, out, rep = run(src, tmp_path)
    assert statuses(rep)["Gradebook_P2.csv"] == "FAIL", res.output
    blob = everything(out, rep, console=res.output)
    for p in picks[:5]:
        assert f"({p})" not in blob, p


@pytest.mark.parametrize("name", ["Bus_Route_214_12_2024", "Room_305_09_2024_Attendance", "Section 101 01 2025 Roster",
                                  "Course_345_10_2024_Grades", "Lunch_Count_105_08_2025", "Rt_214_10_2023",
                                  "PE_101_12_2024", "Class 214 12 2024", "Roster 2024-25 512345",
                                  "Grades_2024-25_612345", "Attendance_2024-25_334455", "ELA_2025_26_101_Roster",
                                  "Roster_512345_2024-25"])
def test_b8_a_number_a_month_and_a_year_or_a_school_year_and_a_number_are_no_ssn_or_phone(name: str) -> None:
    assert not ({"ssn", "phone"} & d34.name_kinds(name)), name


@pytest.mark.parametrize("name, kind", [("SSN_214_12_2024", "ssn"), ("SS# 214 12 2024", "ssn"), ("214-12-2024", "ssn"),
                                        ("Room_214-12-2024", "ssn"), ("123_45_6789", "ssn"), ("Bus 214 13 2024", "ssn"),
                                        ("Route_214_12_2224", "ssn"), ("Parent 2024 25 512345", "phone"),
                                        ("Roster_2024_512_555_0199", "phone"), ("Cell_2024-25_512345", "phone")])
def test_b8_ssns_and_phones_stay_shapes(name: str, kind: str) -> None:
    assert kind in d34.name_kinds(name), name


def test_b8_files_named_by_route_room_or_section_month_and_year_are_written(tmp_path: Path) -> None:
    src = tmp_path / "in"
    rng = random.Random(1)
    names = ["Bus_Route_214_12_2024.csv", "Room_305_09_2024_Attendance.csv", "Section 101 01 2025 Roster.csv",
             "Grades_2024-25_612345.csv"]
    for name in names:
        write_csv(src / name, ["Day", "Count"], [[d, rng.randint(10, 40)] for d in range(1, 21)])
    res, out, rep = run(src, tmp_path)
    assert res.exit_code == 0, res.output
    assert sorted(p.name for p in out.glob("*.csv")) == sorted(names)
