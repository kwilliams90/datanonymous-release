"""Excel reading: .xlsx/.xlsm (openpyxl, read-only) and .xls (xlrd) -> raw tables (AREA E helper).

Each worksheet is streamed row by row (openpyxl ``read_only=True,
data_only=True``: cached values; formulas, layout and extras come from
:mod:`.excel_book`) into a :class:`SheetScan`, which decides:

* **header row** (design 6): within the first 30 rows, the first row that is
  >= 80% as wide as the widest row, all text and unique, with a wide row
  (>= 50% of the widest) among its next 5 non-blank rows (the first data row
  may be sparse) that shares its columns (a title cell in column A is not the
  header of a table in column B). Widths count only the main block of
  columns (adjacent occupied columns holding the most cells), so a note two
  columns beside the table or a title out to the side does not count. A
  one-cell candidate can be a title above a one-column table
  (``_title_above_single_column``: titles above numbers/dates/IDs, dated
  title lines, a title block then a short label). Fallbacks, in order: the
  same test with the most common row width as the reference; the first
  header-like row without the "wide row after it" condition; a row of text
  labels with duplicates (``date, date``) and a wide row after it; the first
  row >= 80% as wide as the widest with a wide row after it (number headers
  such as years); the first non-blank row. ``--header-row N`` and ``--range
  [Sheet!]A5:H200`` override it. Rows above the header are title/notes rows
  (kept in place by the writer with placeholder text).
* **footer rows**: a trailing block (at most 20 non-blank rows, blank rows
  inside it are layout) whose first row is a total row or a note row. A total
  row has an aggregate formula (SUM/COUNT/AVERAGE/SUBTOTAL ...) over a range
  above it (a per-row ``=SUM(B5:D5)`` in row 5 is data), or a first cell such
  as "Total" / "Grand Total" / "Subtotal". A row whose first cell is only a
  summary word ("Average", "Max", "Count" ... which can also be a data value,
  e.g. the first name Max) or a single plain text cell (a note; not a
  data-shaped value) starts the footer only after a blank row (or inside a
  block that a total row started); when every row is total-like ("Total
  enrolled", "Total staff" summary rows) only a blank row separates a footer.
  In a one-column table a text line is a note only when the column is mostly
  data-like (numbers, dates, zero-padded IDs ...) or the line has >= 3 words
  below short labels. A run of 20 total-like rows that continues above the
  block is data (row formulas), not a footer. Rows of the block are not
  data; the writer recomputes SUM/COUNT/AVERAGE from the synthetic rows and
  leaves other static totals blank. Rows whose cells are only formulas with
  an empty cached value (``=IF(B9="","",...)`` filled down past the data)
  are blank; trailing blank rows are not data (pandas drops them too).
* **columns**: from the first to the last column holding a header or data
  cell; interior empty columns are kept (header ``''``). Headerless blocks of
  cells separated from the table by an empty column (side notes) are not
  columns: they are written in place as placeholders (:func:`side_columns`).
* **notes**: a single-column sheet whose first row reads like a sentence, or
  a short block (<= 5 rows) of multi-word lines (:func:`notes_like`), is notes,
  not a table: written back as placeholders.

Cell values become canonical strings (``RawColumn.values``; an empty cell is
``''``, a blank as in the text formats: Excel and pandas do not distinguish an
empty cell from an empty string, and the foundation's missing-value paths
expect strings): text as stored; numbers in
shortest positional form (integer-valued numbers without ``.0``, as pandas
reads them); booleans ``TRUE``/``FALSE``; error cells as their token
(``#N/A``); date cells as ISO 8601 per column: ``YYYY-MM-DD`` (date formats
with no time of day anywhere in the column), ``YYYY-MM-DD HH:MM:SS``
(datetimes, rounded to the second), ``HH:MM:SS`` (time-only cells) and
``[-]H:MM:SS`` (``[h]:mm`` durations, hours may exceed 24). Per-column cell
storage goes to ``RawColumn.native`` (PRIVATE): type counts (s/n/d/b/e), the
main type, the most common number format per type, the date kind (also as
AREA A's ``temporal`` description when date cells are the main type, so the
date detector keeps the kind: ``[h]:mm`` cells under 24 hours are durations),
error tokens, formula count, the full mix of cell types and number formats
(``storage``, reproduced by exact quota: :mod:`.excel_storage`), and for
columns with more than one type/format combination the combinations of each
distinct value.

A notes sheet (see **notes** above) is kept as notes rows with no table. A
detected header row that looks like data (``sniff.header_problems``) fails the file with the
fix ``--header-row`` / ``--range``, because headers are copied verbatim.
Anything that may quote real data (title/notes/footer cells) lives in
``_``-prefixed layout keys.

``--range [Sheet!]A5:H200``: row 5 is the header and rows 6-200 the data
(no footer detection inside an explicit range); rows above the range are
title rows; cells below or beside it are not written (warning with a count).
Without ``Sheet!`` the range applies to every sheet read.

Reading is always streamed (openpyxl read-only; the writer is write-only),
whatever the size. The self-check re-read of a workbook written by this
package uses the header row, data rows and column span the writer recorded
for that file (:func:`register_written`), then compares everything else.
"""

from __future__ import annotations

import datetime as dt
import io
import os
import re
import warnings
import zipfile
import zlib
from collections import Counter, deque
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any, Iterator

import numpy as np

from ..errors import ReadError, UsageError
from ..model import FileLayout, Message, RawColumn, RawTable, ReadResult
from . import excel_book, sniff
from .excel_book import SheetXml
from .numfmt import safe_number_format

if TYPE_CHECKING:  # pragma: no cover
    from .registry import ReadRequest

HEADER_SCAN_ROWS = 30
#: a single-column block of at most this many rows below its first, all multi-word text, is notes
NOTES_MAX_ROWS = 5
NOTES_MIN_WORDS = 3
#: the wide row under a header candidate may be any of its next few non-blank rows (sparse first data rows)
WIDE_AFTER_ROWS = 5
FOOTER_MAX_ROWS = 20
VALUE_TYPES_CAP = 5000
TYPE_PRIORITY = "ndbs"  # main-type tie break

#: first-cell labels of footer summary rows (also the footer words the writer keeps verbatim)
SUMMARY_WORDS = r"sum|count|average|avg|mean|median|min|max|minimum|maximum|overall"
#: "Total"-family first cells: a total row wherever it is
STRONG_TOTAL_RE = re.compile(r"^\s*(?:(?:grand|sub)[\s-]*)?totals?\b", re.I)
#: a summary word alone: a total row only after a blank row or inside a footer block (it can be data: "Max")
WEAK_TOTAL_RE = re.compile(r"^\s*(?:" + SUMMARY_WORDS + r")\s*:?\s*$", re.I)
TOTAL_LABEL_RE = re.compile(STRONG_TOTAL_RE.pattern + "|" + WEAK_TOTAL_RE.pattern, re.I)
#: same-sheet A1 references in a formula: B5, $B$5, B5:B44, B:B (sheet-qualified references are skipped)
_REF_RE = re.compile(r"(?<![A-Za-z0-9_.!$:])\$?([A-Za-z]{1,3})\$?(\d+)"
                     r"(?::\$?([A-Za-z]{1,3})\$?(\d+))?(?![A-Za-z0-9_(!:])"
                     r"|(?<![A-Za-z0-9_.!$:])\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})(?![A-Za-z0-9_(!:])")
#: an aggregate function anywhere in a formula (``=ROUND(AVERAGE(E5:E44),2)`` is a total too)
AGGREGATE_RE = re.compile(r"(?<![A-Za-z0-9_.])(?:_xlfn\.)?(?:SUM|COUNT|COUNTA|AVERAGE|SUBTOTAL|AGGREGATE|MIN|MAX|"
                          r"MEDIAN|SUMIF|SUMIFS|COUNTIF|COUNTIFS|AVERAGEIF|AVERAGEIFS|PRODUCT)\s*\(", re.I)

#: written-workbook layouts registered by the writer for the self-check re-read (path -> sheet -> facts)
_WRITTEN: dict[str, dict[str, dict[str, Any]]] = {}


def path_key(path: Path) -> str:
    return os.path.normcase(os.path.abspath(str(path)))


def register_written(path: Path, sheets: dict[str, dict[str, Any]]) -> None:
    """Writer -> self-check: the header row/data rows/column span written for each sheet."""
    _WRITTEN[path_key(path)] = sheets


# --------------------------------------------------------------------------- cells and scanning
@dataclass
class XCell:
    """One non-empty cell: Python value, type ('s','n','d','b','e'; '' = formula without cached value),
    number format, and formula text (title/footer cells only)."""

    v: Any
    t: str
    fmt: str
    f: str | None = None


Row = dict[int, XCell]


class ColBuf:
    """Column storage for the data rows (values, type codes, number formats)."""

    __slots__ = ("values", "types", "fmts")

    def __init__(self) -> None:
        self.values: list[Any] = []
        self.types = bytearray()
        self.fmts: list[str | None] = []

    def put(self, idx: int, x: XCell) -> None:
        n = len(self.values)
        if n < idx:
            self.pad(idx)
        self.values.append(x.v)
        self.types.append(ord(x.t) if x.t else 0)
        self.fmts.append(x.fmt)

    def pad(self, n: int) -> None:
        k = n - len(self.values)
        if k > 0:
            self.values.extend([None] * k)
            self.types.extend(bytes(k))
            self.fmts.extend([None] * k)

    def truncate(self, n: int) -> None:
        del self.values[n:]
        del self.types[n:]
        del self.fmts[n:]


@dataclass
class Region:
    min_col: int
    min_row: int
    max_col: int | None
    max_row: int | None


def _is_text(x: XCell) -> bool:
    return x.t == "s" and isinstance(x.v, str)


def detect_header_row(rows: list[tuple[int, Row]]) -> int | None:
    """Header heuristic (see module docstring); ``rows`` = the buffered top rows."""
    cand = [(i, r, cells) for i, (r, cells) in enumerate(rows) if r <= HEADER_SCAN_ROWS]
    cols = _table_columns(rows)
    widths = [_width(cells, cols) for _, _, cells in cand]
    if not widths or max(widths) == 0:
        return next((r for r, cells in rows if cells), None)
    nonblank = [w for w in widths if w]
    mode = Counter(nonblank).most_common(1)[0][0]
    # the widest row first (design 6); then the most common width, so that a stray wide row (a note
    # beside the table) does not hide a header that is as wide as the table itself
    for maxw in [max(widths)] + ([mode] if 2 <= mode < max(widths) and mode >= 0.5 * max(widths) else []):
        found = _header_pass(rows, cand, maxw, cols)
        if found is not None:
            return found
    maxw = max(widths)
    h = _Heuristic(rows, maxw, cols)
    for _, r, cells in cand:
        if h.headerish(cells):
            return r
    for i, r, cells in cand:  # duplicate labels ('date', 'date') are still a header row
        if h.headerish(cells, unique=False) and h.wide_after(i, cells):
            return r
    for i, r, cells in cand:  # number/date labels (years): the first wide row with data under it
        if h.wide(cells) and h.wide_after(i, cells):
            return r
    return next((r for r, cells in rows if cells), None)


def _table_columns(rows: list[tuple[int, Row]]) -> set[int]:
    """Columns of the main block of the buffered rows: the run of adjacent occupied columns holding the
    most cells (blocks are separated by columns that are empty in every row).

    Row widths count only these columns, so a note beside the table (two
    columns to its right) does not make its data row the "widest row" of the
    heuristic, and a title out to the side does not count either.
    """
    counts: Counter = Counter(c for _, cells in rows for c in cells)
    if not counts:
        return set()
    blocks: list[list[int]] = []
    for c in sorted(counts):
        if blocks and c == blocks[-1][-1] + 1:
            blocks[-1].append(c)
        else:
            blocks.append([c])
    best = max(blocks, key=lambda b: sum(counts[c] for c in b))  # first (leftmost) on a tie
    return set(best)


def _width(cells: Row, cols: set[int]) -> int:
    return sum(1 for c in cells if c in cols)


class _Heuristic:
    """Row tests of the header heuristic for one reference width ``maxw`` (widths over ``cols``)."""

    def __init__(self, rows: list[tuple[int, Row]], maxw: int, cols: set[int]) -> None:
        self.rows = rows
        self.maxw = maxw
        self.cols = cols

    def wide(self, cells: Row) -> bool:
        return bool(cells) and _width(cells, self.cols) >= 0.8 * self.maxw

    def headerish(self, cells: Row, unique: bool = True) -> bool:
        if not self.wide(cells):
            return False
        texts = []
        for x in cells.values():
            if not _is_text(x) or not x.v.strip():
                return False
            texts.append(x.v.strip())
        return not unique or len(set(texts)) == len(texts)

    def wide_after(self, i: int, cells: Row) -> bool:
        """A wide row among the next few rows after row ``i`` that shares its columns (data under the
        labels; the first data row may be sparse)."""
        seen = 0
        for j in range(i + 1, len(self.rows)):
            nxt = self.rows[j][1]
            if not nxt:
                continue
            if _width(nxt, self.cols) >= 0.5 * self.maxw and min(nxt) <= max(cells) and min(cells) <= max(nxt):
                return True
            seen += 1
            if seen >= WIDE_AFTER_ROWS:
                return False
        return False


def _header_pass(rows: list[tuple[int, Row]], cand: list[tuple[int, int, Row]], maxw: int,
                 cols: set[int]) -> int | None:
    """The design heuristic for one reference width: the first header-like row with a wide row after it."""
    h = _Heuristic(rows, maxw, cols)
    for i, r, cells in cand:
        if h.headerish(cells) and h.wide_after(i, cells):
            return _title_above_single_column(rows, i, h)
    return None


def _title_above_single_column(rows: list[tuple[int, Row]], i: int, h: "_Heuristic") -> int:
    """Row number of the header when candidate ``rows[i]`` may be a one-cell title above a one-column table.

    Only a one-cell candidate can be a title as wide as the header (the table
    has one column, ``tc``: the fullest column of the main block). Looking down
    that column from the candidate:

    * data-like cells (numbers, dates, booleans, or text with a data shape such
      as a zero-padded ID) start the data when the cells directly under the
      first one are data-like too; the header is the nearest text cell above
      it (the rows above that are titles);
    * otherwise (text data): after a dated title line ("Printed 10/01/2024 by
      ..."), the header is the next cell that is not one; after an isolated
      title line (>= 3 words, a blank row under it) or a block of such lines,
      it is the first cell of the next block of two or more cells (past any
      dated lines on top of it). A short label above a blank row stays the
      header.

    A label is chosen only if it is not data-shaped, has data under it and does
    not recur below (a header names its column once; a data value usually
    repeats). Otherwise the candidate itself.
    """
    r, cells = rows[i]
    if len(cells) != 1 or not h.cols:
        return r
    counts = Counter(c for _, row in rows for c in row if c in h.cols)
    tc = max(sorted(counts), key=lambda c: counts[c])
    col = [(j, rows[j][1].get(tc)) for j in range(i, len(rows))]
    found = _header_above_data_like(rows, col)
    if found is not None:
        return found if found > r else r
    return _header_after_titles(rows, col, r)


def _data_like(x: XCell | None) -> bool:
    if x is None:
        return False
    if not _is_text(x):
        return x.v is not None
    prob = sniff.header_problems([x.v], [])
    return prob is not None and not prob.prose


def _label_ok(x: XCell | None) -> bool:
    return x is not None and _is_text(x) and bool(x.v.strip()) and sniff.header_problems([x.v], []) is None


def _header_above_data_like(rows: list[tuple[int, Row]], col: list[tuple[int, XCell | None]]) -> int | None:
    """See :func:`_title_above_single_column` (data-like cells); None when the signal is absent."""
    k = next((p for p, (_, x) in enumerate(col) if _data_like(x)), None)
    if k is None or k == 0:
        return None
    for _, x in col[k + 1:k + 1 + WIDE_AFTER_ROWS]:
        if x is None or (_is_text(x) and TOTAL_LABEL_RE.match(x.v)):
            break  # the run of data cells ends (a footer such as "Total" may follow)
        if not _data_like(x):
            return None  # text among the data: no clear run of data
    above = next((p for p in range(k - 1, -1, -1) if col[p][1] is not None), None)
    if above is None or not _label_ok(col[above][1]):
        return None
    return rows[col[above][0]][0]


_DATED_RE = re.compile(r"(?<!\d)(?:\d{1,2}[/.-]\d{1,2}[/.-](?:\d{4}|\d{2})|\d{4}-\d{2}-\d{2})(?!\d)")


def _dated_line(x: XCell | None) -> bool:
    """A title line: >= 3 words with a date in it ("Printed 10/01/2024 by jdoe"); a column label has none."""
    return (x is not None and _is_text(x) and len(x.v.split()) >= NOTES_MIN_WORDS
            and _DATED_RE.search(x.v) is not None)


def _header_after_titles(rows: list[tuple[int, Row]], col: list[tuple[int, XCell | None]], r: int) -> int:
    """See :func:`_title_above_single_column` (text data)."""
    runs: list[tuple[int, int]] = []  # (start, end) positions of consecutive non-blank cells
    p = 0
    while p < len(col):
        if col[p][1] is None:
            p += 1
            continue
        q = p
        while q + 1 < len(col) and col[q + 1][1] is not None:
            q += 1
        runs.append((p, q))
        p = q + 1
    head = col[0][1]
    if _dated_line(head):
        # a dated title line: the header is the first later cell that is not one
        p = next((k for k in range(1, len(col)) if col[k][1] is not None and not _dated_line(col[k][1])), None)
        return rows[col[p][0]][0] if p is not None and _good_label(col, p) else r
    short_only = False
    later = runs
    if head is not None:
        a0, b0 = runs[0]
        if b0 > a0:
            # cells directly under the candidate: it is the header, unless that block is all title-like
            # lines (>= 3 words each) followed, after a blank row, by a short label
            lines = [x for _, x in col[a0:b0 + 1]]
            if len(runs) < 2 or not all(_is_text(x) and len(x.v.split()) >= NOTES_MIN_WORDS for x in lines):
                return r
            short_only = True
        elif len(head.v.split()) < NOTES_MIN_WORDS:
            return r  # a short label alone above a blank row is the header of data further down
        later = runs[1:]
    run = next(((a, b) for a, b in later if b > a), None)  # the next block of two or more cells
    if run is None:
        return r
    p = run[0]
    while p < run[1] and _dated_line(col[p][1]):  # dated title lines stacked on the header
        p += 1
    if short_only and len(col[p][1].v.split()) >= NOTES_MIN_WORDS:
        return r
    return rows[col[p][0]][0] if _good_label(col, p) else r


def _good_label(col: list[tuple[int, XCell | None]], p: int) -> bool:
    """``col[p]`` can be the header: a plain label with data under it that does not recur below."""
    x = col[p][1]
    if not _label_ok(x) or not any(y is not None for _, y in col[p + 1:]):
        return False
    text = x.v.strip()
    return not any(y is not None and _is_text(y) and y.v.strip() == text for _, y in col[p + 1:])


def _first_cell(cells: Row) -> XCell | None:
    return cells[min(cells)] if cells else None


def is_column_aggregate(formula: str | None, row: int) -> bool:
    """A formula using an aggregate (SUM/COUNT/AVERAGE ...) over a range above ``row`` (or whole columns).

    A row formula such as ``=SUM(B5:D5)`` in row 5 (or a running total that
    includes its own row) is data, not a footer total.
    """
    if not formula or not AGGREGATE_RE.search(formula):
        return False
    for m in _REF_RE.finditer(formula):
        if m.group(5):  # whole columns: B:B
            return True
        r1 = int(m.group(2))
        r2 = int(m.group(4)) if m.group(4) else r1
        if max(r1, r2) < row:
            return True
    return False


def total_kind(cells: Row, formulas: dict[int, str] | None, row: int) -> str | None:
    """'strong' (column aggregate formula or a "Total" label), 'weak' (a summary word alone) or None."""
    if formulas and any(is_column_aggregate(f, row) for f in formulas.values()):
        return "strong"
    first = _first_cell(cells)
    if first is None or not _is_text(first):
        return None
    if STRONG_TOTAL_RE.match(first.v):
        return "strong"
    if WEAK_TOTAL_RE.match(first.v):
        return "weak"
    return None


def is_total_row(cells: Row, formulas: dict[int, str] | None, row: int | None = None) -> bool:
    """A strong or weak total row (see :func:`total_kind`); ``row`` defaults to below any reference."""
    return total_kind(cells, formulas, row if row is not None else 2 ** 31) is not None


def is_note_row(cells: Row) -> bool:
    """One plain text cell (not a data-shaped value such as a zero-padded ID or a date)."""
    if len(cells) != 1:
        return False
    x = next(iter(cells.values()))
    return _is_text(x) and not _data_like(x)


class SheetScan:
    """Streams one sheet's rows and splits them into title / header / data / footer (module docstring)."""

    def __init__(self, name: str, header_override: int | None = None, region: Region | None = None,
                 known: dict[str, Any] | None = None) -> None:
        self.name = name
        self.region = region
        # layout written by our own writer (self-check re-read)
        self.known_notes = bool(known and known.get("notes_only"))
        self.known = known if known is not None and not self.known_notes else None
        self.any_nonblank = False
        if self.known is not None:
            self.header_override = self.known.get("header_row")
        elif region is not None:
            self.header_override = region.min_row
        else:
            self.header_override = header_override
        self.pre: list[tuple[int, Row]] = []
        self.decided = False
        self.header_row: int | None = None
        self.header_cells: Row = {}
        self.title_rows: list[tuple[int, Row]] = []
        self.cols: dict[int, ColBuf] = {}
        self.n_data = 0
        self.nonblank_data_rows = 0
        self.collect_dtypes = False
        self.inferred_dtypes: list[str] | None = None
        self.tail: deque[tuple[int, Row]] = deque(maxlen=FOOTER_MAX_ROWS + 1)
        self.after_known: list[tuple[int, Row]] = []  # rows after the known data end (re-read)
        self.last_nonblank: int | None = None
        self.last_row = 0
        self.dropped_below = 0
        self.dropped_side = 0

    # ------------------------------------------------------------------ feeding
    def feed(self, r: int, cells: Row) -> None:
        self.last_row = r
        rg = self.region
        if rg is not None and self.known is None:
            if rg.max_row is not None and r > rg.max_row:
                self.dropped_below += len(cells)
                return
            if r >= rg.min_row and cells:
                keep = {c: x for c, x in cells.items() if c >= rg.min_col and (rg.max_col is None or c <= rg.max_col)}
                self.dropped_side += len(cells) - len(keep)
                cells = keep
        if not self.decided:
            self.pre.append((r, cells))
            if cells:
                self.any_nonblank = True
            if self._can_decide(r):
                self._decide()
            return
        self._append(r, cells)

    def _can_decide(self, r: int) -> bool:
        if self.header_override is not None:
            return r >= self.header_override
        return r >= HEADER_SCAN_ROWS + 1 and self.any_nonblank

    def _decide(self) -> None:
        if self.header_override is not None:
            h = self.header_override
        else:
            h = detect_header_row(self.pre)
        self.decided = True
        if h is None:  # nothing but blank rows
            self.header_row = None
            self.title_rows = [(r, c) for r, c in self.pre if c]
            self.pre = []
            return
        self.header_row = h
        pre, self.pre = self.pre, []
        for r, cells in pre:
            if r < h:
                if cells:
                    self.title_rows.append((r, cells))
            elif r == h:
                self.header_cells = cells
            else:
                self._append(r, cells)

    def _append(self, r: int, cells: Row) -> None:
        assert self.header_row is not None
        if self.known is not None and r > self.known["data_end"]:
            if cells:
                self.after_known.append((r, cells))
            return
        idx = r - self.header_row - 1
        if idx < self.n_data:  # never happens with increasing rows
            return
        for c, x in cells.items():
            buf = self.cols.get(c)
            if buf is None:
                buf = self.cols[c] = ColBuf()
            buf.put(idx, x)
        self.n_data = idx + 1
        if any(x.v is not None for x in cells.values()):
            # a row of formulas whose cached value is empty ('=IF(B9="","",...)' filled down past the
            # data) is blank, as pandas reads it: it does not extend the data or start a footer
            self.last_nonblank = r
            self.nonblank_data_rows += 1
            self.tail.append((r, cells))

    # ------------------------------------------------------------------ finishing
    def finish(self, sx: SheetXml | None) -> "ScannedSheet":
        if not self.decided:
            if self.header_override is not None and self.header_override > self.last_row and self.any_nonblank:
                raise ReadError(f"sheet {self.name!r} has {self.last_row} row(s); the header row "
                                f"{self.header_override} is beyond its end",
                                fix="check --header-row / --range (or --sheet NAME to read one sheet)")
            if self.header_override is not None:
                self.header_row = self.header_override
                self.decided = True
                pre, self.pre = self.pre, []
                for r, cells in pre:
                    if r < self.header_row and cells:
                        self.title_rows.append((r, cells))
                    elif r == self.header_row:
                        self.header_cells = cells
            else:
                self._decide()
        out = ScannedSheet(self.name)
        out.title_rows = self.title_rows
        out.dropped_below = self.dropped_below
        out.dropped_side = self.dropped_side
        if self.header_row is None:
            return out
        h = self.header_row
        out.header_row = h
        out.header_cells = self.header_cells
        formulas = {r: cols for r, cols in sx.formula_tail} if sx is not None else {}
        footer: list[tuple[int, Row]] = []
        if self.known is not None:
            data_end = self.known["data_end"]
            footer = self.after_known
        else:
            data_end = self.last_nonblank if self.last_nonblank is not None else h
            if self.region is None or self.region.max_row is None:
                start = self._footer_start(formulas)
                if start is not None:
                    tail = list(self.tail)
                    footer = tail[start:]
                    data_end = tail[start - 1][0]
            elif self.region.max_row is not None:
                data_end = max(h, min(data_end, self.region.max_row))
        n = data_end - h
        for buf in self.cols.values():
            buf.truncate(n)
            buf.pad(n)
        out.n_rows = n
        out.data_end = data_end
        out.footer_rows = footer
        out.footer_formulas = {r: formulas.get(r, {}) for r, _ in footer}
        heads = set(self.header_cells)
        span = set(heads)
        for c, buf in self.cols.items():
            if any(buf.types) or any(v is not None for v in buf.values):
                span.add(c)
        if self.known is None and self.region is None and heads:
            out.side_cols = side_columns(span, heads)
            span -= set(out.side_cols)
        if self.known is not None:
            out.first_col, out.last_col = self.known["first_col"], self.known["last_col"]
        elif span:
            out.first_col, out.last_col = min(span), max(span)
        else:
            out.first_col = out.last_col = None
        out.cols = self.cols
        return out

    def _single_column_note(self, cells: Row) -> bool:
        """In a one-column table, a single text cell (after a blank row) is a note when the column holds
        mostly data-like cells (numbers, dates, zero-padded IDs ...), or when it is a multi-word line
        (>= 3 words) below short labels. Its last 500 cells are looked at, the candidate note excluded."""
        if len(self.header_cells) != 1:
            return False
        buf = self.cols.get(next(iter(self.header_cells)))
        x = next(iter(cells.values()))
        if buf is None:
            return False
        typed = 0
        words: list[int] = []
        for v, t in zip(buf.values[-501:], buf.types[-501:]):
            if not t or v is None or t == ord("e") or v is x.v:
                continue
            if _data_like(XCell(v, chr(t), "General")):
                typed += 1
            elif isinstance(v, str):
                words.append(len(v.split()))
        if typed > 2 * len(words):
            return True
        return (len(x.v.split()) >= NOTES_MIN_WORDS and bool(words)
                and sorted(words)[len(words) // 2] < NOTES_MIN_WORDS)

    def _footer_start(self, formulas: dict[int, dict[int, str]]) -> int | None:
        tail = list(self.tail)
        if not tail:
            return None
        width = len(self.header_cells)

        def kind(k: int) -> str | None:
            r, cells = tail[k]
            tk = total_kind(cells, formulas.get(r), r)
            if tk is None and is_note_row(cells) and (width >= 2 or self._single_column_note(cells)):
                return "note"
            return tk

        kinds: dict[int, str] = {}
        i = len(tail) - 1
        while i >= 0 and len(tail) - i <= FOOTER_MAX_ROWS:
            k = kind(i)
            if k is None:
                break
            kinds[i] = k
            i -= 1
        if not kinds:
            return None
        if len(kinds) >= FOOTER_MAX_ROWS and i >= 0 and kind(i) is not None:
            return None  # total-like rows all the way up: row formulas / data, not a footer block
        for start in range(max(i + 1, 1), len(tail)):
            gap = tail[start][0] - tail[start - 1][0] > 1
            # when every row is total-like ("Total A", "Total B" summary rows) only a blank row
            # separates data from a footer
            if gap or (kinds[start] == "strong" and i >= 0):
                return start
        return None


@dataclass
class ScannedSheet:
    name: str
    header_row: int | None = None
    header_cells: Row = field(default_factory=dict)
    title_rows: list[tuple[int, Row]] = field(default_factory=list)
    footer_rows: list[tuple[int, Row]] = field(default_factory=list)
    footer_formulas: dict[int, dict[int, str]] = field(default_factory=dict)
    cols: dict[int, ColBuf] = field(default_factory=dict)
    n_rows: int = 0
    data_end: int = 0
    first_col: int | None = None
    last_col: int | None = None
    dropped_below: int = 0
    dropped_side: int = 0
    side_cols: list[int] = field(default_factory=list)  # note columns beside the table (side_columns)


def side_columns(occupied: set[int], heads: set[int]) -> list[int]:
    """Columns of headerless blocks beside the table: notes, not data.

    ``occupied``: columns holding a header or a data cell; ``heads``: columns
    with a header cell. Occupied columns form blocks separated by fully empty
    columns. A block with no header cell that lies left or right of every block
    with one (``Note: ... withdrew 9/14`` two columns right of the table) is
    side notes: its cells are written in place as same-shape placeholders and
    are not profiled as a column (a one-value column would otherwise be
    reproduced as a real label).
    """
    blocks: list[list[int]] = []
    for c in sorted(occupied):
        if blocks and c == blocks[-1][-1] + 1:
            blocks[-1].append(c)
        else:
            blocks.append([c])
    with_head = [k for k, b in enumerate(blocks) if heads.intersection(b)]
    if not with_head:
        return []
    lo, hi = with_head[0], with_head[-1]
    return [c for k, b in enumerate(blocks) if k < lo or k > hi for c in b]


# --------------------------------------------------------------------------- canonical values
def _fmt_has_time(fmt: str | None) -> bool:
    if not fmt:
        return False
    f = re.sub(r'"[^"]*"|\\.', "", fmt)
    f = re.sub(r"\[(?![hHmMsS]+\])[^\]]*\]", "", f)
    low = f.lower()
    return "h" in low or "s" in low or "am/pm" in low or "a/p" in low


def _fmt_has_date(fmt: str | None) -> bool:
    if not fmt:
        return False
    f = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", fmt).lower()
    return "y" in f or "d" in f


def canon_number(v: Any) -> str:
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return str(v)
        if v == 0:
            return "0"
        if v.is_integer() and abs(v) < 2 ** 63:
            return str(int(v))
        return np.format_float_positional(v, unique=True, trim="-")
    return str(v)


def _round_dt(v: dt.datetime) -> dt.datetime:
    if v.microsecond >= 500000:
        v = v + dt.timedelta(microseconds=1000000 - v.microsecond)
    return v.replace(microsecond=0)


def canon_date(v: Any, kind: str | None) -> str:
    if isinstance(v, dt.timedelta):
        secs = int(round(v.total_seconds()))
        sign = "-" if secs < 0 else ""
        secs = abs(secs)
        h, rem = divmod(secs, 3600)
        m, s = divmod(rem, 60)
        return f"{sign}{h}:{m:02d}:{s:02d}"
    if isinstance(v, dt.time):
        return f"{v.hour:02d}:{v.minute:02d}:{v.second:02d}"
    if isinstance(v, dt.datetime):
        v = _round_dt(v)
        d = f"{v.year:04d}-{v.month:02d}-{v.day:02d}"
        if kind == "date":
            return d
        return f"{d} {v.hour:02d}:{v.minute:02d}:{v.second:02d}"
    if isinstance(v, dt.date):
        d = f"{v.year:04d}-{v.month:02d}-{v.day:02d}"
        return d if kind != "datetime" else d + " 00:00:00"
    if isinstance(v, (int, float)):
        return canon_number(v)
    return str(v)


def canon(v: Any, t: str, kind: str | None) -> str | None:
    if v is None:
        return None
    if t == "s" or isinstance(v, str):
        return v if v != "" else None
    if t == "b" or isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if t == "d" or isinstance(v, (dt.date, dt.time, dt.timedelta)):
        return canon_date(v, kind)
    return canon_number(v)


def _date_kind(values: list[Any], types: bytearray, fmt: str | None) -> str | None:
    dvals = [v for v, t in zip(values, types) if t == 100 and v is not None]  # 'd'
    if not dvals:
        return None
    if all(isinstance(v, dt.timedelta) for v in dvals):
        return "duration"
    if all(isinstance(v, dt.time) for v in dvals):
        return "time"
    for v in dvals:
        if isinstance(v, dt.datetime) and (v.hour or v.minute or v.second or v.microsecond >= 500000):
            return "datetime"
    return "datetime" if _fmt_has_time(fmt) else "date"


def finalize_column(buf: ColBuf | None, n: int, formulas: int = 0,
                    cell_types: bool = False) -> tuple[np.ndarray, dict[str, Any]]:
    """Canonical value strings and the PRIVATE native description of one column.

    ``storage``: non-empty cells per cell type and number format (the full
    mix, kept by exact quota: :mod:`.excel_storage`); ``formats``: the most
    common format per type; ``value_storage`` (columns with more than one
    type/format combination): the combinations of each distinct value.
    ``cell_types`` (self-check re-read): the per-cell type codes are kept too.
    """
    values = np.empty(n, dtype=object)
    native: dict[str, Any] = {"kind": "excel", "types": {}, "main": None, "formats": {}, "date_kind": None,
                              "errors": [], "formulas": formulas, "storage": {}}
    if buf is None or n == 0:
        values[:] = ""
        if cell_types:
            native["cell_types"] = bytes(n)
        return values, native
    buf.truncate(n)
    buf.pad(n)
    counts = Counter(buf.types)
    counts.pop(0, None)
    types = {chr(k): v for k, v in counts.items()}
    native["types"] = types
    fmt_by_type: dict[str, Counter] = {}
    for t, f in zip(buf.types, buf.fmts):
        if t:
            fmt_by_type.setdefault(chr(t), Counter())[f or "General"] += 1
    native["formats"] = {t: c.most_common(1)[0][0] for t, c in fmt_by_type.items()}
    native["storage"] = {t: dict(c) for t, c in fmt_by_type.items()}
    # the storage of the first non-empty cell (readers such as pandas type a bool/text mix by it)
    first = next((k for k, (v, t) in enumerate(zip(buf.values, buf.types)) if t and v is not None), None)
    if first is not None:
        native["first"] = [chr(buf.types[first]), buf.fmts[first] or "General"]
    if cell_types:
        native["cell_types"] = bytes(buf.types)
    main_counts = {t: c for t, c in types.items() if t in "sndb"}
    if main_counts:
        best = max(main_counts.values())
        native["main"] = next(t for t in TYPE_PRIORITY if main_counts.get(t) == best)
        native["main_share"] = round(best / sum(main_counts.values()), 6)
    kind = _date_kind(buf.values, buf.types, native["formats"].get("d"))
    native["date_kind"] = kind
    if native["main"] == "d" and kind:
        # AREA A native-storage contract (generate/dates.py): the temporal kind of a date-cell column, so
        # the date detector keeps it (an [h]:mm duration under 24 h is a DURATION, not a time of day)
        native["temporal"] = {"kind": kind, "tz": None, "unit": None}
    errs: set[str] = set()
    # PRIVATE: the storage combinations of each distinct value, when the column has more than one
    vs: dict[str, Counter] | None = {} if sum(len(c) for c in fmt_by_type.values()) >= 2 else None
    for i, (v, t, f) in enumerate(zip(buf.values, buf.types, buf.fmts)):
        tc = chr(t) if t else ""
        s = canon(v, tc, kind)
        if s is None:
            values[i] = ""  # empty cell: blank, as in the text formats
            continue
        values[i] = s
        if tc == "e":
            errs.add(s)
        if vs is not None and tc in "sndbe" and (s in vs or len(vs) < VALUE_TYPES_CAP):
            vs.setdefault(s, Counter())[(tc, f or "General")] += 1
    native["errors"] = sorted(errs)
    if vs is not None:
        native["value_storage"] = vs
        if len(main_counts) >= 2:
            vt: dict[str, str] = {}
            for s, c in vs.items():
                tcount: Counter = Counter()
                for (tc, _), k in c.items():
                    if tc in "sndb":
                        tcount[tc] += k
                if tcount:
                    vt[s] = tcount.most_common(1)[0][0]
            native["value_types"] = vt
    return values, native


# --------------------------------------------------------------------------- options
def parse_ranges(entries: list[str]) -> dict[str | None, Region]:
    """``--range`` values -> {sheet name or None (every sheet): Region}."""
    from openpyxl.utils.cell import range_boundaries

    out: dict[str | None, Region] = {}
    for e in entries:
        sheet, sep, ref = e.rpartition("!")
        if not sep:
            sheet, ref = "", e
        sheet = sheet.strip()
        if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
            sheet = sheet[1:-1].replace("''", "'")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref.strip().replace("$", ""))
        except (ValueError, TypeError):
            raise UsageError(f"--range {e!r}: expected [Sheet!]A5:H200", fix="e.g. --range \"Sheet1!A5:H200\"") from None
        if min_col is None or min_row is None:
            raise UsageError(f"--range {e!r}: the range needs a first column and row (e.g. A5:H200)")
        out[sheet or None] = Region(min_col, min_row, max_col, max_row)
    return out


# --------------------------------------------------------------------------- sheet sources
def xlsx_rows(ws: Any, sx: SheetXml | None) -> Iterator[tuple[int, Row]]:
    """(row number, non-empty cells) of an openpyxl read-only worksheet, rows 1..last."""
    ws.reset_dimensions()
    fmt_cache: dict[int, str] = {}
    masked = sx.extras if sx is not None else Counter()
    uncached = sx.uncached if sx is not None else set()
    for r, row in enumerate(ws.iter_rows(), 1):
        cells: Row = {}
        for c in row:
            v = c.value
            if v is None or (v == "" and isinstance(v, str)):
                if uncached and (r, getattr(c, "column", 0)) in uncached:
                    cells[c.column] = XCell(None, "", _fmt(c, fmt_cache, masked), "=")
                continue
            cells[c.column] = XCell(v, c.data_type if c.data_type in ("s", "n", "d", "b", "e") else "s",
                                    _fmt(c, fmt_cache, masked))
        yield r, cells


#: ``extras`` key counting cell styles whose number format held literal text (masked, never copied).
MASKED_FORMAT_EXTRA = "number-format literal text (masked)"


def _fmt(c: Any, cache: dict[int, str], masked: Counter | None = None) -> str:
    """The cell's number format with its literal text masked (``io.numfmt``: a format can quote a name)."""
    sid = getattr(c, "_style_id", 0) or 0
    f = cache.get(sid)
    if f is None:
        try:
            raw = c.number_format or "General"
        except Exception:  # noqa: BLE001 - a broken style table falls back to General
            raw = "General"
        f = safe_number_format(raw) or "General"
        if f != raw and masked is not None:
            masked[MASKED_FORMAT_EXTRA] += 1
        cache[sid] = f
    return f


# --------------------------------------------------------------------------- building tables
def build_table(scan: SheetScan, sx: SheetXml | None, state: str, index: int,
                messages: list[Message], purpose: str, header_checked: bool) -> tuple[RawTable, dict[str, Any]]:
    """RawTable + public sheet summary from a finished scan."""
    s = scan.finish(sx)
    sx = sx or SheetXml()
    layout: dict[str, Any] = {"sheet_index": index, "state": state, "freeze_panes": sx.freeze_panes,
                              "autofilter": sx.autofilter, "widths": sx.widths, "default_width": sx.default_width,
                              "_title_rows": s.title_rows}
    summary: dict[str, Any] = {"name": scan.name, "state": state}
    if s.dropped_below or s.dropped_side:
        messages.append(Message("warning", "excel_range_dropped",
                                f"sheet {scan.name!r}: {s.dropped_below + s.dropped_side} non-empty cell(s) outside "
                                "--range were not written"))
    notes_only = s.header_row is None or s.first_col is None
    if not notes_only and header_checked and scan.known is None:
        notes_only = _prose_sheet(s)
    if scan.known_notes:
        notes_only = True
    if notes_only:
        rows = list(s.title_rows)
        if s.header_row is not None:
            rows.append((s.header_row, s.header_cells))
            rows.extend(_rows_from_cols(s))
            rows.extend(s.footer_rows)
        rows = [(r, c) for r, c in sorted(rows, key=lambda rc: rc[0]) if c]
        layout.update({"notes_only": True, "header_row": None, "_title_rows": rows, "merged": sx.merged,
                       "columns": [], "n_title_rows": len(rows)})
        summary.update({"notes_only": True, "n_title_rows": len(rows), "n_rows": 0})
        if rows and purpose != "selfcheck" and not scan.known_notes and s.header_row is not None:
            messages.append(Message("warning", "excel_notes_sheet",
                                    f"sheet {scan.name!r} holds notes/prose, not a table: written back as same-shape "
                                    "placeholder text (--keep-titles copies it verbatim)"))
        return RawTable(scan.name, [], [], 0, layout), summary
    h = s.header_row
    assert h is not None and s.first_col is not None and s.last_col is not None
    header: list[str] = []
    columns: list[RawColumn] = []
    col_layout: list[list[Any]] = []
    header_cells: dict[int, list[Any]] = {}
    for c in range(s.first_col, s.last_col + 1):
        hx = s.header_cells.get(c)
        hkind = _date_kind([hx.v], bytearray([100]), hx.fmt) if hx is not None and hx.t == "d" else None
        name = canon(hx.v, hx.t, hkind) if hx is not None and hx.v is not None else None
        name = name if name is not None else ""
        if hx is not None:
            header_cells[c] = [hx.t, hx.fmt, hkind]  # the writer restores number/date/bool header cells
        nf = excel_book.formula_count_between(sx, c, h + 1, s.data_end)
        values, native = finalize_column(s.cols.get(c), s.n_rows, nf, cell_types=purpose == "selfcheck")
        header.append(name)
        columns.append(RawColumn(name, values, native))
        # the full storage mix (aggregate counts, masked formats) for the self-check's exact comparison
        col_layout.append([name, native["main"], native["formats"].get(native["main"] or "", None),
                           native.get("main_share"),
                           {"cells": native["storage"], "date_kind": native["date_kind"],
                            "errors": native["errors"]}])
    merged_keep, merged_drop = _classify_merges(sx.merged, h, s.data_end)
    if merged_drop and purpose != "selfcheck":
        messages.append(Message("warning", "excel_merged_dropped",
                                f"sheet {scan.name!r}: {merged_drop} merged range(s) inside the data rows were not "
                                "re-merged (the synthetic values differ row by row)"))
    side = _side_cells(s)
    layout.update({
        "notes_only": False, "header_row": h, "first_col": s.first_col, "last_col": s.last_col,
        "n_title_rows": len(s.title_rows), "data_end": s.data_end, "n_footer_rows": len(s.footer_rows),
        "merged": merged_keep, "columns": col_layout, "_header_cells": header_cells,
        "_footer_rows": s.footer_rows, "_footer_formulas": s.footer_formulas, "_side_cells": side,
    })
    if side and purpose != "selfcheck":
        n_side = sum(len(c) for c in side.values())
        messages.append(Message("warning", "excel_side_notes",
                                f"sheet {scan.name!r}: {n_side} cell(s) in {len(s.side_cols)} headerless column(s) "
                                "beside the table (separated from it by an empty column) are treated as notes: "
                                "written in place as same-shape placeholders, not profiled as data",
                                fix=f"--range \"{scan.name}!{_ref(min(s.first_col, *s.side_cols), h)}:"
                                    f"{_ref(max(s.last_col, *s.side_cols), s.data_end)}\" to read them as data "
                                    "columns"))
    summary.update({"header_row": h, "first_col": s.first_col, "last_col": s.last_col, "n_rows": s.n_rows,
                    "n_title_rows": len(s.title_rows), "n_footer_rows": len(s.footer_rows)})
    # Native values are still available here, before ColBuf is released. The
    # read_xlsx caller moves only dtype strings into private file-layout extras.
    if scan.collect_dtypes:
        dtypes = _scanned_dtypes(scan, s, sx)
        if dtypes is not None:
            scan.inferred_dtypes = dtypes
    if purpose != "selfcheck":
        bits = [f"header row {h}"]
        if s.title_rows:
            bits.append(f"{len(s.title_rows)} title/notes row(s) above it (written as placeholders)")
        if s.footer_rows:
            bits.append(f"{len(s.footer_rows)} footer row(s) below the data (totals recomputed or blank)")
        messages.append(Message("info", "excel_sheet_layout", f"sheet {scan.name!r}: " + ", ".join(bits)))
    return RawTable(scan.name, header, columns, s.n_rows, layout), summary


def dtype_file_key(path: Path) -> tuple[int, int, int, int] | None:
    """Rename-stable identity of the workbook read, for the private dtype cache.

    The self-check reads the temporary output before its atomic rename. Device,
    inode, size and modification time identify that same file at its destination;
    the path and ctime cannot (POSIX rename changes ctime). Filesystems without
    an inode use the full pandas read.
    """
    try:
        stat = path.stat()
    except OSError:
        return None
    if not stat.st_ino:
        return None
    return stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns


def dtype_window(info: dict[str, Any]) -> tuple[Any, ...]:
    return tuple(info.get(k) for k in ("header_row", "first_col", "last_col", "n_rows"))


def _native_dtype_values(buf: ColBuf | None, n: int) -> list[Any]:
    """First-appearance native values as pandas' OpenpyxlReader converts them.

    Typed keys keep True/1 and numeric/text spellings distinct. Error cells
    become the one NaN sentinel, blanks become empty strings, and integral
    numeric cells become Python ints even beyond int64 (matching pandas).
    """
    if buf is None:
        return [""] if n else []
    distinct: dict[tuple[type, Any], Any] = {}
    for value, typ in zip(buf.values, buf.types):
        if value is None:
            value = ""
        elif typ == ord("e"):
            value = np.nan
        elif typ == ord("n"):
            integer = int(value)
            value = integer if integer == value else float(value)
        distinct.setdefault((type(value), value), value)
    return list(distinct.values())


def _native_dtype(values: list[Any]) -> str:
    """Run the same Python parser used by read_excel on one distinct column."""
    from pandas.io.parsers import TextParser

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with TextParser([["value"], *[[v] for v in values]], header=0, skip_blank_lines=False) as parser:
            frame = parser.read()
    return str(frame.dtypes.iloc[0])


def _scanned_dtypes(scan: SheetScan, s: ScannedSheet, sx: SheetXml) -> list[str] | None:
    """Infer only a plain XLSX window whose native cells exactly match pandas'."""
    if (s.header_row != 1 or s.first_col != 1 or not s.n_rows or scan.region is not None
            or s.title_rows or s.footer_rows or sx.merged or s.side_cols
            or scan.nonblank_data_rows != s.n_rows or s.last_col is None
            or any(c not in s.header_cells or s.header_cells[c].t != "s"
                   for c in range(1, s.last_col + 1))):
        return None
    try:
        return [_native_dtype(_native_dtype_values(s.cols.get(c), s.n_rows)) for c in range(1, s.last_col + 1)]
    except Exception:
        # pandas may reject malformed native values, unsupported cell mixtures,
        # or parser-version differences; the existing read_excel hook decides.
        return None


def _ref(col: int, row: int) -> str:
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(col)}{row}"


def _side_cells(s: ScannedSheet) -> dict[int, Row]:
    """PRIVATE: data-row index -> {column: cell} of the side-note columns (written as placeholders)."""
    out: dict[int, Row] = {}
    for c in s.side_cols:
        buf = s.cols.get(c)
        if buf is None:
            continue
        for i in range(min(s.n_rows, len(buf.values))):
            t, v = buf.types[i], buf.values[i]
            if t and v is not None:
                out.setdefault(i, {})[c] = XCell(v, chr(t), buf.fmts[i] or "General")
    return out


def _rows_from_cols(s: ScannedSheet) -> list[tuple[int, Row]]:
    assert s.header_row is not None
    out: list[tuple[int, Row]] = []
    for i in range(s.n_rows):
        cells: Row = {}
        for c, buf in s.cols.items():
            t = buf.types[i]
            if t or buf.values[i] is not None:
                cells[c] = XCell(buf.values[i], chr(t) if t else "", buf.fmts[i] or "General")
        if cells:
            out.append((s.header_row + 1 + i, cells))
    return out


def notes_like(header: str, values: list[str]) -> bool:
    """A single-column block that is notes, not a table: ``header`` = its first row, ``values`` the rest.

    Either the first row reads like a sentence (``sniff.header_problems``'
    prose rule), or the block is short (at most :data:`NOTES_MAX_ROWS` non-blank
    rows below the first) and every row is multi-word text ("Class roster - Mr.
    Okafor" / "Printed 10/01/2024 by jdoe"). Its text is written back as
    same-shape placeholders instead of being copied as a header and levels.
    """
    prob = sniff.header_problems([header], [])
    if prob is not None and prob.prose:
        return True
    rest = [v for v in values if v.strip()]
    return len(rest) <= NOTES_MAX_ROWS and all(len(v.split()) >= NOTES_MIN_WORDS for v in [header, *rest])


def _prose_sheet(s: ScannedSheet) -> bool:
    """A single-column sheet that is notes, not a table (:func:`notes_like`)."""
    if s.first_col != s.last_col:
        return False
    hx = s.header_cells.get(s.first_col) if s.first_col is not None else None
    if hx is None or not _is_text(hx):
        return False
    buf = s.cols.get(s.first_col)
    values: list[str] = []
    typed = False  # a number/date/bool cell below the first row: a data column, unless the header is prose
    if buf is not None:
        for v, t in zip(buf.values[:s.n_rows], buf.types[:s.n_rows]):
            if v is not None:
                typed = typed or t != ord("s") or not isinstance(v, str)
                values.append(v if isinstance(v, str) else "")
    if typed:
        prob = sniff.header_problems([hx.v], [])
        return bool(prob and prob.prose)
    return notes_like(hx.v, values)


def check_header(t: RawTable, sheet: str) -> None:
    """Refuse a header row that looks like data (its text would be copied verbatim)."""
    if not t.columns:
        return
    cols = [[v for v in c.values.tolist() if v is not None][:5000] for c in t.columns]
    prob = sniff.header_problems(t.header, cols)
    if prob is not None and not prob.prose:
        raise ReadError(f"sheet {sheet!r}: the detected header row {t.layout.get('header_row')} does not look like a "
                        f"header: " + "; ".join(prob.reasons),
                        fix=f"--header-row N or --range \"{sheet}!A5:H200\" (or --sheet NAME to read one sheet)")


def _classify_merges(merged: list[str], header_row: int, data_end: int) -> tuple[list[str], int]:
    from openpyxl.utils.cell import range_boundaries

    keep: list[str] = []
    drop = 0
    for ref in merged:
        try:
            _, r1, _, r2 = range_boundaries(ref)
        except (ValueError, TypeError):
            continue
        if r2 is None or r1 is None:
            drop += 1
        elif r2 <= header_row or r1 > data_end:
            keep.append(ref)
        else:
            drop += 1
    return keep, drop


def _extras_message(extras: Counter, messages: list[Message]) -> None:
    items = [(k, v) for k, v in sorted(extras.items()) if v]
    if items:
        messages.append(Message("warning", "excel_extras_dropped",
                                "Excel extras not copied to the fresh synthetic workbook: "
                                + ", ".join(f"{k} ({v})" for k, v in items)))


def _select(names: list[str], settings: Any) -> list[str]:
    want = settings.sheet
    if want is None:
        return names
    if want not in names:  # the sheet names are never listed: a sheet can be named after a person (D34)
        raise ReadError(f"no sheet named {want!r} (the workbook has {len(names)} sheet(s))",
                        fix="check --sheet NAME")
    return [want]


# --------------------------------------------------------------------------- xlsx
def _load_workbook(fh: Any) -> Any:
    """``openpyxl.load_workbook`` (read-only, cached values). openpyxl fails on a chart sheet that has no
    relationships part (an empty chart sheet: ``AttributeError`` in ``read_chartsheet``); chart sheets hold no
    table, so the workbook is read again with such a chart sheet left empty."""
    import openpyxl

    try:
        return openpyxl.load_workbook(fh, read_only=True, data_only=True, keep_links=False)
    except AttributeError:
        from openpyxl.reader import excel as ox

        original = ox.ExcelReader.read_chartsheet

        def tolerant(self: Any, sheet: Any, rel: Any) -> None:
            try:
                original(self, sheet, rel)
            except AttributeError:
                return None

        ox.ExcelReader.read_chartsheet = tolerant
        try:
            fh.seek(0)
            return openpyxl.load_workbook(fh, read_only=True, data_only=True, keep_links=False)
        finally:
            ox.ExcelReader.read_chartsheet = original


def read_xlsx(req: "ReadRequest") -> ReadResult:
    import openpyxl

    path = Path(req.path)
    dtype_key = dtype_file_key(path)
    dtype_sheets: dict[str, dict[str, Any]] = {}
    settings = req.settings
    messages: list[Message] = []
    known_book = _WRITTEN.pop(path_key(path), None) if req.purpose == "selfcheck" else None
    ranges = parse_ranges(list(getattr(settings, "range", []) or []))
    try:
        zf = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as e:
        raise ReadError(f"not a readable .xlsx workbook ({type(e).__name__})",
                        fix="open it in Excel and save it as .xlsx") from None
    with zf:
        try:
            book = excel_book.read_book(zf)
        except (KeyError, SyntaxError, ValueError) as e:
            raise ReadError(f"not a readable .xlsx workbook (package part missing or invalid: {type(e).__name__})",
                            fix="open it in Excel and save it as .xlsx") from None
        worksheets = [s for s in book.sheets if s.kind == "worksheet"]
        selected = _select([s.name for s in worksheets], settings)
        extras = Counter(book.extras)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with open(path, "rb") as fh:
                wb = _load_workbook(fh)
                try:
                    tables, summaries = [], []
                    n_formulas = n_uncached = 0
                    for idx, part in enumerate(worksheets):
                        if part.name not in selected:
                            continue
                        known = known_book.get(part.name) if known_book else None
                        region = ranges.get(part.name, ranges.get(None))
                        scan = SheetScan(part.name, settings.header_row, region, known)
                        scan.collect_dtypes = True
                        try:
                            sx = excel_book.read_sheet_xml(zf, part.path) if part.path else SheetXml()
                            for r, cells in xlsx_rows(wb[part.name], sx):
                                scan.feed(r, cells)
                        except (KeyError, SyntaxError, zipfile.BadZipFile, EOFError, zlib.error) as e:
                            raise ReadError(f"sheet {part.name!r} is damaged ({type(e).__name__})",
                                            fix="open the workbook in Excel and save it as .xlsx") from None
                        extras.update(sx.extras)
                        n_formulas += sx.n_formulas
                        n_uncached += sx.n_formulas_uncached
                        checked = settings.header_row is None and region is None
                        t, summ = build_table(scan, sx, part.state, idx, messages, req.purpose, checked)
                        if scan.inferred_dtypes is not None:
                            dtype_sheets[part.name] = {"window": dtype_window(summ), "dtypes": scan.inferred_dtypes}
                        if checked and req.purpose != "selfcheck" and known is None:
                            check_header(t, part.name)
                        tables.append(t)
                        summaries.append(summ)
                finally:
                    wb.close()
    _unused_ranges(ranges, [s.name for s in worksheets], messages)
    kind = "xlsm" if req.rel_path.lower().endswith(".xlsm") or book.has_macros else "xlsx"
    active = book.sheets[book.active_tab].name if 0 <= book.active_tab < len(book.sheets) else None
    extra = _file_extra(kind, 1904 if book.date1904 else 1900, active, selected, summaries, extras, n_formulas)
    if dtype_sheets and dtype_key is not None and dtype_file_key(path) == dtype_key:
        import pandas as pd

        extra["_xlsx_dtypes"] = {"file": dtype_key, "infer_string": pd.options.future.infer_string,
                                 "sheets": dtype_sheets}
    if req.purpose != "selfcheck":
        _book_messages(messages, kind, extras, n_formulas, n_uncached, summaries, len(worksheets), selected,
                       settings)
    _ensure_visible(tables, summaries, messages, req.purpose)
    return ReadResult(FileLayout(format="xlsx", extra=extra), tables, messages)


def _unused_ranges(ranges: dict[str | None, Region], names: list[str], messages: list[Message]) -> None:
    for k in ranges:
        if k is not None and k not in names:
            messages.append(Message("warning", "excel_range_unused", f"--range for sheet {k!r}: no such sheet here"))


def _file_extra(kind: str, date_system: int, active_name: str | None, selected: list[str],
                summaries: list[dict[str, Any]], extras: Counter, n_formulas: int) -> dict[str, Any]:
    return {
        "workbook": kind,
        "date_system": date_system,
        "active_sheet": active_name if active_name in selected else None,
        "sheets": summaries,
        "extras_dropped": {k: int(v) for k, v in sorted(extras.items()) if v},
        "formula_cells": n_formulas,
    }


def _book_messages(messages: list[Message], kind: str, extras: Counter, n_formulas: int, n_uncached: int,
                   summaries: list[dict[str, Any]], n_sheets: int, selected: list[str], settings: Any) -> None:
    if kind == "xlsm":
        messages.append(Message("warning", "xlsm_macros",
                                "macro-enabled workbook: data read, macros (VBA) never copied; written as .xlsx"))
    _extras_message(Counter({k: v for k, v in extras.items() if k != "macros (VBA)"}), messages)
    if n_formulas:
        messages.append(Message("info", "excel_formulas",
                                f"{n_formulas} formula cell(s): cached values read, written as static values "
                                "(footer SUM/COUNT/AVERAGE recomputed from the synthetic rows)"))
        if n_uncached == n_formulas:
            messages.append(Message("warning", "excel_formulas_uncached",
                                    "no formula cell has a cached value (the workbook was never calculated): "
                                    "those cells read as empty",
                                    fix="open the workbook in Excel, save it, and re-run"))
    hidden = [s["name"] for s in summaries if s["state"] != "visible"]
    if hidden:
        messages.append(Message("info", "excel_hidden_sheets",
                                "hidden sheet(s) processed and written back hidden: "
                                + ", ".join(repr(h) for h in hidden)))
    if settings.sheet is not None and n_sheets > len(selected):
        messages.append(Message("warning", "excel_sheets_not_written",
                                f"--sheet: only sheet {settings.sheet!r} was read and written; the other "
                                f"{n_sheets - len(selected)} sheet(s) are not in the output"))
    empty = [s["name"] for s in summaries if s.get("notes_only") and not s.get("n_title_rows")]
    if empty:
        messages.append(Message("info", "excel_empty_sheets",
                                "empty sheet(s) written back empty: " + ", ".join(repr(e) for e in empty)))


def _ensure_visible(tables: list[RawTable], summaries: list[dict[str, Any]], messages: list[Message],
                    purpose: str) -> None:
    """A workbook needs one visible sheet: with only hidden sheets selected, the first is shown."""
    if tables and all(t.layout.get("state") != "visible" for t in tables):
        tables[0].layout["state"] = "visible"
        summaries[0]["state"] = "visible"
        if purpose != "selfcheck":
            messages.append(Message("warning", "excel_hidden_selected",
                                    f"only hidden sheets were selected; sheet {tables[0].name!r} is written visible"))


# --------------------------------------------------------------------------- xls (xlrd)
_XLRD_ERRORS = {0x00: "#NULL!", 0x07: "#DIV/0!", 0x0F: "#VALUE!", 0x17: "#REF!", 0x1D: "#NAME?", 0x24: "#NUM!",
                0x2A: "#N/A"}


def read_xls(req: "ReadRequest") -> ReadResult:
    import xlrd

    path = Path(req.path)
    settings = req.settings
    messages: list[Message] = []
    ranges = parse_ranges(list(getattr(settings, "range", []) or []))
    data = path.read_bytes()
    log = io.StringIO()  # xlrd writes its warnings to stdout by default (would corrupt --json output)
    try:
        try:
            book = xlrd.open_workbook(file_contents=data, formatting_info=True, logfile=log)
        except NotImplementedError:
            book = xlrd.open_workbook(file_contents=data, formatting_info=False, logfile=log)
    except Exception as e:  # noqa: BLE001 - xlrd raises many types for broken files
        raise ReadError(f"not a readable .xls workbook ({type(e).__name__})",
                        fix="open it in Excel and save it as .xlsx or .csv") from None
    names = book.sheet_names()
    selected = _select(names, settings)
    extras: Counter = Counter()
    for nm in getattr(book, "name_obj_list", []) or []:
        if not getattr(nm, "builtin", 0):
            extras["named ranges"] += 1
    tables, summaries = [], []
    for idx, name in enumerate(names):
        if name not in selected:
            continue
        sh = book.sheet_by_index(idx)
        sx = _xls_layout(sh)
        extras.update(sx.extras)
        state = {0: "visible", 1: "hidden", 2: "veryHidden"}.get(getattr(sh, "visibility", 0), "visible")
        region = ranges.get(name, ranges.get(None))
        scan = SheetScan(name, settings.header_row, region, None)
        fmt_cache: dict[int, str] = {}
        masked: Counter = Counter()
        for r in range(sh.nrows):
            scan.feed(r + 1, _xls_row(book, sh, r, fmt_cache, masked))
        extras.update(masked)
        checked = settings.header_row is None and region is None
        t, summ = build_table(scan, sx, state, idx, messages, req.purpose, checked)
        if checked and req.purpose != "selfcheck":
            check_header(t, name)
        tables.append(t)
        summaries.append(summ)
    _unused_ranges(ranges, names, messages)
    extra = _file_extra("xls", 1904 if book.datemode == 1 else 1900, names[0] if names else None, selected,
                        summaries, extras, 0)
    if req.purpose != "selfcheck":
        _book_messages(messages, "xls", extras, 0, 0, summaries, len(names), selected, settings)
        messages.append(Message("info", "xls_formulas_unknown",
                                ".xls formulas are not visible to the reader: cached values are used and footer "
                                "totals are left blank"))
    _ensure_visible(tables, summaries, messages, req.purpose)
    return ReadResult(FileLayout(format="xls", extra=extra), tables, messages)


def _xls_fmt(book: Any, xf_index: int, cache: dict[int, str], masked: Counter | None = None) -> str:
    f = cache.get(xf_index)
    if f is None:
        f = "General"
        try:
            xf = book.xf_list[xf_index]
            fo = book.format_map.get(xf.format_key)
            if fo is not None and fo.format_str:
                f = safe_number_format(fo.format_str) or "General"  # literal text masked (io.numfmt)
                if f != fo.format_str and masked is not None:
                    masked[MASKED_FORMAT_EXTRA] += 1
        except (IndexError, AttributeError):
            pass
        cache[xf_index] = f
    return f


def _xls_row(book: Any, sh: Any, r: int, cache: dict[int, str], masked: Counter | None = None) -> Row:
    import xlrd

    cells: Row = {}
    for c, cell in enumerate(sh.row(r)):
        ct = cell.ctype
        if ct in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            continue
        fmt = _xls_fmt(book, cell.xf_index, cache, masked) if cell.xf_index is not None else "General"
        v = cell.value
        if ct == xlrd.XL_CELL_TEXT:
            if v == "":
                continue
            cells[c + 1] = XCell(v, "s", fmt)
        elif ct == xlrd.XL_CELL_NUMBER:
            cells[c + 1] = XCell(int(v) if float(v).is_integer() and abs(v) < 2 ** 63 else v, "n", fmt)
        elif ct == xlrd.XL_CELL_DATE:
            try:
                if v < 1 and not _fmt_has_date(fmt):
                    d = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    pv: Any = d.time()
                elif fmt.startswith("[h]") or fmt.startswith("[H]"):
                    pv = dt.timedelta(days=v)
                else:
                    pv = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                cells[c + 1] = XCell(pv, "d", fmt)
            except Exception:  # noqa: BLE001 - an out-of-range date serial is kept as a number
                cells[c + 1] = XCell(v, "n", fmt)
        elif ct == xlrd.XL_CELL_BOOLEAN:
            cells[c + 1] = XCell(bool(v), "b", fmt)
        elif ct == xlrd.XL_CELL_ERROR:
            cells[c + 1] = XCell(_XLRD_ERRORS.get(v, "#N/A"), "e", fmt)
    return cells


def _xls_layout(sh: Any) -> SheetXml:
    sx = SheetXml()
    for (rlo, rhi, clo, chi) in getattr(sh, "merged_cells", []) or []:
        from openpyxl.utils import get_column_letter

        sx.merged.append(f"{get_column_letter(clo + 1)}{rlo + 1}:{get_column_letter(chi)}{rhi}")
    for colx, info in sorted((getattr(sh, "colinfo_map", {}) or {}).items()):
        width = getattr(info, "width", None)
        hidden = bool(getattr(info, "hidden", 0))
        if width or hidden:
            sx.widths.append([colx + 1, colx + 1, round(width / 256, 4) if width else None, hidden])
    if getattr(sh, "panes_are_frozen", 0):
        from openpyxl.utils import get_column_letter

        rows = int(getattr(sh, "horz_split_pos", 0) or 0)
        cols = int(getattr(sh, "vert_split_pos", 0) or 0)
        if rows or cols:
            sx.freeze_panes = f"{get_column_letter(cols + 1)}{rows + 1}"
    if getattr(sh, "hyperlink_list", None):
        sx.extras["hyperlinks"] += len(sh.hyperlink_list)
    if getattr(sh, "cell_note_map", None):
        sx.extras["comments"] += len(sh.cell_note_map)
    return sx
