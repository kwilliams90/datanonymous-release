"""Round-1 review fixes, ids/structure cluster (R1-03, R1-08, R1-09, R1-10).

* R1-03: synthetic IDs never equal a real ID of ANY namespace of the run (D9), and the leak scan checks the
  same run-wide set after writing.
* R1-08: ID columns with different headers whose keys are >= 99% contained in another table's ID column
  share one synthetic mapping, so cross-file and cross-sheet joins are kept (D5, D8).
* R1-09: a within-ID-constant numeric/date column keeps the column's exact min/max even when it is held only
  by a minority row of one ID (D14), and a run checks that every exact extreme was written.
* R1-10: within-ID constants follow whole-row blank/non-blank patterns too (D6): jointly per entity, and
  coupled to the row-level patterns.
"""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

from conftest import read_dicts, run_cli
from datanonymous.detect.ids import normalize_id
from datanonymous.generate import ids as ids_mod
from datanonymous.generate.ids import IdRegistry, build_mapping
from datanonymous.privacy.leaks import SensitiveIndex, _Namespace
from datanonymous.profile import structure as st
from datanonymous.profile import table as table_mod


def _write(path: Path, header: list[str], rows: list[list[object]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return path


def _file_report(rep: Path, rel: str) -> dict:
    return json.loads((rep / "files" / (rel + ".json")).read_text(encoding="utf-8"))


def _column(rep: Path, rel: str, name: str) -> dict:
    return next(c for c in _file_report(rep, rel)["tables"][0]["columns"] if c["name"] == name)


# --------------------------------------------------------------------------- R1-03: run-wide "never a real ID"
def test_synthetic_ids_avoid_real_ids_of_other_namespaces(tmp_path: Path) -> None:
    """The real staff IDs are exactly the synthetic student IDs a run without staff.csv writes."""
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[50000 + i, i] for i in range(20)])
    res = run_cli(d, "--out-dir", tmp_path / "o1", "--report-dir", tmp_path / "r1", "--seed", "7")
    assert res.exit_code == 0, res.output
    first = [r["student_id"] for r in read_dicts(tmp_path / "o1" / "people.csv")]
    _write(d / "staff.csv", ["staff_id", "score"], [[v, i] for i, v in enumerate(first)])
    res = run_cli(d, "--out-dir", tmp_path / "o2", "--report-dir", tmp_path / "r2", "--seed", "7")
    assert res.exit_code == 0, res.output
    people = {r["student_id"] for r in read_dicts(tmp_path / "o2" / "people.csv")}
    staff = {r["staff_id"] for r in read_dicts(tmp_path / "o2" / "staff.csv")}
    real = {str(50000 + i) for i in range(20)} | set(first)
    assert len(people) == 20 and len(staff) == 20
    assert not people & real and not staff & real
    assert all(len(v) == 5 and v.isdigit() for v in people | staff)  # same shape (D9)


def test_registry_excludes_every_namespace_but_keeps_mappings_separate() -> None:
    reg = IdRegistry()
    reg.register("a", "a.csv", None, 0, "a_id", np.array(["3", "5", "3"], dtype=object))
    reg.register("b", "b.csv", None, 0, "b_id", np.array(["1", "2", "4", "6"], dtype=object))
    for seed in range(20):
        reg.build(seed)
        a = set(reg.namespaces["a"].mapping.values())
        b = set(reg.namespaces["b"].mapping.values())
        assert len(a) == 2 and a <= {7, 8, 9}, seed  # only 7-9 are free in the whole run (0 is a last resort)
        assert b == {0, 7, 8, 9}, seed  # mappings stay separate: b may reuse a's synthetic values
    # the run-wide exclusion can make a shape too small: unpadded integers take another digit count (D9, noted;
    # K-02 harden 2 A2-6), zero-padded ones fail clearly
    reg.register("c", "c.csv", None, 0, "c_id", np.array(["8", "9"], dtype=object))
    reg.build(0)
    b_ns = reg.namespaces["b"]
    b = set(b_ns.mapping.values())
    assert b_ns.error is None and b_ns.notes and len(b) == 4 and not b & {1, 2, 3, 4, 5, 6, 8, 9}
    assert sum(len(str(v)) == 2 for v in b) == 2
    reg.register("d", "d.csv", None, 0, "d_id", np.array([f"{i:02d}" for i in range(100)], dtype=object))
    reg.build(0)
    assert reg.namespaces["d"].error is not None and "cannot avoid all real IDs" in reg.namespaces["d"].error.message


def test_row_number_namespace_is_not_excluded_and_stays_identity() -> None:
    """Kept row numbers 1..n are not person IDs: they neither shrink other namespaces' space nor get remapped."""
    reg = IdRegistry()
    reg.register("id@a.csv", "a.csv", None, 0, "id", np.array([str(i) for i in range(1, 10)], dtype=object))
    reg.register("staffid", "b.csv", None, 0, "staff_id", np.array(["3", "5"], dtype=object))
    reg.build(1)
    assert reg.namespaces["id@a.csv"].identity and reg.namespaces["id@a.csv"].error is None
    ns = reg.namespaces["staffid"]
    assert ns.error is None and not set(ns.mapping.values()) & {3, 5}


def test_pattern_ids_avoid_other_namespaces() -> None:
    """S00-S29 and S30-S59 are real: both namespaces draw from S60-S99 only (a nearly full shape)."""
    reg = IdRegistry()
    reg.register("a", "a.csv", None, 0, "a_id", np.array([f"S{i:02d}" for i in range(30)], dtype=object))
    reg.register("b", "b.csv", None, 0, "b_id", np.array([f"S{i:02d}" for i in range(30, 60)], dtype=object))
    free = {f"S{i:02d}" for i in range(60, 100)}
    for seed in range(5):
        reg.build(seed)
        for key in ("a", "b"):
            ns = reg.namespaces[key]
            assert ns.error is None and len(set(ns.mapping.values())) == 30
            assert set(ns.mapping.values()) <= free, (seed, key)


def test_build_mapping_without_run_keys_is_unchanged() -> None:
    """A single namespace (no other real IDs) maps exactly as before the run-wide exclusion."""
    reg = IdRegistry()
    reg.register("a", "a.csv", None, 0, "a_id", np.array(["0012345", "0023456", "0034567"], dtype=object))
    ns = reg.namespaces["a"]
    build_mapping(ns, np.random.default_rng(5))
    alone = dict(ns.mapping)
    build_mapping(ns, np.random.default_rng(5), run_keys=set(ns.representatives),
                  run_int=np.asarray(sorted(int(k) for k in ns.representatives), dtype=np.int64))
    assert ns.mapping == alone


def test_leak_scan_checks_real_ids_of_every_namespace() -> None:
    idx = SensitiveIndex()
    idx.namespaces["studentid"] = _Namespace(exact={"50001"}, normalized={"50001"}, origins={0})
    idx.namespaces["staffid"] = _Namespace(exact={"0016192"}, normalized={"16192"}, origins={1})
    idx.namespaces["id@a.csv"] = _Namespace(exact=set(), normalized=set(), origins=set(), identity=True)
    n, origins = idx.real_id_hits({"16192", "77777", "50001"})
    assert n == 2 and origins == {0, 1}
    assert idx.real_id_hits({"77777"}) == (0, set())


def test_leak_scan_fails_a_file_whose_ids_equal_another_namespaces_real_ids(tmp_path: Path, monkeypatch) -> None:
    """Defence in depth: with the allocation fix undone, the final scan FAILs the file instead of writing it."""
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[50000 + i, i] for i in range(20)])
    res = run_cli(d, "--out-dir", tmp_path / "o1", "--report-dir", tmp_path / "r1", "--seed", "7")
    assert res.exit_code == 0, res.output
    first = [r["student_id"] for r in read_dicts(tmp_path / "o1" / "people.csv")]
    _write(d / "staff.csv", ["staff_id", "score"], [[v, i] for i, v in enumerate(first)])
    orig = ids_mod.build_mapping
    monkeypatch.setattr(ids_mod, "build_mapping",
                        lambda ns, rng, id_check=None, run_keys=None, run_int=None, **_: orig(ns, rng, id_check))
    res = run_cli(d, "--out-dir", tmp_path / "o2", "--report-dir", tmp_path / "r2", "--seed", "7")
    assert res.exit_code == 3, res.output
    assert not (tmp_path / "o2" / "people.csv").exists()
    assert _file_report(tmp_path / "r2", "people.csv")["status"] == "FAIL"
    text = (tmp_path / "r2" / "REPORT.md").read_text(encoding="utf-8")
    assert not any(v in text for v in first)


# --------------------------------------------------------------------------- R1-08: key discovery across headers
def _keys_fixture(d: Path) -> None:
    d.mkdir(parents=True, exist_ok=True)
    (d / "people.csv").write_text("student_id,first_name\n0012345,Zelmira\n0023456,Balthazar\n0034567,Mirelda\n",
                                  encoding="utf-8")
    (d / "events.csv").write_text("StudentNumber,score\n12345,11\n12345,12\n23456,13\n34567,14\n", encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["student_id", "score"])
    for i, v in enumerate(["0012345", "0023456", "0034567"]):
        ws.append([v, i])
    ws = wb.create_sheet("Events")
    ws.append(["StudentNumber", "score"])
    for i, v in enumerate(["12345", "12345", "23456", "34567"]):
        ws.append([v, i])
    wb.save(d / "book.xlsx")


def test_differently_named_id_columns_keep_cross_file_and_cross_sheet_joins(tmp_path: Path) -> None:
    d = tmp_path / "in"
    _keys_fixture(d)
    res = run_cli(d, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "123")
    assert res.exit_code == 0, res.output
    real = {"12345", "23456", "34567"}
    people = [r["student_id"] for r in read_dicts(tmp_path / "o" / "people.csv")]
    events = [r["StudentNumber"] for r in read_dicts(tmp_path / "o" / "events.csv")]
    assert all(len(v) == 7 and v.startswith("00") for v in people)  # each file keeps its own raw form (D8)
    assert {normalize_id(v) for v in events} == {normalize_id(v) for v in people}  # 100% join, as in the source
    assert Counter(Counter(map(normalize_id, events)).values()) == Counter({1: 2, 2: 1})
    wb = load_workbook(tmp_path / "o" / "book.xlsx")
    sp = {normalize_id(str(r[0])) for r in wb["People"].iter_rows(min_row=2, values_only=True)}
    se = {normalize_id(str(r[0])) for r in wb["Events"].iter_rows(min_row=2, values_only=True)}
    assert sp == se == {normalize_id(v) for v in people}
    assert not ({normalize_id(v) for v in people} & real)
    cov = _file_report(tmp_path / "r", "events.csv")["tables"][0]["structure"]["cross_file_keys"]
    assert cov and cov[0]["column"] == "StudentNumber" and cov[0]["key_containment"] == 1.0
    msgs = [m for m in _column(tmp_path / "r", "events.csv", "StudentNumber")["messages"] if m["code"] == "id_key_link"]
    assert msgs and "people.csv::student_id" in msgs[0]["text"]


def test_key_link_can_be_disabled(tmp_path: Path) -> None:
    d = tmp_path / "in"
    _keys_fixture(d)
    cfg = tmp_path / "c.yaml"
    cfg.write_text("structure:\n  disable: [cross_file_keys]\n", encoding="utf-8")
    res = run_cli(d, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "123", "--config", cfg)
    assert res.exit_code == 0, res.output
    people = {normalize_id(r["student_id"]) for r in read_dicts(tmp_path / "o" / "people.csv")}
    events = {normalize_id(r["StudentNumber"]) for r in read_dicts(tmp_path / "o" / "events.csv")}
    assert not people & events


def test_id_col_fix_for_an_unlinked_key_column_now_keeps_the_join(tmp_path: Path) -> None:
    """The fix the unlinked-key warning prints (--id-col) now also joins the files."""
    rng = np.random.default_rng(9)
    ids = rng.choice(np.arange(1_000_000, 9_999_999), 80, replace=False).tolist()
    d = tmp_path / "in"
    _write(d / "students.csv", ["student_id", "grade"], [[i, rng.choice(["01", "02"])] for i in ids])
    _write(d / "visits.csv", ["learner_ref", "score"],
           [[int(rng.choice(ids)), int(rng.integers(40, 101))] for _ in range(300)])
    res = run_cli(d, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "3",
                  "--id-col", "visits.csv::learner_ref")
    assert res.exit_code == 0, res.output
    stu = {r["student_id"] for r in read_dicts(tmp_path / "o" / "students.csv")}
    vis = [r["learner_ref"] for r in read_dicts(tmp_path / "o" / "visits.csv")]
    assert set(vis) <= stu and not stu & {str(i) for i in ids}
    real_counts = Counter(Counter(r["learner_ref"] for r in read_dicts(d / "visits.csv")).values())
    assert Counter(Counter(vis).values()) == real_counts


def test_row_number_ids_are_never_linked_by_content(tmp_path: Path) -> None:
    """A person-ID column whose values fall inside another file's kept row numbers is not linked (its real IDs
    would be written as they are)."""
    d = tmp_path / "in"
    _write(d / "courses.csv", ["id", "name_code"], [[i, f"C{i % 4}"] for i in range(1, 41)])
    _write(d / "visits.csv", ["student_id", "score"], [[(i * 7) % 30 + 3, i % 50] for i in range(120)])
    res = run_cli(d, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "4")
    assert res.exit_code == 0, res.output
    assert [r["id"] for r in read_dicts(tmp_path / "o" / "courses.csv")] == [str(i) for i in range(1, 41)]
    real = {r["student_id"] for r in read_dicts(d / "visits.csv")}
    syn = {r["student_id"] for r in read_dicts(tmp_path / "o" / "visits.csv")}
    assert not real & syn
    msgs = [m for m in _column(tmp_path / "r", "visits.csv", "student_id")["messages"] if m["code"] == "id_key_link"]
    assert not msgs


def test_same_table_id_columns_are_not_linked_by_content(tmp_path: Path) -> None:
    """Two ID columns of ONE table are aliases/nesting/unrelated (find_aliases_and_nesting), never key-linked."""
    rows = [[100000 + 7 * i, 100000 + 7 * ((i * 3) % 40)] for i in range(40)]
    src = _write(tmp_path / "in" / "mentors.csv", ["student_id", "mentor_ref"], rows)
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "2",
                  "--id-col", "mentor_ref")
    assert res.exit_code == 0, res.output
    msgs = [m for m in _column(tmp_path / "r", "mentors.csv", "mentor_ref")["messages"] if m["code"] == "id_key_link"]
    assert not msgs


def test_contained_share_stops_early_and_is_exact() -> None:
    keys = {str(i) for i in range(1000)}
    assert st._contained(keys, [keys], 0.99) == 1.0
    assert st._contained(keys, [{str(i) for i in range(990)}], 0.99) == 0.99
    assert st._contained(keys, [{str(i) for i in range(989)}], 0.99) is None
    assert st._contained(keys, [{str(i) for i in range(500)}, {str(i) for i in range(500, 995)}], 0.99) == 0.995
    assert st._contained(set(), [keys], 0.99) is None and st._contained(keys, [], 0.99) is None


# --------------------------------------------------------------------------- R1-09: entity reduction keeps extremes
def _extreme_fixture(path: Path) -> Path:
    rows = []
    for i in range(100):
        for j in range(2):
            odd = i == 99 and j == 1
            rows.append([f"S{i:05d}", "999.99" if odd else f"{i + 1}.25",
                         "2025-12-31" if odd else f"2024-01-{i % 28 + 1:02d}",
                         "$9,999.50" if odd else f"${i + 10:,}.50"])
    return _write(path, ["student_id", "measurement", "observed_date", "fee"], rows)


def test_within_id_constant_keeps_row_level_extremes(tmp_path: Path) -> None:
    src = _extreme_fixture(tmp_path / "in" / "entity_extreme.csv")
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "123")
    assert res.exit_code == 0, res.output
    syn = read_dicts(tmp_path / "o" / "entity_extreme.csv")
    real = read_dicts(src)
    for col in ("measurement", "observed_date", "fee"):
        assert _column(tmp_path / "r", "entity_extreme.csv", col)["role"] == "entity", col
        per = {}
        for s in syn:
            per.setdefault(s["student_id"], set()).add(s[col])
        assert all(len(v) == 1 for v in per.values()), col  # still constant within each ID (D4)
        assert _column(tmp_path / "r", "entity_extreme.csv", col)["flags"]["extreme_value_adjusted"], col
    vals = {s["measurement"] for s in syn}
    assert "999.99" in vals and "1.25" in vals
    assert max(s["observed_date"] for s in syn) == "2025-12-31" == max(r["observed_date"] for r in real)
    assert min(s["observed_date"] for s in syn) == "2024-01-01"
    assert "$9,999.50" in {s["fee"] for s in syn} and "$10.50" in {s["fee"] for s in syn}
    codes = [m["code"] for m in _file_report(tmp_path / "r", "entity_extreme.csv")["messages"]]
    assert "extremes_not_emitted" not in codes


def test_missing_extremes_are_reported(tmp_path: Path, monkeypatch) -> None:
    """The run checks the written values for the exact min/max, independently of how they were generated."""
    src = _extreme_fixture(tmp_path / "in" / "entity_extreme.csv")
    monkeypatch.setattr(table_mod, "floor_entity_extremes", lambda p, emodel: 0)
    from datanonymous.generate import numeric as numeric_gen
    monkeypatch.setattr(numeric_gen, "_floor_row_extremes", lambda profile, model, exact: (model, False))
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "123")
    assert res.exit_code == 0, res.output
    rep = _file_report(tmp_path / "r", "entity_extreme.csv")
    flagged = {m["column"] for m in rep["messages"] if m["code"] == "extremes_not_emitted"}
    assert flagged == {"measurement", "observed_date", "fee"}
    assert rep["status"] == "WARN"


def test_no_exact_extremes_turns_the_entity_floor_off(tmp_path: Path) -> None:
    src = _extreme_fixture(tmp_path / "in" / "entity_extreme.csv")
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "123",
                  "--no-exact-extremes", "measurement")
    assert res.exit_code == 0, res.output
    assert "999.99" not in {s["measurement"] for s in read_dicts(tmp_path / "o" / "entity_extreme.csv")}
    codes = [m["code"] for m in _file_report(tmp_path / "r", "entity_extreme.csv")["messages"]]
    assert "extremes_not_emitted" not in codes


# --------------------------------------------------------------------------- R1-10: whole-row missingness with ID constants
def _pairs(rows: list[dict], a: str, b: str) -> Counter:
    return Counter((bool(r[a].strip()), bool(r[b].strip())) for r in rows)


def test_within_id_constants_stay_blank_together(tmp_path: Path) -> None:
    src = _write(tmp_path / "in" / "entity_missing.csv", ["student_id", "gender", "birth_date", "score"],
                 [[f"S{i:05d}", "F" if i < 10 else "", f"200{i % 10}-01-01" if i < 10 else "", j]
                  for i in range(20) for j in range(3)])
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "123")
    assert res.exit_code == 0, res.output
    syn = read_dicts(tmp_path / "o" / "entity_missing.csv")
    assert _pairs(syn, "gender", "birth_date") == _pairs(read_dicts(src), "gender", "birth_date") \
        == Counter({(True, True): 30, (False, False): 30})
    mm = _file_report(tmp_path / "r", "entity_missing.csv")["tables"][0]["structure"]["missingness"]
    assert mm["entity_columns"] == ["gender", "birth_date"] and mm["n_entity_patterns"] == 2


def _cross_fixture(path: Path, sizes: list[int], seed: int = 0) -> Path:
    """Entity constants (gender, dob) blank together; when they are blank, the row-level score is blank too."""
    rng = np.random.default_rng(seed)
    rows = []
    for i, k in enumerate(sizes):
        blank = i % 2 == 0
        g = "" if blank else str(rng.choice(["F", "M"]))
        dob = "" if blank else f"2010-0{int(rng.integers(1, 10))}-1{int(rng.integers(0, 10))}"
        for _ in range(k):
            rows.append([f"{700000 + 37 * i}", g, dob, "" if blank else int(rng.integers(200, 400)),
                         str(rng.choice(["A", "B", ""]))])
    return _write(path, ["student_id", "gender", "dob", "score", "period"], rows)


def test_entity_and_row_missingness_are_coupled(tmp_path: Path) -> None:
    src = _cross_fixture(tmp_path / "in" / "cross.csv", [3] * 60)
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "5")
    assert res.exit_code == 0, res.output
    real, syn = read_dicts(src), read_dicts(tmp_path / "o" / "cross.csv")
    assert _column(tmp_path / "r", "cross.csv", "gender")["role"] == "entity"
    assert _column(tmp_path / "r", "cross.csv", "score")["role"] == "row"
    assert _pairs(syn, "gender", "dob") == _pairs(real, "gender", "dob")
    assert _pairs(syn, "gender", "score") == _pairs(real, "gender", "score")  # equal block sizes: exact
    for col in ("score", "period"):  # row-level per-column missing counts stay exact
        assert sum(not s[col] for s in syn) == sum(not r[col] for r in real), col


def test_coupling_with_uneven_blocks_keeps_totals_and_most_pairs(tmp_path: Path) -> None:
    sizes = [int(k) for k in np.random.default_rng(3).integers(1, 7, 120)]
    src = _cross_fixture(tmp_path / "in" / "cross.csv", sizes, seed=3)
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "8")
    assert res.exit_code == 0, res.output
    real, syn = read_dicts(src), read_dicts(tmp_path / "o" / "cross.csv")
    assert sum(not s["score"] for s in syn) == sum(not r["score"] for r in real)
    ents = lambda rows: Counter(bool(r["gender"]) for r in {r["student_id"]: r for r in rows}.values())  # noqa: E731
    assert ents(syn) == ents(real)
    broken = sum(bool(s["gender"]) != bool(s["score"]) for s in syn)
    assert broken <= 0.1 * len(syn)  # independent draws would break about half of them


def test_rare_entity_patterns_are_pooled(tmp_path: Path) -> None:
    rows = []
    for i in range(60):
        g = "" if i < 3 else "F"  # 3 entities (< min-cell 10) with gender blank: a pooled pattern
        dob = "" if 3 <= i < 6 else "2011-02-03"
        rows += [[f"{900000 + i}", g, dob, j] for j in range(2)]
    src = _write(tmp_path / "in" / "pool.csv", ["student_id", "gender", "dob", "visit"], rows)
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "2")
    assert res.exit_code == 0, res.output
    syn = read_dicts(tmp_path / "o" / "pool.csv")
    per = {s["student_id"]: s for s in syn}
    assert sum(not s["gender"] for s in per.values()) == 3 and sum(not s["dob"] for s in per.values()) == 3
    mm = _file_report(tmp_path / "r", "pool.csv")["tables"][0]["structure"]["missingness"]
    assert mm["n_pooled_entity_patterns"] == 2 and mm["pooled_entities"] == 6


def test_na_token_held_by_a_minority_row_is_floored_through_the_entity_mask(tmp_path: Path) -> None:
    rows = []
    for i in range(100):
        for j in range(2):
            rows.append([f"S{i:05d}", "NA" if (i == 42 and j == 1) else f"2011-03-{i % 28 + 1:02d}", j])
    src = _write(tmp_path / "in" / "na.csv", ["student_id", "dob", "visit"], rows)
    res = run_cli(src, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "6")
    assert res.exit_code == 0, res.output
    syn = read_dicts(tmp_path / "o" / "na.csv")
    assert _column(tmp_path / "r", "na.csv", "dob")["role"] == "entity"
    assert "NA" in {s["dob"] for s in syn}  # D14: every NA token kept at least once
    assert _column(tmp_path / "r", "na.csv", "dob")["flags"]["rare_category_adjusted"]


def test_coupled_codes_keep_exact_totals() -> None:
    for seed in range(30):
        rng = np.random.default_rng(seed)
        g, r = int(rng.integers(2, 5)), int(rng.integers(2, 6))
        cross = rng.integers(0, 30, (g, r)).astype(float)
        cross[cross < 10] = 0
        groups = rng.integers(0, g, int(rng.integers(5, 400)))
        alloc = np.bincount(rng.integers(0, r, len(groups)), minlength=r)
        codes = st._coupled_codes(cross, groups, alloc, np.random.default_rng(seed))
        assert np.array_equal(np.bincount(codes, minlength=r), alloc), seed


def test_runs_with_entity_patterns_are_deterministic(tmp_path: Path) -> None:
    sizes = [int(k) for k in np.random.default_rng(4).integers(1, 5, 50)]
    src = _cross_fixture(tmp_path / "in" / "cross.csv", sizes, seed=4)
    outs = []
    for k in range(2):
        res = run_cli(src, "--out-dir", tmp_path / f"o{k}", "--report-dir", tmp_path / f"r{k}", "--seed", "11")
        assert res.exit_code == 0, res.output
        outs.append((tmp_path / f"o{k}" / "cross.csv").read_bytes())
    assert outs[0] == outs[1]


def test_entity_mask_without_entity_columns_is_empty() -> None:
    from datanonymous.model import MissingnessModel

    codes, mask = st.sample_entity_mask(MissingnessModel(), 4, np.random.default_rng(0))
    assert codes is None and mask.shape == (4, 0)
