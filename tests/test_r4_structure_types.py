"""Round 4, structure and types: discrete numbers as composite keys (R4-04), header rules over native typing (R4-05).

R4-04: a unique (student_id, visit) key with ``visit`` a small discrete integer was never found, because
INTEGER was missing from the composite-key candidate types; the output repeated (student, visit) pairs without
a warning. Candidates are now every column whose synthetic values keep each real value's count (labels,
booleans, years, codes, discrete integers: the generator's exact-quota test), plus masked codes, invented
labels and dates; continuous numbers never are. Large tables reject non-keys on the first IDs' rows (two
stages) before any full pass, so the extra candidates cost little; pair candidates keep labels, codes and dates
first so discrete numbers never push a (course, term) pair out of the pair budget.

R4-05: a native Arrow ``decimal128(10, 0)`` ``school_id`` was typed DECIMAL before the D10 org-code header
rule ran, so it was interpolated (new school codes). Native storage now decides only what the values cannot
show: numbers of every storage (int, float, decimal, Excel, Stata) go through the value rules, header rules
first, and a native factor (Arrow dictionary, Stata value labels) under an org-code header is an ORG_CODE as it
is in text. ID headers stay overrides and PII headers the D12 guard, for native and text columns alike. The
writer restores the native storage (decimal, dictionary, value labels) from the values.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
from collections import Counter
from decimal import Decimal
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.feather as pf
import pyarrow.parquet as pq
import pytest
from click.testing import CliRunner

from datanonymous.cli import main
from datanonymous.config import Settings
from datanonymous.detect.context import DetectContext
from datanonymous.detect.types import detect_column
from datanonymous.io import csvio
from datanonymous.model import ColumnType
from datanonymous.profile import structure as st
from datanonymous.profile.table import profile_table

T = ColumnType
NAMES = ["Althea Voss", "Ysabel Quispe", "Ravi Kline"]


# --------------------------------------------------------------------------- helpers
def _write_csv(path: Path, header: list[str], rows: list[list]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    return path


def _run(src: Path, tmp_path: Path, *flags: str, seed: int = 42) -> tuple[Path, dict]:
    out, rep = tmp_path / "out", tmp_path / "rep"
    res = CliRunner().invoke(main, ["--mode", "fidelity", str(src), "--out-dir", str(out), "--report-dir", str(rep), "--seed", str(seed),
                                    "--force", *flags], catch_exceptions=False)
    assert res.exit_code == 0, res.output
    report = json.loads((rep / "files" / f"{src.name}.json").read_text(encoding="utf-8"))
    return out / src.name, report


def _types(report: dict) -> dict[str, str]:
    return {c["name"]: c["type"] for c in report["tables"][0]["columns"]}


def _keys(report: dict) -> list[tuple[str, ...]]:
    return [tuple(k["columns"]) for k in report["tables"][0]["structure"].get("composite_keys", [])]


def _profile(text: str, name: str = "t.csv"):
    r = csvio.parse_text_bytes(text.encode("utf-8"))
    return profile_table(r.tables[0], name, "csv", Settings(mode="fidelity"))


def _csv_text(header: list[str], rows: list[list]) -> str:
    return "\n".join(",".join(str(v) for v in r) for r in [header, *rows]) + "\n"


def _dups(rows: list[dict], cols: tuple[str, ...]) -> int:
    filled = [tuple(r[c] for c in cols) for r in rows if all(r[c].strip() for c in cols)]
    return len(filled) - len(set(filled))


# =========================================================================== R4-04: discrete numbers as keys
def test_codex_numeric_visit_reproduction(tmp_path: Path) -> None:
    """60 students x visits 1-4 (each once) plus a score: (student_id, visit) stays unique in the output."""
    src = _write_csv(tmp_path / "in" / "numeric_visit.csv", ["student_id", "visit", "score"],
                     [[f"S{i + 1000}", v, 30 + (i * 17 + v * 3) % 60] for i in range(60) for v in range(1, 5)])
    out, report = _run(src, tmp_path)
    assert ("student_id", "visit") in _keys(report)
    assert _types(report)["visit"] == "INTEGER"
    with out.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 240 and _dups(rows, ("student_id", "visit")) == 0
    assert Counter(r["visit"] for r in rows) == Counter({str(v): 60 for v in range(1, 5)})  # exact counts kept


def _visit_rows(form: str) -> tuple[list[str], list[list]]:
    """(header, rows) of one discrete-number key variant; the key columns are the ID plus the rest of ``keys``."""
    rows: list[list] = []
    if form == "float_text":  # numbers exported through a float dtype: 1.0, 2.0 ...
        rows = [[f"S{i + 1000}", f"{v}.0", (i * 7 + v) % 40] for i in range(50) for v in range(1, 5)]
        return ["student_id", "visit", "score"], rows
    if form == "blanks":  # a key over the filled rows: some visits are not recorded
        rows = [[f"S{i + 1000}", "" if (i + v) % 9 == 0 else v, (i * 7 + v) % 40] for i in range(50)
                for v in range(1, 5)]
        return ["student_id", "visit", "score"], rows
    if form == "period_pair":  # attendance by (date, period): neither alone is a key
        rows = [[f"S{i + 1000}", f"2025-01-{d:02d}", p, "P" if (i + d + p) % 5 else "A"]
                for i in range(40) for d in range(6, 11) for p in range(1, 4)]
        return ["student_id", "date", "period", "status"], rows
    if form == "term_pair":  # grades by (course, semester): semester numbers 1-2 (no org-code header), course labels
        rows = [[f"S{i + 1000}", ["Algebra", "Biology", "Chemistry"][c], t, "ABCDF"[(i + c) % 5]]
                for i in range(40) for c in range(3) for t in (1, 2)]
        return ["student_id", "course", "semester", "mark"], rows
    if form == "boolean":  # at most one primary and one secondary enrolment per student
        rows = [[f"S{i + 1000}", ["Y", "N"][k], "ABCDF"[(i * 3 + k * 2) % 5]] for i in range(60) for k in range(2)]
        return ["student_id", "is_primary", "mark"], rows
    raise AssertionError(form)


@pytest.mark.parametrize("form, key", [
    ("float_text", ("student_id", "visit")),
    ("blanks", ("student_id", "visit")),
    ("period_pair", ("student_id", "date", "period")),
    ("term_pair", ("student_id", "course", "semester")),
    ("boolean", ("student_id", "is_primary")),
])
def test_discrete_number_keys_found_and_kept(tmp_path: Path, form: str, key: tuple[str, ...]) -> None:
    header, rows = _visit_rows(form)
    src = _write_csv(tmp_path / "in" / f"{form}.csv", header, rows)
    out, report = _run(src, tmp_path)
    assert key in _keys(report), _keys(report)
    with out.open(newline="", encoding="utf-8") as fh:
        syn = list(csv.DictReader(fh))
    assert len(syn) == len(rows)
    col = key[-1]
    assert Counter(r[col] for r in syn) == Counter(str(r[header.index(col)]) for r in rows)  # marginal unchanged
    if len(key) == 2:
        assert _dups(syn, key) == 0
    else:  # (ID, A, B): A and B are drawn independently (D3), so a repeat the tuple moves cannot fix is warned
        codes = [m["code"] for t in report["tables"] for m in t.get("messages", [])] + [
            m["code"] for m in report.get("messages", [])]
        assert _dups(syn, key) == 0 or "composite_key_duplicates" in codes


@pytest.mark.parametrize("fmt", ["parquet", "feather", "xlsx", "dta"])
def test_native_integer_storage_visit_key(tmp_path: Path, fmt: str) -> None:
    """The same key whatever stores the numbers: Arrow int64, Excel number cells, Stata int."""
    ids = [f"S{i + 1000}" for i in range(50) for _ in range(4)]
    visits = [v for _ in range(50) for v in range(1, 5)]
    scores = [(i * 13) % 37 for i in range(200)]
    src = tmp_path / "in" / f"visits.{fmt}"
    src.parent.mkdir(parents=True)
    if fmt in ("parquet", "feather"):
        tab = pa.table({"student_id": ids, "visit": pa.array(visits, pa.int64()),
                        "score": pa.array(scores, pa.int16())})
        (pq.write_table if fmt == "parquet" else pf.write_feather)(tab, src)
    elif fmt == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["student_id", "visit", "score"])
        for row in zip(ids, visits, scores):
            ws.append(list(row))
        wb.save(src)
    else:
        pd.DataFrame({"student_id": ids, "visit": visits, "score": scores}).to_stata(src, write_index=False,
                                                                                     version=118)
    out, report = _run(src, tmp_path)
    assert _types(report)["visit"] == "INTEGER" and ("student_id", "visit") in _keys(report)
    if fmt == "parquet":
        syn = pq.read_table(out).to_pandas()
    elif fmt == "feather":
        syn = pf.read_table(out).to_pandas()
    elif fmt == "xlsx":
        syn = pd.read_excel(out)
    else:
        syn = pd.read_stata(out)
    assert len(syn) == 200 and not syn.duplicated(["student_id", "visit"]).any()


def test_continuous_numbers_are_never_keys() -> None:
    """Interpolated numbers (floats, continuous integers) are not key candidates even when unique within IDs."""
    rows = [[f"S{i + 1000}", f"{(i * 4 + v) * 1.37:.2f}", (i * 4 + v) * 1009 + 3, f"2025-02-{v + 1:02d}"]
            for i in range(30) for v in range(4)]
    tp = _profile(_csv_text(["student_id", "amount", "big_count", "visit_date"], rows))
    types = {p.name: (p.type, p.subtype) for p in tp.columns}
    assert types["amount"][0] is T.FLOAT and types["big_count"] == (T.INTEGER, "continuous")
    keys = [tuple(k["columns"]) for k in tp.structure.composite_keys]
    assert keys == [("student_id", "visit_date")]


def test_measurement_integers_are_not_keys_but_long_visit_numbers_are() -> None:
    """Scores (more values than factor-max and any ID's rows) unique within IDs by coincidence are no key; visit
    numbers 1..30 (more values than factor-max, as many as the longest ID's rows) are."""
    rng = np.random.default_rng(3)
    rows = [[f"S{i + 1000}", v, int(rng.integers(0, 101))] for i in range(40) for v in range(1, 4)]
    rows = [r for r in rows if len({x[2] for x in rows if x[0] == r[0]}) == 3]  # every score unique within its ID
    tp = _profile(_csv_text(["student_id", "visit", "score"], rows))
    assert {p.name: p.type for p in tp.columns}["score"] is T.INTEGER
    assert [tuple(k["columns"]) for k in tp.structure.composite_keys] == [("student_id", "visit")]
    long_rows = [[f"S{i + 1000}", v, ["a", "b"][(i + v) % 2]] for i in range(12) for v in range(1, 31)]
    tp = _profile(_csv_text(["student_id", "session", "mark"], long_rows))
    assert [tuple(k["columns"]) for k in tp.structure.composite_keys] == [("student_id", "session")]


def test_codex_pair_with_a_score_is_not_a_key(tmp_path: Path) -> None:
    """privacy/school/enrollments.csv: (course, score) is unique within students only by coincidence."""
    src = Path(__file__).parent / "fixtures" / "privacy" / "school" / "enrollments.csv"
    r = csvio.parse_text_bytes(src.read_bytes())
    tp = profile_table(r.tables[0], "enrollments.csv", "csv", Settings(mode="fidelity"))
    assert not any("score" in k["columns"] for k in tp.structure.composite_keys)


def test_candidate_types_follow_the_generators_exact_quota() -> None:
    """Every level-generated type is a candidate; continuous numbers and never-row columns are not."""
    def p(t: ColumnType, sub: str | None = None, **params) -> st.ColumnProfile:
        return st.ColumnProfile(index=0, name="c", type=t, subtype=sub, params=params, role="row")

    yes = [p(T.FACTOR), p(T.ORG_CODE), p(T.ORDINAL), p(T.SCHOOL_YEAR), p(T.BOOLEAN), p(T.INTEGER_YEAR),
           p(T.INTEGER, "discrete"), p(T.DIGIT_CODE, levels_mode=True), p(T.DIGIT_CODE, levels_mode=False),
           p(T.HIGH_CARD_CATEGORICAL), p(T.DATE), p(T.DATETIME), p(T.TIME)]
    no = [p(T.INTEGER, "continuous"), p(T.FLOAT), p(T.DECIMAL), p(T.NUMERIC_FORMATTED), p(T.NUMERIC_WITH_TOKENS),
          p(T.LAT), p(T.LON), p(T.FREE_TEXT), p(T.EMAIL), p(T.PERSON_NAME),
          p(T.INTEGER, "discrete", id_candidate="unique sequential integers")]
    assert all(st._composite_candidate(c) for c in yes)
    assert not any(st._composite_candidate(c) for c in no)
    entity = p(T.INTEGER, "discrete")
    entity.role = "entity"  # constant within the ID: never part of a key
    assert not st._composite_candidate(entity)


def test_pairs_keep_label_candidates_first() -> None:
    """Twelve non-key discrete-number columns before (course, term) never push the label pair out of the budget."""
    nums = [f"n{k}" for k in range(st.MAX_PAIR_CANDIDATES)]
    rows = [[f"S{i + 1000}", *[(i + k + c) % 3 for k in range(len(nums))], ["Art", "Bio"][c], ["Fall", "Spring"][t]]
            for i in range(30) for c in range(2) for t in range(2)]
    tp = _profile(_csv_text(["student_id", *nums, "course", "term"], rows))
    assert all(tp.columns[i + 1].type is T.INTEGER for i in range(len(nums)))
    assert ("student_id", "course", "term") in [tuple(k["columns"]) for k in tp.structure.composite_keys]


def _random_table(seed: int) -> str:
    rng = np.random.default_rng(seed)
    n_ids = int(rng.integers(20, 60))
    rows = []
    for i in range(n_ids):
        k = int(rng.integers(1, 6))
        for v in range(k):
            rows.append([f"S{i + 1000}", v + 1 if rng.random() > 0.02 else int(rng.integers(1, 6)),
                         int(rng.integers(0, 4)), ["a", "b", "c", "d", "e"][int(rng.integers(0, 5))],
                         int(rng.integers(0, 50)), (v + i) % 7])
    order = rng.permutation(len(rows))
    return _csv_text(["student_id", "visit", "slot", "label", "score", "cycle"], [rows[j] for j in order])


@pytest.mark.parametrize("seed", range(12))
def test_staged_prechecks_find_the_same_keys(monkeypatch, seed: int) -> None:
    """The first-IDs prechecks only reject what a full pass rejects too (same keys with and without them)."""
    text = _random_table(seed)
    monkeypatch.setattr(st, "KEY_PRECHECK_ROWS", 10**9)
    full = [k["columns"] for k in _profile(text).structure.composite_keys]
    monkeypatch.setattr(st, "KEY_PRECHECK_ROWS", 10)
    monkeypatch.setattr(st, "KEY_PRECHECK_IDS", 2)
    staged = [k["columns"] for k in _profile(text).structure.composite_keys]
    assert staged == full


def test_large_table_rejects_non_keys_without_a_full_pass(monkeypatch) -> None:
    """Columns that repeat within the first IDs never get a pass over every row (runtime stays flat)."""
    rng = np.random.default_rng(5)
    rows = [[f"S{i + 1000}", v, int(rng.integers(0, 3)), int(rng.integers(0, 2)), "ab"[int(rng.integers(0, 2))]]
            for i in range(400) for v in range(1, 5)]
    monkeypatch.setattr(st, "KEY_PRECHECK_ROWS", 100)
    monkeypatch.setattr(st, "KEY_PRECHECK_IDS", 20)
    full_passes: list[int] = []
    orig = st._KeyColumns._data

    def spy(self, i, stage):  # noqa: ANN001
        if stage < 0 and (i, stage) not in self._data_of:
            full_passes.append(i)
        return orig(self, i, stage)

    monkeypatch.setattr(st._KeyColumns, "_data", spy)
    tp = _profile(_csv_text(["student_id", "visit", "score_a", "score_b", "half"], rows))
    names = {p.index: p.name for p in tp.columns}
    assert [tuple(k["columns"]) for k in tp.structure.composite_keys] == [("student_id", "visit")]
    assert {names[i] for i in full_passes} == {"visit"}


# =========================================================================== R4-05: header rules over native typing
def test_codex_decimal_org_reproduction(tmp_path: Path) -> None:
    """school_id stored as decimal128(10, 0): the 30 real codes with their exact counts, still decimal128."""
    vals = [Decimal((i % 30 + 1) * 100) for i in range(300)]
    src = tmp_path / "in" / "schools.parquet"
    src.parent.mkdir(parents=True)
    pq.write_table(pa.table({"student_id": [f"S{i + 1000}" for i in range(300)],
                             "school_id": pa.array(vals, type=pa.decimal128(10, 0)),
                             "score": [(i * 17) % 91 for i in range(300)]}), src)
    out, report = _run(src, tmp_path)
    assert _types(report)["school_id"] == "ORG_CODE"
    syn = pq.read_table(out)
    assert syn.schema.field("school_id").type == pa.decimal128(10, 0)
    assert Counter(syn["school_id"].to_pylist()) == Counter(vals)


def _typed_table(n: int = 120) -> pa.Table:
    """The same header semantics in every native storage (decimal, int, float, dictionary, ordered dictionary)."""
    return pa.table({
        "student_id": [f"S{1000 + i // 2}" for i in range(n)],
        "school_id": pa.array([Decimal((i % 30 + 1) * 100) for i in range(n)], type=pa.decimal128(10, 0)),
        "course_code": pa.array([Decimal(f"{i % 25 + 101}.00") for i in range(n)], type=pa.decimal128(8, 2)),
        "site": pa.array([(i % 30 + 1) * 100 for i in range(n)], type=pa.int32()),
        "campus_no": pa.array([float((i % 30 + 1) * 10) for i in range(n)], type=pa.float64()),
        "Section": pa.DictionaryArray.from_arrays(pa.array([i % 25 for i in range(n)], type=pa.int8()),
                                                  pa.array([f"Sec {k:02d} - room {k}" for k in range(25)])),
        "Term": pa.DictionaryArray.from_arrays(pa.array([i % 3 for i in range(n)], type=pa.int8()),
                                               pa.array(["Fall", "Winter", "Spring"]), ordered=True),
        "school_year": pa.array([Decimal(2020 + i % 4) for i in range(n)], type=pa.decimal128(4, 0)),
        "grade": pa.array([Decimal(i % 13) for i in range(n)], type=pa.decimal128(2, 0)),
        "dob": pa.array([Decimal(20100101 + i % 28) for i in range(n)], type=pa.decimal128(8, 0)),
        "phone": pa.array([Decimal(3125550100 + i) for i in range(n)], type=pa.decimal128(10, 0)),
        "Notes": pa.array([Decimal(i % 7) for i in range(n)], type=pa.decimal128(3, 0)),
        "flag01": pa.array([Decimal(i % 2) for i in range(n)], type=pa.decimal128(1, 0)),
        "balance": pa.array([Decimal(f"{i * 3 % 97}.25") for i in range(n)], type=pa.decimal128(10, 2)),
        "Student_Name": pa.DictionaryArray.from_arrays(pa.array([i % 3 for i in range(n)], type=pa.int8()),
                                                       pa.array(NAMES)),
        "Mentor": pa.DictionaryArray.from_arrays(pa.array([(i + 1) % 3 for i in range(n)], type=pa.int8()),
                                                 pa.array(NAMES)),
    })


def _canonical(v) -> str:  # noqa: ANN001
    return "" if v is None else (format(v, "f") if isinstance(v, Decimal) else str(v))


def _inspect_types(path: Path) -> dict[str, str]:
    res = CliRunner().invoke(main, ["inspect", "--mode", "fidelity", str(path), "--json"], catch_exceptions=False)
    data = json.loads(res.output)
    return {c["name"]: c["type"] for c in data["files"][0]["tables"][0]["columns"]}


def test_native_storage_and_text_get_the_same_header_rules(tmp_path: Path) -> None:
    """Uniformity (R4-05): the Parquet columns and the same canonical texts in a CSV get the same type, except
    that a plain native decimal keeps its storage type (DECIMAL where the text is FLOAT)."""
    tab = _typed_table()
    src = tmp_path / "in" / "typed.parquet"
    src.parent.mkdir(parents=True)
    pq.write_table(tab, src)
    text = _write_csv(tmp_path / "in" / "typed.csv", tab.column_names,
                      [[_canonical(tab[c][i].as_py()) for c in tab.column_names] for i in range(tab.num_rows)])
    native, plain = _inspect_types(src), _inspect_types(text)
    assert native == {**plain, "balance": "DECIMAL"}
    assert native["balance"] == "DECIMAL" and plain["balance"] == "FLOAT" and native["Term"] == "ORDINAL"
    assert {native[c] for c in ("school_id", "course_code", "site", "campus_no", "Section")} == {"ORG_CODE"}
    assert native["school_year"] == "SCHOOL_YEAR" and native["grade"] == "ORDINAL" and native["flag01"] == "BOOLEAN"
    assert native["dob"] == "DATE_AS_INT" and native["phone"] == "PHONE"
    assert native["Student_Name"] == native["Mentor"] == "PERSON_NAME"


@pytest.mark.parametrize("fmt", ["parquet", "feather"])
def test_native_columns_written_back_in_their_storage(tmp_path: Path, fmt: str) -> None:
    """Org codes keep their real values and counts in decimal/int/float/dictionary storage; PII never copied."""
    tab = _typed_table()
    src = tmp_path / "in" / f"typed.{fmt}"
    src.parent.mkdir(parents=True)
    (pq.write_table if fmt == "parquet" else pf.write_feather)(tab, src)
    out, report = _run(src, tmp_path)
    syn = pq.read_table(out) if fmt == "parquet" else pf.read_table(out)
    assert syn.schema.equals(tab.schema, check_metadata=False)
    for c in ("school_id", "course_code", "site", "campus_no", "Section", "school_year", "grade", "flag01"):
        assert Counter(syn[c].to_pylist()) == Counter(tab[c].to_pylist()), c
    assert not set(syn["phone"].to_pylist()) & set(tab["phone"].to_pylist())
    assert not set(syn["Student_Name"].to_pylist()) & set(NAMES)
    assert not set(syn["Mentor"].to_pylist()) & set(NAMES)
    assert syn["Term"].type.ordered and syn["Term"].to_pylist().count("Winter") == tab["Term"].to_pylist().count(
        "Winter")
    dob = [int(v) for v in syn["dob"].to_pylist()]
    assert min(dob) == 20100101 and max(dob) == 20100128  # a date (D14 exact extremes), not masked digits


def test_native_dates_and_ordered_levels_stay_before_the_org_code_rule(tmp_path: Path) -> None:
    """Dates and ordinal vocabularies come before the org-code rule in the value order: native ones too."""
    src = tmp_path / "in" / "terms.parquet"
    src.parent.mkdir(parents=True)
    n = 60
    pq.write_table(pa.table({
        "student_id": [f"S{1000 + i}" for i in range(n)],
        "term": pa.array([dt.date(2025, 1, 6) + dt.timedelta(days=7 * (i % 10)) for i in range(n)], pa.date32()),
        "section": pa.DictionaryArray.from_arrays(pa.array([i % 3 for i in range(n)], type=pa.int8()),
                                                  pa.array(["Low", "Medium", "High"]), ordered=True),
    }), src)
    types = _inspect_types(src)
    assert types["term"] == "DATE" and types["section"] == "ORDINAL"


def test_id_headers_on_native_decimals_get_synthetic_ids(tmp_path: Path) -> None:
    """ID header rules are overrides for every storage: decimal student and family numbers are synthesised."""
    n = 80
    real_ids = [Decimal(5000000 + 37 * i) for i in range(n)]
    real_fam = [Decimal(900000 + 11 * (i // 2)) for i in range(n)]
    src = tmp_path / "in" / "ids.parquet"
    src.parent.mkdir(parents=True)
    tab = pa.table({"student_id": pa.array(real_ids, pa.decimal128(12, 0)),
                    "Family_Number": pa.array(real_fam, pa.decimal128(12, 0)),
                    "score": [(i * 7) % 50 for i in range(n)]})
    pq.write_table(tab, src)
    out, report = _run(src, tmp_path)
    types = _types(report)
    assert types["student_id"] == types["Family_Number"] == "IDENTIFIER"
    syn = pq.read_table(out)
    assert syn.schema.equals(tab.schema, check_metadata=False)
    assert not set(syn["student_id"].to_pylist()) & set(real_ids)
    assert not set(syn["Family_Number"].to_pylist()) & set(real_fam)


def test_stata_value_labels_under_org_header(tmp_path: Path) -> None:
    """A value-labelled org code is ORG_CODE with its labels kept; labelled names are never written."""
    n = 120
    src = tmp_path / "in" / "labels.dta"
    src.parent.mkdir(parents=True)
    pd.DataFrame({"student_id": [f"S{1000 + i // 2}" for i in range(n)],
                  "school_code": [i % 30 + 1 for i in range(n)],
                  "status": [i % 3 + 1 for i in range(n)],
                  "Guardian": [i % 3 + 1 for i in range(n)]}).to_stata(
        src, write_index=False, version=118,
        value_labels={"school_code": {k: f"Campus {k:02d}" for k in range(1, 31)},
                      "status": {1: "Enrolled", 2: "Withdrawn", 3: "Graduated"},
                      "Guardian": dict(enumerate(NAMES, start=1))})
    out, report = _run(src, tmp_path)
    types = _types(report)
    assert types["school_code"] == "ORG_CODE" and types["status"] == "FACTOR" and types["Guardian"] == "PERSON_NAME"
    real = pd.read_stata(src)
    syn = pd.read_stata(out)
    assert Counter(syn["school_code"].astype(str)) == Counter(real["school_code"].astype(str))
    with pd.io.stata.StataReader(out) as r:
        labels = r.value_labels()
    assert labels["school_code"] == {k: f"Campus {k:02d}" for k in range(1, 31)}
    assert not set(labels["Guardian"].values()) & set(NAMES)


def test_excel_number_cells_under_org_header(tmp_path: Path) -> None:
    """Excel number cells under an org-code header keep their real codes and stay number cells."""
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["student_id", "school_id", "score"])
    for i in range(90):
        ws.append([f"S{1000 + i}", (i % 30 + 1) * 100, i % 41])
    src = tmp_path / "in" / "book.xlsx"
    src.parent.mkdir(parents=True)
    wb.save(src)
    out, report = _run(src, tmp_path)
    assert _types(report)["school_id"] == "ORG_CODE"
    cells = [row[1] for row in load_workbook(out).active.iter_rows(min_row=2)]
    assert all(c.data_type == "n" for c in cells)
    assert Counter(c.value for c in cells) == Counter((i % 30 + 1) * 100 for i in range(90))


def _fees(n: int = 120) -> tuple[list[Decimal], list[Decimal]]:
    owed = [Decimal(f"{50 + (i * 37) % 400}.{i % 100:02d}") for i in range(n)]
    return owed, [o - Decimal(f"{(i * 13) % 50}.00") for i, o in enumerate(owed)]


def test_declared_rule_on_arrow_decimals_is_built(tmp_path: Path) -> None:
    """Rules compare numbers in every storage: Parquet decimal128 columns are no longer refused (the same data
    in a CSV always worked); the rule holds on every output row and the decimal type is kept."""
    owed, paid = _fees()
    src = tmp_path / "in" / "fees.parquet"
    src.parent.mkdir(parents=True)
    tab = pa.table({"student_id": [f"S{1000 + i}" for i in range(len(owed))],
                    "owed": pa.array(owed, pa.decimal128(10, 2)), "paid": pa.array(paid, pa.decimal128(10, 2))})
    pq.write_table(tab, src)
    cfg = tmp_path / "rules.yaml"
    cfg.write_text('rules: ["paid <= owed"]\n', encoding="utf-8")
    out, report = _run(src, tmp_path, "--config", str(cfg))
    structure = report["tables"][0]["structure"]
    assert structure["declared_rules"][0]["status"] == "applied"
    assert [r["rule"] for r in structure["suggested_rules"]] == ["paid <= owed"]  # suggested for decimals too
    syn = pq.read_table(out)
    assert syn.schema.equals(tab.schema, check_metadata=False)
    assert all(p <= o for p, o in zip(syn["paid"].to_pylist(), syn["owed"].to_pylist()))


def test_rule_view_reads_decimals() -> None:
    owed, _ = _fees()
    view = st.rule_view(np.asarray([str(v) for v in owed], dtype=object),
                        st.ColumnProfile(index=0, name="owed", type=T.DECIMAL, params={"scale": 2, "decimals": 2},
                                         role="row"))
    assert view is not None and view.unit == "number" and view.ok.all()
    assert view.vals.tolist() == [float(v) for v in owed]


def _ctx(values: list, native: dict | None, name: str) -> DetectContext:
    arr = np.empty(len(values), dtype=object)
    arr[:] = values
    s = Settings(mode="fidelity")
    return DetectContext(index=0, name=name, values=arr, file_rel="t.parquet", table=None, format="parquet",
                         options=s.column_options("t.parquet", None, name), settings=s, native=native)


def test_native_factor_turned_org_code_keeps_its_dictionary_facts() -> None:
    from datanonymous.io import arrowio

    field = pa.field("Section", pa.dictionary(pa.int8(), pa.string()))
    native = {"family": "arrow", "kind": "dictionary", "field": field,
              "dictionary_values": ["A1", "B2", "C3", "Unused label"]}
    ctx = _ctx(["A1", "B2", "C3", "A1"] * 5, native, "Section")
    raw = arrowio.detect_native(ctx)
    assert raw.type is T.FACTOR  # the handler's view: a dictionary column
    det = detect_column(ctx)
    assert det.type is T.ORG_CODE and det.source == "native"
    assert det.params["dictionary_size"] == 4 and det.params["n_levels"] == 3
    assert "D10" in det.reason
