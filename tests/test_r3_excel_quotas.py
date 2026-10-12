"""Round-3 review fixes, Excel storage quotas (R3-05).

R3-05: the reader kept only the most common number format per cell type and one type per distinct value; the
writer then typed new values by guessing, so minority number formats vanished (50 ``0%`` + 50 ``0.00`` cells
became 100 ``0%``) and the text/number mix drifted (50/50 became 56/44); the self-check compared only the
dominant type and format, so both passed. Now (``io.excel_storage``):

* the reader keeps each column's full mix of (cell type, number format) counts, the combinations of each
  distinct value, and the storage of the first non-empty cell;
* the writer gives every type ``min(target, cells that can take it)`` cells, target = the source count
  scaled to the written cells, and deals each type's number formats by exact quota (a value that is also a
  source value takes its source storage first; the first cell keeps the source's first storage);
* the self-check compares the per-type counts and the per-type format counts exactly, so a writer that
  falls back to the dominant type or format FAILs the file.

Variants covered: minority number format (``0%``/``0.00``), minority date format, text format ``@`` vs
General on text cells, numbers stored as text (scores, IDs), a grade level stored as text and as a number,
booleans mixed with TRUE/FALSE text, error cells, ISO text dates beside date cells, NA tokens in a number
column, fewer written cells than source cells (scaled quota), values that cannot take a type (no false
refusal), footer totals over numbers stored as text, determinism.
"""

from __future__ import annotations

import collections
import datetime as dt
import io
import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from conftest import run_cli
from datanonymous.io import excel_html
from datanonymous.io import excel_storage as xs
from datanonymous.io import excel_write as xw
from datanonymous.io.excel_read import ColBuf, XCell, finalize_column

N = 120
HEADER = ["student_id", "pct_or_dec", "score_mixed", "entry_date", "notes_fmt", "flag", "grade", "local_id",
          "test_score", "when_mixed", "attendance"]
GRADES = ["KG"] * 12 + [str(g) for g in range(1, 6) for _ in range(20)] + ["PK"] * 8


def _run(inp: Path, tmp: Path, *flags: str, seed: str = "5", tag: str = "") -> tuple[int, Path, Path]:
    out, rep = tmp / f"out{tag}", tmp / f"rep{tag}"
    res = run_cli(inp, "--out-dir", out, "--report-dir", rep, "--seed", seed, *flags)
    return res.exit_code, out, rep


def _file_report(rep: Path, name: str) -> dict:
    return json.loads((rep / "files" / f"{name}.json").read_text(encoding="utf-8"))


def _checks(rep: Path, name: str) -> dict[str, str]:
    files = json.loads((rep / "report.json").read_text(encoding="utf-8"))["verdict"]["files"]
    f = next(x for x in files if x["source"] == name)
    return {k: v["result"] for k, v in f["checks"].items()}


def _mix(path: Path, header: list[str], first_row: int = 2, n: int = N) -> dict[str, collections.Counter]:
    ws = openpyxl.load_workbook(path).active
    out = {}
    for j, h in enumerate(header, 1):
        out[h] = collections.Counter((ws.cell(r, j).data_type, ws.cell(r, j).number_format)
                                     for r in range(first_row, first_row + n) if ws.cell(r, j).value is not None)
    return out


def _levels(path: Path, col: int) -> collections.Counter:
    ws = openpyxl.load_workbook(path).active
    return collections.Counter((str(ws.cell(r, col).value), ws.cell(r, col).data_type) for r in range(2, N + 2))


def _variants_book(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws.append(HEADER)
    for i in range(N):
        r = i + 2
        ws.cell(r, 1, f"S{i:05d}")
        ws.cell(r, 2, (i % 37) + 0.25).number_format = "0%" if i % 3 == 0 else "0.00"
        ws.cell(r, 3, (i / 10) if i % 2 else str(round(i / 10 + 0.05, 1)))  # numbers stored as text
        ws.cell(r, 3).number_format = "0.00" if i % 2 else "@"
        ws.cell(r, 4, dt.datetime(2024, 1, 1) + dt.timedelta(days=i)).number_format = "m/d/yyyy" if i % 4 else \
            "yyyy-mm-dd"
        ws.cell(r, 5, f"note {i % 30}").number_format = "@" if i % 5 == 0 else "General"
        ws.cell(r, 6, (i % 2 == 0) if i % 3 else ("TRUE" if i % 2 == 0 else "FALSE"))
        g = GRADES[i]
        ws.cell(r, 7, g if g in ("KG", "PK") or (g == "3" and i % 2) else int(g))  # level 3: half text
        ws.cell(r, 8, 100000 + i * 7 if i % 4 else str(100000 + i * 7))  # an ID stored as text for 1/4
        if i % 8 == 0:
            ws.cell(r, 9, "#N/A").data_type = "e"
        else:
            ws.cell(r, 9, 50 + i % 50)
        when = dt.datetime(2023, 9, 1) + dt.timedelta(days=i)
        if i % 5 == 0:
            ws.cell(r, 10, when.strftime("%Y-%m-%d"))  # ISO text beside date cells
        else:
            ws.cell(r, 10, when).number_format = "mm/dd/yyyy"
        ws.cell(r, 11, "N/A" if i % 10 == 0 else 80 + i % 20).number_format = "0" if i % 10 else "General"
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def variants(tmp_path_factory: pytest.TempPathFactory):
    base = tmp_path_factory.mktemp("r3_quota")
    src = _variants_book((base / "in").mkdir() or base / "in" / "roster.xlsx")
    code, out, rep = _run(src.parent, base)
    assert code == 0
    return src, out / "roster.xlsx", rep


# =========================================================================== the reproduction
def _codex_book(path: Path, kind: str) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    if kind == "formats":
        ws.append(["student_id", "score", "grade"])
        for i in range(100):
            ws.append([f"S{i:06}", i + 0.25, i % 4 + 5])
            ws.cell(i + 2, 2).number_format = "0.00" if i % 2 else "0%"
    else:
        ws.append(["student_id", "score", "first_name"])
        for i in range(100):
            ws.append([f"S{i:05}", i / 10 if i % 2 else str(i / 10), ["Zelmira", "Althea"][i % 2]])
            ws.cell(i + 2, 2).number_format = "0.00" if i % 2 else "@"
    wb.save(path)
    return path


@pytest.mark.parametrize("kind, want", [
    ("formats", {("n", "0%"): 50, ("n", "0.00"): 50}),
    ("mixed", {("s", "@"): 50, ("n", "0.00"): 50}),
])
def test_codex_reproduction_keeps_the_exact_mix(tmp_path: Path, kind: str, want: dict) -> None:
    src = _codex_book((tmp_path / "in").mkdir() or tmp_path / "in" / "book.xlsx", kind)
    code, out, rep = _run(src.parent, tmp_path, seed="123")
    assert code == 0
    assert _mix(src, ["id", "score"], n=100)["score"] == want
    assert _mix(out / "book.xlsx", ["id", "score"], n=100)["score"] == want
    assert _checks(rep, "book.xlsx")["layout_selfcheck"] == "PASS"


# =========================================================================== every variant, exact
@pytest.mark.parametrize("col", HEADER)
def test_every_variant_keeps_its_exact_mix(variants, col: str) -> None:
    src, out, _ = variants
    assert _mix(out, HEADER)[col] == _mix(src, HEADER)[col]


def test_levels_keep_their_source_storage(variants) -> None:
    src, out, _ = variants
    want = _levels(src, 7)
    assert want[("3", "s")] == want[("3", "n")] == 10 and want[("KG", "s")] == 12
    assert _levels(out, 7) == want  # KG/PK text, 1-5 numbers, level 3 half text and half numbers


def test_variants_pass_every_check_and_pandas_dtypes(variants) -> None:
    src, out, rep = variants
    checks = _checks(rep, "roster.xlsx")
    assert checks["layout_selfcheck"] == "PASS" and checks["dtypes"] == "PASS", checks
    assert pd.read_excel(src).dtypes.to_dict() == pd.read_excel(out).dtypes.to_dict()
    # the first cell keeps the source's first storage (pandas types a bool/TRUE-text mix by it)
    s_ws, o_ws = openpyxl.load_workbook(src).active, openpyxl.load_workbook(out).active
    assert s_ws.cell(2, 6).data_type == o_ws.cell(2, 6).data_type == "s"


def test_variants_are_deterministic(variants, tmp_path: Path) -> None:
    src, out, _ = variants
    code, out2, _ = _run(src.parent, tmp_path)
    assert code == 0 and (out2 / "roster.xlsx").read_bytes() == out.read_bytes()


# =========================================================================== the self-check
def _dominant_plan(values, nat, rng):
    """The old writer: every value typed alone, every type with its most common format."""
    strings = [None if v in (None, "") else str(v) for v in values]
    codes = np.full(len(strings), -1, dtype=np.int32)
    plan = xs.ColumnPlan(strings, codes, [], (nat or {}).get("date_kind"))
    for i, s in enumerate(strings):
        if s is None:
            continue
        _, t = xw.typed_value(s, nat)
        combo = (t, ((nat or {}).get("formats") or {}).get(t))
        if combo not in plan.combos:
            plan.combos.append(combo)
        codes[i] = plan.combos.index(combo)
    return plan


@pytest.mark.parametrize("kind", ["formats", "mixed"])
def test_self_check_fails_a_dominant_only_writer(tmp_path: Path, monkeypatch, kind: str) -> None:
    src = _codex_book((tmp_path / "in").mkdir() or tmp_path / "in" / "book.xlsx", kind)
    monkeypatch.setattr(xw, "plan_column", _dominant_plan)
    code, out, rep = _run(src.parent, tmp_path, seed="123")
    assert code == 3 and not (out / "book.xlsx").exists()
    reasons = " ".join(json.loads((rep / "report.json").read_text(encoding="utf-8"))["failures"][0]["reasons"])
    assert "self-check" in reasons
    if kind == "formats":
        assert "'score': number formats of the number cells {'0%': 100}, expected {'0%': 50, '0.00': 50}" in reasons
    else:
        assert "'score':" in reasons and "number cell(s) of 100, expected 50" in reasons


def _nat(storage: dict, **kw) -> dict:
    types = {t: sum(f.values()) for t, f in storage.items()}
    return {"kind": "excel", "types": types, "storage": storage, "date_kind": kw.get("date_kind"),
            "errors": kw.get("errors", []), "formats": {t: max(f, key=f.get) for t, f in storage.items()},
            **({"value_storage": kw["value_storage"]} if "value_storage" in kw else {})}


def _plan_counts(plan: xs.ColumnPlan) -> collections.Counter:
    return collections.Counter(plan.combos[c] for c in plan.codes.tolist() if c >= 0)


def test_plan_scales_the_quota_to_the_written_cells() -> None:
    nat = _nat({"n": {"0%": 50, "0.00": 50}})
    values = [str(i) for i in range(30)] + [""] * 70
    plan = xs.plan_column(np.array(values, dtype=object), nat, np.random.default_rng(1))
    assert _plan_counts(plan) == {("n", "0%"): 15, ("n", "0.00"): 15}
    nat = _nat({"s": {"@": 30}, "n": {"General": 70}})
    plan = xs.plan_column(np.array([str(i) for i in range(10)], dtype=object), nat, np.random.default_rng(1))
    assert _plan_counts(plan) == {("s", "@"): 3, ("n", "General"): 7}


def test_plan_never_changes_a_values_text() -> None:
    """In a column that also holds text cells, '02134', '1.0' or '+5' stay text whatever the quota says."""
    nat = _nat({"s": {"General": 1}, "n": {"General": 9}})
    values = np.array(["02134", "1.0", "+5", "7", "8"], dtype=object)
    plan = xs.plan_column(values, nat, np.random.default_rng(3))
    cells = [plan.cell(i) for i in range(5)]
    assert [t for _, t, _ in cells] == ["s", "s", "s", "n", "n"]
    assert [v for v, _, _ in cells] == ["02134", "1.0", "+5", 7, 8]


def test_plan_without_storage_mix_matches_the_single_cell_typing() -> None:
    """A one-type column (the common case) and an older native description type every value as before."""
    nat = {"kind": "excel", "types": {"n": 5, "e": 1}, "main": "n", "formats": {"n": "0.00"},
           "date_kind": None, "errors": ["#N/A"]}
    values = ["12", "3.25", "#N/A", "*", "", None]
    plan = xs.plan_column(np.array(values, dtype=object), nat, np.random.default_rng(0))
    assert [plan.cell(i)[:2] for i in range(6)] == [xw.typed_value(v, nat) for v in values]
    assert plan.cell(0)[2] == "0.00"


def test_value_storage_is_kept_in_proportion() -> None:
    vs = {"5": collections.Counter({("s", "General"): 10, ("n", "General"): 20}),
          "KG": collections.Counter({("s", "General"): 10})}
    nat = _nat({"s": {"General": 20}, "n": {"General": 20}}, value_storage=vs)
    values = np.array(["5"] * 30 + ["KG"] * 10, dtype=object)
    plan = xs.plan_column(values, nat, np.random.default_rng(9))
    got = collections.Counter((values[i], plan.combos[c][0]) for i, c in enumerate(plan.codes.tolist()))
    assert got == {("5", "s"): 10, ("5", "n"): 20, ("KG", "s"): 10}


def test_reader_keeps_the_full_storage_mix() -> None:
    buf = ColBuf()
    cells = [XCell(0.25, "n", "0%"), XCell(0.5, "n", "0.00"), XCell("x", "s", "@"), XCell(0.75, "n", "0.00"),
             XCell("x", "s", "General")]
    for i, c in enumerate(cells):
        buf.put(i, c)
    _, nat = finalize_column(buf, 6)
    assert nat["storage"] == {"n": {"0%": 1, "0.00": 2}, "s": {"@": 1, "General": 1}}
    assert nat["formats"] == {"n": "0.00", "s": "@"} and nat["first"] == ["n", "0%"]
    assert nat["value_storage"]["x"] == {("s", "@"): 1, ("s", "General"): 1}
    assert "cell_types" not in nat
    _, nat2 = finalize_column(buf, 6, cell_types=True)
    assert nat2["cell_types"] == b"nnsns\x00"


def test_storage_issues_compare_types_and_formats_exactly() -> None:
    src = {"cells": {"s": {"@": 50}, "n": {"0.00": 50}}, "date_kind": None, "errors": []}
    values = [str(i) for i in range(100)]
    ok = {"storage": {"s": {"@": 50}, "n": {"0.00": 50}}, "cell_types": b"s" * 50 + b"n" * 50}
    assert xs.storage_issues("score", src, values, ok) == []
    all_text = {"storage": {"s": {"@": 100}}, "cell_types": b"s" * 100}
    assert xs.storage_issues("score", src, values, all_text) == [
        "column 'score': 0 number cell(s) of 100, expected 50 (the source holds 50 of 100)"]
    wrong_fmt = {"storage": {"s": {"@": 50}, "n": {"General": 50}}, "cell_types": b"s" * 50 + b"n" * 50}
    assert xs.storage_issues("score", src, values, wrong_fmt) == [
        "column 'score': number formats of the number cells {'General': 50}, expected {'0.00': 50}"]
    invented = {"storage": {"s": {"@": 50}, "n": {"0.00": 49}, "d": {"yyyy-mm-dd": 1}},
                "cell_types": b"s" * 50 + b"n" * 49 + b"d"}
    assert any("1 date cell(s), a type the source column does not hold" in i
               for i in xs.storage_issues("score", src, values, invented))


def test_values_that_cannot_take_a_type_are_not_a_layout_failure() -> None:
    """Free-text labels where 10% of the source cells were numbers: the writer cannot store text as numbers,
    so the self-check does not fail the file for it (no false refusal); the dominant type is unchanged."""
    src = {"cells": {"s": {"General": 90}, "n": {"General": 10}}, "date_kind": None, "errors": []}
    values = [f"lorem ipsum {i}" for i in range(100)]
    nat = _nat(src["cells"])
    plan = xs.plan_column(np.array(values, dtype=object), nat, np.random.default_rng(0))
    assert _plan_counts(plan) == {("s", "General"): 100}
    out = {"storage": {"s": {"General": 100}}, "cell_types": b"s" * 100}
    assert xs.storage_issues("note", src, values, out) == []


def test_cell_option() -> None:
    strict, loose = {"s": 1, "n": 1, "d": 1, "b": 1, "e": 1}, {"n": 1, "d": 1, "b": 1}
    assert xs.cell_option("#N/A", strict, True, "date", {"#N/A"}) == "e"
    assert xs.cell_option("2134", strict, True, "date", set()) == "n"
    assert xs.cell_option("02134", strict, True, "date", set()) is None
    assert xs.cell_option("02134", loose, False, "date", set()) == "n"
    assert xs.cell_option("2024-01-05", strict, True, "date", set()) == "d"
    assert xs.cell_option("2024-01-05", strict, True, "datetime", set()) is None  # reads back with a time
    assert xs.cell_option("2024-01-05 08:30:00", strict, True, "datetime", set()) == "d"
    assert xs.cell_option("20240105", {"s": 1, "d": 1}, True, "date", set()) is None  # a text ID, not a date
    assert xs.cell_option("2024-01-05", loose, False, "datetime", set()) == "d"
    assert xs.cell_option("true", strict, True, None, set()) is None
    assert xs.cell_option("true", loose, False, None, set()) == "b"
    assert xs.cell_option("Kindergarten", strict, True, None, set()) is None


def test_scaled_is_exact_and_deterministic() -> None:
    assert xs.scaled({"a": 50, "b": 50}, 100) == {"a": 50, "b": 50}
    assert xs.scaled({"a": 1, "b": 1, "c": 1}, 2) == {"a": 1, "b": 1, "c": 0}
    assert sum(xs.scaled({"0%": 40, "0.00": 80}, 77).values()) == 77


# =========================================================================== footer totals
def test_footer_total_skips_numbers_stored_as_text(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["student_id", "points"])
    for i in range(40):
        ws.append([f"S{i:04d}", (i % 9) + 1 if i % 4 else str((i % 9) + 1)])
    ws.append([])
    ws.append(["Total", "=SUM(B2:B41)"])
    src = (tmp_path / "in").mkdir() or tmp_path / "in" / "points.xlsx"
    wb.save(src)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0 and _checks(rep, "points.xlsx")["layout_selfcheck"] == "PASS"
    o = openpyxl.load_workbook(out / "points.xlsx").active
    cells = [o.cell(r, 2) for r in range(2, 42)]
    assert collections.Counter(c.data_type for c in cells) == {"n": 30, "s": 10}
    assert o.cell(43, 2).value == sum(c.value for c in cells if c.data_type == "n")


# =========================================================================== no false refusals
def test_plain_single_type_workbook_is_unchanged_in_behaviour(tmp_path: Path) -> None:
    """A typical SIS export (one type and one format per column) still passes with the same storage."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Student_Number", "Grade_Level", "Entry_Date", "GPA", "Zip"])
    for i in range(60):
        ws.append([4100 + i * 13, 9 + i % 4, dt.datetime(2024, 8, 1) + dt.timedelta(days=i % 30),
                   round(1.5 + (i * 37 % 250) / 100, 2), f"{2100 + i % 50:05d}"])
        ws.cell(i + 2, 3).number_format = "mm/dd/yyyy"
        ws.cell(i + 2, 4).number_format = "0.00"
        ws.cell(i + 2, 5).number_format = "@"
    src = (tmp_path / "in").mkdir() or tmp_path / "in" / "sis.xlsx"
    wb.save(src)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0
    hdr = ["Student_Number", "Grade_Level", "Entry_Date", "GPA", "Zip"]
    assert _mix(out / "sis.xlsx", hdr, n=60) == _mix(src, hdr, n=60)
    assert _checks(rep, "sis.xlsx")["layout_selfcheck"] == "PASS"


def test_excel_fixtures_still_pass(tmp_path: Path) -> None:
    import shutil

    fix = Path(__file__).resolve().parent / "fixtures" / "excel"
    inp = tmp_path / "in"
    inp.mkdir()
    for name in ("multi_sheet.xlsx", "adversarial.xlsx", "real_biff.xls"):
        shutil.copy(fix / name, inp / name)
    code, out, rep = _run(inp, tmp_path, seed="7")
    assert code == 0
    files = json.loads((rep / "report.json").read_text(encoding="utf-8"))["verdict"]["files"]
    assert {f["source"]: f["checks"]["layout_selfcheck"]["result"] for f in files} == {
        "multi_sheet.xlsx": "PASS", "adversarial.xlsx": "PASS", "real_biff.xls": "PASS"}
    # the error cells beside the numbers of multi_sheet.xlsx keep their exact count
    o = openpyxl.load_workbook(out / "multi_sheet.xlsx")["Enrollment"]
    s = openpyxl.load_workbook(inp / "multi_sheet.xlsx")["Enrollment"]

    def errors(ws) -> int:
        return sum(1 for row in ws.iter_rows() for c in row if c.data_type == "e")

    assert errors(o) == errors(s) > 0


def test_readable_by_pandas(variants) -> None:
    _, out, _ = variants
    df = pd.read_excel(io.BytesIO(out.read_bytes()))
    assert len(df) == N and list(df.columns) == HEADER


# =========================================================================== HTML "Excel" exports (same class)
def _html_export(path: Path) -> Path:
    """An SIS HTML table saved as .xls whose columns mix cell templates (class/style = Excel number format)."""
    rows = []
    for i in range(80):
        sid = (f"<td style=\"mso-number-format:'\\@'\">{4100 + i * 13:06d}</td>" if i % 4
               else f"<td>{4100 + i * 13}</td>")
        gpa = (f'<td class="num">{1.5 + (i * 37 % 250) / 100:.2f}</td>' if i % 5
               else f'<td class="pct">{(i % 9) / 10:.1f}</td>')
        grade = f"<td>{9 + i % 4}</td>" if i % 3 else '<td class="blank"></td>'
        rows.append(f"<tr>{sid}<td>Name{i}</td>{grade}{gpa}</tr>")
    path.write_text('<html><head><style>.num { mso-number-format:"0.00"; } .pct { mso-number-format:"0%"; }'
                    "</style></head><body><table>\n<tr><th>Student_Number</th><th>LastFirst</th>"
                    "<th>Grade_Level</th><th>GPA</th></tr>\n" + "\n".join(rows) + "\n</table></body></html>",
                    encoding="utf-8")
    return path


def _html_templates(path: Path) -> list[collections.Counter]:
    import re

    text = path.read_text(encoding="utf-8")
    cols: list[collections.Counter] = [collections.Counter() for _ in range(4)]
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.S)[1:]:
        for j, (tag, body) in enumerate(re.findall(r"(<td[^>]*>)(.*?)</td>", tr, re.S)):
            cols[j][(tag, excel_html.cell_class(body))] += 1
    return cols


def test_html_cell_templates_keep_their_mix_per_content_class(tmp_path: Path) -> None:
    src = _html_export((tmp_path / "in").mkdir() or tmp_path / "in" / "export.xls")
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 0 and _checks(rep, "export.xls")["layout_selfcheck"] == "PASS"
    s, o = _html_templates(src), _html_templates(out / "export.xls")
    # zero-padded IDs keep the text-format style, unpadded ones the plain cell (the old writer gave every
    # cell the most common template); blank grade cells keep their own class
    assert o[0] == s[0] == {("<td style=\"mso-number-format:'\\@'\">", "number_text"): 60, ("<td>", "number"): 20}
    assert o[2] == s[2]
    # GPA: both styles kept (0.00 and 0%); within each content class the exact source mix
    styles = collections.Counter(tag for (tag, _), c in o[3].items() for _ in range(c))
    assert set(styles) == {'<td class="num">', '<td class="pct">'} and abs(styles['<td class="pct">'] - 16) <= 2
    code, out2, _ = _run(src.parent, tmp_path, tag="2")
    assert code == 0 and (out2 / "export.xls").read_bytes() == (out / "export.xls").read_bytes()  # D26


def test_html_self_check_fails_a_dominant_template_writer(tmp_path: Path, monkeypatch) -> None:
    src = _html_export((tmp_path / "in").mkdir() or tmp_path / "in" / "export.xls")

    def dominant(ct, values, rng):
        return [ct.overall.most_common(1)[0][0]] * len(values)

    monkeypatch.setattr(excel_html, "plan_templates", dominant)
    code, out, rep = _run(src.parent, tmp_path)
    assert code == 3 and not (out / "export.xls").exists()
    reasons = " ".join(json.loads((rep / "report.json").read_text(encoding="utf-8"))["failures"][0]["reasons"])
    assert ("'Student_Number': 20 number cell(s) written with 1 cell style(s) [20], expected 1 [20] "
            "(the source mix of cell styles; not the source's styles)") in reasons
    assert "'GPA':" in reasons and "'Grade_Level': 27 empty cell(s)" in reasons
    assert "mso-number-format" not in reasons and "class=" not in reasons  # templates are never quoted


def test_html_uniform_columns_are_written_as_before() -> None:
    """A column with one template keeps it on every row (the SIS fixture: one style per column)."""
    ct = excel_html.ColumnTemplates()
    ct.overall["td"] = 5
    assert excel_html.plan_templates(ct, ["a", "", None, "1", "2"], np.random.default_rng(0)) == ["td"] * 5
