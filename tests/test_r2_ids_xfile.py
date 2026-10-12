"""Round-2 review fixes, ids/cross-file cluster (R2-01, R2-10).

* R2-01: a column that holds the real IDs of an ID column of the run (another file, another sheet, or its own
  table) but is not an ID column itself was written with its real values. It is now treated as an ID column
  (detection source ``key-ref``): synthetic IDs from the ID column's mapping, so no real ID is written and
  joins are kept. The bug class covered: every column type that writes real values (factor, discrete and
  continuous integers, a constant, an org-code-hinted column), within-ID-constant columns, self references,
  references with orphans, few distinct values, header matches with dense IDs, workbooks; explicit flags and
  PII/text types keep their type (warned); coincidental overlaps (scores, rooms) are not evidence; a column
  that could not be promoted fails its file instead of being written.
* R2-10: within-ID constants (D4) and their blank patterns (D6) were dealt to entities by exact quota without
  regard to ID block sizes, so their row counts changed. Patterns, levels and NA tokens are now dealt so their
  row counts stay as in the source (labels held by fewer than --min-cell IDs are not tied to block sizes), and
  a per-row fidelity check warns when they drift.
"""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook

from conftest import read_dicts, run_cli
from datanonymous.detect.ids import normalize_id
from datanonymous.profile import entity as entity_mod
from datanonymous.profile import structure as st

PAT = [f"S{100000 + i * 113}" for i in range(40)]
DIG = [int(v) for v in np.random.default_rng(5).choice(np.arange(1_000_000, 9_999_999), 120, replace=False)]


def _write(path: Path, header: list[str], rows: list[list[object]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return path


def _run(src: Path, tmp_path: Path, *opts: object, seed: str = "123", tag: str = "") -> tuple[Path, Path]:
    out, rep = tmp_path / f"o{tag}", tmp_path / f"r{tag}"
    res = run_cli(src, "--out-dir", out, "--report-dir", rep, "--seed", seed, *opts)
    assert res.exit_code == 0, res.output
    return out, rep


def _columns(rep: Path, rel: str, table: int = 0) -> dict[str, dict]:
    data = json.loads((rep / "files" / (rel + ".json")).read_text(encoding="utf-8"))
    return {c["name"]: c for c in data["tables"][table]["columns"]}


def _codes(col: dict) -> list[str]:
    return [m["code"] for m in col["messages"]]


def _keys(values: object) -> set[str]:
    return {normalize_id(str(v)) for v in values if str(v).strip()}


def _no_real_anywhere(out: Path, rep: Path, real: list[object]) -> None:
    """No real ID in any written data file, out-dir doc or report file (whole tokens of the CSV cells)."""
    realk = _keys(real)
    for path in [q for q in out.rglob("*.csv")]:
        for row in read_dicts(path):
            assert not _keys(row.values()) & realk, path.name
    texts = [q for q in list(rep.rglob("*")) + list(out.glob("*")) if q.is_file() and q.suffix in (".md", ".json")]
    for path in texts:
        text = path.read_text(encoding="utf-8")
        assert not any(str(v) in text for v in real), path.name


# --------------------------------------------------------------------------- R2-01: the reproduced case
def test_reference_column_holding_real_ids_gets_the_parent_mapping(tmp_path: Path) -> None:
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT[:30])])
    _write(d / "visits.csv", ["subject", "value"], [[PAT[i % 10], i] for i in range(40)])
    out, rep = _run(d, tmp_path)
    people = {r["student_id"] for r in read_dicts(out / "people.csv")}
    refs = [r["subject"] for r in read_dicts(out / "visits.csv")]
    assert not set(refs) & set(PAT)  # no real ID is written
    assert set(refs) <= people  # the join is kept: every reference finds its synthetic person
    assert Counter(Counter(refs).values()) == Counter({4: 10})  # the same rows per referenced ID
    assert all(len(v) == 7 and v.startswith("S") for v in refs)  # same shape (D9)
    col = _columns(rep, "visits.csv")["subject"]
    assert col["type"] == "IDENTIFIER" and col["role"] == "id"
    assert "id_key_ref" in _codes(col) and "id_key_link" in _codes(col)
    _no_real_anywhere(out, rep, PAT)


# --------------------------------------------------------------------------- R2-01: the bug class
def _variant(name: str, d: Path) -> tuple[str, str, list[object], tuple[str, str] | None]:
    """(file, column, its real IDs, (parent file, parent column) or None for a sheet) of one variant."""
    if name == "discrete_integer":  # <= 50 distinct: exact quota over the real values
        _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(DIG)])
        _write(d / "visits.csv", ["visitor", "value"], [[DIG[i % 30], i % 7] for i in range(200)])
        return "visits.csv", "visitor", DIG, ("people.csv", "student_id")
    if name == "continuous_integer":  # many distinct: interpolated, with the exact min and max (D14)
        _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(DIG)])
        _write(d / "visits.csv", ["visitor", "value"], [[DIG[i % 110], i % 7] for i in range(400)])
        return "visits.csv", "visitor", DIG, ("people.csv", "student_id")
    if name == "within_id_constant":  # a D4 column of another ID's table
        staff = [f"T{5000 + i * 37}" for i in range(15)]
        _write(d / "staff.csv", ["staff_id", "dept"], [[s, "A" if i % 2 else "B"] for i, s in enumerate(staff)])
        _write(d / "visits.csv", ["student_id", "advisor", "day"],
               [[PAT[i], staff[i % 15], k] for i in range(30) for k in range(3)])
        return "visits.csv", "advisor", staff, ("staff.csv", "staff_id")
    if name == "self_reference":  # its own table's ID column
        _write(d / "mentors.csv", ["student_id", "mentor", "score"],
               [[DIG[i], DIG[(i * 7) % 40], i % 5] for i in range(40)])
        return "mentors.csv", "mentor", DIG[:40], ("mentors.csv", "student_id")
    if name == "one_value":  # a constant column holding one real ID
        _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT)])
        _write(d / "visits.csv", ["subject", "value"], [[PAT[3], i % 7] for i in range(30)])
        return "visits.csv", "subject", PAT, ("people.csv", "student_id")
    if name == "org_code_header":  # an org-code-hinted header does not make real person IDs a factor
        _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT)])
        _write(d / "visits.csv", ["section", "value"], [[PAT[i % 6], i % 7] for i in range(30)])
        return "visits.csv", "section", PAT, ("people.csv", "student_id")
    if name == "header_match_dense":  # same header as a (file-qualified) ID column, dense IDs 100..139
        _write(d / "a.csv", ["buddy", "x"], [[100 + i, i % 3] for i in range(40)])
        _write(d / "b.csv", ["buddy", "y"], [[100 + i % 30, i % 2] for i in range(50)])
        return "b.csv", "buddy", list(range(100, 140)), ("a.csv", "buddy")
    raise AssertionError(name)


@pytest.mark.parametrize("name", ["discrete_integer", "continuous_integer", "within_id_constant", "self_reference",
                                  "one_value", "org_code_header", "header_match_dense"])
def test_reference_variants_never_write_real_ids_and_keep_the_join(tmp_path: Path, name: str) -> None:
    d = tmp_path / "in"
    rel, col, real, parent = _variant(name, d)
    opts = ("--id-col", "a.csv::buddy") if name == "header_match_dense" else ()
    out, rep = _run(d, tmp_path, *opts)
    vals = [r[col] for r in read_dicts(out / rel)]
    assert not _keys(vals) & _keys(real), name
    assert parent is not None
    par = {r[parent[1]] for r in read_dicts(out / parent[0])}
    assert set(vals) <= par, name  # every reference joins its synthetic parent
    src_counts = Counter(Counter(r[col] for r in read_dicts(d / rel)).values())
    assert Counter(Counter(vals).values()) == src_counts, name  # rows per referenced ID kept (D8)
    info = _columns(rep, rel)[col]
    assert info["type"] == "IDENTIFIER" and "id_key_ref" in _codes(info), name


def test_workbook_sheet_referencing_another_sheet(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["student_id", "score"])
    for i, v in enumerate(PAT):
        ws.append([v, i])
    ws = wb.create_sheet("Visits")
    ws.append(["subject", "value"])
    for i in range(60):
        ws.append([PAT[i % 12], i % 7])
    src = tmp_path / "in" / "book.xlsx"
    src.parent.mkdir(parents=True)
    wb.save(src)
    out, rep = _run(src.parent, tmp_path)
    book = load_workbook(out / "book.xlsx")
    people = {r[0] for r in book["People"].iter_rows(min_row=2, values_only=True)}
    refs = [r[0] for r in book["Visits"].iter_rows(min_row=2, values_only=True)]
    assert not set(refs) & set(PAT) and set(refs) <= people
    assert _columns(rep, "book.xlsx", 1)["subject"]["type"] == "IDENTIFIER"


def test_orphan_references_are_synthetic_and_the_rest_still_join(tmp_path: Path) -> None:
    """70% of the references are known IDs (the rest are people missing from the roster): no real value is
    written, known references join, orphans stay orphans."""
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT[:20])])
    refs = PAT[:14] + [f"S{900000 + i}" for i in range(6)]
    _write(d / "visits.csv", ["subject", "value"], [[refs[i % 20], i % 7] for i in range(60)])
    out, rep = _run(d, tmp_path)
    people = {r["student_id"] for r in read_dicts(out / "people.csv")}
    vals = [r["subject"] for r in read_dicts(out / "visits.csv")]
    assert not _keys(vals) & _keys(refs)
    assert sum(v in people for v in vals) == 42  # 14 known IDs x 3 rows, as in the source
    cov = json.loads((rep / "files" / "visits.csv.json").read_text(encoding="utf-8"))["tables"][0]
    entry = cov["structure"]["cross_file_keys"][0]
    assert entry["column"] == "subject" and entry["orphan_rows"] == 18


def test_explicitly_typed_reference_keeps_its_type_with_a_warning(tmp_path: Path) -> None:
    """--force-factor is the user's explicit choice: the real values are written, and the run says so."""
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT)])
    _write(d / "visits.csv", ["subject", "value"], [[PAT[i % 10], i % 7] for i in range(60)])
    out, rep = _run(d, tmp_path, "--force-factor", "visits.csv::subject")
    assert Counter(r["subject"] for r in read_dicts(out / "visits.csv")) == Counter(PAT[i % 10] for i in range(60))
    col = _columns(rep, "visits.csv")["subject"]
    assert col["type"] == "FACTOR"
    msg = next(m for m in col["messages"] if m["code"] == "unlinked_key_column")
    assert "real values are written" in msg["text"] and msg["fix"] == "--id-col visits.csv::subject"


def test_pii_typed_reference_keeps_its_generator(tmp_path: Path) -> None:
    """An SSN used as a staff ID stays an SSN column elsewhere (invalid area range), never a same-shape ID that
    could be another real SSN."""
    rng = np.random.default_rng(3)
    ssns = [f"{rng.integers(100, 666)}-{rng.integers(10, 99)}-{rng.integers(1000, 9999)}" for _ in range(30)]
    d = tmp_path / "in"
    _write(d / "staff.csv", ["staff_id", "dept"], [[s, "A" if i % 2 else "B"] for i, s in enumerate(ssns)])
    _write(d / "payroll.csv", ["ssn", "amount"], [[ssns[i % 30], 100 + i] for i in range(60)])
    out, rep = _run(d, tmp_path)
    col = _columns(rep, "payroll.csv")["ssn"]
    assert col["type"] == "SSN" and "unlinked_key_column" in _codes(col)
    vals = [r["ssn"] for r in read_dicts(out / "payroll.csv")]
    assert not set(vals) & set(ssns) and all(v[0] == "9" for v in vals)


def test_coincidental_numbers_are_not_references(tmp_path: Path) -> None:
    """Scale scores and room numbers that fall inside a dense ID range are not evidence (chance guard)."""
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(range(100, 400, 3))])
    _write(d / "tests.csv", ["student_id", "scale", "room"],
           [[100 + 3 * (i % 100), 200 + (i * 7) % 200, 1 + i % 80] for i in range(300)])
    _, rep = _run(d, tmp_path)
    cols = _columns(rep, "tests.csv")
    assert cols["scale"]["type"] == "INTEGER" and cols["room"]["type"] == "INTEGER"
    assert not {"id_key_ref", "unlinked_key_column"} & set(_codes(cols["scale"]) + _codes(cols["room"]))


def test_key_ref_share_rules() -> None:
    sparse = {str(v) for v in DIG}
    ints = np.asarray(sorted(DIG), dtype=np.int64)
    # 1-2 values of a sparse 7-digit namespace are evidence; of a dense one they are not
    assert st._key_ref_share({str(DIG[0])}, sparse, ints, False) == 1.0
    dense = {str(v) for v in range(100, 1000)}
    dints = np.arange(100, 1000, dtype=np.int64)
    assert st._key_ref_share({"150", "870"}, dense, dints, False) is None
    assert st._key_ref_share({str(v) for v in range(150, 160)}, dense, dints, False) is None
    # half or more of the values must be keys; a header match needs nothing else
    assert st._key_ref_share({"150", "x1", "x2"}, dense, dints, True) is None
    assert st._key_ref_share({"150", "151", "x2"}, dense, dints, True) == pytest.approx(2 / 3)
    # words: a lone 'unknown' that is also an ID value is no evidence; IDs with digits are
    assert st._key_ref_share({"unknown"}, {"unknown", "s100113"}, np.zeros(0, np.int64), False) is None
    assert st._key_ref_share({"s100113"}, {"unknown", "s100113"}, np.zeros(0, np.int64), False) == 1.0


def test_promotion_without_cross_file_links_still_writes_no_real_id(tmp_path: Path) -> None:
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT[:30])])
    _write(d / "visits.csv", ["subject", "value"], [[PAT[i % 10], i] for i in range(40)])
    cfg = tmp_path / "c.yaml"
    cfg.write_text("structure:\n  disable: [cross_file_keys]\n", encoding="utf-8")
    out, rep = _run(d, tmp_path, "--config", cfg)
    refs = {r["subject"] for r in read_dicts(out / "visits.csv")}
    assert not refs & set(PAT)
    assert not refs & {r["student_id"] for r in read_dicts(out / "people.csv")}  # not linked, as configured
    assert _columns(rep, "visits.csv")["subject"]["type"] == "IDENTIFIER"


def test_inspect_shows_the_reference_as_an_id_column(tmp_path: Path) -> None:
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT[:30])])
    _write(d / "visits.csv", ["subject", "value"], [[PAT[i % 10], i] for i in range(40)])
    res = run_cli("inspect", d, "--json")
    assert res.exit_code == 0, res.output
    data = json.loads(res.output)
    visits = next(f for f in data["files"] if f["source"] == "visits.csv")
    col = next(c for c in visits["tables"][0]["columns"] if c["name"] == "subject")
    assert col["type"] == "IDENTIFIER" and col["role"] == "id"
    assert not any(v in res.output for v in PAT)


def test_a_reference_that_cannot_be_promoted_fails_its_file(tmp_path: Path, monkeypatch) -> None:
    """Defence in depth: if the column could not be made an ID column, its file is not written (D19)."""
    monkeypatch.setattr(entity_mod, "key_ref_detection", lambda ctx, det, ref: det)
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT[:30])])
    _write(d / "visits.csv", ["subject", "value"], [[PAT[i % 10], i] for i in range(40)])
    res = run_cli(d, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "1")
    assert res.exit_code == 3, res.output
    assert not (tmp_path / "o" / "visits.csv").exists() and (tmp_path / "o" / "people.csv").exists()
    entry = json.loads((tmp_path / "r" / "files" / "visits.csv.json").read_text(encoding="utf-8"))
    assert entry["status"] == "FAIL"
    _no_real_anywhere(tmp_path / "o", tmp_path / "r", PAT)


def test_lists_of_real_ids_fail_their_file_with_the_fix(tmp_path: Path) -> None:
    """A multi-valued cell of person IDs ('S1;S2;S3') cannot be one ID per cell; its items would be written as
    they are, so the file is not written (D19) and the fix is named; --force-text then writes invented labels."""
    d = tmp_path / "in"
    _write(d / "people.csv", ["student_id", "score"], [[v, i] for i, v in enumerate(PAT)])
    _write(d / "groups.csv", ["members", "value"],
           [[";".join(PAT[(i + j) % 12] for j in range(3)), i % 5] for i in range(60)])
    res = run_cli(d, "--out-dir", tmp_path / "o", "--report-dir", tmp_path / "r", "--seed", "1")
    assert res.exit_code == 3, res.output
    assert not (tmp_path / "o" / "groups.csv").exists() and (tmp_path / "o" / "people.csv").exists()
    entry = json.loads((tmp_path / "r" / "files" / "groups.csv.json").read_text(encoding="utf-8"))
    assert entry["status"] == "FAIL"
    assert "--force-text groups.csv::members" in json.dumps(entry)
    _no_real_anywhere(tmp_path / "o", tmp_path / "r", PAT)
    out, _ = _run(d, tmp_path, "--force-text", "groups.csv::members", tag="2")
    items = {x for r in read_dicts(out / "groups.csv") for x in r["members"].split(";")}
    assert not items & set(PAT)


def test_reference_runs_are_deterministic(tmp_path: Path) -> None:
    d = tmp_path / "in"
    _variant("within_id_constant", d)
    outs = [(_run(d, tmp_path, tag=str(k))[0] / "visits.csv").read_bytes() for k in range(2)]
    assert outs[0] == outs[1]


# --------------------------------------------------------------------------- R2-10: entity patterns per row
def _pairs(rows: list[dict], a: str, b: str) -> Counter:
    return Counter((r[a] == "", r[b] == "") for r in rows)


def _unequal(path: Path) -> Path:
    """Ten 2-row IDs with blank gender/race and ten 5-row IDs with values (the reproduced case)."""
    return _write(path, ["student_id", "gender", "race", "score"],
                  [[f"S{i + 100}", "" if i < 10 else "F", "" if i < 10 else "groupA", j * 3 + i]
                   for i in range(20) for j in range(2 if i < 10 else 5)])


def test_entity_patterns_keep_their_row_counts_with_unequal_blocks(tmp_path: Path) -> None:
    src = _unequal(tmp_path / "in" / "data.csv")
    out, rep = _run(src, tmp_path)
    syn = read_dicts(out / "data.csv")
    assert _pairs(syn, "gender", "race") == _pairs(read_dicts(src), "gender", "race") \
        == Counter({(True, True): 20, (False, False): 50})
    per = {r["student_id"]: r["gender"] for r in syn}
    assert Counter(v == "" for v in per.values()) == Counter({True: 10, False: 10})  # entity counts too
    codes = [m["code"] for m in json.loads((rep / "files" / "data.csv.json").read_text(encoding="utf-8"))["messages"]]
    assert "fidelity_drift" not in codes


def test_row_level_fidelity_check_catches_unbalanced_blocks(tmp_path: Path, monkeypatch) -> None:
    """With the balancing undone, the per-row check reports the drift the entity-level check cannot see."""
    monkeypatch.setattr(st, "balance_entity_rows", lambda labels, sizes, target, rng, min_entities=1, real=None, **kw: labels)
    src = _unequal(tmp_path / "in" / "data.csv")
    seen = False
    for seed in ("1", "2", "3", "4", "5", "6"):
        out, rep = _run(src, tmp_path, seed=seed, tag=seed)
        blank = sum(r["gender"] == "" for r in read_dicts(out / "data.csv"))
        msgs = json.loads((rep / "files" / "data.csv.json").read_text(encoding="utf-8"))["messages"]
        flagged = any(m["code"] == "fidelity_drift" and m.get("column") == "gender" and "per row" in m["text"]
                      for m in msgs)
        if abs(blank - 20) / 70 > 0.11:  # beyond the limit (2 standard errors at 28.6% of 70 rows)
            assert flagged, seed
            seen = True
        elif blank == 20:
            assert not flagged, seed
    assert seen


def _sized(path: Path, seed: int = 1) -> Path:
    """Random block sizes; blank gender/race mostly on short blocks; F/M and NA tokens tied to sizes too."""
    rng = np.random.default_rng(seed)
    rows = []
    for i, k in enumerate(rng.integers(1, 12, 300)):
        blank = k <= 4
        g = "" if blank else ("F" if k > 6 else "M")
        tok = "NA" if k % 3 == 0 else ""
        dob = tok if k <= 3 else f"20{10 + i % 9}-0{1 + i % 9}-1{i % 9}"
        for j in range(k):
            rows.append([f"S{1000 + i}", g, "" if blank else "grpA", dob, j])
    return _write(path, ["student_id", "gender", "race", "dob", "visit"], rows)


@pytest.mark.parametrize("seed", [1, 2, 3])
def test_levels_tokens_and_patterns_keep_row_counts_with_random_blocks(tmp_path: Path, seed: int) -> None:
    src = _sized(tmp_path / "in" / "data.csv", seed)
    out, rep = _run(src, tmp_path, seed=str(seed))
    real, syn = read_dicts(src), read_dicts(out / "data.csv")
    cols = _columns(rep, "data.csv")
    assert {cols[c]["role"] for c in ("gender", "race", "dob")} == {"entity"}
    for col in ("gender", "race"):
        assert Counter(r[col] for r in syn) == Counter(r[col] for r in real), col  # levels and blanks per row
    for tok in ("NA", ""):
        assert sum(r["dob"] == tok for r in syn) == sum(r["dob"] == tok for r in real), tok  # NA tokens per row
    ents = lambda rows, c: Counter(r[c] for r in {r["student_id"]: r for r in rows}.values())  # noqa: E731
    assert ents(syn, "gender") == ents(real, "gender")  # and per entity (exact quota, D4/D6)


def test_a_rare_level_is_not_tied_to_its_unique_block(tmp_path: Path) -> None:
    """One ID with 50 rows holds the only 'X' (fewer than --min-cell IDs): it is not forced onto the 50-row
    block (that would reproduce a real person's attribute); the per-row drift is reported instead."""
    rows = []
    for i in range(60):
        g = "X" if i == 0 else ("F" if i % 2 else "M")
        rows += [[f"S{1000 + i}", g, j] for j in range(50 if i == 0 else 3)]
    src = _write(tmp_path / "in" / "data.csv", ["student_id", "gender", "visit"], rows)
    forced = 0
    for seed in ("1", "2", "3", "4"):
        out, rep = _run(src, tmp_path, seed=seed, tag=seed)
        syn = read_dicts(out / "data.csv")
        forced += sum(r["gender"] == "X" for r in syn) == 50
    assert forced <= 1


def test_balancing_keeps_entity_counts_and_reaches_feasible_targets() -> None:
    for seed in range(40):
        rng = np.random.default_rng(seed)
        n, k = int(rng.integers(20, 300)), int(rng.integers(2, 6))
        sizes = rng.integers(1, 15, n)
        real = rng.integers(0, k, n)
        target = np.bincount(real, weights=sizes, minlength=k)
        start = rng.permutation(real)
        out = st.balance_entity_rows(start, sizes, target, np.random.default_rng(seed), 1)
        assert np.array_equal(np.bincount(out, minlength=k), np.bincount(real, minlength=k)), seed
        assert np.array_equal(np.bincount(out, weights=sizes, minlength=k), target), seed  # reachable: exact
    # labels tied to block sizes, many entities and levels: still exact, and quick
    rng = np.random.default_rng(1)
    sizes = rng.integers(1, 30, 50_000)
    real = np.minimum(19, (np.argsort(np.argsort(sizes + rng.integers(0, 30, 50_000))) * 20) // 50_000)
    target = np.bincount(real, weights=sizes, minlength=20)
    out = st.balance_entity_rows(rng.permutation(real), sizes, target, np.random.default_rng(2), 10)
    assert np.array_equal(np.bincount(out, weights=sizes, minlength=20), target)


def test_balancing_uses_no_randomness_when_already_exact_and_freezes_rare_labels() -> None:
    sizes = np.full(30, 3)
    labels = np.repeat([0, 1, 2], 10)
    rng = np.random.default_rng(0)
    state = rng.bit_generator.state
    assert np.array_equal(st.balance_entity_rows(labels, sizes, [30, 30, 30], rng, 10), labels)
    assert rng.bit_generator.state == state
    sizes = np.asarray([50] + [3] * 29)
    labels = np.asarray([1, 0] + [1] * 14 + [2] * 14)  # label 0 (1 entity) is rare: it never moves
    out = st.balance_entity_rows(labels, sizes, [50, 45, 42], np.random.default_rng(1), 10)
    assert out[1] == 0 and np.array_equal(np.bincount(out), np.bincount(labels))


def test_equal_blocks_generate_as_before(tmp_path: Path, monkeypatch) -> None:
    """With equal block sizes nothing is re-dealt: the output is byte-identical to a run without balancing."""
    rows = [[f"S{1000 + i}", "" if i % 4 == 0 else ("F" if i % 3 else "M"), "" if i % 4 == 0 else "A", j]
            for i in range(60) for j in range(3)]
    src = _write(tmp_path / "in" / "data.csv", ["student_id", "gender", "race", "visit"], rows)
    with_balance = (_run(src, tmp_path, tag="a")[0] / "data.csv").read_bytes()
    monkeypatch.setattr(st, "balance_entity_rows", lambda labels, sizes, target, rng, min_entities=1, real=None, **kw: labels)
    without = (_run(src, tmp_path, tag="b")[0] / "data.csv").read_bytes()
    assert with_balance == without


def test_unequal_block_runs_are_deterministic(tmp_path: Path) -> None:
    src = _sized(tmp_path / "in" / "data.csv", 4)
    outs = [(_run(src, tmp_path, seed="9", tag=str(k))[0] / "data.csv").read_bytes() for k in range(2)]
    assert outs[0] == outs[1]
